"""Daily Excel interchange for the AI reader and standalone processes.

One Referrals row per extraction; Providers rows preserve repeated records.
All field cells are written as literal text, including identifiers and formulas.
"""
from contextlib import contextmanager
from datetime import datetime
import os
from pathlib import Path
import sys
import tempfile
import time
import uuid
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

try:
    from .referral_schema import REQUIRED_FIELDS, PROVIDER_FIELDS, ReferralFields, completeness, split_claims_manager_name
except ImportError:
    from referral_schema import REQUIRED_FIELDS, PROVIDER_FIELDS, ReferralFields, completeness, split_claims_manager_name

VERSION = "3"
HEADERS = ["Record ID", "Extracted At", "Source File", "NEXT STEP", "Missing Required Fields"] + REQUIRED_FIELDS
PROVIDER_HEADERS = ["Record ID", "Provider Index"] + PROVIDER_FIELDS + ['Date Only', 'Appointment Confirmation']
NEW_FIELDS = {"Compensable Body/Part(s)", "Additional Diagnosis Codes", "Doctor First Name", "Doctor Last Name",
              "Provider Address Line 2", "Employer First Name", "Employer Last Name", "Employer Mobile", "Language", "Special Instructions",
              "Claims Case Manager First Name", "Claims Case Manager Last Name"}
V2_HEADERS = [h for h in HEADERS if h not in ('Claims Case Manager First Name', 'Claims Case Manager Last Name')]
LEGACY_HEADERS = [h for h in HEADERS if h not in NEW_FIELDS]
LEGACY_PROVIDER_HEADERS = [h for h in ['Record ID', 'Provider Index'] + PROVIDER_FIELDS if h not in NEW_FIELDS]
REVIEW_SHEETS = ('Address Review', 'Appointment Review', 'Name Inference')


def output_directory():
    configured = os.getenv("FCM_AI_OUTPUT_DIR", "").strip()
    base = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parents[1]
    return Path(configured).expanduser().resolve() if configured else base / "Output" / "AI"


@contextmanager
def workbook_lock(path):
    lock = path.with_suffix(".xlsx.lock")
    deadline = time.monotonic() + 5
    while True:
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            break
        except FileExistsError:
            if time.monotonic() >= deadline:
                raise OSError("Daily workbook is busy. Retry after the other writer finishes.")
            time.sleep(.1)
    try:
        os.close(fd)
        yield
    finally:
        lock.unlink(missing_ok=True)


def append_text(sheet, values):
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        text = str(value if value is not None else "")
        if len(text) > 32767:
            raise ValueError("An output value exceeds Excel's 32,767-character cell limit; workbook was not changed.")
        cell = sheet.cell(row, column, text)
        cell.data_type = "s"
        cell.number_format = "@"


def check_layout(book):
    if "Schema" not in book or book["Schema"]["B1"].value not in ('1', '2', VERSION):
        raise ValueError("Unsupported daily workbook schema.")
    legacy = book['Schema']['B1'].value == '1'
    referral_headers = LEGACY_HEADERS if legacy else V2_HEADERS if book['Schema']['B1'].value == '2' else HEADERS
    for name, headers in [("Referrals", referral_headers),
                          ("Providers", LEGACY_PROVIDER_HEADERS if legacy else PROVIDER_HEADERS)]:
        if name not in book or [c.value for c in book[name][1]] != headers:
            raise ValueError(f"Unexpected {name} columns; use the AI daily output workbook without renaming headers.")


def save_daily_output(fields, source_file, *, directory=None, record_id=None, extracted_at=None, update_existing=False):
    timestamp = extracted_at or datetime.now().astimezone()
    record_id = record_id or str(uuid.uuid4())
    folder = Path(directory) if directory is not None else output_directory()
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"AI-FCM-Output-{timestamp:%Y-%m-%d}.xlsx"
    providers = getattr(fields, "providers", None)
    if providers is None:
        raise ValueError("Saving requires structured extraction results, including provider records.")
    assessment = completeness(fields)
    with workbook_lock(path):
        book = load_workbook(path) if path.exists() else Workbook()
        temp_path = None
        try:
            if path.exists():
                check_layout(book)
                if book['Schema']['B1'].value != VERSION:
                    for name, headers in [('Referrals', HEADERS), ('Providers', PROVIDER_HEADERS)]:
                        sheet = book[name]
                        old = [cell.value for cell in sheet[1]]
                        rows = [dict(zip(old, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
                        sheet.delete_rows(1, sheet.max_row)
                        append_text(sheet, headers)
                        for row in rows:
                            if name == 'Referrals':
                                split_claims_manager_name(row)
                            append_text(sheet, [row.get(key, 'Not found') for key in headers])
                    book['Schema']['B1'] = VERSION
                    # Refresh old status cells under the new requirements too.
                    provider_rows = {}
                    for row in book['Providers'].iter_rows(min_row=2, values_only=True):
                        values = dict(zip(PROVIDER_HEADERS, row))
                        provider_rows.setdefault(values['Record ID'], []).append(
                            {k: values.get(k, 'Not found') for k in PROVIDER_FIELDS})
                    for cells in book['Referrals'].iter_rows(min_row=2):
                        row = dict(zip(HEADERS, [c.value for c in cells]))
                        previous = ReferralFields({k: row.get(k, 'Not found') for k in REQUIRED_FIELDS},
                                                  provider_rows.get(row['Record ID'], []))
                        _load_reviews(book, row['Record ID'], previous)
                        checked = completeness(previous)
                        cells[3].value = checked['status']
                        cells[4].value = '\n'.join(checked['missing_fields'] + checked['confirmation_required'])
            else:
                book.active.title = "Referrals"
                append_text(book.active, HEADERS)
                append_text(book.create_sheet("Providers"), PROVIDER_HEADERS)
                append_text(book.create_sheet("Schema"), ["Version", VERSION])
            # Retrying a save of the same extraction never appends a duplicate.
            existing = [row[0] for row in book["Referrals"].iter_rows(min_row=2, values_only=True)]
            if record_id in existing and not update_existing:
                return path, record_id
            if record_id in existing:
                for name in ["Referrals", "Providers", *REVIEW_SHEETS]:
                    if name in book:
                        sheet = book[name]
                        for row in range(sheet.max_row, 1, -1):
                            if sheet.cell(row, 1).value == record_id:
                                sheet.delete_rows(row)
            append_text(book["Referrals"], [record_id, timestamp.isoformat(), Path(source_file).name,
                        assessment["status"], "\n".join(assessment["missing_fields"] + assessment['confirmation_required'])]
                        + [fields.get(key, "Not found") for key in REQUIRED_FIELDS])
            for index, provider in enumerate(providers, 1):
                check = assessment['appointments'][index - 1]
                status = 'Confirmed' if check['confirmed'] else 'Required' if check['reasons'] else 'Not required'
                append_text(book["Providers"], [record_id, index] + [provider.get(k, "Not found") for k in PROVIDER_FIELDS]
                            + ['Yes' if check['date_only'] else 'No', status])
            for name in REVIEW_SHEETS:
                reviews = getattr(fields, name.lower().replace(' ', '_'), [])
                if reviews:
                    if name not in book:
                        append_text(book.create_sheet(name), ['Record ID', 'Review JSON'])
                    for review in reviews:
                        append_text(book[name], [record_id, json.dumps(review, ensure_ascii=False)])
            for name in ["Referrals", "Providers"]:
                sheet = book[name]
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24486C")
                    sheet.column_dimensions[cell.column_letter].width = min(45, max(20, len(str(cell.value)) + 2))
            with tempfile.NamedTemporaryFile(dir=folder, suffix=".xlsx", delete=False) as temp:
                temp_path = Path(temp.name)
            book.save(temp_path)
            os.replace(temp_path, path)
        finally:
            book.close()
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)
    return path, record_id


def _load_reviews(book, record_id, fields):
    for name in REVIEW_SHEETS:
        if name in book:
            reviews = []
            for cells in book[name].iter_rows(min_row=2):
                if cells[0].value == record_id:
                    if any(cell.data_type == 'f' for cell in cells):
                        raise ValueError('Formula cells are not supported in review input.')
                    reviews.append(json.loads(cells[1].value))
            setattr(fields, name.lower().replace(' ', '_'), reviews)


def read_record(path, record_id):
    """Shared standalone reader; never executes Excel formulas."""
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        check_layout(book)
        headers = [cell.value for cell in book['Referrals'][1]]
        provider_headers = [cell.value for cell in book['Providers'][1]][2:]
        rows = []
        for cells in book["Referrals"].iter_rows(min_row=2):
            if cells[0].value == record_id:
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Formula cells are not supported in process input.")
                rows.append(dict(zip(headers, [cell.value for cell in cells])))
        if len(rows) != 1:
            raise ValueError("Record ID must identify exactly one Referrals row.")
        record = rows[0]
        providers = []
        for cells in book["Providers"].iter_rows(min_row=2):
            if cells[0].value == record_id:
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Formula cells are not supported in provider input.")
                values = dict(zip(provider_headers, [c.value or 'Not found' for c in cells[2:]]))
                providers.append({key: values.get(key, 'Not found') for key in PROVIDER_FIELDS})
        fields = ReferralFields({key: record.get(key) or "Not found" for key in REQUIRED_FIELDS}, providers)
        split_claims_manager_name(fields)
        _load_reviews(book, record_id, fields)
        return fields, completeness(fields)
    finally:
        book.close()


def reopen_input(fields):
    def value(key):
        text = str(fields.get(key) or "").strip()
        return "" if text.casefold() == "not found" else text
    return {"claimNumber": value("Claim Number"), "customer": value("Employer Name"),
            "claimID": value("Claim ID"), "claimantFull": " ".join(filter(None, [value("First Name"), value("Last Name")])),
            "referralType": value("Referral Type")}
