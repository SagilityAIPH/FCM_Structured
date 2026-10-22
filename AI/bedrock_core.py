"""Shared document extraction and Bedrock logic; no UI dependencies."""
import os
import re
import time
from datetime import date
from io import BytesIO
import pymupdf as fitz
import pandas as pd
from docx import Document
from openai import OpenAI

DEFAULT_BEDROCK_REGION = "us-east-2"
DEFAULT_BEDROCK_MODEL = "openai.gpt-oss-120b"
try:
    from .inference_settings import DEFAULT_MAX_TOKENS, MAX_OUTPUT_TOKENS, check_output_limit
except ImportError:
    from inference_settings import DEFAULT_MAX_TOKENS, MAX_OUTPUT_TOKENS, check_output_limit
DEFAULT_TEMP = 0.0
DEFAULT_MAX_DOC_CHARS = 100000
DEFAULT_REQUEST_TIMEOUT = 180.0


def build_bedrock_base_url(region: str) -> str:
    region = (region or DEFAULT_BEDROCK_REGION).strip()
    return f"https://bedrock-mantle.{region}.api.aws/v1"


def resolve_bedrock_api_key(manual_key: str = "") -> str:
    """
    Priority:
    1. Key entered in the application
    2. OPENAI_API_KEY environment variable
    """
    manual_key = (manual_key or "").strip()
    if manual_key:
        return manual_key

    env_key = os.getenv("OPENAI_API_KEY", "").strip()
    if env_key:
        return env_key

    return ""


def create_bedrock_client(api_key: str, base_url: str) -> OpenAI:
    if not api_key:
        raise ValueError(
            "Bedrock API key is missing. Enter it in the application or set OPENAI_API_KEY."
        )

    return OpenAI(
        api_key=api_key,
        base_url=base_url,
        timeout=DEFAULT_REQUEST_TIMEOUT,
        max_retries=2,
    )


def test_bedrock_connection(client: OpenAI, model_id: str) -> str:
    """Send a tiny inference request so access and model permissions are truly tested."""
    response = client.chat.completions.create(
        model=model_id,
        messages=[
            {
                "role": "user",
                "content": "Reply with exactly: BEDROCK_OK",
            }
        ],
        max_completion_tokens=32,
        temperature=0.0,
    )

    check_output_limit(getattr(response.choices[0], "finish_reason", None), 32)
    return (response.choices[0].message.content or "").strip()


try:
    from .referral_schema import *
except ImportError:
    from referral_schema import *

try:
    from .appointment_rules import reconcile_model_appointment_years
except ImportError:
    from appointment_rules import reconcile_model_appointment_years

# The schema retains its source-literal rule; the live prompt explains advance
# inference. Final stored years still come from appointment_rules after merging.
_APPOINTMENT_SOURCE_RULE = (
    "Return dates as YYYY-MM-DD with an explicit\n"
    "  documented year; do not guess a year. If only month/day is documented, preserve\n"
    "  that literal date (e.g. 3/23); never discard it as Not found. Python resolves a\n"
    "  missing appointment year after source corroboration: current calendar year\n"
    "  unless month/day has passed, then next calendar year (today stays this year)."
)
_APPOINTMENT_ADVANCE_RULE = (
    "Return appointment dates with a documented year as YYYY-MM-DD; never change\n"
    "  that year. A yearless appointment month/day is an ADVANCE date: relative to\n"
    "  the current date, choose the current calendar year if the month/day is today\n"
    "  or still ahead, otherwise the next calendar year. Return that inferred date\n"
    "  as YYYY-MM-DD. Python checks the source and resolves the final stored year\n"
    "  after provider corroboration; do not apply this rule to DOB, DOI or other dates."
)

def extract_pages_from_pdf(file_bytes: bytes):
    pages = []

    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            pages.append({
                "page": page_num,
                "text": text
            })

    return pages


def extract_pages_from_docx(file_bytes: bytes):
    doc = Document(BytesIO(file_bytes))
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"\n===== TABLE {table_index} =====")

        for row in table.rows:
            row_values = []

            for cell in row.cells:
                value = cell.text.strip().replace("\n", " ")
                row_values.append(value)

            parts.append(" | ".join(row_values))

    return [{
        "page": 1,
        "text": "\n".join(parts).strip()
    }]


def extract_pages_from_txt(file_bytes: bytes):
    try:
        text = file_bytes.decode("utf-8", errors="ignore").strip()
    except Exception:
        text = file_bytes.decode("latin-1", errors="ignore").strip()

    return [{
        "page": 1,
        "text": text
    }]


def extract_document_pages(uploaded_file):
    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if file_name.endswith(".pdf"):
        return extract_pages_from_pdf(file_bytes)

    if file_name.endswith(".docx"):
        return extract_pages_from_docx(file_bytes)

    if file_name.endswith(".txt"):
        return extract_pages_from_txt(file_bytes)

    raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")


def join_pages(pages):
    output = []

    for item in pages:
        output.append(f"\n===== PAGE {item['page']} =====\n{item['text']}")

    return "\n".join(output).strip()

def call_llm_once(
    client: OpenAI,
    model_id: str,
    prompt: str,
    max_tokens: int,
    temperature: float,
):
    response = client.chat.completions.create(
        model=model_id,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a strict document field extraction engine. "
                    "Return only the requested JSON object. "
                    "Do not explain. Do not provide reasoning. "
                    "Do not copy blank templates."
                ),
            },
            {
                "role": "user",
                "content": prompt,
            },
        ],
        max_completion_tokens=max_tokens,
        temperature=temperature,
        top_p=0.4,
    )

    check_output_limit(getattr(response.choices[0], "finish_reason", None), max_tokens)
    return (response.choices[0].message.content or "").strip()


def fields_to_table_df(fields):
    return pd.DataFrame(field_rows(fields), columns=["Section", "Field", "Value", "Optional"])


def table_df_to_csv_text(df):
    return df.to_csv(index=False)


def run_reasoning(client, model_id, full_document_text, validation_source_text="",
                  max_tokens=DEFAULT_MAX_TOKENS, temperature=0.0, llm_call=None):
    """Retry malformed/truncated output once; optional missing data is valid."""
    today = date.today()
    rules = FIELD_ONLY_PROMPT
    if _APPOINTMENT_SOURCE_RULE not in rules:
        raise ValueError("Appointment extraction rule is missing from the schema prompt")
    rules = rules.replace(_APPOINTMENT_SOURCE_RULE, _APPOINTMENT_ADVANCE_RULE)
    rules = rules.replace("DOCUMENT:\n{DOCUMENT_TEXT}",
                          f"Today's date for appointment-year inference: {today.isoformat()}.\n\n"
                          "DOCUMENT:\n{DOCUMENT_TEXT}")
    prompt = rules.replace("{DOCUMENT_TEXT}", full_document_text)
    start = time.time()
    invoke = llm_call or call_llm_once
    for attempt in range(2):
        raw = invoke(client=client, model_id=model_id, prompt=prompt,
                     max_tokens=max_tokens, temperature=temperature)
        try:
            source = validation_source_text or full_document_text
            checked = reconcile_model_appointment_years(raw, source, today)
            result, fields = force_exact_field_output(checked, source)
            return result, fields, raw, round(time.time() - start, 2)
        except (ValueError, TypeError) as error:
            if attempt:
                raise ValueError("Bedrock returned invalid or incomplete JSON. Increase output tokens and retry.") from error
            prompt += "\nReturn valid complete JSON with every requested key and provider array. Use Not found for missing facts."
