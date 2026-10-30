"""Local CMS Registration workbook input; no lookup or portal side effects."""

from copy import copy, deepcopy
from datetime import date, datetime, time
from itertools import chain
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .daily_output import HEADERS as DAILY_HEADERS, REVIEW_SHEETS, CLASSIFICATION_SHEET


class ExcelValidationError(ValueError):
    """Operator-safe validation message containing no cell values or file paths."""


SHEET = "CMS Input"
SOURCE_SHEET = "Subject Line Input"
SOURCE_HEADERS = (
    "Claim - Employer / customer", "Claim - Company / market",
    "Claim - Claim number", "Claim - Referral type", "Claim - Date of injury",
    "Claim - Body part", "Claim - Injury type", "Claim - Injury cause",
    "Claimant - First name", "Claimant - Last name", "Claimant - Date of birth",
    "Claimant - Gender", "Claimant - Address line 1", "Claimant - Address line 2",
    "Claimant - City", "Claimant - State", "Claimant - ZIP", "Claimant - Phone",
    "Provider - Facility name", "Provider - Doctor first name",
    "Provider - Doctor last name", "Provider - Address", "Provider - City",
    "Provider - State", "Provider - ZIP", "Provider - Phone",
    "Provider - Appointment date", "Provider - Appointment time",
    "Adjuster - First name", "Adjuster - Last name", "Adjuster - Email",
    "Adjuster - Phone", "Adjuster - Requested nurse", "Adjuster - Referral source",
    "Attorney - First name", "Attorney - Last name", "Attorney - Firm",
    "Attorney - Phone", "Attorney - Address line 1", "Attorney - Address line 2",
    "Attorney - City", "Attorney - State", "Attorney - ZIP",
)
CMS_HEADERS = (
    "CMS - Case ID", "CMS - Referral type", "CMS - Referral number",
    "CMS - CEM employer request", "CMS - Diagnosis code", "CMS - Office number",
    "CMS - Special instructions", "CMS - Claim Identifier", "CMS - Other Notes",
    "CMS - Customer email", "CMS - Insured address line 1",
)
LEGACY_HEADERS = SOURCE_HEADERS + CMS_HEADERS
AI_SOURCE_HEADERS = ("Accurate",) + tuple(DAILY_HEADERS)
AI_CMS_HEADERS = CMS_HEADERS + ("CMS - Body part", "CMS - Referral source")
HEADERS = AI_SOURCE_HEADERS + AI_CMS_HEADERS
_KNOWN = {header.casefold(): header for header in HEADERS + LEGACY_HEADERS}
_SETUP_FIELDS = {
    "dxCode": "CMS - Diagnosis code",
    "bodyParts": "Claim - Body part",
    "officeNumber": "CMS - Office number",
    "xlCustomerName": "Claim - Employer / customer",
    "xlReferralSource": "Adjuster - Referral source",
    "xlProviderPhone": "Provider - Phone",
    "xlCustomerEmail": "CMS - Customer email",
    "xlInsuredAddr1": "CMS - Insured address line 1",
    "xlAttorneyAdd1": "Attorney - Address line 1",
    "xlAttorneyAdd2": "Attorney - Address line 2",
    "xlAttorneyCity": "Attorney - City",
    "xlAttorneyState": "Attorney - State",
    "xlAttorneyZip": "Attorney - ZIP",
    "xlAttorneyPhone": "Attorney - Phone",
    "apptDate": "Provider - Appointment date",
    "Special_Instructions": "CMS - Special instructions",
    "reviewClaimIdentifier": "CMS - Claim Identifier",
    "reviewOtherNotes": "CMS - Other Notes",
    "reviewAppointmentDate": "Provider - Appointment date",
    "reviewAppointmentTime": "Provider - Appointment time",
}

EDITABLE_SETUP_FIELDS = tuple(
    {"key": key, "label": label, **options}
    for key, label, options in (
        ("dxCode", "Diagnosis code", {}),
        ("bodyParts", "CMS body part", {}),
        ("officeNumber", "Office number", {}),
        ("xlCustomerName", "Customer / employer name", {}),
        ("xlReferralSource", "Referral source", {}),
        ("xlProvider", "Primary provider name", {}),
        ("provType", "Primary provider type", {"choices": ["", "doctor", "facility"]}),
        ("xlProviderPhone", "Primary provider phone", {}),
        ("xlCustomerEmail", "Customer email", {}),
        ("xlInsuredAddr1", "Insured address line 1", {}),
        ("xlAttorney", "Attorney name", {}),
        ("xlAttorneyAdd1", "Attorney address line 1", {}),
        ("xlAttorneyAdd2", "Attorney address line 2", {}),
        ("xlAttorneyCity", "Attorney city", {}),
        ("xlAttorneyState", "Attorney state", {}),
        ("xlAttorneyZip", "Attorney ZIP", {}),
        ("xlAttorneyPhone", "Attorney phone", {}),
        ("apptDate", "Appointment date (MM/DD/YYYY)", {}),
        ("Special_Instructions", "Special instructions", {"multiline": True}),
        ("reviewClaimIdentifier", "Claim Identifier (manual review)", {}),
        ("reviewOtherNotes", "Other Notes (manual review)", {"multiline": True}),
        ("reviewAppointmentTime", "Appointment time (manual review)", {}),
    )
)
# apptDate is canonical; reviewAppointmentDate is an import-time review alias.
EDITABLE_PROVIDER_FIELDS = tuple(
    {"key": key, "label": label, **options}
    for key, label, options in (
        ("name", "Provider name", {}),
        ("type", "Provider type", {"choices": ["", "doctor", "facility"]}),
        ("first_name", "Doctor first name", {}),
        ("last_name", "Doctor last name", {}),
        ("phone", "Phone", {}),
        ("address", "Address line 1", {}),
        ("address2", "Address line 2", {}),
        ("city", "City", {}),
        ("state", "State", {}),
        ("zip", "ZIP", {}),
    )
)


def _validate_strings(values, descriptors):
    allowed = {field["key"]: field for field in descriptors}
    if not isinstance(values, dict):
        raise ExcelValidationError("Setup edits must be an object.")
    if any(not isinstance(key, str) or key not in allowed for key in values):
        raise ExcelValidationError("Unsupported setup field.")
    for key, value in values.items():
        if not isinstance(value, str):
            raise ExcelValidationError(f"{key}: enter a text value.")
        if "choices" in allowed[key] and value not in allowed[key]["choices"]:
            raise ExcelValidationError(f"{key}: choose a listed value.")
    return deepcopy(values)


def validate_setup_edits(values):
    """Validate partial edits without allowing request or provenance overwrites."""
    if not isinstance(values, dict):
        raise ExcelValidationError("Setup edits must be an object.")
    root = {key: value for key, value in values.items() if key != "providers"}
    result = _validate_strings(root, EDITABLE_SETUP_FIELDS)
    if "providers" in values:
        providers = values["providers"]
        if not isinstance(providers, list) or not providers:
            raise ExcelValidationError("Providers must be a nonempty list of objects.")
        result["providers"] = [_validate_strings(provider, EDITABLE_PROVIDER_FIELDS)
                               for provider in providers]
    return result


def _effective(value):
    return "" if value.strip().casefold() in ("", "not found", "unknown", "n/a", "none", "null") else value


def _text(cell, *, preserve_whitespace=False):
    if cell.data_type in ("f", "e"):
        raise ExcelValidationError(f"{cell.coordinate}: replace formulas or Excel errors with plain values.")
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, time):
        return value.strftime("%I:%M %p")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value).is_integer():
            digits = str(int(value))
            if re.fullmatch(r"0+", cell.number_format):
                return digits.zfill(len(cell.number_format))
            return digits
    return str(value) if preserve_whitespace else str(value).strip()


def _open(path):
    if Path(path).suffix.lower() != ".xlsx":
        raise ExcelValidationError("Choose an .xlsx workbook.")
    try:
        return load_workbook(path, read_only=True, data_only=False)
    except Exception:
        raise ExcelValidationError("Could not read the Excel workbook. Check the file and try again.") from None


def _sheet(book, name):
    if name not in book.sheetnames:
        raise ExcelValidationError(f'Missing "{name}" sheet. Use the matching input template.')
    sheet = book[name]
    if sheet.max_row > 10001 or sheet.max_column > 200:
        raise ExcelValidationError("Use a workbook with up to 10,000 input rows and 200 columns.")
    return sheet


def _headers(cells):
    headers = [_text(cell) for cell in cells]
    seen = set()
    for cell, header in zip(cells, headers):
        if header and header.casefold() in seen:
            raise ExcelValidationError(f"{cell.coordinate}: duplicate column header.")
        if header:
            seen.add(header.casefold())
    return headers


def _record(number, supplied, coordinates):
    def get(header):
        value = _effective(supplied.get(header, ""))
        return value if header in ("CMS - Special instructions", "CMS - Other Notes") else value.strip()
    explicit_type = get("CMS - Referral type")
    referral_type = explicit_type or get("Claim - Referral type")
    referral_type = {"medical": "Medical", "vocational": "Vocational"}.get(
        referral_type.casefold(), "")
    if explicit_type and not referral_type:
        raise ExcelValidationError(
            f'{coordinates["CMS - Referral type"]}: CMS referral type must be Medical or Vocational.')
    cem = get("CMS - CEM employer request").casefold() or "auto"
    if cem not in ("yes", "no", "auto"):
        raise ExcelValidationError(
            f'{coordinates["CMS - CEM employer request"]}: CEM employer request must be yes, no or auto.')
    form = {
        "claim_number": get("Claim - Claim number"),
        "claimant_first_name": get("Claimant - First name"),
        "claimant_last_name": get("Claimant - Last name"),
        "referral_type": referral_type,
        "cms_case_id": get("CMS - Case ID"),
        "referral_number": get("CMS - Referral number"),
        "cem_employer_request": cem,
    }
    data = {key: get(header) for key, header in _SETUP_FIELDS.items()}
    facility = get("Provider - Facility name")
    data["xlProvider"] = facility or " ".join(filter(None, (
        get("Provider - Doctor first name"), get("Provider - Doctor last name"))))
    data["provType"] = "facility" if facility else ""
    data["xlAttorney"] = " ".join(filter(None, (
        get("Attorney - First name"), get("Attorney - Last name")))) or get("Attorney - Firm")
    data["refType"] = referral_type
    data["review_fields"] = {header: supplied.get(header, "") for header in SOURCE_HEADERS}
    if cem != "auto":
        data["cem_employer_request"] = cem == "yes"
    warnings = []
    for key, message in (
        ("claim_number", "Claim number is required."),
        ("claimant_first_name", "Claimant first name is required."),
        ("claimant_last_name", "Claimant last name is required."),
        ("referral_type", "Select CMS referral type Medical or Vocational."),
    ):
        if not form[key]:
            warnings.append(message)
    if cem == "auto":
        warnings.append("CEM employer request is unknown; select yes or no before proceeding.")
    warnings.append("Review blank setup fields, Claim Identifier, Other Notes and appointment details; no RRS lookup or blank-field fill is performed.")
    name = " ".join(filter(None, (form["claimant_first_name"], form["claimant_last_name"])))
    return {
        "row": number,
        "label": f'Row {number} | {name or "Unnamed claimant"} | {form["claim_number"] or "No claim number"}',
        "form": form,
        "data": data,
        "warning": " ".join(warnings),
    }


_AI_LEGACY_MAP = {
    "Claim - Claim number": "Claim Number",
    "Claimant - First name": "First Name", "Claimant - Last name": "Last Name",
    "Claim - Referral type": "Referral Type",
    "Claim - Employer / customer": "Employer Name",
    "Provider - Phone": "Provider Phone",
    "Provider - Appointment date": "Appointment Date",
    "Provider - Appointment time": "Appointment Time",
    "Attorney - Address line 1": "Attorney Address",
    "Attorney - Address line 2": "Attorney Address-line-2",
    "Attorney - City": "Attorney City", "Attorney - State": "Attorney State",
    "Attorney - ZIP": "Attorney Zip", "Attorney - Phone": "Attorney Phone Number",
}
_PROVIDER_MAP = {
    "name": "Provider Name (First Name / Last Name)",
    "first_name": "Doctor First Name", "last_name": "Doctor Last Name",
    "phone": "Provider Phone", "address": "Provider Address",
    "address2": "Provider Address Line 2", "city": "Provider City",
    "state": "Provider State", "zip": "Provider Zip",
}


def _provider(values):
    result = {key: _effective(values.get(header, "")).strip()
              for key, header in _PROVIDER_MAP.items()}
    kind = _effective(values.get("Determining if Doctor or Provider Name", "")).strip().casefold()
    result["type"] = {"doctor": "doctor", "provider": "facility",
                      "facility": "facility"}.get(kind, "")
    return result


def _ai_record(number, supplied, coordinates, related):
    legacy = {key: supplied.get(header, "") for key, header in _AI_LEGACY_MAP.items()}
    legacy.update({header: supplied.get(header, "") for header in CMS_HEADERS})
    for override, source in (
        ("CMS - Diagnosis code", "Diagnosis Code"),
        ("CMS - Office number", "Claims Office Number"),
        ("CMS - Special instructions", "Special Instructions"),
    ):
        legacy[override] = _effective(legacy.get(override, "")) or supplied.get(source, "")
    record_id = _effective(supplied.get("Record ID", "")).strip()
    classification = related.get(CLASSIFICATION_SHEET, {}).get(record_id, [])
    if len(classification) > 1:
        raise ExcelValidationError("Repeated RRS Classification Record ID; resolve before importing.")
    body = classification[0].get("Body Part", "") if classification else ""
    legacy["Claim - Body part"] = _effective(supplied.get("CMS - Body part", "")) or body
    legacy["Adjuster - Referral source"] = supplied.get("CMS - Referral source", "")
    result = _record(number, legacy, coordinates)
    data = result["data"]
    data["xlAttorney"] = _effective(supplied.get("Attorney Name", "")).strip()
    primary = _provider(supplied)
    data["xlProvider"] = primary["name"] or " ".join(
        filter(None, (primary["first_name"], primary["last_name"])))
    data["provType"] = primary["type"]
    providers = related.get("Providers", {}).get(record_id, [])
    if providers:
        data["providers"] = [_provider(provider) for provider in providers]
    data["review_fields"] = dict(supplied)
    data["source_reviews"] = {
        name: deepcopy(records[record_id]) for name, records in related.items()
        if record_id in records
    }
    if not record_id:
        result["warning"] += " Source Record ID is missing; related records cannot be associated."
    if _effective(supplied.get("Compensable Body/Part(s)", "")) and not data["bodyParts"]:
        result["warning"] += " Raw body part has no mapped CMS body part; review manually."
    if (providers and any(not item["type"] for item in data["providers"])) or (
            not providers and primary["name"] and not primary["type"]):
        result["warning"] += " Provider type is unknown; review before proceeding."
    return result


def _related_rows(book, selected):
    related = {}
    for name in ("Providers", CLASSIFICATION_SHEET, *REVIEW_SHEETS):
        if name not in book:
            continue
        sheet = _sheet(book, name)
        rows = sheet.iter_rows()
        try:
            headers = _headers(next(rows, ()))
            if "Record ID" not in headers:
                raise ExcelValidationError(f'{name}: missing Record ID column.')
            index = headers.index("Record ID")
            grouped = {}
            for row in rows:
                record_id = _effective(_text(row[index])).strip()
                if record_id not in selected:
                    continue
                values = [_text(cell, preserve_whitespace=True) for cell in row]
                grouped.setdefault(record_id, []).append(dict(zip(headers, values)))
            related[name] = grouped
        finally:
            rows.close()
    return related


def read_workbook(path):
    """Read editable rows, retaining blanks and requiring explicit CMS decisions."""
    book = _open(path)
    rows = None
    try:
        sheet = _sheet(book, SHEET)
        rows = sheet.iter_rows()
        headers = _headers(next(rows, ()))
        mapping = {index: _KNOWN[header.casefold()] for index, header in enumerate(headers)
                   if header.casefold() in _KNOWN}
        if not mapping:
            raise ExcelValidationError("No recognized columns. Use the CMS Registration input template.")
        is_ai = any(header.casefold() in ("record id", "claim number", "first name")
                    for header in headers)
        pending = []
        for number, row in enumerate(rows, 2):
            # Inspect even unrecognized columns: formulas must never hide in an imported row.
            values = [_text(cell, preserve_whitespace=True) for cell in row]
            supplied = {
                header: values[index] if is_ai or header in ("CMS - Special instructions", "CMS - Other Notes")
                else values[index].strip()
                for index, header in mapping.items()
            }
            if not any(value.strip() for value in supplied.values()):
                continue
            coordinates = {header: f"{get_column_letter(index + 1)}{number}"
                           for index, header in mapping.items()}
            pending.append((number, supplied, coordinates))
        if not pending:
            raise ExcelValidationError("No filled recognized rows found. Enter CMS input starting on row 2.")
        if not is_ai:
            return [_record(*entry) for entry in pending]
        identifiers = [_effective(entry[1].get("Record ID", "")).strip() for entry in pending]
        nonblank = [value for value in identifiers if value]
        if len(nonblank) != len(set(nonblank)):
            raise ExcelValidationError("Repeated referral Record ID; resolve before importing.")
        related = _related_rows(book, set(nonblank))
        return [_ai_record(*entry, related) for entry in pending]
    finally:
        if rows is not None:
            rows.close()
        book.close()


def _new_book(headers=HEADERS):
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(list(headers))
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}501"
    sheet.row_dimensions[1].height = 44
    required = {"Claim - Claim number", "Claimant - First name", "Claimant - Last name",
                "Claim Number", "First Name", "Last Name", "CMS - Referral type",
                "CMS - CEM employer request"}
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="2563EB" if header in required else "475569")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(
            "Required before starting." if header in required else
            "Leave unknown values blank. No RRS lookup or automatic blank-field fill.", "FCM")
        column = get_column_letter(index)
        sheet.column_dimensions[column].width = 25
        for row in range(2, 502):
            sheet.cell(row, index).number_format = "@"
        choices = {"CMS - Referral type": "Medical,Vocational",
                   "CMS - CEM employer request": "yes,no,auto"}.get(header)
        if choices:
            rule = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
            rule.errorTitle = "Choose a listed value"
            rule.error = "Use the dropdown choices."
            rule.showErrorMessage = True
            sheet.add_data_validation(rule)
            rule.add(f"{column}2:{column}10001")
    guide = book.create_sheet("Instructions")
    for note in (
        f"CMS Registration input - {len(headers)} columns. Blank templates contain no samples; converted files retain source data.",
        "Fill CMS Input starting at row 2. One referral per row; keep headers. Providers preserves repeated providers by Record ID.",
        "AI source columns (or original Subject Line columns in converted legacy files) are preserved for review.",
        "Choose a row in the desktop app and Load row. Review request and setup data, then Start explicitly. Loading does not start CMS.",
        "The selected Excel row supplies request AND setup data. No RRS DB lookup and no silent RRS fill of blank fields.",
        "CMS referral type must be Medical or Vocational. Original Subject Line categories are NOT mapped to these types.",
        "CEM employer request: yes/no; blank or auto means UNKNOWN, not a legacy, daily-workbook or database fallback.",
        "Referral number is explicit CMS input; Adjuster - Referral source is NOT a referral number. Office number is explicit, never inferred from the claim.",
        "Customer email and insured address are explicit CMS columns; adjuster email and claimant address are preserved separately, not assumed equivalent.",
        "Facility is preferred for provider setup; doctor and facility source fields remain preserved. Attorney names take precedence over firm.",
        "Review Claim Identifier, Other Notes and first appointment date/time manually; their presence does not automate those CMS steps.",
        "Required request columns: Claim Number, First Name, Last Name, CMS - Referral type and CMS - CEM employer request (or legacy equivalents). Case ID / referral number are explicit CMS identifiers, never Claim ID or Record ID.",
        "CMS setup overrides take precedence when nonblank. Diagnosis Code, Claims Office Number and Special Instructions map directly. Referral Instructions remains separate source metadata.",
        "CMS - Body part overrides the RRS Classification Body Part; raw Compensable Body/Part(s) is never treated as a dropdown choice. Review all remaining blank setup fields before Start.",
        "Appointment date uses the single editable apptDate field; reviewAppointmentDate is a review alias. Provider entries, when present, are canonical instead of primary-provider controls.",
        "Identifiers, ZIPs and phones: text to preserve leading zeros. Dates: MM/DD/YYYY; times: 08:30 AM. No formulas or Excel errors.",
        "Unknown values stay blank. Incomplete rows can load for correction. Keep live preview enabled; this workbook proves no live integration or deployment.",
        "Save as .xlsx. Protect populated workbooks as sensitive local input; do not commit them.",
    ):
        guide.append([note])
        guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[guide.max_row].height = 42
    guide.column_dimensions["A"].width = 115
    guide["A1"].font = Font(size=16, bold=True, color="2563EB")
    return book


def _save_new(book, path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ExcelValidationError("Choose an .xlsx destination.")
    try:
        stream = path.open("xb")
    except FileExistsError:
        raise ExcelValidationError("Destination already exists; choose a separate new workbook.") from None
    try:
        with stream:
            book.save(stream)
    except Exception:
        path.unlink()
        raise
    return str(path)


def create_template(path):
    """Write a blank CMS template without replacing an existing workbook."""
    book = _new_book()
    try:
        return _save_new(book, path)
    finally:
        book.close()


def convert_subject_line_workbook(source, destination):
    """Copy the original 43 columns losslessly into a separate CMS workbook."""
    if Path(source).resolve() == Path(destination).resolve():
        raise ExcelValidationError("Source and destination must be separate workbooks.")
    if Path(destination).exists():
        raise ExcelValidationError("Destination already exists; choose a separate new workbook.")
    source_book = _open(source)
    output = None
    rows = None
    try:
        source_sheet = _sheet(source_book, SOURCE_SHEET)
        rows = source_sheet.iter_rows()
        header_cells = next(rows, ())
        headers = _headers(header_cells)
        if len(headers) != len(SOURCE_HEADERS) or set(headers) != set(SOURCE_HEADERS):
            raise ExcelValidationError("Source must contain the original 43 Subject Line input columns.")
        output = _new_book(tuple(headers) + CMS_HEADERS)
        target = output[SHEET]
        for row_number, row in enumerate(chain((header_cells,), rows), 1):
            for column_number, cell in enumerate(row, 1):
                _text(cell)
                copied = target.cell(row_number, column_number, cell.value)
                copied.data_type = cell.data_type
                if isinstance(cell, ReadOnlyCell):
                    for attribute in ("font", "fill", "border", "alignment", "protection"):
                        setattr(copied, attribute, copy(getattr(cell, attribute)))
                    copied.number_format = cell.number_format
        target.auto_filter.ref = f"A1:{get_column_letter(len(headers) + len(CMS_HEADERS))}{max(501, source_sheet.max_row)}"
        return _save_new(output, destination)
    finally:
        if rows is not None:
            rows.close()
        source_book.close()
        if output is not None:
            output.close()


def _copy_row(row, target, number):
    for column, cell in enumerate(row, 1):
        _text(cell)
        copied = target.cell(number, column, cell.value)
        copied.data_type = cell.data_type
        if isinstance(cell, ReadOnlyCell):
            for attribute in ("font", "fill", "border", "alignment", "protection"):
                setattr(copied, attribute, copy(getattr(cell, attribute)))
            copied.number_format = cell.number_format
        else:
            copied.number_format = "General"


def convert_ai_workbook(source, destination, *, limit=5):
    """Copy the first populated AI referrals and their related records, locally."""
    if not isinstance(limit, int) or isinstance(limit, bool) or limit < 1:
        raise ExcelValidationError("Referral limit must be a positive integer.")
    if Path(source).resolve() == Path(destination).resolve():
        raise ExcelValidationError("Source and destination must be separate workbooks.")
    if Path(destination).exists():
        raise ExcelValidationError("Destination already exists; choose a separate new workbook.")
    book = _open(source)
    output = None
    try:
        sheet = _sheet(book, "Referrals")
        rows = sheet.iter_rows()
        try:
            header_cells = next(rows, ())
            headers = _headers(header_cells)
            if tuple(headers) not in (AI_SOURCE_HEADERS, tuple(DAILY_HEADERS)):
                raise ExcelValidationError("Source must contain the AI output Referrals columns, optionally with Accurate.")
            selected = []
            seen = set()
            id_column = headers.index("Record ID")
            for row in rows:
                # Selection is by physical source order, not completeness or category.
                if not any(cell.value is not None and str(cell.value).strip() for cell in row):
                    continue
                record_id = _effective(_text(row[id_column])).strip()
                if not record_id:
                    raise ExcelValidationError("Populated referral has no Record ID; resolve before conversion.")
                if record_id in seen:
                    raise ExcelValidationError("Repeated referral Record ID; related rows cannot be associated safely.")
                seen.add(record_id)
                if len(selected) < limit:
                    selected.append((record_id, row))
        finally:
            rows.close()
        if len(selected) < limit:
            raise ExcelValidationError("Source has fewer populated referrals than requested.")
        selected_ids = {record_id for record_id, _ in selected}
        related = _related_rows(book, selected_ids)
        if any(len(entries) > 1 for entries in related.get(CLASSIFICATION_SHEET, {}).values()):
            raise ExcelValidationError("Repeated RRS Classification Record ID; resolve before conversion.")
        output = _new_book(tuple(headers) + AI_CMS_HEADERS)
        target = output[SHEET]
        _copy_row(header_cells, target, 1)
        for index, (_, row) in enumerate(selected, 2):
            _copy_row(row, target, index)
        for index, header in enumerate(headers, 1):
            if header in ("Claim Number", "First Name", "Last Name"):
                target.cell(1, index).comment = Comment("Required before starting.", "FCM")
        for name in ("Providers", CLASSIFICATION_SHEET, *REVIEW_SHEETS, "Schema"):
            if name not in book:
                continue
            source_sheet = _sheet(book, name)
            rows = source_sheet.iter_rows()
            try:
                header = next(rows, ())
                page = output.create_sheet(name)
                _copy_row(header, page, 1)
                names = _headers(header)
                id_column = names.index("Record ID") if name != "Schema" else None
                target_row = 2
                for row in rows:
                    if id_column is None or _text(row[id_column]).strip() in selected_ids:
                        _copy_row(row, page, target_row)
                        target_row += 1
            finally:
                rows.close()
        return _save_new(output, destination)
    finally:
        book.close()
        if output is not None:
            output.close()
