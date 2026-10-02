"""CMSCustomerSearch desktop form and isolated live worker."""
import contextlib
import json
import multiprocessing as mp
import os
from pathlib import Path
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText

from run import HERE, ROOT, load_core

for folder in [ROOT, ROOT / "src"]:
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

from AI.daily_output import read_record, reopen_input, check_layout
from openpyxl import load_workbook

STAGES = {"Both: re-open + customer": "all", "Re-open only": "reopen", "Customer only": "customer"}
INPUTS = [("claimNumber", "Claim Number"), ("customer", "Customer Name"),
          ("claimID", "Claim ID"), ("claimantFull", "Claimant Full Name"), ("referralType", "Referral Type")]


def workbook_records(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        check_layout(book)
        records = [str(row[0]) for row in book["Referrals"].iter_rows(min_row=2, values_only=True) if row[0]]
        if len(records) != len(set(records)):
            raise ValueError("Workbook contains duplicate Record IDs.")
        return records
    finally:
        book.close()


def validate_input(data, stage):
    required = ["claimNumber"] if stage in {"all", "reopen"} else []
    if stage in {"all", "customer"}:
        required += ["customer", "claimID", "claimantFull"]
    missing = [label for key, label in INPUTS if key in required and not data.get(key, "").strip()]
    if missing:
        raise ValueError("Missing input: " + ", ".join(missing))


class DesktopPrompts:
    """Create the Tk root lazily, on first prompt, not at worker startup.

    Creating a Tk root puts this process's main thread into a COM apartment
    with no message pump. If that happens before CMS login, it permanently
    deadlocks the certificate-warning handler's background UI Automation
    thread (it never returns from window enumeration). Login never needs a
    prompt, so deferring root creation until a prompt is actually requested
    avoids the deadlock while keeping dialogs available afterward.
    """

    def __init__(self):
        self._root = None

    def _ensure_root(self):
        if self._root is None:
            self._root = tk.Tk()
            self._root.withdraw()
        return self._root

    def ask_yes_no(self, title, message):
        return messagebox.askyesno(title, message, parent=self._ensure_root())

    def ask_text(self, title, message):
        from tkinter.simpledialog import askstring
        return askstring(title, message, parent=self._ensure_root())

    def notify(self, title, message):
        messagebox.showinfo(title, message, parent=self._ensure_root())

    def destroy(self):
        if self._root is not None:
            self._root.destroy()
            self._root = None


def execute_request(request):
    """Same process logic for offline scenarios and live CMS; no UI duplication."""
    if request["mode"] == "offline":
        scenario = json.loads((HERE / "scenarios" / request["scenario"]).read_text(encoding="utf-8"))
        notices = []
        result = load_core().run_flow(scenario["input"], stage=request["stage"],
            search_cases=lambda *_: scenario.get("cases", []),
            check_customer=lambda *_: scenario.get("selected_customer"),
            notify=lambda title, text: notices.append({"title": title, "message": text}),
            confirm=lambda *_: scenario.get("proceed", False))
        return {"result": result, "messages": notices, "mode": "offline"}
    data = request["input"]
    if request.get("excel"):
        fields, assessment = read_record(request["excel"], request["record_id"])
        if assessment["status"] != "Passed":
            raise ValueError("Selected Excel record is Failed. Review missing required information first.")
        data = reopen_input(fields)
    validate_input(data, request["stage"])
    from fcm_intake.cms import session
    from fcm_intake.workflows.reopen_flow import run_live
    prompts = DesktopPrompts()
    session.set_credentials(request["username"], request["password"])
    try:
        result = run_live(data, stage=request["stage"], app=prompts,
                          notify=prompts.notify, confirm=prompts.ask_yes_no)
        return {"result": result, "mode": "live"}
    finally:
        session.close_shared_driver()
        prompts.destroy()


def worker_main(connection, request):
    """Keep legacy Selenium/Tk work on its own main thread; UI stays responsive."""
    stage = 'Preparing process'
    def progress(message):
        nonlocal stage
        stage = message
        connection.send({'progress': message})
    try:
        if request['mode'] == 'live':
            from fcm_intake.cms import session
            session.set_status_callback(progress)
        with open(os.devnull, "w") as sink, contextlib.redirect_stdout(sink), contextlib.redirect_stderr(sink):
            response = execute_request(request)
    except SystemExit:
        response = {"result": {"status": "stopped", "reason": "legacy_process_stopped"}}
    except Exception as error:
        response = {"error": f"{type(error).__name__}: {error}", 'stage': stage}
    try:
        connection.send(response)
    finally:
        connection.close()


class App:
    def __init__(self, root):
        self.root = root
        root.title("CMSCustomerSearch")
        root.geometry("940x820")
        root.minsize(850, 700)
        self.worker = None
        self.pipe = None
        self.result = None
        self.controls = []
        body = ttk.Frame(root, padding=16)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="CMSCustomerSearch", font=("Segoe UI", 20, "bold")).pack(anchor="w")
        ttk.Label(body, text="Test re-open checking and customer validation independently of the intake bot.").pack(anchor="w", pady=(0, 12))
        options = ttk.Frame(body)
        options.pack(fill="x")
        self.mode = tk.StringVar(value="Offline scenario")
        self.stage = tk.StringVar(value=next(iter(STAGES)))
        self.combo(options, self.mode, ["Offline scenario", "Live CMS"], 23).pack(side="left", padx=(0, 12))
        self.combo(options, self.stage, list(STAGES), 28).pack(side="left")
        self.scenario = tk.StringVar(value="happy-path.json")
        self.combo(options, self.scenario, sorted(p.name for p in (HERE / "scenarios").glob("*.json")), 25).pack(side="right")
        ttk.Label(body, text="Offline mode uses synthetic scenario inputs. Live mode uses the Excel record or manual fields below.").pack(anchor="w", pady=8)
        files = ttk.LabelFrame(body, text="Daily Excel input", padding=8)
        files.pack(fill="x")
        self.filename = tk.StringVar()
        ttk.Label(files, textvariable=self.filename, wraplength=780).pack(anchor="w")
        row = ttk.Frame(files)
        row.pack(fill="x", pady=5)
        self.button(row, "Choose daily Excel", self.browse).pack(side="left")
        self.record = tk.StringVar()
        self.records = self.combo(row, self.record, [], 48)
        self.records.pack(side="left", padx=8)
        self.records.bind("<<ComboboxSelected>>", lambda _: self.load_record())
        self.button(row, "Reload record", self.load_record).pack(side="left")
        self.button(row, "Use manual input", self.manual).pack(side="left", padx=8)
        self.assessment = tk.StringVar(value="No workbook selected. Manual input is available for live runs.")
        ttk.Label(files, textvariable=self.assessment, wraplength=850).pack(anchor="w")
        form = ttk.LabelFrame(body, text="Process input", padding=8)
        form.pack(fill="x", pady=10)
        self.inputs = {}
        self.entries = []
        for index, (key, label) in enumerate(INPUTS):
            ttk.Label(form, text=label).grid(row=index, column=0, sticky="w", padx=(0, 12))
            variable = tk.StringVar()
            self.inputs[key] = variable
            entry = ttk.Entry(form, textvariable=variable)
            entry.grid(row=index, column=1, sticky="ew", pady=2)
            self.entries.append(entry)
            self.controls.append((entry, "normal"))
        form.columnconfigure(1, weight=1)
        auth = ttk.Frame(body)
        auth.pack(fill="x", pady=4)
        self.username = tk.StringVar(value=os.getenv("FCM_CMS_USERNAME", ""))
        self.password = tk.StringVar(value=os.getenv("FCM_CMS_PASSWORD", ""))
        for label, variable, show in [("CMS username", self.username, ""), ("Password", self.password, "*")]:
            ttk.Label(auth, text=label).pack(side="left", padx=(0, 6))
            entry = ttk.Entry(auth, textvariable=variable, show=show, width=28)
            entry.pack(side="left", padx=(0, 15))
            self.controls.append((entry, "normal"))
        actions = ttk.Frame(body)
        actions.pack(fill="x", pady=8)
        self.button(actions, "Run selected process", self.start).pack(side="left")
        self.button(actions, "Save result JSON", self.save_result).pack(side="left", padx=8)
        self.status = tk.StringVar(value="Ready — choose offline testing or configure live CMS.")
        ttk.Label(body, textvariable=self.status, wraplength=880).pack(anchor="w")
        self.output = ScrolledText(body, height=12, wrap="word")
        self.output.pack(fill="both", expand=True, pady=(8, 0))
        self.output.configure(state="disabled")
        root.protocol("WM_DELETE_WINDOW", self.close)

    def combo(self, parent, variable, values, width):
        widget = ttk.Combobox(parent, textvariable=variable, values=values, state="readonly", width=width)
        self.controls.append((widget, "readonly"))
        return widget

    def button(self, parent, label, command):
        widget = ttk.Button(parent, text=label, command=command)
        self.controls.append((widget, "normal"))
        return widget

    def display(self, data):
        self.output.configure(state="normal")
        self.output.delete("1.0", "end")
        self.output.insert("1.0", json.dumps(data, indent=2, default=str))
        self.output.configure(state="disabled")

    def browse(self):
        path = filedialog.askopenfilename(parent=self.root, filetypes=[("AI daily workbook", "*.xlsx")])
        if path:
            try:
                ids = workbook_records(path)
                self.filename.set(path)
                self.records.configure(values=ids)
                self.record.set(ids[0] if ids else "")
                self.load_record()
            except Exception as error:
                self.manual()
                messagebox.showerror("Cannot load workbook", str(error), parent=self.root)

    def load_record(self):
        self.result = None
        for variable in self.inputs.values():
            variable.set("")
        if not self.filename.get() or not self.record.get():
            self.assessment.set("Select a workbook containing a Referrals record.")
            return
        try:
            fields, assessment = read_record(self.filename.get(), self.record.get())
            for key, value in reopen_input(fields).items():
                self.inputs[key].set(value)
            self.assessment.set("NEXT STEP: " + assessment["status"] + " — Excel input is read-only; live runs recheck the file.")
            self.display({"record_id": self.record.get(), "next_step": assessment})
        except Exception as error:
            self.assessment.set("Workbook record could not be loaded.")
            messagebox.showerror("Cannot load record", str(error), parent=self.root)
        finally:
            for entry in self.entries:
                entry.configure(state="readonly")

    def manual(self):
        self.filename.set("")
        self.record.set("")
        self.records.configure(values=[])
        for key, variable in self.inputs.items():
            variable.set("")
        for entry in self.entries:
            entry.configure(state="normal")
        self.result = None
        self.display({})
        self.assessment.set("Manual input selected. Required fields depend on the selected stage.")

    def request(self):
        request = {"mode": "offline" if self.mode.get() == "Offline scenario" else "live",
                   "stage": STAGES[self.stage.get()], "scenario": self.scenario.get()}
        if request["mode"] == "live":
            data = {key: value.get().strip() for key, value in self.inputs.items()}
            if self.filename.get():
                fields, assessment = read_record(self.filename.get(), self.record.get())
                if assessment["status"] != "Passed":
                    raise ValueError("Selected Excel record is Failed. Review the missing required fields first.")
                data = reopen_input(fields)
                request.update(excel=self.filename.get(), record_id=self.record.get())
            validate_input(data, request["stage"])
            if not self.username.get().strip() or not self.password.get():
                raise ValueError("Enter CMS username and password for live processing.")
            request.update(input=data, username=self.username.get().strip(), password=self.password.get())
        return request

    def start(self):
        try:
            request = self.request()
            self.result = None
            self.display({})
            context = mp.get_context("spawn")
            self.pipe, child = context.Pipe(duplex=False)
            self.worker = context.Process(target=worker_main, args=(child, request))
            self.worker.start()
            child.close()
        except Exception as error:
            if self.pipe is not None:
                self.pipe.close()
                self.pipe = None
            self.worker = None
            messagebox.showerror("Cannot start process", str(error), parent=self.root)
            return
        for control, _ in self.controls:
            control.configure(state="disabled")
        self.status.set("Running " + request["mode"] + " process. Respond to any CMS prompts that appear.")
        self.root.after(100, self.poll)

    def poll(self):
        response = None
        if self.pipe.poll():
            try:
                response = self.pipe.recv()
            except EOFError:
                response = {"error": "Process ended without returning a result."}
        elif not self.worker.is_alive():
            response = {"error": f"Process exited unexpectedly (code {self.worker.exitcode})."}
        if response is None:
            self.root.after(100, self.poll)
            return
        if 'progress' in response:
            self.status.set(response['progress'])
            self.display({'stage': response['progress']})
            self.root.after(100, self.poll)
            return
        self.worker.join(timeout=1)
        self.pipe.close()
        self.pipe = None
        self.worker = None
        self.result = response
        self.display(response)
        self.status.set("Error — see details below." if "error" in response else
                        "Process " + response.get("result", {}).get("status", "finished") + ".")
        for control, state in self.controls:
            control.configure(state=state)
        if self.filename.get():
            for entry in self.entries:
                entry.configure(state="readonly")

    def save_result(self):
        if self.result is None:
            messagebox.showinfo("No result", "Run a process first.", parent=self.root)
            return
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".json", initialfile="CMSCustomerSearch-result.json")
        if path:
            try:
                Path(path).write_text(json.dumps(self.result, indent=2, default=str), encoding="utf-8")
            except OSError as error:
                messagebox.showerror("Cannot save result", str(error), parent=self.root)

    def close(self):
        if self.worker is not None:
            messagebox.showinfo("Process running", "Finish the current CMS process and its prompts before closing.", parent=self.root)
        else:
            self.root.destroy()
