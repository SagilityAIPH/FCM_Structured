"""Daily CEM compatibility workbook: explicit path, locked atomic writes."""
from contextlib import contextmanager
from datetime import datetime
import os
from pathlib import Path
import tempfile
import time

from openpyxl import Workbook, load_workbook

FILENAME = "FCM CEM Checher.xlsx"
# Legacy ensure_headers rewrites A:C to these strings on every write. B is semantically
# claim NUMBER in new rows despite the legacy heading "Claim ID"; keep legacy compatible.
HEADERS = ("Claimant", "Claim ID", "Status", "Claim ID (PDF)", "Employer",
           "Request Reference", "Submitted At", "Environment")


def _text(value):
    if not isinstance(value, str):
        raise ValueError("Tracking values must be text")
    value = value.strip()
    if value.startswith(("=", "+", "-", "@")) or any(ord(c) < 32 for c in value):
        raise ValueError("Unsafe tracking text")
    return value


def _location(path):
    if not isinstance(path, (str, os.PathLike)) or not path or Path(path).name != FILENAME or not Path(path).is_absolute():
        raise ValueError("Explicit absolute CEM workbook path required")
    return Path(path)


def _clock(now):
    value = now() if callable(now) else now
    if not isinstance(value, datetime):
        raise TypeError("Clock must return datetime")
    return value


@contextmanager
def _lock(path, timeout=10):
    """Windows advisory sidecar lock across check/load/save/replace (not legacy-aware)."""
    import msvcrt  # Only reached during an actual workbook operation on Windows.
    lock_path = str(path) + ".lock"
    with open(lock_path, "a+b") as handle:
        handle.seek(0)
        if not handle.read(1):
            handle.seek(0)
            handle.write(b"0")
            handle.flush()
        deadline = time.monotonic() + timeout
        while True:
            try:
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                break
            except OSError:
                if time.monotonic() >= deadline:
                    raise TimeoutError("CEM workbook lock unavailable") from None
                time.sleep(.05)
        try:
            yield
        finally:
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)


def _rows(ws, claimant, claim_number):
    target = (_text(claimant).casefold(), _text(claim_number).casefold())
    for row in range(2, ws.max_row + 1):
        values = tuple(str(ws.cell(row, column).value or "").strip().casefold() for column in (1, 2))
        if values == target:
            yield row


def find_duplicate(request, settings, *, now=datetime.now):
    """Mirror CMS.py claim_exists: today's A/B, strip+lower, no status check."""
    path = _location(settings.cem_log_path)
    claimant = f"{request.claimant_first_name} {request.claimant_last_name}"
    day = _clock(now).strftime("%m%d%Y")
    with _lock(path):
        if not path.exists():
            return None
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if day not in wb:
                return None
            row = next(_rows(wb[day], claimant, request.claim_number), None)
            return {"sheet": day, "row": row} if row else None
        finally:
            wb.close()


def record(request, reference, settings, *, now=datetime.now):
    """Idempotently append a verified receipt; return exact sheet/row after atomic replace."""
    path = _location(settings.cem_log_path)
    if not path.parent.is_dir():
        raise FileNotFoundError("CEM workbook directory does not exist")
    moment = _clock(now)
    day = moment.strftime("%m%d%Y")
    values = (f"{request.claimant_first_name} {request.claimant_last_name}",
              request.claim_number, "Submitted", request.claim_id, request.employer_name,
              reference, moment.isoformat(timespec="seconds"), settings.environment)
    values = tuple(_text(v) for v in values)
    if not values[5]:
        raise ValueError("Verified request reference required")
    with _lock(path):
        wb = load_workbook(path) if path.exists() else Workbook()
        temp = None
        try:
            if day in wb:
                ws = wb[day]
            elif not path.exists() and wb.active is not None and wb.active.title == "Sheet":
                ws = wb.active
                ws.title = day
            else:
                ws = wb.create_sheet(day)
            if ws is None:
                raise ValueError("CEM worksheet unavailable")
            # Do not rewrite a legacy sheet's headings or rows: its ensure_headers
            # will restore A:C next time the old writer runs. New headers D:H only.
            for column, header in enumerate(HEADERS, 1):
                if ws.cell(1, column).value is None:
                    ws.cell(1, column, _text(header))
            for row in _rows(ws, values[0], values[1]):
                existing = ws.cell(row, 6).value
                if (not existing or str(existing).strip() != reference or
                        ws.cell(row, 3).value != "Submitted" or
                        str(ws.cell(row, 4).value or "").strip() != request.claim_id or
                        str(ws.cell(row, 5).value or "").strip().casefold() != request.employer_name.strip().casefold()):
                    raise ValueError("Existing claim lacks this verified request identity")
                return {"status": "existing", "row": {"sheet": day, "row": row}}
            # Append only after max_row, never recycle a hole with D:H metadata.
            row = max(2, ws.max_row + 1)
            for column, value in enumerate(values, 1):
                ws.cell(row, column, value=value).data_type = "s"
            fd, temp = tempfile.mkstemp(prefix=".cem-", suffix=".xlsx", dir=path.parent)
            os.close(fd)
            wb.save(temp)
            os.replace(temp, path)
            temp = None
        finally:
            wb.close()
            if temp and os.path.exists(temp):
                os.unlink(temp)
    # Confirm exact row was persisted, not merely accepted by save().
    with _lock(path):
        check = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = check[day]
            if tuple(ws.cell(row, col).value for col in range(1, 9)) != values:
                raise IOError("CEM tracking readback mismatch")
        finally:
            check.close()
    return {"status": "written", "row": {"sheet": day, "row": row}}
