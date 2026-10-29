"""Synthetic workbook tests; CMS.py import has live dependency side effects, replicate claim_exists."""
from dataclasses import replace
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from contract import CEMSettings
from test_contract import BASE
import cem_log as L

NOW = lambda: datetime(2026, 1, 2, 9, 30)
DAY = "01022026"


def cms_claim_exists(path, claimant, claim_number):
    """CMS.py:1930-1965 exact strip/lower A/B, today's sheet only, ignore C."""
    if not path.exists():
        return False
    wb = load_workbook(path, data_only=True)
    try:
        if DAY not in wb:
            return False
        norm = lambda v: str(v).strip().lower() if v is not None else ""
        return any(norm(wb[DAY].cell(r, 1).value) == norm(claimant) and
                   norm(wb[DAY].cell(r, 2).value) == norm(claim_number)
                   for r in range(2, wb[DAY].max_row + 1))
    finally:
        wb.close()

class LogTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / L.FILENAME
        self.settings = CEMSettings(cem_log_path=str(self.path))

    def test_lookup_uses_case_insensitive_exact_claimant_claim_pair(self):
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        self.assertIsNotNone(L.find_duplicate(replace(BASE, claimant_first_name="TEST",
            claimant_last_name="PERSON", claim_number=BASE.claim_number.lower()), self.settings, now=NOW))
        self.assertIsNone(L.find_duplicate(replace(BASE, claim_number=BASE.claim_id), self.settings, now=NOW))
        self.assertIsNone(L.find_duplicate(replace(BASE, claimant_last_name="Someone"), self.settings, now=NOW))

    def test_cem_status_not_approved_in_log(self):
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY]["C2"].value, "Submitted")
            self.assertNotEqual(wb[DAY]["C2"].value, "Approved")
        finally:
            wb.close()

    def test_repeated_daily_write_is_idempotent_and_no_extra_row(self):
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY].max_row, 2)
        finally:
            wb.close()

    def test_write_failure_does_not_replace_existing_workbook(self):
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        original = self.path.read_bytes()
        with self.assertRaises(ValueError):
            L.record(replace(BASE, claim_id="DIFFERENT"), "SYN-REF-3", self.settings, now=NOW)
        self.assertEqual(self.path.read_bytes(), original)

    def test_claim_number_vs_claim_id_and_idempotency(self):
        self.assertIsNone(L.find_duplicate(BASE, self.settings, now=NOW))
        saved = L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        self.assertEqual(saved["status"], "written")
        self.assertTrue(cms_claim_exists(self.path, " test person ", BASE.claim_number.lower()))
        self.assertFalse(cms_claim_exists(self.path, "Test Person", BASE.claim_id))
        self.assertEqual(L.find_duplicate(BASE, self.settings, now=NOW), saved["row"])
        self.assertEqual(L.record(BASE, "SYN-REF-3", self.settings, now=NOW)["status"], "existing")
        with self.assertRaises(ValueError):
            L.record(replace(BASE, claim_id="SYN-DIFFERENT-ID"), "SYN-REF-3", self.settings, now=NOW)
        with self.assertRaises(ValueError):
            L.record(BASE, "SYN-OTHER", self.settings, now=NOW)
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY]["B2"].value, BASE.claim_number)
            self.assertEqual(wb[DAY]["D2"].value, BASE.claim_id)
            self.assertEqual(wb[DAY]["C2"].value, "Submitted")
            self.assertEqual(wb[DAY]["B2"].data_type, "s")
        finally:
            wb.close()

    def test_legacy_sheet_preserved_and_extended(self):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = DAY
        for col, text in enumerate(("Claimant", "Claim ID", "Status"), 1):
            ws.cell(1, col, text)
        ws.append(("Older Synthetic", "OLD-CLAIM", ""))
        wb.save(self.path)
        wb.close()
        saved = L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        self.assertEqual(saved["row"]["row"], 3)
        wb = load_workbook(self.path)
        try:
            ws = wb[DAY]
            self.assertEqual([ws.cell(1, c).value for c in range(1, 4)], ["Claimant", "Claim ID", "Status"])
            self.assertEqual([ws.cell(2, c).value for c in range(1, 4)], ["Older Synthetic", "OLD-CLAIM", None])
            self.assertEqual(ws["D1"].value, "Claim ID (PDF)")
        finally:
            wb.close()
        # Legacy ensure_headers changes only A:C and therefore preserves D:H.

    def test_existing_sheet_header_rewritten_by_legacy_unchanged_extra_columns(self):
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        wb = load_workbook(self.path)
        ws = wb[DAY]
        for col, header in enumerate(("Claimant", "Claim ID", "Status"), 1):
            ws.cell(1, col, header)
        wb.save(self.path)
        wb.close()
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY]["D1"].value, "Claim ID (PDF)")
            self.assertEqual(wb[DAY]["H2"].value, "uat")
        finally:
            wb.close()

    def test_legacy_key_mismatch_is_visible(self):
        legacy = Workbook()
        ws = legacy.active
        assert ws is not None
        ws.title = DAY
        ws.append(("Claimant", "Claim ID", "Status"))
        ws.append(("Test Person", BASE.claim_id, ""))
        legacy.save(self.path)
        legacy.close()
        self.assertFalse(cms_claim_exists(self.path, "Test Person", BASE.claim_number))
        self.assertIsNone(L.find_duplicate(BASE, self.settings, now=NOW))
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        self.assertTrue(cms_claim_exists(self.path, "Test Person", BASE.claim_number))
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY]["B2"].value, BASE.claim_id)
            self.assertEqual(wb[DAY]["B3"].value, BASE.claim_number)
        finally:
            wb.close()

    def test_sheet_header_preservation_even_if_unusual(self):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = DAY
        ws.append(("Claimant", "Claim Number", "Status", "Claim ID (PDF)"))
        wb.save(self.path)
        wb.close()
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY]["B1"].value, "Claim Number")
            self.assertEqual(wb[DAY]["D1"].value, "Claim ID (PDF)")
        finally:
            wb.close()

    def test_legacy_row_blocks_same_claim_without_verified_reference(self):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = DAY
        ws.append(("Claimant", "Claim ID", "Status"))
        ws.append(("Test Person", BASE.claim_number, None))
        wb.save(self.path)
        wb.close()
        self.assertEqual(L.find_duplicate(BASE, self.settings, now=NOW), {"sheet": DAY, "row": 2})
        with self.assertRaises(ValueError):
            L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb[DAY].max_row, 2)
        finally:
            wb.close()

    def test_run_controller_with_real_log_and_synthetic_portal(self):
        # Exercise the actual injected tracking callback, not only a fake writer.
        import cem_flow as F
        import cem_rules as R
        from contract import CEMStages
        settings = CEMSettings(cem_log_path=str(self.path),
                               employer_search_url="https://example.invalid/search",
                               employer_add_url="https://example.invalid/add")
        stages = CEMStages(search_employer=lambda *a: {"state": "no_results", "names": []},
                           find_existing_request=lambda *a: False,
                           find_log_duplicate=lambda r, s: L.find_duplicate(r, s, now=NOW),
                           fill_form=lambda r, address, s: dict(address, employer_name=r.employer_name,
                                        bill_to_name=R.BILL_TO_NAME, bill_to_code=R.BILL_TO_CODE),
                           review=lambda *a: True, confirm_submit=lambda *a: True,
                           submit=lambda *a: True,
                           read_back_request=lambda r, s: dict(unique=True, status="submitted",
                                        reference="SYN-REF-3", employer_name=r.employer_name,
                                        claim_number=r.claim_number),
                           record=lambda r, ref, s: L.record(r, ref, s, now=NOW))
        first = F.run(BASE, stages=stages, settings=settings)
        self.assertEqual((first["outcome"], first["tracking_status"]), ("completed", "written"))
        second = F.run(BASE, stages=stages, settings=settings)
        self.assertEqual((second["outcome"], second["reason"]), ("stopped", "duplicate_request"))
        self.assertTrue(cms_claim_exists(self.path, "Test Person", BASE.claim_number))

    def test_find_duplicate_rejects_invalid_filename_without_read(self):
        wrong = CEMSettings(cem_log_path=str(self.path.with_name("elsewhere.xlsx")))
        with self.assertRaises(ValueError):
            L.find_duplicate(BASE, wrong, now=NOW)

    def test_insert_does_not_touch_other_day(self):
        yesterday = lambda: datetime(2026, 1, 1)
        first = L.record(BASE, "SYN-REF-A", self.settings, now=yesterday)
        second = L.record(BASE, "SYN-REF-B", self.settings, now=NOW)
        self.assertEqual(first["row"]["sheet"], "01012026")
        self.assertEqual(second["row"]["sheet"], DAY)
        wb = load_workbook(self.path)
        try:
            self.assertEqual(wb["01012026"]["F2"].value, "SYN-REF-A")
            self.assertEqual(wb[DAY]["F2"].value, "SYN-REF-B")
        finally:
            wb.close()

    def test_existing_log_row_detected_from_actual_workbook(self):
        import cem_flow as F
        from contract import CEMStages
        L.record(BASE, "SYN-REF-3", self.settings, now=NOW)
        settings = CEMSettings(cem_log_path=str(self.path),
                               employer_search_url="https://example.invalid/search",
                               employer_add_url="https://example.invalid/add")
        stages = CEMStages(search_employer=lambda *a: {"state": "no_results", "names": []},
                           find_log_duplicate=lambda r, s: L.find_duplicate(r, s, now=NOW),
                           find_existing_request=lambda *a: False,
                           fill_form=lambda *a: self.fail("Existing log must prevent fill"),
                           review=lambda *a: self.fail("No review"),
                           confirm_submit=lambda *a: self.fail("No submit confirmation"),
                           submit=lambda *a: self.fail("No submit"),
                           read_back_request=lambda *a: self.fail("No receipt"),
                           record=lambda *a: self.fail("No write"))
        result = F.run(BASE, stages=stages, settings=settings)
        self.assertEqual((result["outcome"], result["reason"]), ("stopped", "duplicate_request"))

    def test_real_log_write_failure_after_verified_submit(self):
        import cem_flow as F
        import cem_rules as R
        from contract import CEMStages
        invalid_path = str(Path(self.tmp.name) / "missing" / L.FILENAME)
        settings = CEMSettings(cem_log_path=invalid_path,
                               employer_search_url="https://example.invalid/search",
                               employer_add_url="https://example.invalid/add")
        stages = CEMStages(search_employer=lambda *a: {"state": "no_results", "names": []},
                           find_existing_request=lambda *a: False,
                           find_log_duplicate=lambda *a: False,
                           fill_form=lambda r, address, s: dict(address, employer_name=r.employer_name,
                                         bill_to_name=R.BILL_TO_NAME, bill_to_code=R.BILL_TO_CODE),
                           review=lambda *a: True, confirm_submit=lambda *a: True,
                           submit=lambda *a: True,
                           read_back_request=lambda r, s: dict(unique=True, status="submitted",
                                         reference="SYN-REF-3", employer_name=r.employer_name,
                                         claim_number=r.claim_number),
                           record=lambda r, ref, s: L.record(r, ref, s, now=NOW))
        result = F.run(BASE, stages=stages, settings=settings)
        self.assertEqual((result["outcome"], result["reason"]), ("failed", "tracking_failed_after_submit"))
        self.assertEqual(result["request_status"], "submitted")
        self.assertEqual(result["request_reference"], "SYN-REF-3")
        self.assertIn("do not resubmit", result["warning"])
        self.assertNotIn(invalid_path, str(result))

    def test_explicit_absolute_path_required(self):
        with self.assertRaises(ValueError):
            L.record(BASE, "SYN-REF", CEMSettings(cem_log_path="relative/FCM CEM Checher.xlsx"), now=NOW)
        with self.assertRaises(ValueError):
            L.find_duplicate(BASE, CEMSettings(cem_log_path=str(self.path.with_name("other.xlsx"))), now=NOW)

    def test_formula_and_missing_sheet(self):
        with self.assertRaises(ValueError):
            L.record(replace(BASE, claim_number="=1+2"), "SYN-REF", self.settings, now=NOW)
        with self.assertRaises(ValueError):
            L.record(replace(BASE, employer_name="+malicious"), "SYN-REF", self.settings, now=NOW)
        with self.assertRaises(ValueError):
            L.record(BASE, "@bad", self.settings, now=NOW)
        other_day = lambda: datetime(2026, 1, 3)
        L.record(BASE, "SYN-REF", self.settings, now=NOW)
        self.assertIsNone(L.find_duplicate(BASE, self.settings, now=other_day))

if __name__ == "__main__":
    unittest.main()
