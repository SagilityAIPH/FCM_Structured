"""Attended live CMS/RRS adapters. Nothing live is imported or opened on import.

Explicit-driver DOM helpers replace CMS.py's module-global driver helpers. The
stored-procedure mapper alone is reused lazily. No credentials are returned.
"""
from dataclasses import asdict
from datetime import datetime
import importlib.util
import os
import logging
from pathlib import Path
import re
import sys
import time
from types import SimpleNamespace
from difflib import SequenceMatcher

SETTINGS_KEYS = (
    "cms_db_connection", "reference_dir", "cms_user", "cms_password",
    "ie_driver_path", "notify", "confirm", "office_matrix_filename",
    "supervisor_mapping_filename", "cem_log_path", "reference_fallback_dirs", "preview_only", "load_data",
)
OFFICE_MATRIX = "Liberty Mutual Claims Office Matrix_08-29-19.xlsx"
SUPERVISOR_MAPPING = "Liberty Mutual Medical Dedicated Supervisor Mapping 04-07-2025.xlsx"


def _sibling(filename, name):
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, Path(__file__).with_name(filename))
        if spec is None or spec.loader is None:
            raise ImportError("CMS registration module unavailable")
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    return sys.modules[name]


def _contract():
    return _sibling("contract.py", "_cms_registration_contract")


def _rules():
    return _sibling("cms_rules.py", "_cms_registration_pure_rules")


# The new CEM app's daily log reader (Processes/CEM/Deploy-Ready/cem_log.py).
# Loaded by path, lazily, under a module name unique to this adapter.
CEM_LOG_MODULE = "_cmsreg_cem_log"
CEM_LOG_SOURCE = Path(__file__).resolve().parents[2] / "CEM" / "Deploy-Ready" / "cem_log.py"
_cem_clock = datetime.now  # one reading fixes "today" for probe and lookup


def _cem_log():
    if CEM_LOG_MODULE not in sys.modules:
        spec = importlib.util.spec_from_file_location(CEM_LOG_MODULE, CEM_LOG_SOURCE)
        if spec is None or spec.loader is None:
            raise ImportError("CEM log module unavailable")
        module = importlib.util.module_from_spec(spec)
        sys.modules[CEM_LOG_MODULE] = module
        try:
            spec.loader.exec_module(module)
        except BaseException:
            sys.modules.pop(CEM_LOG_MODULE, None)
            raise
    return sys.modules[CEM_LOG_MODULE]


def _local_cem_log_path():
    """[cem] cem_log_path / FCM_CEM_LOG_PATH, same convention as the CEM app."""
    try:
        from fcm_intake.config import get_local_setting
    except ImportError:
        return os.environ.get("FCM_CEM_LOG_PATH", "")
    try:
        return get_local_setting("cem", "cem_log_path", "FCM_CEM_LOG_PATH")
    except KeyError:
        # No built-in default: absent means "not configured".
        return ""


def _configured_cem_log_path(settings):
    """Return the configured CEM log path, or '' when not configured.

    A non-blank ``settings['cem_log_path']`` wins over local settings/env.
    A non-text value counts as configured-but-invalid (the reader rejects it).
    """
    value = settings.get("cem_log_path")
    if value is None or (isinstance(value, str) and not value.strip()):
        value = _local_cem_log_path()
    if isinstance(value, os.PathLike):
        value = os.fspath(value)
    if not isinstance(value, str):
        return value if value else ""
    return value.strip()


def _legacy():
    # CMS imports selenium, pyodbc, openpyxl and pywinauto: NEVER at import time.
    from fcm_intake.legacy import CMS
    return CMS


def _element(driver, by, selector):
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    return WebDriverWait(driver, 20).until(EC.presence_of_element_located((by, selector)))


def _type(driver, by, selector, value):
    element = _element(driver, by, selector)
    element.clear()
    element.send_keys(str(value or ""))
    return element


def _click(driver, by, selector):
    element = _element(driver, by, selector)
    element.click()
    return element


def _manual(settings, step, message):
    try:
        approved = settings["confirm"]("CMS Registration", message)
    except (Exception, SystemExit):
        raise RuntimeError("Operator confirmation failed") from None
    return _contract().StepResult(step, "manual" if approved else "failed",
                                  "Operator confirmed" if approved else "Operator declined")


def _safe_step(step, callback):
    try:
        return callback()
    except (_CaseSearchError, _CaseSelectionError, _ClaimSetupError, _ClaimantSetupError,
            _CaseSetupError, _CustomerSetupError, _ProviderSetupError,
            _AttorneySetupError, _BillingSetupError) as error:
        return _contract().StepResult(step, "failed", str(error))
    except (Exception, SystemExit):
        # DOM / driver exceptions can echo typed values or passwords: omit them.
        return _contract().StepResult(step, "failed", "Portal action could not be verified")


def _load_reference(settings, filename_key, default):
    """FCM-1.1-105: never silently use CWD or expose missing reference path."""
    filename = Path(settings.get(filename_key, default))
    if (filename.name != str(filename) or filename.suffix.lower() != ".xlsx"
            or "\\" in str(filename) or ":" in str(filename)):
        raise ValueError("Invalid reference workbook filename")
    for folder in [settings["reference_dir"], *settings.get("reference_fallback_dirs", ())]:
        path = Path(folder) / filename
        if path.is_file():
            break
    else:
        return None
    # Workbook content stays local/in memory and is never included in errors.
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (Exception, SystemExit):
        if filename_key == "supervisor_mapping_filename":
            raise
        raise ValueError("Reference workbook unreadable") from None
    try:
        if filename_key == "supervisor_mapping_filename":
            return list(workbook.active.values)
        from AI.cms_office_matrix import SHEETS
        return {name: list(workbook[name].values) for name in SHEETS
                if name in workbook.sheetnames}
    except (Exception, SystemExit):
        if filename_key == "supervisor_mapping_filename":
            raise
        raise ValueError("Reference workbook unreadable") from None
    finally:
        try:
            workbook.close()
        except (Exception, SystemExit):
            if filename_key == "supervisor_mapping_filename":
                raise
            raise ValueError("Reference workbook unreadable") from None


class _CaseSearchError(RuntimeError):
    """Only fixed phase labels and exception types may cross the CASE boundary."""


class _CaseSelectionError(RuntimeError):
    """Construct selection diagnostics without copying portal exception text."""

    PHASES = (
        "frame reset", "search frame", "search input", "search submission",
        "search results", "claimant profile opening", "case row inspection",
        "case activation", "case identity readback", "referral type readback",
    )

    def __init__(self, phase, error, *, matching_claim_rows=None, selectable_result_rows=None):
        if phase not in self.PHASES:
            raise ValueError("Invalid case selection diagnostic phase")
        diagnostic = type(error).__name__
        if matching_claim_rows is not None or selectable_result_rows is not None:
            if (phase != "search results"
                    or any(type(count) is not int or count < 0
                           for count in (matching_claim_rows, selectable_result_rows))):
                raise ValueError("Invalid case selection diagnostic counts")
            diagnostic += (f"; matching claim rows: {matching_claim_rows}"
                           f"; selectable result rows: {selectable_result_rows}")
        super().__init__(f"CMS case selection {phase} failed ({diagnostic})")


class _ClaimSetupError(RuntimeError):
    """Keep claim diagnostics limited to fixed phases and exception types."""

    PHASES = (
        "claim tab opening", "diagnosis dialog opening", "diagnosis input",
        "diagnosis save", "body part dialog opening", "body part selection",
        "body part save",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid claim setup diagnostic phase")
        super().__init__(f"CMS claim setup {phase} failed ({type(error).__name__})")


class _ClaimantSetupError(RuntimeError):
    """Keep claimant navigation diagnostics limited to fixed phases and types."""

    PHASES = ("claimant tab lookup", "claimant tab opening", "claimant tab readiness")

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid claimant setup diagnostic phase")
        super().__init__(f"CMS claimant setup {phase} failed ({type(error).__name__})")


class _CaseSetupError(RuntimeError):
    """Keep Case setup diagnostics limited to fixed phases and types."""

    PHASES = (
        "case tab lookup", "case tab opening", "case tab readiness",
        "referral source readback", "referral source input", "referral type readback",
        "inventory selection/readback", "case save", "saved referral source readback",
        "manual referral verification",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid case setup diagnostic phase")
        super().__init__(f"CMS case setup {phase} failed ({type(error).__name__})")


class _ProviderSetupError(RuntimeError):
    """Keep Providers diagnostics limited to fixed phases and exception types."""

    PHASES = (
        "initialization", "providers tab lookup", "providers tab opening",
        "providers tab readiness", "existing-row readback", "existing-provider confirmation",
        "provider planning", "search-link lookup", "search-link opening",
        "search readiness", "facility input", "first name input", "last name input",
        "phone input", "city input", "state input", "zip input",
        "search submission", "search result readiness", "search result readback",
        "operator confirmation", "return-tab lookup", "return-tab opening",
        "return-tab readiness", "final provider verification",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid provider setup diagnostic phase")
        super().__init__(f"CMS provider setup {phase} failed ({type(error).__name__})")


class _SupervisorReadbackError(ValueError):
    def __init__(self, reason):
        super().__init__()
        self.reason = reason


class _AttorneySetupError(RuntimeError):
    """Only allowlisted phases and exception types cross the attorney boundary."""

    PHASES = (
        "initialization", "owned context", "tab lookup", "tab opening",
        "tab readiness", "tab verification", "name resolution", "existing-row readback",
        "search opening", "search readiness", "search input", "search submission",
        "search result readiness", "search results", "association",
        "create opening", "create readiness", "create input", "state selection",
        "create type", "create save", "create save readiness", "final readback",
        "manual confirmation", "manual readback",
        "manual return owned context", "manual return tab lookup",
        "manual return tab opening", "manual return tab readiness",
        "manual return tab verification",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid attorney setup diagnostic phase")
        super().__init__(f"CMS attorney setup {phase} failed ({type(error).__name__})")


class _BillingSetupError(RuntimeError):
    """Only allowlisted phases and exception types cross the billing boundary."""

    PHASES = (
        "initialization", "owned context", "tab lookup", "tab opening",
        "tab readiness", "tab verification", "referral readback", "activity readback",
        "activity opening", "activity readiness", "activity input",
        "note completion", "activity save", "activity save readiness", "final readback",
        "manual confirmation", "manual readback",
        "manual return owned context", "manual return tab lookup",
        "manual return tab opening", "manual return tab readiness",
        "manual return tab verification",
        "case owned context", "case tab lookup", "case tab opening",
        "case tab readiness", "case tab verification",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid billing setup diagnostic phase")
        kind = _billing_case_error_kind(error)
        self.diagnostic = f"CMS billing setup {phase} failed ({kind})"
        super().__init__(self.diagnostic)


def _billing_case_error_kind(error):
    kind = type(error).__name__
    return kind if kind in (
        "TimeoutError", "TimeoutException", "JavascriptException", "WebDriverException",
        "NoSuchElementException", "StaleElementReferenceException", "NoSuchFrameException",
        "NoSuchWindowException", "ElementNotInteractableException", "RuntimeError",
        "ValueError", "TypeError", "KeyError", "AttributeError", "ImportError",
        "ModuleNotFoundError", "SystemExit",
    ) else "OtherError"


class _BillingInspectionStop(_BillingSetupError):
    """Failed tab verification cannot enter Billing's corrective recovery."""

    def __init__(self, phase, error, detail):
        super().__init__(phase, error)
        self.args = (f"{self}; {detail}",)


class _BillingCaseInspectionStop(_BillingInspectionStop):
    def __init__(self, error, detail):
        super().__init__("case tab verification", error, detail)


_BILLING_CASE_OBSERVATION_SCRIPT = r"""
    var stage = arguments[0], element = arguments[1], referralId = arguments[2];
    if (stage === "target") {
        if (!element) { return null; }
        var text = element.innerText, label = "unavailable";
        if (typeof text === "string") {
            text = text.replace(/^\s+|\s+$/g, "").toLowerCase();
            label = text === "" ? "blank" :
                (text === "case" ? "Case" :
                (text === "customer" ? "Customer" : (text === "billing" ? "Billing" : "other")));
        }
        return {
            label: label,
            owner_document: element.ownerDocument === document ? "same" : "different",
            visible: typeof element.offsetWidth === "number" && typeof element.offsetHeight === "number" ?
                (element.offsetWidth > 0 && element.offsetHeight > 0 ? "yes" : "no") : "unavailable",
            enabled: element.disabled === true ? "no" : "yes"
        };
    }
    if (stage === "failure") {
        var ready = document.readyState;
        return {
            document: document ? "available" : "unavailable",
            mode: typeof document.documentMode === "number" ?
                (document.documentMode === 8 ? "8" : "other") : "unavailable",
            frame: window === window.top ? "top" : "child",
            readiness: ready === "loading" || ready === "interactive" || ready === "complete" ?
                ready : (typeof ready === "string" ? "other" : "unavailable"),
            referral_control: document.getElementById(referralId) ? "present" : "absent"
        };
    }
    return null;
"""


def _billing_case_observation(driver, stage, *, target=None, referral_id=None):
    fields = {
        "target": {
            "label": ("Case", "Customer", "Billing", "other", "blank", "unavailable"),
            "owner_document": ("same", "different", "unavailable"),
            "visible": ("yes", "no", "unavailable"),
            "enabled": ("yes", "no", "unavailable"),
        },
        "failure": {
            "document": ("available", "unavailable"),
            "mode": ("8", "other", "unavailable"),
            "frame": ("top", "child", "unavailable"),
            "readiness": ("loading", "interactive", "complete", "other", "unavailable"),
            "referral_control": ("present", "absent", "unavailable"),
        },
    }
    try:
        schema = fields[stage]
        payload = driver.execute_script(_BILLING_CASE_OBSERVATION_SCRIPT, stage, target, referral_id)
        if (type(payload) is not dict or any(type(key) is not str for key in payload)
                or set(payload) != set(schema)
                or any(type(payload[key]) is not str or payload[key] not in choices
                       for key, choices in schema.items())):
            raise ValueError()
        return ", ".join(f"{key}={payload[key]}" for key in schema)
    except (Exception, SystemExit) as error:
        return f"unavailable ({_billing_case_error_kind(error)})"


def _billing_inspection(settings, diagnostic, context):
    try:
        notify = settings.get("notify")
        if not callable(notify):
            return "unavailable (missing or invalid callback); stopped"
        notify("CMS Registration", f"{diagnostic}. "
               "Leave this UI prompt and the browser open for read-only inspection of the "
               f"current failed {context} context. Acknowledgement ends inspection and stops this run "
               "without retry, Save, corrective confirmation or downstream continuation. "
               "Earlier setup or activity work may already have changed CMS.")
    except (Exception, SystemExit) as error:
        return f"failed ({_billing_case_error_kind(error)}); stopped"
    return "acknowledged; stopped"


def _billing_case_inspection(settings, diagnostic):
    return _billing_inspection(settings, diagnostic, "Case")


class _CustomerSetupError(RuntimeError):
    """Keep Customer diagnostics limited to fixed phases and exception types."""

    PHASES = (
        "owned context", "customer tab lookup", "customer tab opening",
        "customer tab readiness", "expected employer", "matrix resolution",
        "employer popup opening", "employer popup readiness", "employer search",
        "employer assignment", "employer readback", "bill to guard",
        "location opening", "location readiness", "location search",
        "location matching", "location assignment", "bill to readback",
        "frame return", "manual confirmation", "manual readback",
        "employer equivalence input", "employer equivalence readback",
        "employer equivalence confirmation",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid customer setup diagnostic phase")
        self.phase = phase
        kind = _billing_case_error_kind(error)
        self.diagnostic = f"CMS customer setup {phase} failed ({kind})"
        super().__init__(self.diagnostic)


def _customer_lookup_inspection(settings, diagnostic):
    try:
        notify = settings.get("notify")
        if not callable(notify):
            return "unavailable (missing or invalid callback); stopped"
        notify("CMS Registration", f"{diagnostic}. "
               "Leave this UI prompt and the browser open for read-only inspection of the "
               "current failed Customer lookup context. Acknowledgement ends inspection and "
               "stops this run without retry, Save, corrective confirmation or downstream "
               "continuation. It does not attest employer identity or any business decision. "
               "Earlier setup or lookup work may already have changed CMS.")
    except (Exception, SystemExit) as error:
        return f"failed ({_billing_case_error_kind(error)}); stopped"
    return "acknowledged; stopped"


class _SupervisorSetupError(RuntimeError):
    """Expose only fixed supervisor phases, exception types and readback reasons."""

    PHASES = (
        "case context", "mapping file", "mapping lookup", "current supervisor readback",
        "popup opening", "popup readiness", "popup result matching", "assignment",
        "frame return", "save", "saved supervisor readback", "manual confirmation",
        "manual supervisor readback",
    )

    def __init__(self, phase, error):
        if phase not in self.PHASES:
            raise ValueError("Invalid supervisor setup diagnostic phase")
        reason = ""
        error_type = type(error).__name__
        if isinstance(error, _SupervisorReadbackError):
            error_type = "ValueError"
            if phase in ("current supervisor readback", "saved supervisor readback",
                         "manual supervisor readback"):
                for allowed_reason in (
                        "unreadable field", "blank required field",
                        "full-name mismatch (commas/whitespace only)", "full-name mismatch"):
                    if error.reason == allowed_reason:
                        reason = f": {allowed_reason}"
                        break
        super().__init__(f"CMS supervisor setup {phase} failed ({error_type}){reason}")


def _search_case_claim(driver, request, *, case_id=""):
    phase = "menu readiness"
    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        from fcm_intake.cms.session import element_click, legacy_safe_type

        driver.switch_to.default_content()
        WebDriverWait(driver, 20).until(
            lambda d: d.execute_script(
                "return document.readyState === 'complete' && typeof window.h$ === 'function';"
            )
        )
        # Preserve the legacy IE-mode menu settling intervals.
        time.sleep(2)
        phase = "menu navigation"
        driver.execute_script("h$(0);")
        time.sleep(1)
        driver.execute_script("h$(10);")
        time.sleep(1)
        element_click(_element(driver, By.ID, "mmlink1"), driver=driver)
        phase = "search frame"
        WebDriverWait(driver, 20).until(EC.frame_to_be_available_and_switch_to_it(
            (By.XPATH, "/html/body/form/div[5]/iframe")
        ))
        phase = "search input"
        for selector, value in (
            ("ctl00_Body_CaseNumber", case_id),
            ("ctl00_Body_ClaimNumber", request.claim_number),
        ):
            if not legacy_safe_type(By.ID, selector, value, driver=driver):
                raise RuntimeError("Search input readback failed")
        phase = "search submission"
        element_click(_element(driver, By.ID, "ctl00_Body_SearchButton"), driver=driver)
        phase = "search results"
        return _element(driver, By.ID, "ctl00_Body_SearchResultsTable_Table")
    except (Exception, SystemExit) as error:
        raise _CaseSearchError(
            f"CMS case search {phase} failed ({type(error).__name__})"
        ) from None


def find_cases(driver, request, settings):
    """CMS.py:981-1140: search then read real case candidates (FCM-1.1-98).

    Search only by claim, not an unverified Unity case ID, to enumerate all
    candidates. Each profile case is opened to read assigned date/type before
    controller selection. The winning case is explicitly reopened during setup.
    Search UI result traversal is legacy-intent based and must be checked in UAT.
    Each Select Case leaves the profile; ClaimantProfileLink restores it before
    another profile-row inspection (CMS.py:1097-1100).
    """
    from selenium.webdriver.common.by import By
    from fcm_intake.cms.session import element_click
    results = _search_case_claim(driver, request)
    phase = "result row inspection"
    try:
        row_count = len(results.find_elements(By.TAG_NAME, "tr"))
        candidates = []
        for row_index in range(1, row_count):
            phase = "result row inspection"
            results = _element(driver, By.ID, "ctl00_Body_SearchResultsTable_Table")
            rows = results.find_elements(By.TAG_NAME, "tr")
            if row_index >= len(rows):
                raise ValueError("CMS search results changed during inspection")
            phase = "claimant profile opening"
            element_click(rows[row_index].find_element(By.TAG_NAME, "a"), driver=driver)
            profile = _element(driver, By.ID, "ctl00_Body_ClaimantProfileTable_Table")
            phase = "claim section inspection"
            profile_rows = profile.find_elements(By.TAG_NAME, "tr")
            in_claim = False
            for profile_index in range(len(profile_rows)):
                phase = "claim section inspection"
                # Selecting a case navigates away from the profile. Re-enter via
                # ClaimantProfileLink before inspecting another row (CMS.py:1097).
                profile = _element(driver, By.ID, "ctl00_Body_ClaimantProfileTable_Table")
                profile_rows = profile.find_elements(By.TAG_NAME, "tr")
                if profile_index >= len(profile_rows):
                    raise ValueError("CMS profile rows changed during inspection")
                row = profile_rows[profile_index]
                text = " ".join(row.text.split())
                if request.claim_number in text and not text.lower().startswith("case"):
                    in_claim = True
                    continue
                if in_claim and not text.lower().startswith("case"):
                    break
                if not in_claim:
                    continue
                parts = text.split()
                case_id = parts[2] if len(parts) > 2 else ""
                if not case_id:
                    continue
                phase = "case selection"
                element_click(
                    row.find_element(By.XPATH, ".//a[normalize-space()='Select Case']"),
                    driver=driver,
                )
                phase = "assigned date readback"
                raw_date = _element(driver, By.ID, "ctl00_Body_AssignedDateField_DatePickerTextBox").get_attribute("value") or ""
                try:
                    assigned = datetime.strptime(raw_date.strip(), "%d-%b-%Y").date().isoformat()
                except ValueError:
                    try:
                        assigned = datetime.strptime(raw_date.strip(), "%m/%d/%Y").date().isoformat()
                    except ValueError:
                        assigned = ""
                phase = "referral type readback"
                from selenium.webdriver.support.ui import Select
                referral_label = Select(_element(driver, By.ID, "ctl00_Body_ReferralTypeField")).first_selected_option.text.strip()
                referral = ("Medical" if "medical" in referral_label.casefold() else
                            "Vocational" if "vocational" in referral_label.casefold() else referral_label)
                phase = "candidate validation"
                candidates.append(_contract().CaseCandidate(case_id, assigned, referral))
                phase = "profile return"
                element_click(
                    _element(driver, By.ID, "ctl00_Body_ClaimantProfileLink"),
                    driver=driver,
                )
                profile = _element(driver, By.ID, "ctl00_Body_ClaimantProfileTable_Table")
                profile_rows = profile.find_elements(By.TAG_NAME, "tr")
            phase = "search results return"
            element_click(
                _element(driver, By.PARTIAL_LINK_TEXT, "Return to Claimant Listing"),
                driver=driver,
            )
        phase = "candidate validation"
        matched_referral = [c for c in candidates if c.referral_type.casefold() == request.referral_type.casefold()]
        if any(not c.case_id.strip() for c in matched_referral):
            raise ValueError("CMS case ID missing in search results")
        return candidates
    except (Exception, SystemExit) as error:
        raise _CaseSearchError(
            f"CMS case search {phase} failed ({type(error).__name__})"
        ) from None


def _select_case(driver, request, case):
    phase = "search results"
    try:
        from selenium.webdriver.common.by import By
        from fcm_intake.cms.session import element_click
        # Re-search by claim, then locate the exact confirmed case in the profile.
        results = _search_case_claim(driver, request)
        rows = results.find_elements(By.TAG_NAME, "tr")
        matches = [row for row in rows if request.claim_number in row.text]
        if len(matches) != 1:
            raise _CaseSelectionError(
                phase, ValueError(), matching_claim_rows=len(matches),
                selectable_result_rows=sum(
                    bool(row.find_elements(By.XPATH, ".//td//a")) for row in rows),
            )
        phase = "claimant profile opening"
        element_click(matches[0].find_element(By.TAG_NAME, "a"), driver=driver)
        profile = _element(driver, By.ID, "ctl00_Body_ClaimantProfileTable_Table")
        phase = "case row inspection"
        matches = [row for row in profile.find_elements(By.TAG_NAME, "tr")
                   if case.case_id in row.text.split() and row.text.lower().startswith("case")]
        if len(matches) != 1:
            raise ValueError("CMS case row not unique")
        phase = "case activation"
        # Selection is deliberately by the exact confirmed ID, not date only.
        element_click(
            matches[0].find_element(By.XPATH, ".//a[normalize-space()='Select Case']"),
            driver=driver,
        )
        phase = "case identity readback"
        case_field = _element(driver, By.ID, "ctl00_Body_CaseNumberField")
        displayed = (case_field.get_attribute("value") or case_field.text).strip()
        if displayed != case.case_id:
            raise ValueError("CMS case identity not verified")
        phase = "referral type readback"
        # Recheck referral type after the portal switches to the chosen case.
        from selenium.webdriver.support.ui import Select
        referral_label = Select(_element(driver, By.ID, "ctl00_Body_ReferralTypeField")).first_selected_option.text.strip()
        if request.referral_type.casefold() not in referral_label.casefold():
            raise ValueError("CMS case referral type changed")
    except (_CaseSearchError, _CaseSelectionError):
        raise
    except (Exception, SystemExit) as error:
        raise _CaseSelectionError(phase, error) from None


def _return_frame(driver, main):
    from selenium.webdriver.common.by import By
    driver.switch_to.window(main)
    driver.switch_to.default_content()
    driver.switch_to.frame(_element(driver, By.XPATH, "/html/body/form/div[5]/iframe"))


def _address_cell(driver, selector):
    from selenium.webdriver.common.by import By
    return " ".join(_element(driver, By.ID, selector).text.casefold().split())


def _office_address(driver, data, settings):
    """CMS.py:1214-1279,616-687; select an exact, unique market/location."""
    from selenium.webdriver.common.by import By
    try:
        rows = _load_reference(settings, "office_matrix_filename", OFFICE_MATRIX)
        physical, mailing = _rules().resolve_office(rows, data.get("officeNumber", ""))
        address = _rules().split_address(physical)
        if (not physical or not mailing or not address or not address["Address1"]
                or not address["City"] or not address["State"] or not address["Zip"]):
            raise ValueError("Unresolved office")
        physical_id = "ctl00_Body_ClaimAdjuster_ctl00_PhysicalAddressCell"
        mailing_id = "ctl00_Body_ClaimAdjuster_ctl00_MailingAddressCell"
        if (SequenceMatcher(None, physical, _address_cell(driver, physical_id)).ratio() < .9
                or SequenceMatcher(None, mailing, _address_cell(driver, mailing_id)).ratio() < .9):
            main = driver.current_window_handle
            try:
                _click(driver, By.ID, "ctl00_Body_ClaimAdjuster_ctl00_EditLink")
                old = set(driver.window_handles)
                _click(driver, By.NAME, "ctl00$Body$ctl00$ctl23")
                from selenium.webdriver.support.ui import WebDriverWait
                WebDriverWait(driver, 15).until(lambda d: len(d.window_handles) > len(old))
                driver.switch_to.window(next(h for h in driver.window_handles if h not in old))
                _type(driver, By.ID, "ctl00_Body__Name", "Liberty Mutual Commercial Market")
                _click(driver, By.XPATH, "//input[@type='submit' and @value='Search']")
                result = _element(driver, By.ID, "ctl00_Body__ResultTable_Table")
                matches = [cells for row in result.find_elements(By.XPATH, ".//tbody/tr")
                           if len(cells := row.find_elements(By.TAG_NAME, "td")) > 7
                           and cells[1].text.strip() == "Liberty Mutual Commercial Market"]
                if len(matches) != 1:
                    raise ValueError("Ambiguous office market")
                matches[0][7].find_element(By.TAG_NAME, "a").click()
                _type(driver, By.ID, "ctl00_Body__FilterCity", address["City"])
                _click(driver, By.XPATH, "//input[@type='submit' and @value='Filter List']")
                result = _element(driver, By.ID, "ctl00_Body__LocationTable_Table")
                matches = [cells for row in result.find_elements(By.XPATH, ".//tbody/tr")
                           if len(cells := row.find_elements(By.TAG_NAME, "td")) > 5
                           and cells[0].text.strip().casefold() in mailing.casefold()
                           and cells[0].text.strip()]
                if len(matches) != 1:
                    raise ValueError("Ambiguous office location")
                # Never take a first fuzzy hit without a unique address match.
                matches[0][5].find_element(By.TAG_NAME, "a").click()
            finally:
                _return_frame(driver, main)
            _click(driver, By.XPATH, "//input[@value='Save']")
        # Read the visible saved cells; durable cross-session persistence needs UAT.
        if (SequenceMatcher(None, physical, _address_cell(driver, physical_id)).ratio() < .9
                or SequenceMatcher(None, mailing, _address_cell(driver, mailing_id)).ratio() < .9):
            raise ValueError("Office readback mismatch")
        return _contract().StepResult("adjuster_address", "done")
    except (Exception, SystemExit):
        # Never claim the popup changed both address cells without readback.
        return _manual(settings, "adjuster_address", "Verify and correct adjuster physical and mailing addresses against the controlled office matrix.")


def step_claim(driver, request, case, data, settings):
    """CMS.py:1141-1211: claim diagnosis and body part; office is its own step."""
    phase = "claim tab opening"
    manual_body_part = False
    try:
        from selenium.webdriver.common.by import By
        from fcm_intake.cms.session import element_click, legacy_safe_type

        element_click(_element(driver, By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[1]/td[2]/table[2]/tbody/tr/td[2]/a"), driver=driver)
        if data.get("dxCode"):
            phase = "diagnosis dialog opening"
            element_click(_element(driver, By.ID, "ctl00_Body_ClaimDiagnosisTable_ctl00_AddLinkButton"), driver=driver)
            phase = "diagnosis input"
            if not legacy_safe_type(By.ID, "ctl00_Body_ClaimDiagnosis1_DiagnosisCode", data["dxCode"], driver=driver):
                raise RuntimeError("Diagnosis input readback failed")
            phase = "diagnosis save"
            element_click(_element(driver, By.XPATH, "/html/body/form/div[3]/div[3]/div/input[1]"), driver=driver)
        if data.get("bodyParts"):
            phase = "body part dialog opening"
            element_click(_element(driver, By.ID, "ctl00_Body_ClaimBodyParts1_Section_AddLinkButton"), driver=driver)
            phase = "body part selection"
            from selenium.common.exceptions import ElementNotInteractableException, NoSuchElementException
            from selenium.webdriver.support.ui import Select
            from fcm_intake.cms.session import ie_select_by_visible_text
            dropdown = _element(driver, By.ID, "ctl00_Body_BodyPartDropDown")
            body_part = Select(dropdown)
            if body_part.is_multiple or not dropdown.is_enabled() or not dropdown.is_displayed():
                raise ValueError("Body part dropdown is not available")
            options = body_part.options
            if any(not isinstance(option.text, str) for option in options):
                raise ValueError("Body part options could not be verified")
            try:
                # The shared helper rejects absent targets as ValueError; only
                # an actually missing option belongs in the attended gate.
                if not any(option.text == data["bodyParts"] for option in options):
                    raise NoSuchElementException("Body part option is unavailable")
                if ie_select_by_visible_text(
                        By.ID, "ctl00_Body_BodyPartDropDown", data["bodyParts"], driver=driver) is not True:
                    raise ValueError("Body part selection was not verified")
            except (NoSuchElementException, ElementNotInteractableException):
                resolution = _manual(settings, "claim",
                    "CMS could not select the body-part option automatically. "
                    "Select and verify the correct actual option in the currently open "
                    "CMS body-part dropdown. Do not click Save. Then Confirm to let "
                    "the bot verify and save, or decline to stop.")
                if resolution.status == "failed":
                    return resolution
                # The attended pause may replace the dropdown DOM.
                dropdown = _element(driver, By.ID, "ctl00_Body_BodyPartDropDown")
                body_part = Select(dropdown)
                selected = [option for option in body_part.options if option.is_selected()]
                if len(selected) != 1 or not dropdown.is_enabled():
                    raise ValueError("Body part selection readback failed")
                option = selected[0]
                text, value = option.text, option.get_attribute("value")
                placeholders = {
                    "", "none", "null", "select", "please select", "none selected",
                    "select body part", "select a body part", "please select body part",
                    "please select a body part", "select an option", "please select an option",
                }
                if (not all(isinstance(item, str) and item.strip() for item in (text, value))
                        or any(" ".join(item.split()).casefold().strip(" -_()[]<>.:")
                               in placeholders for item in (text, value))
                        or not option.is_enabled()
                        or body_part.first_selected_option != option):
                    raise ValueError("Body part selection readback failed")
                manual_body_part = True
            if not manual_body_part:
                body_part = Select(_element(driver, By.ID, "ctl00_Body_BodyPartDropDown"))
                if body_part.first_selected_option.text != data["bodyParts"]:
                    raise ValueError("Body part selection readback failed")
            phase = "body part save"
            element_click(_element(driver, By.ID, "ctl00_Body_ButtonSave"), driver=driver)
    except (Exception, SystemExit) as error:
        raise _ClaimSetupError(phase, error) from None
    if manual_body_part:
        return _contract().StepResult("claim", "manual", "Operator-selected CMS body part verified")
    return _contract().StepResult("claim", "done")


def step_claimant(driver, request, case, data, settings):
    """CMS.py:1281-1287: open claimant; ordinary referrals remain unchanged."""
    phase = "claimant tab lookup"
    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        from fcm_intake.cms.session import element_click, wait_page_ready

        tab = WebDriverWait(driver, 60).until(EC.element_to_be_clickable((
            By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[2]/td[2]/table[1]/tbody/tr/td[2]/a"
        )))
        phase = "claimant tab opening"
        element_click(tab, driver=driver)
        phase = "claimant tab readiness"
        time.sleep(5)
        # Document readiness is not a claimant identity or field readback.
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError("Claimant document readiness not verified")
    except (Exception, SystemExit) as error:
        raise _ClaimantSetupError(phase, error) from None
    return _contract().StepResult("claimant", "done")


def _supervisor_name_key(value):
    """Normalize the portal's Last, First display without inferring name parts."""
    parts = value.split(",")
    if len(parts) == 2 and all(part.strip() for part in parts):
        value = parts[1] + " " + parts[0]
    elif len(parts) != 1:
        return None
    return " ".join(value.split()).casefold() or None


def _supervisor(driver, case, settings):
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from fcm_intake.cms.session import element_click, wait_page_ready

    field = "ctl00_Body_CaseSupervisorNameField"
    main = None
    supervisor = None
    phase = "case context"

    def readback():
        _element(driver, By.ID, field)
        value = driver.execute_script(
            "var field = document.getElementById(arguments[0]); return field ? field.value : null;",
            field)
        if not isinstance(value, str):
            raise _SupervisorReadbackError("unreadable field")
        return value.strip()

    def verify(value):
        if not value:
            raise _SupervisorReadbackError("blank required field")
        if supervisor is not None and value.casefold() != supervisor.casefold():
            actual_key = _supervisor_name_key(value)
            if actual_key is not None and actual_key == _supervisor_name_key(supervisor):
                return
            actual_format = " ".join(value.replace(",", " ").split()).casefold()
            expected_format = " ".join(supervisor.replace(",", " ").split()).casefold()
            reason = ("full-name mismatch (commas/whitespace only)"
                      if actual_format == expected_format else "full-name mismatch")
            raise _SupervisorReadbackError(reason)

    try:
        main = driver.current_window_handle
        phase = "mapping file"
        rows = _load_reference(settings, "supervisor_mapping_filename", SUPERVISOR_MAPPING)
        phase = "mapping lookup"
        prefix = case.case_id[:2]
        if rows is None or not prefix or sum(1 for row in rows[5:] if len(row) > 1 and str(row[0] or "").strip().casefold() == prefix.casefold()) != 1:
            raise ValueError("Supervisor mapping ambiguous")
        mapped = _rules().supervisor_lookup(rows[5:], prefix)
        if not isinstance(mapped, str) or len(mapped.split()) < 2:
            raise ValueError("Supervisor unresolved")
        supervisor = mapped.strip()
        phase = "current supervisor readback"
        current = readback()
        if not current:
            try:
                phase = "popup opening"
                old = set(driver.window_handles)
                element_click(_element(driver, By.XPATH,
                    "/html/body/form/div[3]/table/tbody/tr[9]/td[2]/a"), driver=driver)
                phase = "popup readiness"
                time.sleep(3)
                WebDriverWait(driver, 10).until(lambda d: set(d.window_handles) - old)
                new = set(driver.window_handles) - old
                if len(new) != 1:
                    raise ValueError("Supervisor popup ambiguous")
                driver.switch_to.window(new.pop())
                if not wait_page_ready(driver, timeout=20):
                    raise RuntimeError("Supervisor popup not ready")
                phase = "popup result matching"
                table = _element(driver, By.ID, "ctl00_Body_Results_Table")
                tokens = supervisor.split()[:2]
                matches = [cells for row in table.find_elements(By.XPATH, ".//tbody/tr")
                           if len(cells := row.find_elements(By.TAG_NAME, "td")) > 2
                           and all(t.casefold() in cells[0].text.casefold() for t in tokens)]
                if len(matches) != 1:
                    raise ValueError("Ambiguous supervisor")
                # Do not choose the first row if both mapped name tokens recur.
                phase = "assignment"
                time.sleep(1)
                element_click(matches[0][2].find_element(By.TAG_NAME, "a"), driver=driver)
                time.sleep(5)
            finally:
                # Preserve the original failure phase unless restoration itself fails.
                try:
                    _return_frame(driver, main)
                except (Exception, SystemExit):
                    phase = "frame return"
                    raise
            phase = "save"
            element_click(_element(driver, By.ID, "ctl00_Body_SaveButton"), driver=driver)
            time.sleep(5)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError("Supervisor save not ready")
            phase = "saved supervisor readback"
        # Read supervisor from the saved Case tab, not the search popup.
        verify(readback())
        return _contract().StepResult("case_supervisor", "done")
    except (Exception, SystemExit) as error:
        diagnostic = str(_SupervisorSetupError(phase, error))

    def failed(detail):
        return _contract().StepResult("case_supervisor", "failed", detail)

    # Never ask for a correction or accept a readback in an unowned popup/frame.
    if main is None or phase == "frame return":
        return failed(diagnostic)
    try:
        _return_frame(driver, main)
    except (Exception, SystemExit) as error:
        return failed(f"{diagnostic}; {_SupervisorSetupError('frame return', error)}")
    try:
        approved = settings["confirm"]("CMS Registration",
            f"Automatic supervisor assignment failed. {diagnostic}. "
            "Confirm DOES NOT fill or save the supervisor. Manually select the correct "
            "case supervisor against the controlled dedicated mapping and click Save "
            "in CMS first. Then confirm to verify the saved field, or decline to stop.")
    except (Exception, SystemExit) as error:
        return failed(f"{diagnostic}; {_SupervisorSetupError('manual confirmation', error)}")
    if not approved:
        return failed(f"{diagnostic}; operator declined manual correction")
    try:
        phase = "frame return"
        _return_frame(driver, main)
        phase = "manual supervisor readback"
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError("Supervisor document not ready")
        verify(readback())
    except (Exception, SystemExit) as error:
        return failed(f"{diagnostic}; {_SupervisorSetupError(phase, error)}")
    provenance = ("controlled mapping matched" if supervisor is not None
                  else "operator-authorized mapping resolution; nonblank field verified")
    return _contract().StepResult("case_supervisor", "manual",
        f"{diagnostic}; manual correction verified ({provenance})")


def step_case(driver, request, case, data, settings):
    """CMS.py:1289-1383: referral source, inventory and supervisor."""
    phase = "case tab lookup"
    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import Select, WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        from fcm_intake.cms.session import (
            element_click, legacy_safe_type, wait_page_ready, ie_select_by_visible_text,
        )

        tab = WebDriverWait(driver, 60).until(EC.element_to_be_clickable((
            By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[1]/td[2]/table[3]/tbody/tr/td[2]/a"
        )))
        phase = "case tab opening"
        element_click(tab, driver=driver)
        phase = "case tab readiness"
        time.sleep(10)
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError("Case document readiness not verified")
        phase = "referral source readback"
        field = _element(driver, By.ID, "ctl00_Body_ReferralSourceField")
        source_issue = not request.referral_number
        current_source = (field.get_attribute("value") or "").strip()
        if request.referral_number and not current_source:
            phase = "referral source input"
            if legacy_safe_type(By.ID, "ctl00_Body_ReferralSourceField",
                                request.referral_number, driver=driver) is not True:
                raise RuntimeError("Referral source input not verified")
            if (_element(driver, By.ID, "ctl00_Body_ReferralSourceField").get_attribute("value") or "").strip() != request.referral_number:
                raise ValueError("Referral source input readback mismatch")
        elif request.referral_number and current_source != request.referral_number:
            source_issue = True
        phase = "referral type readback"
        displayed_type = Select(_element(driver, By.ID, "ctl00_Body_ReferralTypeField")).first_selected_option.text
        if request.referral_type.casefold() not in displayed_type.casefold():
            raise ValueError("CMS referral type mismatch")
        phase = "inventory selection/readback"
        if ie_select_by_visible_text(
                By.ID, "ctl00_Body_InventoryTypeField",
                _rules().inventory_type(data.get("refType") or request.referral_type),
                driver=driver) is not True:
            raise RuntimeError("Inventory selection not verified")
        phase = "case save"
        element_click(_element(driver, By.ID, "ctl00_Body_SaveButton"), driver=driver)
        # Visible readback is not proof of durable cross-session persistence.
        if not current_source and request.referral_number:
            phase = "saved referral source readback"
            readback_source = (_element(driver, By.ID, "ctl00_Body_ReferralSourceField").get_attribute("value") or "").strip()
            if readback_source != request.referral_number:
                source_issue = True
        if source_issue:
            phase = "manual referral verification"
            return _manual(settings, "case",
                "Verify CMS Referral Source against the referral and correct it manually before continuing.")
    except (Exception, SystemExit) as error:
        raise _CaseSetupError(phase, error) from None
    return _contract().StepResult("case", "done")


def _customer_mailing(mailing):
    if not isinstance(mailing, str) or len(mailing.split(",")) < 2:
        raise ValueError()
    parts = [part.strip() for part in mailing.split(",")]
    street = parts[0].casefold()
    match = re.search(r"([A-Za-z\s]+),\s*[A-Za-z]{2}\s*\d{5}(?:-\d{4})?", mailing, re.I)
    city = match.group(1).strip() if match else (parts[-2] if len(parts) > 2 else parts[1])
    if not street or not city or not re.search(r"[A-Za-z]", city):
        raise ValueError()
    return street, city


def _customer_expectations(data, settings):
    phase = "expected employer"
    try:
        name = data.get("xlCustomerName")
        if not isinstance(name, str) or not name.strip():
            raise ValueError()
        phase = "matrix resolution"
        office = data.get("officeNumber")
        if not isinstance(office, (str, int)) or isinstance(office, bool) or not str(office).strip():
            raise ValueError()
        rows = _load_reference(settings, "office_matrix_filename", OFFICE_MATRIX)
        _, mailing = _rules().resolve_office(rows, office)
        street, _ = _customer_mailing(mailing)
        return name.strip(), mailing, street
    except (Exception, SystemExit) as error:
        raise _CustomerSetupError(phase, error) from None


def _customer_text(driver, selector):
    from selenium.webdriver.common.by import By
    value = _element(driver, By.ID, selector).text
    if not isinstance(value, str) or not value.strip():
        raise ValueError()
    return value.strip()


def _customer_fields(driver):
    """Presence of these fields, not readyState alone, identifies Customer."""
    from selenium.webdriver.common.by import By
    values = [_element(driver, By.ID, selector).text for selector in (
        "EmployerField", "ctl00_Body_BillToDetail_Content_CustomerName",
        "ctl00_Body_BillToDetail_Content_Address1")]
    if not all(isinstance(value, str) for value in values):
        raise ValueError()
    return [value.strip() for value in values]


def _customer_alias_identity(driver, data):
    """Read the exact per-run identity; never substitute the historical on-case ID."""
    from selenium.webdriver.common.by import By
    phase = "employer equivalence input"
    try:
        expected = data.get("xlCustomerName")
        if (not isinstance(expected, str) or not expected.strip()
                or expected.strip().casefold() == "not given / not applicable"):
            raise ValueError()
        phase = "employer equivalence readback"
        actual = _element(driver, By.ID, "EmployerField").text
        selected_id = _element(
            driver, By.ID, "ctl00_Body_CustomerSelector_EmployerIdField").get_attribute("value")
        if (not isinstance(actual, str) or not actual.strip()
                or actual.strip().casefold() == "not given / not applicable"
                or not isinstance(selected_id, str) or not selected_id.strip()):
            raise ValueError()
        return expected, actual, selected_id
    except (Exception, SystemExit) as error:
        raise _CustomerSetupError(phase, error) from None


def _customer_verify(driver, data, settings, *, employer_equivalence=None):
    name, _, street = _customer_expectations(data, settings)
    phase = "employer readback"
    try:
        if employer_equivalence is None:
            employer = _customer_text(driver, "EmployerField")
            if employer.casefold() != name.casefold():
                raise ValueError()
        else:
            phase = "employer equivalence readback"
            if _customer_alias_identity(driver, data) != employer_equivalence:
                raise ValueError()
        phase = "bill to guard"
        bill_to = _customer_text(driver, "ctl00_Body_BillToDetail_Content_CustomerName")
        if bill_to != "Liberty Mutual Commercial Market":
            raise ValueError()
        phase = "bill to readback"
        address = _customer_text(driver, "ctl00_Body_BillToDetail_Content_Address1")
        if street not in address.casefold():
            raise ValueError()
    except _CustomerSetupError:
        raise
    except (Exception, SystemExit) as error:
        raise _CustomerSetupError(phase, error) from None


_CUSTOMER_HANDOFF_SCRIPT = r"""
    var label = document.getElementById(arguments[0]);
    var employer = document.getElementById(arguments[1]);
    var onCase = document.getElementById(arguments[2]);
    var text, visible = "missing", relationship = "unavailable";
    function blank(value) { return value.replace(/^\s+|\s+$/g, "") === ""; }
    if (label) {
        text = label.innerText;
        if (typeof text !== "string") { return null; }
        visible = blank(text) ? "blank" :
            text.replace(/^\s+|\s+$/g, "") === "Not Given / Not Applicable" ?
            "placeholder" : "populated";
    }
    if (employer && onCase && typeof employer.value === "string" &&
            typeof onCase.value === "string" &&
            !blank(employer.value) && !blank(onCase.value)) {
        relationship = employer.value === onCase.value ? "same" : "different";
    }
    return {label: visible, relationship: relationship};
"""


def _customer_handoff(driver):
    """Optional observation: no identifiers or label text leave the document."""
    try:
        state = driver.execute_script(
            _CUSTOMER_HANDOFF_SCRIPT, "EmployerField",
            "ctl00_Body_CustomerSelector_EmployerIdField", "ctl00_Body_onCaseEmployerid")
        if (type(state) is not dict or set(state) != {"label", "relationship"}
                or type(state["label"]) is not str
                or state["label"] not in ("blank", "placeholder", "populated", "missing")
                or type(state["relationship"]) is not str
                or state["relationship"] not in ("same", "different", "unavailable")):
            raise ValueError()
        return f'{state["label"]}/{state["relationship"]}'
    except (Exception, SystemExit) as error:
        return f"unavailable ({_billing_case_error_kind(error)})"


def _customer_popup_state(driver, handle):
    try:
        if handle is None:
            return "unavailable"
        return "open" if handle in driver.window_handles else "closed"
    except (Exception, SystemExit) as error:
        return f"unavailable ({_billing_case_error_kind(error)})"


def _customer_popup(driver, old, timeout):
    from selenium.webdriver.support.ui import WebDriverWait

    def new_window(d):
        new = set(d.window_handles) - old
        if len(new) > 1:
            raise ValueError()
        return next(iter(new)) if len(new) == 1 else False

    handle = WebDriverWait(driver, timeout).until(new_window)
    # Recheck ownership after the wait; never choose an unrelated existing tab.
    if set(driver.window_handles) - old != {handle}:
        raise ValueError()
    driver.switch_to.window(handle)
    return handle


def _customer_location(driver, mailing, main=None):
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from fcm_intake.cms.session import element_click, legacy_safe_type, wait_page_ready
    phase = "location opening"
    try:
        if main is None:
            main = driver.current_window_handle
        try:
            street, city = _customer_mailing(mailing)
            old = set(driver.window_handles)
            element_click(_element(driver, By.ID,
                "ctl00_Body_BillToDetail_Content_AssignPrimaryLocationLink"), driver=driver)
            phase = "location readiness"
            _customer_popup(driver, old, 60)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
            phase = "location search"
            if legacy_safe_type(By.ID, "ctl00_Body_City", city, driver=driver) is not True:
                raise RuntimeError()
            element_click(_element(driver, By.ID, "ctl00_Body_Search"), driver=driver)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
            phase = "location matching"
            locations = WebDriverWait(driver, 30).until(
                EC.presence_of_all_elements_located((By.CLASS_NAME, "Location")))
            matches = []
            for location in locations:
                if street not in location.find_element(By.TAG_NAME, "pre").text.casefold():
                    continue
                contacts = location.find_element(By.XPATH, ".//following::ul[contains(@id,'Contacts')][1]")
                matches.extend(c.find_element(By.TAG_NAME, "a") for c in contacts.find_elements(By.XPATH, ".//li[contains(@class,'Contact')]")
                               if c.find_element(By.TAG_NAME, "pre").text.strip().casefold() == "unknown, unknown")
            if len(matches) != 1:
                raise ValueError()
            phase = "location assignment"
            element_click(matches[0], driver=driver)
        finally:
            try:
                _return_frame(driver, main)
            except (Exception, SystemExit):
                phase = "frame return"
                raise
        time.sleep(5)
    except (Exception, SystemExit) as error:
        raise _CustomerSetupError(phase, error) from None


def _customer_tab(driver):
    from selenium.common.exceptions import JavascriptException
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC

    native = EC.element_to_be_clickable((
        By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[1]/td[2]/table[3]/tbody/tr/td[2]/a"))
    legacy = False

    def lookup(owned):
        nonlocal legacy
        if not legacy:
            try:
                return native(owned)
            except JavascriptException:
                legacy = True
        from selenium.webdriver.remote.webelement import WebElement
        # IE8's native clickability atom can fail; inspect only this document.
        element = owned.execute_script(r"""
            if (document.documentMode !== 8) { return false; }
            var links = document.getElementsByTagName("a");
            var found = null, count = 0, i, text, node, style, value, href, disabled;
            function normalized(value) {
                return value.replace(/\s+/g, " ").replace(/^\s+|\s+$/g, "");
            }
            for (i = 0; i < links.length; i++) {
                text = links[i].innerText;
                if (typeof text !== "string") { return false; }
                if (normalized(text) === "Customer") {
                    found = links[i];
                    count++;
                }
            }
            if (count !== 1 || found.ownerDocument !== document ||
                    !(found.offsetWidth > 0) || !(found.offsetHeight > 0)) {
                return false;
            }
            href = found.getAttribute("href");
            if (typeof href !== "string" || !normalized(href) ||
                    normalized(href) === "#") { return false; }
            for (node = found; node && node.nodeType === 1; node = node.parentNode) {
                disabled = node.getAttributeNode("disabled");
                if (node.disabled || (disabled && disabled.specified !== false)) {
                    return false;
                }
                value = node.getAttribute("aria-disabled");
                if (value !== null && (typeof value !== "string" ||
                        normalized(value).toLowerCase() !== "false")) { return false; }
                style = node.currentStyle;
                if (!style || typeof style.display !== "string" ||
                        typeof style.visibility !== "string" ||
                        style.display.toLowerCase() === "none" ||
                        style.visibility.toLowerCase() === "hidden" ||
                        style.visibility.toLowerCase() === "collapse") { return false; }
                if (node === document.documentElement) { return found; }
            }
            return false;
        """)
        if element is False:
            return False
        if not isinstance(element, WebElement):
            raise ValueError()
        return element

    return WebDriverWait(driver, 60).until(lookup)


def step_customer(driver, request, case, data, settings):
    """CMS.py:1385-1519: employer and matrix-based Bill To selection."""
    main = None
    phase = "owned context"
    selection = None
    employer_popup = None
    employer_frame_restored = False
    expected_employer = data.get("xlCustomerName")
    try:
        from selenium.common.exceptions import TimeoutException
        from selenium.webdriver.common.by import By
        from fcm_intake.cms.session import element_click, legacy_safe_type, wait_page_ready

        main = driver.current_window_handle
        if main not in driver.window_handles:
            raise ValueError()
        phase = "customer tab lookup"
        tab = _customer_tab(driver)
        phase = "customer tab opening"
        element_click(tab, driver=driver)
        phase = "customer tab readiness"
        time.sleep(5)
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError()
        employer = _customer_fields(driver)[0]
        name, mailing, street = _customer_expectations(data, settings)
        phase = "employer readback"
        if employer == "Not Given / Not Applicable":
            selection = {
                "opening": "not completed", "results": "not reached",
                "choice": "not reached", "popup": "unavailable",
                "pre": _customer_handoff(driver), "post": "unavailable (not reached)",
            }
            lookup_failure = None
            try:
                phase = "employer popup opening"
                opening = _element(driver, By.XPATH,
                    "/html/body/form/div[3]/table[1]/tbody/tr[1]/td[2]/a")
                old = set(driver.window_handles)
                if main not in old or driver.current_window_handle != main:
                    raise ValueError()
                timed_out_popup = None
                try:
                    element_click(opening, driver=driver)
                    selection["opening"] = "returned normally"
                except TimeoutException:
                    # IE8 showModalDialog can open synchronously before the click times out.
                    handles = set(driver.window_handles)
                    if not old <= handles or len(handles - old) != 1:
                        raise
                    timed_out_popup = next(iter(handles - old))
                    employer_popup = timed_out_popup
                phase = "employer popup readiness"
                employer_popup = _customer_popup(driver, old, 10)
                if not wait_page_ready(driver, timeout=20):
                    raise RuntimeError()
                if timed_out_popup is not None:
                    if driver.execute_script(
                            "return document.documentMode === 8 && typeof window.dialogArguments !== 'undefined';") is not True:
                        raise ValueError()
                    _element(driver, By.ID, "ctl00_Body_Name")
                    _element(driver, By.ID, "ctl00_Body_Search")
                phase = "employer search"
                clean_name = re.sub(r"[^A-Za-z0-9 &]+", "", name)
                if not clean_name.strip():
                    raise ValueError()
                if timed_out_popup is not None:
                    handles = set(driver.window_handles)
                    if (not old <= handles or handles - old != {timed_out_popup}
                            or driver.current_window_handle != timed_out_popup):
                        raise ValueError()
                    selection["opening"] = "timeout recovered"
                if legacy_safe_type(By.ID, "ctl00_Body_Name", clean_name, driver=driver) is not True:
                    raise RuntimeError()
                element_click(_element(driver, By.ID, "ctl00_Body_Search"), driver=driver)
                if not wait_page_ready(driver, timeout=20):
                    raise RuntimeError()
                _element(driver, By.ID, "ctl00_Body_SearchResults")
                matches = [a for a in driver.find_elements(By.TAG_NAME, "a") if a.text.strip() == "Assign"]
                selection["results"] = "zero" if not matches else "one" if len(matches) == 1 else "multiple"
                phase = "employer assignment"
                if len(matches) == 1:
                    selection["choice"] = "automatic Assign"
                    element_click(matches[0], driver=driver)
                else:
                    selection["choice"] = "operator requested"
                    settings["notify"]("CMS Customer", "Select the exact customer in the search window, then acknowledge.")
                    selection["choice"] = "operator acknowledged (not selection proof)"
            except (Exception, SystemExit) as error:
                failure = error if isinstance(error, _CustomerSetupError) else _CustomerSetupError(phase, error)
                lookup_failure = failure.diagnostic
                # Pause before restoration or any further driver observation can hide the failure.
                inspection = _customer_lookup_inspection(settings, lookup_failure)
            finally:
                try:
                    _return_frame(driver, main)
                    employer_frame_restored = True
                    if lookup_failure is None:
                        selection["popup"] = _customer_popup_state(driver, employer_popup)
                except (Exception, SystemExit) as error:
                    restoration = _CustomerSetupError("frame return", error).diagnostic
                    if lookup_failure is None:
                        lookup_failure = restoration
                        inspection = _customer_lookup_inspection(settings, lookup_failure)
                    else:
                        lookup_failure += f"; {restoration}"
            if lookup_failure is not None:
                summary = ", ".join(f"{key}={value}" for key, value in selection.items())
                return _contract().StepResult("customer", "failed",
                    f"{lookup_failure}; selection/handoff: {summary}; inspection={inspection}")
            time.sleep(5)
            selection["post"] = _customer_handoff(driver)
        phase = "employer readback"
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError()
        if _element(driver, By.ID, "EmployerField").text.strip().casefold() != name.casefold():
            raise ValueError()
        phase = "bill to guard"
        bill_to_name = _element(driver, By.ID, "ctl00_Body_BillToDetail_Content_CustomerName").text.strip()
        if bill_to_name != "Liberty Mutual Commercial Market":
            raise ValueError()
        phase = "bill to readback"
        old_address = _element(driver, By.ID, "ctl00_Body_BillToDetail_Content_Address1").text.strip().casefold()
        if street not in old_address:
            _customer_location(driver, mailing, main)
        phase = "frame return"
        _return_frame(driver, main)
        phase = "bill to readback"
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError()
        _customer_verify(driver, data, settings)
        return _contract().StepResult("customer", "done")
    except (Exception, SystemExit) as error:
        failure = error if isinstance(error, _CustomerSetupError) else _CustomerSetupError(phase, error)
        diagnostic = str(failure)
        if selection is not None:
            if employer_frame_restored and selection["post"] == "unavailable (not reached)":
                selection["post"] = _customer_handoff(driver)
            diagnostic += "; selection/handoff: " + ", ".join(
                f"{key}={value}" for key, value in selection.items())

    def failed(detail):
        return _contract().StepResult("customer", "failed", detail)

    if main is None or failure.phase in ("owned context", "frame return"):
        return failed(diagnostic)
    try:
        _return_frame(driver, main)
    except (Exception, SystemExit) as error:
        return failed(f"{diagnostic}; {_CustomerSetupError('frame return', error)}")
    try:
        approved = settings["confirm"]("CMS Registration",
            f"Automatic Customer setup failed. {diagnostic}. "
            "Confirm DOES NOT correct, fill or save anything. Manually open Customer "
            "and correct the employer against the controlled input and Bill To primary "
            "location/contact against the controlled office matrix first. Then confirm "
            "to verify the fields, or decline to stop. Employer-specific policy remains "
            "a separate gate.")
    except (Exception, SystemExit) as error:
        return failed(f"{diagnostic}; {_CustomerSetupError('manual confirmation', error)}")
    if not approved:
        return failed(f"{diagnostic}; operator declined manual correction")
    equivalence_verified = False
    try:
        phase = "frame return"
        _return_frame(driver, main)
        phase = "manual readback"
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError()
        phase = "employer equivalence input"
        if data.get("xlCustomerName") != expected_employer:
            raise ValueError()
        # Verification only: never click a positional tab or rerun assignments.
        try:
            _customer_verify(driver, data, settings)
        except _CustomerSetupError as error:
            if error.phase != "employer readback":
                raise
            identity = _customer_alias_identity(driver, data)
            phase = "employer equivalence confirmation"
            try:
                alias_approved = settings["confirm"]("CMS Employer Equivalence",
                    "Explicit employer equivalence for this run only:\n\n"
                    f"Expected employer:\n{identity[0]}\n\n"
                    f"Currently assigned employer:\n{identity[1]}\n\n"
                    "Do you explicitly attest that these two names denote the SAME "
                    "employer? Yes accepts only this unchanged selected employer for "
                    "this run. No stops. This does not assign, correct or save anything, "
                    "and does not verify persistence. Bill To checks still apply.")
            finally:
                try:
                    _return_frame(driver, main)
                except (Exception, SystemExit):
                    phase = "frame return"
                    raise
            if alias_approved is not True:
                raise ValueError()
            phase = "manual readback"
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
            _customer_verify(driver, data, settings, employer_equivalence=identity)
            equivalence_verified = True
    except (Exception, SystemExit) as error:
        detail = error if isinstance(error, _CustomerSetupError) else _CustomerSetupError(phase, error)
        return failed(f"{diagnostic}; {detail}")
    return _contract().StepResult("customer", "manual",
        f"{diagnostic}; " + (
            "operator employer equivalence confirmed for this run; unchanged selection "
            "and Bill To verified (not persistence proof)"
            if equivalence_verified else "manual correction verified"))


def _tab_providers(driver):
    from selenium.webdriver.common.by import By
    rows = driver.find_elements(By.XPATH, "/html/body/form/div[3]/table/tbody/tr[2]/td/table/tbody/tr")
    return [{"name": row.find_element(By.XPATH, "./td[2]").text.strip()} for row in rows
            if row.find_elements(By.XPATH, "./td[2]/a")]


def _provider_name(provider):
    return str(provider.get("name") or provider.get("facility") or " ".join(
        (str(provider.get("first_name") or ""), str(provider.get("last_name") or "")))).strip()


def step_providers(driver, request, case, data, settings):
    """CMS.py:1524-1558: BRD 113/114 add only missing providers (FCM-1.1-100)."""
    phase = "initialization"
    try:
        from selenium.webdriver.common.by import By
        from fcm_intake.cms.session import element_click, legacy_safe_type, wait_page_ready
        tab_path = "/html/body/form/div[3]/div[1]/table/tbody/tr[2]/td[2]/table[8]/tbody/tr/td[2]/a"
        fields = {
            "facility": ("ctl00_Body_FacilityName", "facility input"),
            "first_name": ("ctl00_Body_FirstName", "first name input"),
            "last_name": ("ctl00_Body_LastName", "last name input"),
            "phone": ("ctl00_Body_Phone_PNumber", "phone input"),
            "city": ("ctl00_Body_City", "city input"),
            "state": ("ctl00_Body_State", "state input"),
            "zip": ("ctl00_Body_ZipCode", "zip input"),
        }
        providers = data.get("providers") or ([dict(name=data.get("xlProvider", ""),
                           type=data.get("provType", ""), phone=data.get("xlProviderPhone", ""))]
                           if data.get("xlProvider") else [])
        phase = "providers tab lookup"
        tab = _element(driver, By.XPATH, tab_path)
        phase = "providers tab opening"
        element_click(tab, driver=driver)
        phase = "providers tab readiness"
        time.sleep(3)
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError()
        phase = "existing-row readback"
        existing = _tab_providers(driver)
        # Tab rows expose names only: the operator must adjudicate phone/type.
        existing_names = {p["name"].casefold() for p in existing}
        ambiguous = [p for p in providers if _provider_name(p).casefold() in existing_names]
        if ambiguous:
            phase = "existing-provider confirmation"
            if not settings["confirm"]("CMS Registration",
                    "Verify each prepopulated provider identity, including phone and type; correct any mismatch manually before continuing."):
                return _contract().StepResult("providers_existing", "failed", "Operator declined")
        phase = "provider planning"
        missing = [p for p in _rules().providers_to_add(providers, [])
                   if _provider_name(p).casefold() not in existing_names]
        for provider in missing:
            phase = "provider planning"
            if not _provider_name(provider):
                return _contract().StepResult("providers", "failed", "Provider identity missing")
            plan = _rules().provider_search_plan(provider)
            phase = "search-link lookup"
            link = _element(driver, By.XPATH, "/html/body/form/div[3]/table/tbody/tr[1]/td/a")
            phase = "search-link opening"
            element_click(link, driver=driver)
            phase = "search readiness"
            time.sleep(3)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
            rows = []
            for criteria in plan:
                for field, value in criteria.items():
                    phase = "provider planning"
                    control, input_phase = fields[field]
                    phase = input_phase
                    if legacy_safe_type(By.ID, control, value, driver=driver) is not True:
                        raise RuntimeError()
                phase = "search submission"
                element_click(_element(driver, By.ID, "ctl00_Body_SearchButton"), driver=driver)
                phase = "search result readiness"
                time.sleep(3)
                if not wait_page_ready(driver, timeout=20):
                    raise RuntimeError()
                phase = "search result readback"
                rows = _element(driver, By.ID, "ctl00_Body_SearchResultsTable_Table").find_elements(By.TAG_NAME, "tr")[1:]
                if 1 <= len(rows) <= 10:
                    break
            message = ("Review provider search; select the exact Add To Case result manually. "
                       "If no result remains, Provider Entry Rules are unavailable: obtain authorization "
                       "and create the record manually before continuing.")
            if len(rows) > 10:
                message = "Provider search remains ambiguous; select and verify the exact record manually before continuing."
            phase = "operator confirmation"
            if not settings["confirm"]("CMS Registration", message):
                return _contract().StepResult("providers", "failed", "Provider not confirmed")
            phase = "return-tab lookup"
            tab = _element(driver, By.XPATH, tab_path)
            phase = "return-tab opening"
            element_click(tab, driver=driver)
            phase = "return-tab readiness"
            time.sleep(3)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
            # Readiness/acknowledgement never substitute for actual tab presence.
            phase = "final provider verification"
            if _provider_name(provider).casefold() not in {
                    p["name"].casefold() for p in _tab_providers(driver)}:
                return _contract().StepResult("providers", "failed", "Provider not present on tab")
        return _contract().StepResult("providers", "manual" if missing or ambiguous else "skipped")
    except (Exception, SystemExit) as error:
        raise _ProviderSetupError(phase, error) from None


_BILLING_CASE_TARGET_SCRIPT = r"""
    var scope = arguments[0], found = null, count = 0, links, i, text, node, style;
    var disabled, value, inScope = false;
    function normalized(value) {
        return value.replace(/\s+/g, " ").replace(/^\s+|\s+$/g, "");
    }
    if (!scope || scope.nodeType !== 1 || scope.ownerDocument !== document) { return false; }
    for (node = scope; node && node !== document.documentElement; node = node.parentNode) {}
    if (node !== document.documentElement) { return false; }
    links = scope.getElementsByTagName("a");
    for (i = 0; i < links.length; i++) {
        text = links[i].innerText;
        if (typeof text !== "string") { return false; }
        if (normalized(text) === "Case") { found = links[i]; count++; }
    }
    if (count !== 1 || found.nodeType !== 1 || found.ownerDocument !== document ||
            typeof found.offsetWidth !== "number" || typeof found.offsetHeight !== "number" ||
            !(found.offsetWidth > 0) || !(found.offsetHeight > 0)) { return false; }
    value = found.getAttribute("href");
    if (typeof value !== "string" || !normalized(value) || normalized(value) === "#") {
        return false;
    }
    for (node = found; node && node.nodeType === 1; node = node.parentNode) {
        if (node.ownerDocument !== document) { return false; }
        if (node === scope) { inScope = true; }
        disabled = node.getAttributeNode("disabled");
        if (node.disabled || (disabled && disabled.specified !== false)) { return false; }
        value = node.getAttribute("aria-disabled");
        if (value !== null && (typeof value !== "string" ||
                normalized(value).toLowerCase() !== "false")) { return false; }
        style = node.currentStyle;
        if (!style || typeof style.display !== "string" ||
                typeof style.visibility !== "string" ||
                style.display.toLowerCase() === "none" ||
                style.visibility.toLowerCase() === "hidden" ||
                style.visibility.toLowerCase() === "collapse") { return false; }
        if (node === document.documentElement) { return inScope ? found : false; }
    }
    return false;
"""


def _billing_case_tab(driver, tab_path):
    from selenium.webdriver.common.by import By
    from selenium.webdriver.remote.webelement import WebElement
    from selenium.webdriver.support.ui import WebDriverWait

    # Both existing tab paths share this navigation container; never search the body.
    scope = _element(driver, By.XPATH, tab_path.split("/table", 1)[0])
    if not isinstance(scope, WebElement) or scope.parent is not driver:
        raise ValueError()

    def lookup(owned):
        element = owned.execute_script(_BILLING_CASE_TARGET_SCRIPT, scope)
        if element is False:
            return False
        if not isinstance(element, WebElement) or element.parent is not owned:
            raise ValueError()
        return element

    return WebDriverWait(driver, 20).until(lookup)


def _required_case_tab(driver, main, tab_path, control, error_type, *, manual=False, settle=3,
                       observe_target=None, resolve_target=None):
    """Return only after identifying the required tab in this driver's main frame."""
    phase = "owned context"
    prefix = "manual return " if manual else ""
    try:
        from selenium.webdriver.common.by import By
        from fcm_intake.cms.session import element_click, wait_page_ready
        if not main:
            raise ValueError()
        _return_frame(driver, main)
        phase = "tab lookup"
        tab = (_element(driver, By.XPATH, tab_path) if resolve_target is None
               else resolve_target(driver, tab_path))
        if observe_target is not None:
            observe_target(tab)
        phase = "tab opening"
        element_click(tab, driver=driver)
        phase = "tab readiness"
        time.sleep(settle)
        if not wait_page_ready(driver, timeout=20):
            raise RuntimeError()
        phase = "tab verification"
        return _element(driver, *control)
    except (Exception, SystemExit) as error:
        raise error_type(prefix + phase, error) from None


def _required_tab_recovery(step, settings, failure, error_type, visit, verify, message, detail):
    """Recovery only navigates and reads; confirmation never retries automatic writes."""
    diagnostic = str(failure)
    phase = "manual confirmation"
    try:
        visit()
        if not settings["confirm"]("CMS Registration", f"{diagnostic}. {message}"):
            return _contract().StepResult(step, "failed", f"{diagnostic}; operator declined")
        visit()
        phase = "manual readback"
        verify()
    except (Exception, SystemExit) as error:
        safe = error if isinstance(error, error_type) else error_type(phase, error)
        return _contract().StepResult(step, "failed", f"{diagnostic}; {safe}")
    return _contract().StepResult(step, "manual", f"{diagnostic}; {detail}")


def step_attorneys(driver, request, case, data, settings):
    """CMS.py:1614-1687: preserve matching/add policy with mandatory tab readback."""
    phase = "initialization"
    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import Select
        from fcm_intake.cms.session import element_click, legacy_safe_type, wait_page_ready
        phase = "owned context"
        main = driver.current_window_handle
        tab_path = "/html/body/form/div[3]/div[1]/table/tbody/tr[2]/td[2]/table[9]/tbody/tr/td[2]/a"
        control = (By.ID, "ctl00_Body_AttorneyListTable_Table")
        def visit(manual=False):
            return _required_case_tab(driver, main, tab_path, control,
                                      _AttorneySetupError, manual=manual)
        def rows():
            return [r for r in _element(driver, *control).find_elements(By.XPATH, ".//tbody/tr")
                    if last.casefold() in r.text.casefold() and first.casefold() in r.text.casefold()]
        def verify():
            if len(rows()) != 1:
                raise ValueError()
        def settle():
            time.sleep(3)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
        visit()
    except _AttorneySetupError:
        raise
    except (Exception, SystemExit) as error:
        raise _AttorneySetupError(phase, error) from None
    first = last = ""
    try:
        phase = "name resolution"
        name = (data.get("xlAttorney") or "").strip()
        first, last = _rules().split_name(name)
        if not first or not last:
            raise ValueError()
        phase = "existing-row readback"
        existing = rows()
        if len(existing) > 1:
            raise ValueError()
        if not existing:
            phase = "search opening"
            element_click(_element(driver, By.XPATH, "/html/body/form/div[3]/table/tbody/tr[4]/td/a"), driver=driver)
            phase = "search readiness"
            settle()
            phase = "search input"
            for field, value in (("LastName", last), ("FirstName", first)):
                if legacy_safe_type(By.ID, "ctl00_Body_" + field, value, driver=driver) is not True:
                    raise RuntimeError()
            phase = "search submission"
            element_click(_element(driver, By.ID, "ctl00_Body_SearchButton"), driver=driver)
            phase = "search result readiness"
            settle()
            phase = "search results"
            table = _element(driver, By.ID, "ctl00_Body_SearchResultsTable_Table")
            matches = [r for r in table.find_elements(By.XPATH, ".//tbody/tr")
                       if f"{last}, {first}".casefold() in r.text.casefold()]
            if len(matches) > 1:
                raise ValueError()
            if matches:
                phase = "association"
                links = matches[0].find_elements(By.TAG_NAME, "a")
                if not links:
                    raise ValueError()
                element_click(links[-1], driver=driver)
            else:
                phase = "create opening"
                element_click(_element(driver, By.ID, "ctl00_Body_AddLink"), driver=driver)
                phase = "create readiness"
                settle()
                phase = "create input"
                for field, value in (("LastName", last), ("FirstName", first),
                                     ("Address1", data.get("xlAttorneyAdd1")),
                                     ("Address2", data.get("xlAttorneyAdd2")),
                                     ("City", data.get("xlAttorneyCity")),
                                     ("ZipCode", data.get("xlAttorneyZip")),
                                     ("Phone_PNumber", data.get("xlAttorneyPhone"))):
                    if legacy_safe_type(By.ID, "ctl00_Body__" + field, str(value or ""), driver=driver) is not True:
                        raise RuntimeError()
                phase = "state selection"
                Select(_element(driver, By.ID, "ctl00_Body__State_SingleSelectList")).select_by_value(data.get("xlAttorneyState", ""))
                phase = "create type"
                element_click(_element(driver, By.ID, "ctl00_Body__AttorneyType_List_2"), driver=driver)
                phase = "create save"
                element_click(_element(driver, By.XPATH, "/html/body/form/div[3]/div[4]/input[1]"), driver=driver)
                phase = "create save readiness"
                settle()
            visit()
        phase = "final readback"
        verify()
        return _contract().StepResult("attorneys", "done")
    except (Exception, SystemExit) as error:
        failure = error if isinstance(error, _AttorneySetupError) else _AttorneySetupError(phase, error)
    resolved = bool(first and last)
    return _required_tab_recovery(
        "attorneys", settings, failure, _AttorneySetupError, lambda: visit(True),
        verify if resolved else lambda: _element(driver, *control).find_elements(By.XPATH, ".//tbody/tr"),
        "Verify attorney identity and tab entry, or identity/N/A when unresolved, using authorized "
        "Attorney Entry Rules. Correct manually before confirming; Confirm does not add or save.",
        "manual correction verified by unique tab row" if resolved else
        "operator confirmed identity/N/A; Attorneys table accessible; identity not verified")


def step_billing(driver, request, case, data, settings):
    """CMS.py:1688-1724: displayed referral policy, explicit save and tab code readback."""
    phase = "initialization"
    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import Select
        from fcm_intake.cms.session import element_click, legacy_safe_type, wait_page_ready
        phase = "owned context"
        main = driver.current_window_handle
        tab_path = "/html/body/form/div[3]/div[1]/table/tbody/tr[1]/td[2]/table[9]/tbody/tr/td[2]/a"
        control = (By.XPATH, "/html/body/form/div[3]/div[6]/div[2]/div/table/tbody/tr[2]/td[2]/input[3]")
        case_tab = "/html/body/form/div[3]/div[1]/table/tbody/tr[1]/td[2]/table[3]/tbody/tr/td[2]/a"
        referral_control = (By.ID, "ctl00_Body_ReferralTypeField")
        def visit(manual=False, *, context="initial", inspect_failure=True):
            if context not in ("initial", "referral return", "post-save",
                               "final referral return", "manual verification",
                               "manual referral return"):
                raise ValueError()
            target_observation = "unavailable (not observed)"
            def observe_target(tab):
                nonlocal target_observation
                target_observation = _billing_case_observation(driver, "target", target=tab)
            def billing_error(phase, error):
                failure = _BillingSetupError(phase, error)
                if phase not in ("tab verification", "manual return tab verification"):
                    return failure
                observation = (f"Billing visit={context}; Billing target: {target_observation}; "
                               "required code control=unverified")
                if not inspect_failure:
                    failure.diagnostic += f"; {observation}; inspection=suppressed (prior failure)"
                    failure.args = (failure.diagnostic,)
                    return failure
                # Do not query or restore the failed document before the blocking prompt.
                inspection = _billing_inspection(settings, f"{failure}; {observation}", "Billing")
                return _BillingInspectionStop(phase, error, f"{observation}; inspection={inspection}")
            return _required_case_tab(driver, main, tab_path, control,
                                      billing_error, manual=manual, settle=5,
                                      observe_target=observe_target)
        def expected(manual=False, *, context="referral return"):
            target_observation = "unavailable (not observed)"
            pending = None
            def observe_target(tab):
                nonlocal target_observation
                target_observation = _billing_case_observation(driver, "target", target=tab)
            def case_error(phase, error):
                observation = (
                    f"Case target: {target_observation}; failed context: " +
                    _billing_case_observation(driver, "failure", referral_id=referral_control[1]))
                if phase == "tab verification":
                    diagnostic = str(_BillingSetupError("case " + phase, error)) + "; " + observation
                    inspection = _billing_case_inspection(settings, diagnostic)
                    return _BillingCaseInspectionStop(error, f"{observation}; inspection={inspection}")
                failure = _BillingSetupError("case " + phase, error)
                failure.args = (f"{failure}; {observation}",)
                return failure
            try:
                field = _required_case_tab(
                    driver, main, case_tab, referral_control,
                    case_error, settle=10, observe_target=observe_target,
                    resolve_target=_billing_case_tab)
                referral = Select(field).first_selected_option.text
                if not isinstance(referral, str) or not referral.strip():
                    raise ValueError()
            except (Exception, SystemExit) as error:
                pending = error
                raise
            finally:
                # The referral control belongs to Case; code work belongs to Billing.
                try:
                    visit(manual, context=context, inspect_failure=pending is None)
                except (Exception, SystemExit) as error:
                    if pending is None:
                        raise
                    restoration = (error.diagnostic if isinstance(error, _BillingSetupError)
                                   else _billing_case_error_kind(error))
                    if isinstance(pending, _BillingSetupError):
                        pending.args = (f"{pending}; Billing return failed: {restoration}",)
                    else:
                        failure = _BillingSetupError(
                            "manual readback" if manual else "referral readback", pending)
                        failure.args = (f"{failure}; Billing return failed: {restoration}",)
                        raise failure from None
            return (("OPN71", "Case Opening - Task") if referral == "Medical - Task Assignment"
                    else ("OPN01", "Case Opening - Field"))
        def current():
            return (_element(driver, *control).get_attribute("value") or "").strip()
        def verify(manual=False):
            code, _ = expected(manual, context=(
                "manual referral return" if manual else "final referral return"))
            if current() != code:
                raise ValueError()
        def settle(seconds):
            time.sleep(seconds)
            if not wait_page_ready(driver, timeout=20):
                raise RuntimeError()
        visit()
        phase = "referral readback"
        code, note = expected()
    except _BillingSetupError:
        raise
    except (Exception, SystemExit) as error:
        raise _BillingSetupError(phase, error) from None
    try:
        phase = "activity readback"
        actual = current()
        if actual != code:
            if actual:
                raise ValueError()
            phase = "activity opening"
            element_click(_element(driver, By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[2]/td[2]/table[1]/tbody/tr/td[2]/a"), driver=driver)
            phase = "activity readiness"
            settle(3)
            phase = "activity input"
            for field, value in (("Code", code), ("Note", note), ("Description", "Case Opening")):
                if legacy_safe_type(By.ID, "ctl00_Body__Activity__" + field, value, driver=driver) is not True:
                    raise RuntimeError()
            phase = "note completion"
            element_click(_element(driver, By.ID, "ctl00_Body__Activity__NoteCompleted"), driver=driver)
            phase = "activity save"
            element_click(_element(driver, By.ID, "ctl00_Body__Activity_Save"), driver=driver)
            phase = "activity save readiness"
            settle(5)
            visit(context="post-save")
        phase = "final readback"
        verify()
        return _contract().StepResult("billing", "done")
    except _BillingInspectionStop:
        raise
    except (Exception, SystemExit) as error:
        failure = error if isinstance(error, _BillingSetupError) else _BillingSetupError(phase, error)
    return _required_tab_recovery(
        "billing", settings, failure, _BillingSetupError,
        lambda: visit(True, context="manual verification"), lambda: verify(True),
        "Verify the case-opening billing activity, including code, note, description and persistence. "
        "Correct manually before confirming; Confirm does not fill or save.",
        "displayed referral/activity code verified; operator confirmed note, description and persistence")


def step_attachments(driver, request, case, data, settings):
    """CMS.py:1738-1772: select Unity sender and today's attachment only."""
    from selenium.webdriver.common.by import By
    _click(driver, By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[2]/td[2]/table[2]/tbody/tr/td[2]/a")
    table = _element(driver, By.ID, "ctl00_Body__AttachmentTable_Section_Table_Table")
    today = datetime.now().date()
    matches = []
    for row in table.find_elements(By.XPATH, ".//tbody/tr"):
        cells = row.find_elements(By.TAG_NAME, "td")
        if len(cells) < 7 or "UnityApiUser" not in cells[1].text:
            continue
        try:
            sent = datetime.strptime(cells[0].text.strip(), "%m/%d/%Y").date()
        except ValueError:
            continue
        if sent == today:
            matches.append(row)
    if len(matches) != 1 or not request.referral_number:
        return _manual(settings, "attachments", "Verify today's Unity attachment and append the referral number manually; confirm only after saving.")
    # One matching sender/date is required; first-row fallback is unsafe.
    matches[0].find_elements(By.TAG_NAME, "td")[6].find_element(By.TAG_NAME, "a").click()
    subject = _element(driver, By.ID, "ctl00_Body__Subject").get_attribute("value") or ""
    expected = _rules().attachment_subject(subject, request.referral_number)
    _type(driver, By.ID, "ctl00_Body__Subject", expected)
    _click(driver, By.ID, "ctl00_Body__SaveButton")
    return _contract().StepResult("attachments", "done")


def step_notes(driver, request, case, data, settings):
    """CMS.py:1738-1782: missing Other Notes verification (FCM-1.1-102)."""
    return _manual(settings, "notes", "Verify Other Notes transferred and correct any missing notes manually before continuing.")


def step_letters(driver, request, case, data, settings):
    """CMS.py:1776-1841: configure existing ERFS; additional ERFS stays manual."""
    from selenium.webdriver.common.by import By
    try:
        _click(driver, By.XPATH, "/html/body/form/div[3]/div[1]/table/tbody/tr[2]/td[2]/table[7]/tbody/tr/td[2]/a")
        driver.switch_to.default_content()
        frames = driver.find_elements(By.TAG_NAME, "iframe")
        matches = []
        for frame in frames:
            driver.switch_to.default_content()
            driver.switch_to.frame(frame)
            if driver.find_elements(By.XPATH, "//span[normalize-space()='CPD']"):
                matches.append(frame)
        if len(matches) != 1:
            raise ValueError("CPD frame ambiguous")
        # Legacy took the first iframe containing CPD; fail closed when more
        # than one frame could receive the instructions.
        driver.switch_to.default_content()
        driver.switch_to.frame(matches[0])
        _click(driver, By.XPATH, "//span[normalize-space()='CPD']/preceding::a[contains(@class,'exp')][1]")
        time.sleep(.5)
        _click(driver, By.XPATH, "//span[normalize-space()='CPD']/ancestor::div[@id][1]/following-sibling::div[1]//a[normalize-space(.)='CPD ERFS']")
        value = data.get("Special_Instructions", "")
        if not value:
            raise ValueError("ERFS instructions unresolved")
        _type(driver, By.ID, "instructions", value)
        _click(driver, By.XPATH, "//input[@type='submit' and @value='Configure']")
        # Portal availability of the configured instructions after submit is
        # the only exposed readback here; durable document status needs UAT.
        if (_element(driver, By.ID, "instructions").get_attribute("value") or "") != value:
            raise ValueError("ERFS readback mismatch")
    except (Exception, SystemExit):
        # If Configure navigates away or the instructions cannot be read,
        # require an attended verification rather than reporting done.
        return _manual(settings, "letters", "Verify existing CPD ERFS configuration and its saved instructions.")
    return _contract().StepResult("letters", "done")


def setup_cms(driver, request, case, data, settings, on_step):
    """CMS.py:1141-1843: ordered, explicit-driver per-tab work."""
    steps = []
    def select_confirmed_case(*_unused):
        _select_case(driver, request, case)
        return _contract().StepResult("select_case", "done")

    for name, fn in (("select_case", select_confirmed_case),
                     ("claim", step_claim), ("adjuster_address", lambda d, r, c, data, s: _office_address(d, data, s)),
                     ("claimant", step_claimant),
                     ("case", step_case), ("case_supervisor", lambda d, r, c, data, s: _supervisor(d, c, s)),
                     ("customer", step_customer),
                     ("providers", step_providers), ("attorneys", step_attorneys),
                     ("billing", step_billing), ("attachments", step_attachments),
                     ("notes", step_notes), ("letters", step_letters)):
        result = (_safe_step(name, fn) if name == "select_case" else
                  _safe_step(name, lambda fn=fn: fn(driver, request, case, data, settings)))
        steps.append(result)
        on_step(result)
        if result.status == "failed":
            break
    # The remaining business gaps are explicit separate manual gates; a
    # successful legacy-parity write never bypasses their owner decisions.
    for name, message in (
        ("first_appointment_claim_identifier", "Verify Date of First Appointment and Claim Identifier against the referral (FCM-1.1-101)."),
        ("bill_to_policy", "Verify employer-specific Bill To rules and C02283 fallback (FCM-1.1-103)."),
        ("additional_cpd_erfs", "Create and verify an additional CPD ERFS using approved instructions (FCM-1.1-104)."),
    ):
        if steps[-1].status == "failed":
            break
        step = _safe_step(name, lambda m=message, n=name: _manual(settings, n, m))
        steps.append(step)
        on_step(step)
    return steps


def complete_registration(request, case, cem, settings, desktop=None):
    """CMS.py:1843-1897: confirmed case ID and readback (FCM-1.1-99).

    RRS input is irreversible; unverified work-row identity must stop before
    btnReg. The operator confirms the selected claim/referral row, then the
    adapter reads claimant and case prompt. For CEM 10, RRS prompt availability
    requires live validation; absence returns false, never manufactured success.
    """
    if desktop is None:
        from pywinauto import Desktop
        desktop = Desktop(backend="uia")
    main = desktop.window(title_re=r"Referral Routing System*", control_type="Window")
    main.set_focus()
    grid = main.child_window(auto_id="dgWorkWindow", control_type="Table")
    # The legacy row-zero assumption is unsafe: human confirmation binds current
    # selected work row to claim/name/referral before the irreversible click.
    if not case.case_id.strip():
        raise ValueError("Confirmed CMS case ID required")
    # The current workbook lookup is a routing hint, not authorization.
    # For CEM 10, require an independent human confirmation.
    if cem:
        try:
            cem_approved = settings["confirm"]("RRS CEM Route",
                "Confirm a real CEM employer request exists and option 10 is authorized for this referral.")
        except (Exception, SystemExit):
            raise RuntimeError("CEM confirmation failed") from None
        if not cem_approved:
            raise ValueError("CEM employer request not confirmed")
    try:
        approved = settings["confirm"]("RRS Registration",
            "Verify the selected RRS work row matches the claim, claimant, referral type and confirmed CMS case. Continue only after checking the row.")
    except (Exception, SystemExit):
        raise RuntimeError("RRS confirmation failed") from None
    if not approved:
        raise ValueError("RRS work row not confirmed")
    claimant = grid.child_window(title="Claimant Row 0", control_type="Edit").wrapper_object().iface_value.CurrentValue
    expected = f"{request.claimant_first_name} {request.claimant_last_name}"
    if " ".join(claimant.split()).casefold() != " ".join(expected.split()).casefold():
        raise ValueError("RRS claimant does not match")
    grid.child_window(title="Select Row 0", control_type="CheckBox").wrapper_object().click_input()
    main.child_window(auto_id="btnReg", control_type="Button").wrapper_object().click_input()
    menu = main.child_window(auto_id="frmSmallInputBox", control_type="Window")
    menu.wait("exists", timeout=5)
    option = _rules().complete_reg_option(cem)
    option_input = menu.child_window(auto_id="tbInput", control_type="Edit").wrapper_object()
    option_input.iface_value.SetValue(option)
    if option_input.iface_value.CurrentValue != option:
        raise ValueError("RRS registration option readback mismatch")
    menu.child_window(auto_id="btnOk", control_type="Button").wrapper_object().click_input()
    if desktop is not None and hasattr(desktop, "case_prompt"):
        prompt = desktop.case_prompt
    else:
        from pywinauto.application import Application
        prompt = Application(backend="win32").connect(process=main.element_info.process_id).window(title_re=r"Case Id")
    # For CEM 10 this prompt was skipped in legacy; source policy requires
    # readback, so absence fails rather than reporting a fabricated success.
    prompt.wait("ready", timeout=10)
    prompt.Edit.set_edit_text(case.case_id)
    readback = prompt.Edit.window_text()
    entered = readback.strip() == case.case_id
    if entered:
        prompt.OK.click()
        prompt.wait_not("exists", timeout=10)
    return {"option": option, "case_id_entered": entered, "readback": readback}


def _daily_cem_request(legacy, request):
    """Legacy exact claimant+claim daily check; unknown is not False."""
    # Caller supplies the legacy module lazily from the live session; neither
    # the workbook nor legacy imports are touched on cms_steps import.
    try:
        path = Path(legacy.get_excel_path())
        if not path.is_file():
            return None
        # Probe sheet availability separately: legacy claim_exists returns
        # False for both a genuine negative and a missing workbook/sheet.
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            # Missing/unreadable workbook or sheet cannot distinguish a genuine
            # negative from an absent check, so do not route to option 3.
            if legacy.get_today_sheet_name() not in wb.sheetnames:
                return None
        finally:
            wb.close()
        value = legacy.claim_exists(
            f"{request.claimant_first_name} {request.claimant_last_name}", request.claim_number)
        if not isinstance(value, bool):
            raise ValueError("Daily CEM result invalid")
        return value
    except (Exception, SystemExit):
        # The controller then stops with cem_status_unknown unless an explicit
        # request boolean was provided; do not silently choose option 3.
        return None


def _cem_sheet_available(log, location, day):
    """True only when the workbook opens and has the dated sheet.

    Probed under the CEM module's own sidecar lock (the one find_duplicate and
    record take) so this read never holds the file open across the CEM app's
    atomic replace. A held lock times out (TimeoutError) and the caller treats
    that as unknown.
    """
    with log._lock(location):
        if not location.is_file():
            return False
        from openpyxl import load_workbook
        wb = load_workbook(location, read_only=True, data_only=True)
        try:
            return day in wb.sheetnames
        finally:
            wb.close()


def _cem_log_request(request, path):
    """Configured new-CEM-log check: True match, False readable no-match, None unknown.

    Uses cem_log.find_duplicate (today's sheet, column A claimant, column B claim
    NUMBER, strip+casefold, Status ignored). It returns None both for "no row"
    and "no workbook/sheet", so availability is probed before and after the
    lookup; only a sheet present on both sides with no row is a genuine False.
    Any failure (missing, unreadable, locked, invalid path/value, module
    unavailable) is unknown, never False.
    """
    try:
        log = _cem_log()
        moment = _cem_clock()
        if not isinstance(moment, datetime):
            raise TypeError("CEM clock must return datetime")
        day = moment.strftime("%m%d%Y")
        location = log._location(path)  # same absolute-path/filename rule as the writer
        if not _cem_sheet_available(log, location, day):
            return None
        # find_duplicate builds the claimant exactly as the CEM writer
        # (cem_log.record) stores column A, f"{first} {last}", which is also the
        # string load_data passes to legacy claim_exists; both sides then
        # strip+casefold. RegistrationRequest has the same three attribute names.
        found = log.find_duplicate(request, SimpleNamespace(cem_log_path=path), now=lambda: moment)
        if found is None:
            return False if _cem_sheet_available(log, location, day) else None
        if isinstance(found, dict) and found.get("sheet") == day and found.get("row"):
            return True
        raise ValueError("CEM log result invalid")
    except (Exception, SystemExit):
        return None


class LiveSession:
    """One CMS IE-mode session per operator-started run; idempotent cleanup."""
    def __init__(self, driver, stages):
        self._driver = driver
        self.stages = stages
        self._closed = False

    def close(self):
        if self._closed:
            return
        self._closed = True
        if self._driver is not None:
            try:
                self._driver.quit()
            except (Exception, SystemExit):
                raise RuntimeError("CMS browser cleanup failed") from None


def open_live_session(settings: dict) -> LiveSession:
    """Open a live session only on explicit --live / desktop Start invocation.

    Required for the default data loader: cms_db_connection (historical adapter key for the explicitly
    configured RRS stored-procedure connection), reference_dir, cms_user,
    cms_password, notify, confirm. The desktop supplies FCM_RRS_DB_CONNECTION
    for this lookup only; no CMS connection fallback is permitted.
    Optional: ie_driver_path, office_matrix_filename,
    supervisor_mapping_filename, preview_only (defaults to True; blocks writes),
    cem_log_path (else [cem] cem_log_path / FCM_CEM_LOG_PATH; blank selects
    the legacy CMS.claim_exists daily checker for the default DATA loader).
    An optional callable load_data(request, case) replaces the entire default
    DB/CEM loader; no database, CEM log or daily checker is read in that mode.
    """
    C = _contract()
    supplied_loader = settings.get("load_data")
    if "load_data" in settings and not callable(supplied_loader):
        raise ValueError("load_data must be callable")
    required = ("reference_dir", "cms_user", "cms_password", "notify", "confirm")
    if supplied_loader is None:
        required += ("cms_db_connection",)
    if any(not settings.get(key) for key in required):
        raise ValueError("Required CMS session setting missing")
    if not callable(settings["notify"]) or not callable(settings["confirm"]):
        raise ValueError("CMS operator callbacks must be callable")
    preview_only = settings.get("preview_only", True)
    if type(preview_only) is not bool:
        raise ValueError("preview_only must be boolean")
    # CEM status is obtained lazily at DATA from the configured CEM app log or,
    # when unconfigured, the legacy daily checker; the independent option-10
    # operator confirmation remains mandatory either way.
    # Reuse the legacy stored-procedure mapper and daily checker only. Its
    # browser helpers rely on module-global `driver`; keep explicit-driver actions.
    driver = None
    try:
        legacy = _legacy()
        # [cem] cem_log_path / FCM_CEM_LOG_PATH (FCM-1.1-126): when configured,
        # the new CEM app's log is the only CEM hint source; legacy is unused.
        cem_log_path = _configured_cem_log_path(settings)
        from fcm_intake.cms.security_warning import (
            CmsCertificateMonitor, new_browser_window, snapshot_browser_windows,
        )
        from fcm_intake.cms.session import authenticate_cms
        from selenium.common.exceptions import TimeoutException
        from selenium.webdriver.common.by import By
        previous_windows = snapshot_browser_windows()
        if settings.get("ie_driver_path"):
            # Legacy creator uses its module constant; construct a local driver
            # instead of mutating shared state.
            from selenium import webdriver
            from selenium.webdriver.ie.options import Options
            from selenium.webdriver.ie.service import Service
            options = Options()
            options.add_additional_option("ie.edgechromium", True)
            options.add_additional_option("ie.edgepath", legacy.EDGE_PATH)
            options.add_additional_option("ignoreProtectedModeSettings", True)
            options.add_additional_option("requireWindowFocus", True)
            options.add_additional_option("nativeEvents", False)
            options.ensure_clean_session = True
            driver = webdriver.Ie(service=Service(executable_path=settings["ie_driver_path"]), options=options)
        else:
            driver = legacy.create_ie_driver()
        driver._cms_warning_window = new_browser_window(previous_windows)
        driver._cms_browser_handles_before = previous_windows
        driver.set_page_load_timeout(45)
        driver.maximize_window()
        # Fixed helper status labels must not block on operator acknowledgement.
        with CmsCertificateMonitor(driver, legacy.CMS_LOGIN_URL,
                                   logging.getLogger(__name__).info) as certificate:
            try:
                driver.get(legacy.CMS_LOGIN_URL)
            except TimeoutException:
                if not certificate.proceeded:
                    raise
            # Proceed only means the warning was clicked, not that login loaded.
            _element(driver, By.ID, "ctl00_Body_UserName")
        authenticate_cms(driver, settings["cms_user"], settings["cms_password"],
                         status=logging.getLogger(__name__).info)
        if not settings["confirm"]("CMS Login", "Confirm that CMS sign-in succeeded before searching cases."):
            raise ValueError("CMS sign-in not confirmed")
        # Operator verification remains required after automated sign-in checks.
        # Login credentials stay in the caller's in-memory settings; no result or
        # exception text contains them. Never send them to a diagnostic logger.

        def load_data(request, case):
            if supplied_loader is not None:
                fields = supplied_loader(request, case)
                if not isinstance(fields, dict):
                    raise TypeError("load_data must return a dict")
                return fields
            try:
                raw = legacy.fetch_case_data_from_db(settings["cms_db_connection"],
                    f"{request.claimant_first_name} {request.claimant_last_name}", request.claim_number)
            except (Exception, SystemExit) as error:
                if (isinstance(error, ValueError) and error.args
                        and isinstance(error.args[0], str)
                        and error.args[0].startswith("No records found for claimant=")):
                    raise ValueError(
                        "No matching registration referral was returned by the RRS lookup."
                    ) from None
                # ODBC's first argument may be a SQLSTATE; never relay its
                # accompanying driver text, parameters or connection string.
                state = error.args[0] if error.args else None
                detail = (f"SQLSTATE {state}" if isinstance(state, str)
                          and re.fullmatch(r"[A-Z0-9]{5}", state)
                          and state[:2] in {"01", "02", "07", "08", "0A", "21", "22",
                                           "23", "24", "25", "28", "40", "42", "HY", "IM"}
                          else type(error).__name__)
                raise ValueError(
                    f"RRS stored-procedure lookup failed during database execution/mapping ({detail})"
                ) from None
            try:
                fields: dict[str, str | bool] = {key: str(value).strip() if value is not None else ""
                                               for key, value in asdict(raw).items()}
            except (Exception, SystemExit):
                raise ValueError("RRS stored-procedure lookup result mapping failed") from None
            # Configured CEM app log: its reader owns the exact normalized
            # claimant/claim matching. Unconfigured: the legacy date-sheet
            # checker, only when its workbook and sheet can be read. Either way
            # unknown stays unknown; never manufacture a True from a path alone.
            if cem_log_path:
                cem = _cem_log_request(request, cem_log_path)
            else:
                cem = _daily_cem_request(legacy, request)
            if cem is not None:
                fields["cem_employer_request"] = cem
            # Return only a boolean routing field, never the workbook path or
            # the underlying checker record.
            return fields

        def setup(request, case, data, on_step):
            if preview_only:
                raise RuntimeError("Preview only: CMS setup and RRS registration are disabled")
            return setup_cms(driver, request, case, data, settings, on_step)

        def search(request):
            try:
                candidates = find_cases(driver, request, settings)
                if preview_only:
                    case, reason = _rules().select_case(candidates, request)
                    if case is not None and not reason:
                        _select_case(driver, request, case)
                return candidates
            except (_CaseSearchError, _CaseSelectionError):
                raise
            except (Exception, SystemExit):
                raise RuntimeError("CMS case search could not be verified") from None

        def finish(request, case, cem):
            if preview_only:
                raise RuntimeError("Preview only: CMS setup and RRS registration are disabled")
            try:
                return complete_registration(request, case, cem, settings)
            except (Exception, SystemExit):
                raise RuntimeError("RRS registration could not be verified") from None

        def review(request, case, data):
            try:
                return settings["confirm"]("CMS Registration Review",
                    "Review registration facts and case identity before setup.")
            except (Exception, SystemExit):
                raise RuntimeError("CMS review confirmation failed") from None

        stages = C.RegistrationStages(
            find_cases=search,
            load_data=load_data,
            review=review,
            setup_cms=setup,
            complete_registration=finish,
        )
        return LiveSession(driver, stages)
    except (Exception, SystemExit):
        if driver is not None:
            try:
                driver.quit()
            except (Exception, SystemExit):
                pass
        raise RuntimeError("CMS session could not be opened or authenticated") from None
