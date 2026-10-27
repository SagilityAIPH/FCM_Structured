"""Synthetic workbook + threaded bridge smoke. No live systems or PHI."""
import json
from pathlib import Path
from types import SimpleNamespace
import tempfile
import time
from openpyxl import load_workbook
from backend import UnityAPI
from unity_model import defaults, schema, verify_legacy_keys
from unity_excel import create_template, read_workbook, COLUMNS, SHEET


def main():
    import sys
    from unity_model import ROOT
    from AI.unity_browser import UnityBrowserSession
    events = []
    driver = SimpleNamespace(quit=lambda: events.append('quit'))
    with UnityBrowserSession() as session:
        assert session.acquire(lambda: driver) is driver
        assert session.acquire(lambda: None) is driver
        assert not events
    assert events == ['quit']
    if getattr(sys, 'frozen', False):
        required = ['frontend/dist/index.html',
            'Processes/Unity/Stand-Alone/run.py',
            'Processes/Unity/Deploy-Ready/unity_flow.py',
            'src/fcm_intake/legacy/legacy_unity.py',
            'src/fcm_intake/legacy/legacy_fcm.py',
            'src/fcm_intake/legacy/CMS.py',
            'Processes/Reopen-Check/Deploy-Ready/reopen_flow.py',
            'Final Code/finalcode/field_adapter.py',
            'Processes/SubjectLineBuilder/Deploy-Ready/office_matrix.py',
            'Liberty Mutual Claims Office Matrix_08-29-19.xlsx',
            'drivers/IEDriverServer.exe']
        for relative in required:
            assert (ROOT / relative).is_file(), relative
        import titlecase
        import pyodbc
        import edge_auto
        from AI.unity_browser import create_driver, ChromeOptions, ChromeService
        from selenium.webdriver import Chrome
        assert all(callable(item) for item in (create_driver, ChromeOptions, ChromeService, Chrome))
        from AI.rrs_assignment import complete_assignment, RRSCompletionError
        assert callable(complete_assignment) and issubclass(RRSCompletionError, RuntimeError)
        from AI.unity_completion import complete_task
        from AI.unity_errors import UnityTaskCompletionError
        assert callable(complete_task) and issubclass(UnityTaskCompletionError, RuntimeError)
        from AI.unity_ncm import assign_ncm
        from AI.unity_errors import UnityNCMAssignmentError
        assert callable(assign_ncm) and issubclass(UnityNCMAssignmentError, RuntimeError)
        import dotenv
        from fcm_intake.config import UNITY_SCRIPT, FCM_SCRIPT, get_rrs_db_connection
        assert UNITY_SCRIPT.is_file() and FCM_SCRIPT.is_file()
        assert callable(get_rrs_db_connection) and callable(dotenv.load_dotenv)
        from fcm_intake.legacy_loader import _ensure_legacy_compat
        _ensure_legacy_compat(FCM_SCRIPT)
        import importlib.util
        assert importlib.util.find_spec("legacy.CMS") is not None
        # Exercise the shared adapter's original lazy file loader in isolation;
        # do not import legacy_fcm, which binds live systems at import time.
        spec = importlib.util.spec_from_file_location('_packaged_adapter_probe', ROOT / 'Final Code/finalcode/field_adapter.py')
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        assert module._load_office_matrix().MATRIX_NAME
        print('Frozen resource layout and live import prerequisites: OK')
    fields = {field['key']: field for group in schema() for field in group['fields']}
    assert fields['specialInstructions']['maxLength'] == 32767
    assert all(field['maxLength'] == 1000 for key, field in fields.items() if key != 'specialInstructions')
    prefix, suffix = ' \n\tSynthetic instructions\n', '\nEND-SYNTHETIC \t\n'
    instructions = prefix + 'x' * (32767 - len(prefix) - len(suffix)) + suffix
    assert len(instructions) == 32767
    data = defaults() | {'claimNumber': 'SYNTHETIC-ROUNDTRIP-001', 'customer': 'Synthetic Continental Tire',
                        'claimantFirst': 'Invented', 'claimantLast': 'Example', 'claimType': 'Auto', 'zip': '00123',
                        'specialInstructions': instructions, 'ncm': 'Synthetic Nurse'}
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / 'synthetic.xlsx'
        create_template(path)
        book = load_workbook(path)
        for i, (key, _, _) in enumerate(COLUMNS, 1):
            book[SHEET].cell(2, i, data[key])
        book.save(path)
        book.close()
        row = read_workbook(path)[0]
        assert row['data'] == data
        factory_calls, payloads, ncm_events = [], [], []
        def callbacks():
            factory_calls.append(True)
            return {'list_attachments': lambda: ['synthetic-referral.pdf'],
                'run_unity': lambda customer, claimant, payload: payloads.append(payload),
                'record_assignment': lambda *args: None,
                'complete_assignment': lambda *args: ncm_events.append('rrs'),
                'assign_ncm': lambda *args: ncm_events.append('ncm')}
        api = UnityAPI(callbacks)
        oversized = data | {'specialInstructions': instructions + 'x'}
        for rejected in (api.validate_fields(oversized, 'all', 'MI'), api.start(oversized, 'all', True, 'MI')):
            assert not rejected['ok']
            assert 'specialInstructions. Use text up to 32,767 characters' in rejected['error']
        assert api._thread is None
        assert api.get_state()['status'] == 'idle'
        assert not factory_calls and not payloads
        assert api.start(row['data'], 'all', True, 'MI')['ok']
        deadline = time.monotonic() + 5
        while api.get_state()['busy']:
            state = api.get_state()
            if state['prompt']:
                assert api.answer_prompt(state['prompt']['id'], True)['ok']
            if time.monotonic() > deadline:
                raise RuntimeError('Synthetic smoke worker timeout.')
            time.sleep(.01)
        result = api.get_state()['result']
        assert result['status'] == 'completed'
        assert len(payloads) == 1
        assert payloads[0]['specialInstructions'] == instructions
        local_results = {}
        for extension in ('PDF', 'mSg', 'eMl'):
            selected = Path(directory) / f'synthetic-selected.{extension}'
            selected.touch()
            api._window = SimpleNamespace(create_file_dialog=lambda *args, **kwargs: [str(selected)])
            api._dialog = lambda kind: kind
            assert api.choose_referral_document()['ok']
            assert api.start(row['data'], 'unity', True, 'MI')['ok']
            deadline = time.monotonic() + 5
            while api.get_state()['busy']:
                if time.monotonic() > deadline:
                    raise RuntimeError('Synthetic local document worker timeout.')
                time.sleep(.01)
            local_result = api.get_state()['result']
            assert local_result['status'] == 'completed'
            assert local_result['attachment'] == str(selected.resolve())
            local_results[extension.lower()] = local_result
            assert api.use_rrs_export_folder()['ok']
        assert len(payloads) == 4
        ncm_events.clear()
        assert api.start(row['data'], 'assign', True, '', False, True)['ok']
        deadline = time.monotonic() + 5
        while api.get_state()['busy']:
            state = api.get_state()
            if state['prompt']:
                assert not ncm_events
                assert api.answer_prompt(state['prompt']['id'], True)['ok']
            if time.monotonic() > deadline:
                raise RuntimeError('Synthetic NCM worker timeout.')
            time.sleep(.01)
        ncm_result = api.get_state()['result']
        assert ncm_result['status'] == 'completed'
        assert ncm_result['steps'] == ['await_icasemanager_approval', 'assign_ncm', 'complete_assignment']
        assert ncm_events == ['ncm', 'rrs']
        audit = verify_legacy_keys()
        print(json.dumps({'mode': 'synthetic workbook roundtrip + bridge', 'operator_columns': len(COLUMNS),
            'single_all_session_cleanup': events == ['quit'],
            'rows': 1, 'zip_preserved': row['data']['zip'], 'ast_consumed_keys': audit['count'],
            'context_only': audit['context_only'], 'result': result,
            'instruction_length': len(payloads[0]['specialInstructions']), 'oversized_rejected': True,
            'ncm_column_preserved': row['data']['ncm'] == 'Synthetic Nurse',
            'postapproval_ncm_result': ncm_result,
            'local_pdf_result': local_results['pdf'], 'local_msg_result': local_results['msg'],
            'local_eml_result': local_results['eml']}, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
