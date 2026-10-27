"""Standalone workbook contract: one referral and one provider per row."""
from datetime import date, datetime, time
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from unity_model import GROUPS, CHOICES, LENGTH_LIMITS, defaults, validate

SHEET = 'Unity Input'
PREFIXES = ['Claim', 'Claimant', 'Provider', 'Adjuster', 'Attorney', 'NCM']
COLUMNS = [(field[0], f'{prefix} - {field[1]}', len(field) > 2 and field[2])
           for prefix, (_, fields) in zip(PREFIXES, GROUPS) for field in fields]


def create_template(path):
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append([label for _, label, _ in COLUMNS])
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}501'
    sheet.row_dimensions[1].height = 44
    for index, (key, label, required) in enumerate(COLUMNS, 1):
        header = sheet.cell(1, index)
        header.font = Font(color='FFFFFF', bold=True)
        header.fill = PatternFill('solid', fgColor='2563EB' if required else '475569')
        header.alignment = Alignment(wrap_text=True, vertical='center')
        header.comment = Comment('Required for unity/all; claimNumber required for every stage.' if required else 'Optional. Leave unknown values blank.', 'FCM')
        if key == 'specialInstructions':
            header.comment = Comment(
                f'Optional. Up to {LENGTH_LIMITS[key]:,} characters (Excel-compatible). '
                'Line breaks and leading/trailing whitespace are preserved when loaded.', 'FCM')
        if key == 'ncm':
            header.comment = Comment(
                'Exact full NCM name. Required only for Scenario 2 in All/Assign; '
                'claimant first and last names are also required in that scenario.', 'FCM')
        sheet.column_dimensions[get_column_letter(index)].width = 25
        for row in range(2, 502):
            sheet.cell(row, index).number_format = '@'
        if key in CHOICES:
            rule = DataValidation(type='list', formula1='"' + ','.join(v for v in CHOICES[key] if v) + '"', allow_blank=True)
            rule.errorTitle = 'Choose a listed value'
            rule.error = 'Use the dropdown choices.'
            rule.showErrorMessage = True
            sheet.add_data_validation(rule)
            rule.add(f'{get_column_letter(index)}2:{get_column_letter(index)}501')
    guide = book.create_sheet('Instructions')
    notes = [
        f'Unity desktop workbook — {len(COLUMNS)} operator fields; no sample patient data.',
        'One referral and one provider per row. Start at row 2. Keep sheet and headers.',
        'claimNumber is required for every stage; customer and claimant names for unity/all.',
        'Scenario 2 All/Assign also requires NCM full name and claimant first/last names. '
        'The scenario selection is not stored in Excel. Older 35-column workbooks load NCM blank.',
        'Other values are optional. Incomplete rows load with warnings for editing.',
        f'Special instructions: up to {LENGTH_LIMITS["specialInstructions"]:,} characters; line breaks and surrounding whitespace preserved.',
        f'All other fields: up to {LENGTH_LIMITS["claimNumber"]:,} characters. Longer supplied values must be corrected before running.',
        'Dates MM/DD/YYYY; time 08:30 AM. ZIPs, phones and identifiers are text. No formulas.',
        'Blank serviceType derives Medical Full (Medical Task for Auto / Continental Tire / MI jurisdiction).',
        'Claim jurisdiction is separate optional context entered in the app; never inferred from claimant State.',
        'Template has no derived columns. Review defaults and jurisdiction before live work.',
        'Loading a row never starts the portal. Select a stage and run explicitly.',
    ]
    for note in notes:
        guide.append([note])
        guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
        guide.row_dimensions[guide.max_row].height = 42
    guide.column_dimensions['A'].width = 115
    guide['A1'].font = Font(size=16, bold=True, color='2563EB')
    book.save(path)
    book.close()
    return str(path)


def _text(cell, *, preserve_whitespace=False):
    value = cell.value
    if value is None:
        return ''
    if cell.data_type in ('f', 'e'):
        raise ValueError(f'{cell.coordinate}: replace formulas or Excel errors with plain values.')
    if isinstance(value, (datetime, date)):
        return value.strftime('%m/%d/%Y')
    if isinstance(value, time):
        return value.strftime('%I:%M %p')
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if re.fullmatch(r'0+', cell.number_format) and float(value).is_integer():
            return str(int(value)).zfill(len(cell.number_format))
        if float(value).is_integer():
            return str(int(value))
    return str(value) if preserve_whitespace else str(value).strip()


def read_workbook(path):
    if Path(path).suffix.lower() != '.xlsx':
        raise ValueError('Choose an .xlsx workbook made from the Unity template.')
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if SHEET not in book.sheetnames:
            raise ValueError(f'Missing "{SHEET}" sheet. Use the Unity template.')
        sheet = book[SHEET]
        if sheet.max_row > 10001 or sheet.max_column > 200:
            raise ValueError('Use a workbook with up to 10,000 input rows and 200 columns.')
        rows = sheet.iter_rows()
        headers = [str(c.value or '').strip() for c in next(rows, [])]
        known = {label.casefold(): key for key, label, _ in COLUMNS}
        mapping = {}
        for index, label in enumerate(headers):
            key = known.get(label.casefold())
            if key:
                if key in mapping.values():
                    raise ValueError(f'Duplicate column: {label}')
                mapping[index] = key
        if not mapping:
            raise ValueError('No recognized columns. Use the Unity template.')
        records = []
        for number, row in enumerate(rows, 2):
            supplied = {key: _text(row[index], preserve_whitespace=key == 'specialInstructions')
                        for index, key in mapping.items()}
            if not any(supplied.values()):
                continue
            data = defaults()
            data.update({key: value for key, value in supplied.items() if value})
            try:
                validate(data)
                warning = 'Optional values and jurisdiction require review.'
            except ValueError as error:
                warning = str(error)
            name = ' '.join(filter(None, [data['claimantFirst'], data['claimantLast']])) or 'Unnamed claimant'
            records.append({'row': number, 'label': f"Row {number} | {name} | {data['claimNumber'] or 'No claim number'}",
                            'data': data, 'warning': warning})
        if not records:
            raise ValueError('No filled rows found. Enter referral details starting on row 2, save, then reopen the workbook.')
        return records
    finally:
        book.close()
