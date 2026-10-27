"""Synthetic desktop regression tests; unique imports coexist with sibling apps."""
import ast
import importlib.util
import json
from copy import deepcopy
import os
import re
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import MagicMock, Mock, patch
from types import SimpleNamespace

from openpyxl import load_workbook
from AI.unity_errors import UnityCustomerFieldsError, UnityEmployerFieldsError, UnityFlagError, UnityTaskCompletionError
from AI.unity_completion import complete_task
from AI.rrs_assignment import RRSCompletionError

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import unity_model as model
import unity_excel as excel
spec = importlib.util.spec_from_file_location('unity_desktop_backend_tests', HERE / 'backend.py')
backend = importlib.util.module_from_spec(spec)
spec.loader.exec_module(backend)


def synthetic():
    return model.defaults() | {'claimNumber': 'SYNTHETIC-001', 'customer': 'Invented Widgets',
                              'claimantFirst': 'Synthetic', 'claimantLast': 'Person'}


def synthetic_instructions(length):
    prefix, suffix = ' \n\tSynthetic instructions\n', '\nEND-SYNTHETIC \t\n'
    return prefix + 'x' * (length - len(prefix) - len(suffix)) + suffix


def legacy_function(name, filename='legacy_unity.py'):
    """Read one function, stopping before the next definition or sample records."""
    path = model.ROOT / 'src' / 'fcm_intake' / 'legacy' / filename
    lines = []
    with path.open(encoding='utf-8-sig') as source:
        for line in source:
            if not lines:
                if line.startswith(f'def {name}('):
                    lines.append(line)
            elif line.startswith(('def ', 'class ', 'if __name__')):
                break
            else:
                lines.append(line)
    return ast.parse(''.join(lines)).body[0]


class StartupTests(unittest.TestCase):
    def test_direct_entrypoints_self_test_without_pythonpath(self):
        env = os.environ.copy()
        env.pop('PYTHONPATH', None)
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            env.update(TEMP=directory, TMP=directory, TMPDIR=directory)
            for entrypoint in ('run.py', 'launcher.py'):
                with self.subTest(entrypoint=entrypoint):
                    report = Path(directory) / f'{entrypoint}-report.json'
                    completed = subprocess.run(
                        [sys.executable, str(HERE / entrypoint), '--self-test', '--report', str(report)],
                        cwd=model.ROOT, env=env, capture_output=True, text=True, timeout=30)
                    self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
                    result = json.loads(report.read_text(encoding='utf-8'))
                    self.assertEqual(result['mode'], 'synthetic workbook roundtrip + bridge')
                    self.assertEqual(result['result']['status'], 'completed')
                    self.assertEqual(result['rows'], 1)
                    self.assertEqual(result['instruction_length'], 32767)
                    self.assertIs(result['oversized_rejected'], True)
                    for extension in ('pdf', 'msg', 'eml'):
                        local = result[f'local_{extension}_result']
                        self.assertEqual(local['status'], 'completed')
                        self.assertEqual(Path(local['attachment']).suffix.lower(), '.' + extension)

    def test_model_and_backend_import_without_repository_on_path(self):
        env = os.environ.copy()
        env.pop('PYTHONPATH', None)
        for module in ('unity_model', 'backend'):
            with self.subTest(module=module):
                completed = subprocess.run(
                    [sys.executable, '-c', f'import {module}'],
                    cwd=HERE, env=env, capture_output=True, text=True, timeout=30)
                self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)

    def test_spec_covers_raw_loader_shared_ai_imports(self):
        adapter_tree = ast.parse((model.ROOT / 'Final Code' / 'finalcode' / 'field_adapter.py')
                                 .read_text(encoding='utf-8-sig'))
        required = {node.module for node in ast.walk(adapter_tree)
                    if isinstance(node, ast.ImportFrom) and node.module
                    and node.module.startswith('AI.')}
        required.add('AI.referral_attachments')
        required.add('AI.unity_errors')
        required.add('AI.unity_completion')
        required.add('AI.unity_browser')
        required.add('AI.rrs_assignment')
        required.add('legacy.CMS')
        required.add('fcm_intake.config')
        required.add('dotenv')
        spec_tree = ast.parse((HERE / 'unity_assignment.spec').read_text(encoding='utf-8'))
        hiddenimports = next(ast.literal_eval(node.value) for node in spec_tree.body
                             if isinstance(node, ast.Assign)
                             and any(isinstance(target, ast.Name) and target.id == 'hiddenimports'
                                     for target in node.targets))
        self.assertFalse(required - set(hiddenimports),
                         f'Raw field adapter imports missing from hiddenimports: {required - set(hiddenimports)}')

    def test_case_lookup_runtime_package_and_frozen_resources(self):
        with patch.object(sys, 'path', [str(model.ROOT / 'src'), *sys.path]):
            from fcm_intake.legacy_loader import _ensure_legacy_compat
            with patch.dict(sys.modules):
                sys.modules.pop('legacy', None)
                sys.modules.pop('legacy.CMS', None)
                legacy = model.ROOT / 'src' / 'fcm_intake' / 'legacy'
                _ensure_legacy_compat(legacy / 'legacy_fcm.py')
                # Resolve without executing CMS's Windows/browser imports.
                found = importlib.util.find_spec('legacy.CMS')
                self.assertEqual(Path(found.origin), legacy / 'CMS.py')
        tree = ast.parse((HERE / 'unity_assignment.spec').read_text(encoding='utf-8'))
        raw_files = next(ast.literal_eval(node.iter) for node in tree.body
                         if isinstance(node, ast.For) and isinstance(node.iter, ast.Tuple)
                         and 'legacy_fcm.py' in ast.literal_eval(node.iter))
        self.assertIn('CMS.py', raw_files)
        self.assertTrue((legacy / 'CMS.py').is_file())
        smoke_tree = ast.parse((HERE / 'smoke.py').read_text(encoding='utf-8'))
        required = next(ast.literal_eval(node.value) for node in ast.walk(smoke_tree)
                        if isinstance(node, ast.Assign)
                        and any(isinstance(target, ast.Name) and target.id == 'required'
                                for target in node.targets))
        self.assertIn('src/fcm_intake/legacy/CMS.py', required)
        self.assertTrue(any(isinstance(node, ast.ImportFrom) and node.module == 'AI.rrs_assignment'
                            for node in ast.walk(smoke_tree)))
        self.assertTrue(any(isinstance(node, ast.ImportFrom) and node.module == 'AI.unity_browser'
                            for node in ast.walk(smoke_tree)))
        self.assertTrue(any(isinstance(node, ast.ImportFrom) and node.module == 'AI.unity_completion'
                            for node in ast.walk(smoke_tree)))
        self.assertTrue(any(isinstance(node, ast.ImportFrom)
                            and node.module == 'selenium.webdriver' and
                            any(alias.name == 'Chrome' for alias in node.names)
                            for node in ast.walk(smoke_tree)))
        self.assertFalse(any(isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
                             and node.func.id in {'Chrome', 'ChromeService', 'create_driver', 'complete_task'}
                             for node in ast.walk(smoke_tree)))


class ChromeDriverTests(unittest.TestCase):
    def test_shared_chrome_factory_manager_static_path_and_opt_in_detach(self):
        from AI import unity_browser
        from selenium.webdriver.chrome.options import Options

        for path in ('', r'C:\synthetic\chromedriver.exe'):
            for keywords, detached in (({}, False), ({'detach': True}, True), ({'detach': False}, False)):
                with self.subTest(path=path, keywords=keywords), \
                        patch.object(unity_browser, 'webdriver') as webdriver, \
                        patch.object(unity_browser, 'ChromeService') as service:
                    driver = unity_browser.create_driver(path, **keywords)
                    service.assert_called_once_with(executable_path=path or None, log_output=None)
                    options = webdriver.Chrome.call_args.kwargs['options']
                    self.assertIsInstance(options, Options)
                    self.assertEqual(options.binary_location, '')
                    self.assertIs(options.experimental_options['detach'], detached)
                    self.assertEqual(options.experimental_options['excludeSwitches'], ['enable-logging'])
                    self.assertEqual(set(options.arguments), {
                        '--ignore-certificate-errors', '--allow-insecure-localhost',
                        '--disable-popup-blocking', '--disable-gpu', '--disable-software-rasterizer',
                        '--no-sandbox', '--disable-dev-shm-usage', '--log-level=3',
                        '--disable-features=OnDeviceModel,OptimizationHints,OptimizationGuideModelDownloading'})
                    webdriver.Chrome.assert_called_once_with(service=service.return_value, options=options)
                    webdriver.Edge.assert_not_called()
                    driver.set_page_load_timeout.assert_called_once_with(60)
                    driver.set_script_timeout.assert_called_once_with(60)

    def test_chrome_errors_propagate_without_service_or_browser_fallback(self):
        from AI import unity_browser

        for failure_stage in ('service', 'browser'):
            error = RuntimeError('synthetic Chrome failure')
            with self.subTest(stage=failure_stage), \
                    patch.object(unity_browser, 'webdriver') as webdriver, \
                    patch.object(unity_browser, 'ChromeService') as service:
                if failure_stage == 'service':
                    service.side_effect = error
                else:
                    webdriver.Chrome.side_effect = error
                with self.assertRaises(RuntimeError) as raised:
                    unity_browser.create_driver(r'C:\synthetic\missing.exe', log_output='chromedriver.log')
                self.assertIs(raised.exception, error)
                service.assert_called_once_with(
                    executable_path=r'C:\synthetic\missing.exe', log_output='chromedriver.log')
                self.assertEqual(webdriver.Chrome.call_count, failure_stage == 'browser')
                webdriver.Edge.assert_not_called()

    def test_actual_legacy_factories_delegate_and_propagate_failures(self):
        for filename in ('legacy_unity.py', 'CMS.py'):
            path = model.ROOT / 'src' / 'fcm_intake' / 'legacy' / filename
            imports = []
            with path.open(encoding='utf-8-sig') as source:
                for line in source:
                    if line.startswith('def '):
                        break
                    imports.append(line)
            import_tree = ast.parse(''.join(imports))
            if filename == 'CMS.py':
                import_tree = legacy_function('CheckCaseUnity', filename)
            self.assertTrue(any(isinstance(node, ast.ImportFrom)
                                and node.module == 'AI.unity_browser'
                                and any(alias.name == 'create_driver'
                                        and alias.asname == 'create_unity_driver' for alias in node.names)
                                for node in ast.walk(import_tree)), filename)
        assignment = legacy_function('create_driver')
        lookup = next(node for node in legacy_function('CheckCaseUnity', 'CMS.py').body
                      if isinstance(node, ast.FunctionDef) and node.name == 'create_driver')
        cases = [(assignment, keywords, {'detach': detached})
                 for keywords, detached in (({}, False), ({'detach': True}, True), ({'detach': False}, False))]
        cases.append((lookup, {}, {'log_output': 'chromedriver.log'}))
        for factory, keywords, forwarded in cases:
            for failing in (False, True):
                with self.subTest(factory=factory.lineno, keywords=keywords, failing=failing):
                    shared = Mock()
                    error = RuntimeError('synthetic Chrome failure')
                    if failing:
                        shared.side_effect = error
                    env = {'create_unity_driver': shared, 'CHROME_DRIVER_PATH': 'synthetic-driver.exe',
                           'CONFIG_CHROME_DRIVER_PATH': 'synthetic-driver.exe'}
                    with patch.dict(sys.modules, {'AI.unity_browser': SimpleNamespace(create_driver=shared)}):
                        exec(compile(ast.Module(body=[factory], type_ignores=[]), '<actual driver>', 'exec'), env)
                        if failing:
                            with self.assertRaises(RuntimeError) as raised:
                                env['create_driver'](**keywords)
                            self.assertIs(raised.exception, error)
                        else:
                            self.assertIs(env['create_driver'](**keywords), shared.return_value)
                    shared.assert_called_once_with('synthetic-driver.exe', **forwarded)


class TrackingConfigurationTests(unittest.TestCase):
    CONNECTION = 'DRIVER={Synthetic Driver};SERVER=synthetic;PWD={fake;brace}}value}; '

    def setUp(self):
        self.enterContext(patch.dict(os.environ, {'FCM_RRS_DB_CONNECTION': self.CONNECTION}))
        self.enterContext(patch.object(sys, 'path', [str(model.ROOT / 'src'), *sys.path]))
        from fcm_intake import config, legacy_loader
        self.config = config
        self.runner = backend._runner()
        self.driver_error = type('SyntheticDriverError', (Exception,), {})
        self.odbc = MagicMock(Error=self.driver_error)
        self.connection = self.odbc.connect.return_value.__enter__.return_value
        self.cursor = self.connection.cursor.return_value.__enter__.return_value
        namespace = {'pyodbc': self.odbc}
        function = legacy_function('InsertIntoTable', 'legacy_fcm.py')
        exec(compile(ast.Module(body=[function], type_ignores=[]), '<actual tracking>', 'exec'), namespace)
        self.insert = namespace['InsertIntoTable']
        self.loader = self.enterContext(patch.object(legacy_loader, 'load_module_from_path',
            return_value=SimpleNamespace(InsertIntoTable=self.insert, main=Mock())))

    def test_existing_connection_and_exact_single_insert_are_preserved(self):
        with backend.redirect_stdout(backend.StringIO()):
            callbacks = backend.live_callbacks()
            callbacks['record_assignment']('Invented Person', 'SYNTHETIC-001')
        self.odbc.connect.assert_called_once_with(self.CONNECTION)
        self.cursor.execute.assert_called_once_with(
            'INSERT INTO [dbo].[RRS_APP_Sagility_Bot_Assignment_Log] ([Claimant],[ClaimNumber],[isAssigned]) VALUES (?,?,?)',
            ('Invented Person', 'SYNTHETIC-001', 0))
        self.connection.commit.assert_called_once_with()
        self.loader.assert_called_once()

    def test_missing_and_blank_stop_desktop_before_worker_and_callbacks(self):
        for value in (None, '', ' \t '):
            for stage in ('all', 'unity'):
                with self.subTest(value=value, stage=stage):
                    if value is None:
                        os.environ.pop('FCM_RRS_DB_CONNECTION', None)
                    else:
                        os.environ['FCM_RRS_DB_CONNECTION'] = value
                    api = backend.UnityAPI()
                    with patch.object(backend.threading, 'Thread') as thread, \
                            patch.object(backend, 'live_callbacks') as live:
                        result = api.start(synthetic(), stage)
                    self.assertFalse(result['ok'])
                    self.assertIn('FCM_RRS_DB_CONNECTION', result['error'])
                    self.assertIn('beside the EXE', result['error'])
                    self.assertIn('No automatic Unity replay', result['error'])
                    self.assertEqual(api.get_state()['status'], 'idle')
                    self.assertFalse(api.get_state()['busy'])
                    self.assertIsNone(api._thread)
                    thread.assert_not_called()
                    live.assert_not_called()
        self.odbc.connect.assert_not_called()
        self.loader.assert_not_called()

    def test_cli_preflight_stops_before_loading_portal(self):
        for value in ('', ' \t '):
            for stage in ('all', 'unity'):
                with self.subTest(stage=stage, value=value):
                    os.environ['FCM_RRS_DB_CONNECTION'] = value
                    output = backend.StringIO()
                    with backend.redirect_stdout(output), patch('sys.stderr', output):
                        code = self.runner.main(['--live', '--stage', stage])
                    self.assertEqual(code, 1)
                    self.assertIn('FCM_RRS_DB_CONNECTION', output.getvalue())
        self.loader.assert_not_called()
        self.odbc.connect.assert_not_called()

    def test_assign_injected_and_offline_do_not_require_tracking_configuration(self):
        os.environ['FCM_RRS_DB_CONNECTION'] = ''
        # Do not start a real worker: assign must get as far as worker creation.
        with patch.object(backend.threading, 'Thread') as thread:
            self.assertTrue(backend.UnityAPI().start({'claimNumber': 'SYNTHETIC-001'}, 'assign')['ok'])
            thread.assert_called_once()
        with patch.object(backend.threading, 'Thread') as thread:
            factory = Mock()
            self.assertTrue(backend.UnityAPI(factory).start(synthetic(), 'all')['ok'])
            thread.assert_called_once()
            factory.assert_not_called()
        with backend.redirect_stdout(backend.StringIO()), patch.object(
                self.runner, '_preflight_tracking', side_effect=AssertionError('Unexpected preflight')), \
                patch.object(self.runner, 'ask_yes_no', return_value=False):
            self.assertEqual(self.runner.main([]), 0)
            self.assertEqual(self.runner.main(['--live', '--stage', 'assign',
                                             '--claim-number', 'SYNTHETIC-001']),
                             2)  # The operator declines; no completion call is made.
        self.odbc.connect.assert_not_called()

    def test_driver_errors_are_value_free_keep_cause_and_never_retry(self):
        private = 'SYNTHETIC-PRIVATE-PASSWORD-ENDPOINT'
        for boundary in ('connect', 'execute', 'commit'):
            with self.subTest(boundary=boundary):
                self.odbc.reset_mock(return_value=False, side_effect=True)
                self.cursor.execute.side_effect = None
                self.connection.commit.side_effect = None
                error = self.driver_error('28000', private)
                target = {'connect': self.odbc.connect, 'execute': self.cursor.execute,
                          'commit': self.connection.commit}[boundary]
                target.side_effect = error
                with self.assertRaises(RuntimeError) as caught:
                    self.insert('Invented Person', 'SYNTHETIC-001')
                self.assertIs(caught.exception.__cause__, error)
                self.assertNotIn(private, str(caught.exception))
                self.odbc.connect.reset_mock()
                confirm, complete = Mock(), Mock()
                output = backend.StringIO()
                with backend.redirect_stdout(output):
                    result = self.runner.load_core().run_flow(synthetic(),
                        list_attachments=lambda: ['synthetic.pdf'], run_unity=Mock(),
                        record_assignment=self.insert, confirm=confirm, notify=Mock(),
                        complete_assignment=complete)
                self.assertEqual(result['reason'], 'record_assignment_failed')
                self.assertNotIn(private, json.dumps(result) + output.getvalue())
                self.assertIn('No automatic retry', result['error']['message'])
                self.odbc.connect.assert_called_once_with(self.CONNECTION)
                confirm.assert_not_called()
                complete.assert_not_called()


class RRSCompletionTests(unittest.TestCase):
    """Actual CLI/desktop completion delegate and controller, with synthetic UI."""

    class DialogAbsent(Exception):
        pass

    def setUp(self):
        self.events = []
        self.original_error = RuntimeError('SYNTHETIC-PRIVATE-PAGE-CONTENT')
        self.lookup = Mock(side_effect=lambda claim: self.events.append(('lookup', claim)) or 'SYNTHETIC-CASE')
        self.dlg, self.form, self.case = Mock(), Mock(), Mock()
        self.assign, self.input, self.input_ok = Mock(), Mock(), Mock()
        self.first, self.second, self.final = Mock(), Mock(), Mock()
        self.yes_first, self.yes_second, self.final_ok = Mock(), Mock(), Mock()
        self.first.__getitem__ = Mock(return_value=self.yes_first)
        self.second.__getitem__ = Mock(return_value=self.yes_second)
        controls = {'btnAssign': self.assign, 'frmInputBox': self.form,
                    'tbInput': self.input, 'btnOk': self.input_ok}
        for control in controls.values():
            control.wrapper_object.return_value = control
        self.dlg.child_window.side_effect = lambda **kw: controls[kw['auto_id']]
        self.form.child_window.side_effect = lambda **kw: controls[kw['auto_id']]
        self.final.child_window.return_value.wrapper_object.return_value = self.final_ok
        desktop = Mock()
        unity_prompts = iter((self.first, self.second))
        def window(**kw):
            if kw.get('title_re') == r'Unity Complete \?':
                return next(unity_prompts)
            if kw.get('title_re') == 'ReferralRoutingSystem':
                return self.final
            return self.dlg
        desktop.return_value.window.side_effect = window
        application = Mock()
        application.return_value.connect.return_value.window.return_value = self.case
        for control, name in ((self.assign, 'Assign'), (self.input_ok, 'input OK'),
                              (self.yes_first, 'first Yes'), (self.case.OK, 'Case Id OK'),
                              (self.yes_second, 'second Yes'), (self.final_ok, 'RRS OK')):
            control.click_input.side_effect = lambda name=name: self.events.append(name)
        self.input.iface_value.SetValue.side_effect = lambda value: self.events.append(('input', value))
        self.input.iface_value.CurrentValue = '19'
        self.case.Edit.set_edit_text.side_effect = lambda value: self.events.append(('Case Id', value))
        self.case.Edit.window_text.return_value = 'SYNTHETIC-CASE'
        env = {'Desktop': desktop, 'Application': application,
               'time': SimpleNamespace(sleep=lambda _: None),
               'data': {'claimNumber': 'SYNTHETIC-STALE'}}
        function = legacy_function('CompleteAssignment', 'legacy_fcm.py')
        exec(compile(ast.Module(body=[function], type_ignores=[]), '<actual RRS delegate>', 'exec'), env)
        self.legacy = SimpleNamespace(CompleteAssignment=env['CompleteAssignment'])
        self.enterContext(patch.dict(sys.modules, {
            'pywinauto.timings': SimpleNamespace(TimeoutError=self.DialogAbsent),
            'legacy': SimpleNamespace(CMS=SimpleNamespace(CheckCaseUnity=self.lookup))}))
        self.enterContext(patch.object(sys, 'path', [str(model.ROOT / 'src'), *sys.path]))
        from fcm_intake import legacy_loader
        self.loader = self.enterContext(patch.object(
            legacy_loader, 'load_module_from_path', return_value=self.legacy))
        self.runner = backend._runner()
        self.core = self.runner.load_core()

    def run_completion(self, *, stage='all', approved=True):
        # Both live frontends use this callback, not a parallel implementation.
        callbacks = backend.live_callbacks()
        callbacks.update(list_attachments=lambda: self.events.append('attachment') or ['synthetic.pdf'],
                         run_unity=lambda *args: self.events.append('Unity'),
                         record_assignment=lambda *args: self.events.append('tracking'))
        def confirm(*args):
            self.events.append('approval')
            return approved
        return self.core.run_flow(synthetic(), stage=stage, **callbacks,
                                  confirm=confirm, notify=Mock())

    def test_all_order_uses_selected_claim_despite_stale_legacy_global(self):
        result = self.run_completion()
        self.assertEqual(result['status'], 'completed')
        self.assertEqual(self.events, [
            'attachment', 'Unity', 'tracking', 'approval', 'Assign', ('input', '19'),
            'input OK', 'first Yes', ('lookup', 'SYNTHETIC-001'),
            ('Case Id', 'SYNTHETIC-CASE'), 'Case Id OK', 'second Yes', 'RRS OK'])
        self.lookup.assert_called_once_with('SYNTHETIC-001')
        self.loader.assert_called_once()
        self.case.Edit.window_text.assert_called_once()
        self.case.wait_not.assert_called_once_with('visible', timeout=10)

    def test_denied_gate_and_unity_only_never_touch_rrs_or_lookup(self):
        for stage, approved in (('all', False), ('assign', False), ('unity', True)):
            with self.subTest(stage=stage):
                self.events.clear()
                result = self.run_completion(stage=stage, approved=approved)
                self.assertEqual(result['status'], 'completed' if stage == 'unity' else 'stopped')
                self.assign.click_input.assert_not_called()
                self.lookup.assert_not_called()
                self.case.Edit.set_edit_text.assert_not_called()
                self.loader.assert_not_called()

    def test_missing_or_failed_lookup_blocks_case_id_submission_and_retains_cause(self):
        for value in (None, '', ' \t', self.original_error):
            with self.subTest(value_type=type(value).__name__):
                self.doCleanups()
                self.setUp()
                self.lookup.side_effect = self.original_error if isinstance(value, Exception) else None
                self.lookup.return_value = value
                result = self.run_completion(stage='assign')
                self.assertEqual(result['reason'], 'complete_assignment_failed')
                self.assertNotIn('SYNTHETIC', result['error']['message'])
                self.case.Edit.set_edit_text.assert_not_called()
                self.case.OK.click_input.assert_not_called()
        self.lookup.side_effect = self.original_error
        with self.assertRaises(RRSCompletionError) as raised:
            self.legacy.CompleteAssignment('SYNTHETIC-CURRENT')
        self.assertIs(raised.exception.__cause__, self.original_error)

    def test_required_writes_and_clicks_fail_without_retry_or_false_completion(self):
        for action in ('assign', 'input', 'input_ok', 'case_entry', 'case_ok'):
            with self.subTest(action=action):
                self.doCleanups()
                self.setUp()
                methods = {'assign': self.assign.click_input, 'input': self.input.iface_value.SetValue,
                           'input_ok': self.input_ok.click_input,
                           'case_entry': self.case.Edit.set_edit_text, 'case_ok': self.case.OK.click_input}
                methods[action].side_effect = self.original_error
                result = self.run_completion(stage='assign')
                self.assertEqual(result['reason'], 'complete_assignment_failed')
                self.assertNotIn('SYNTHETIC', result['error']['message'])
                methods[action].assert_called_once()
                self.final_ok.click_input.assert_not_called()
                if action in ('assign', 'input', 'input_ok'):
                    self.lookup.assert_not_called()
                if action != 'case_ok':
                    self.case.OK.click_input.assert_not_called()

    def test_optional_absence_is_allowed_but_present_action_or_readiness_failure_is_not(self):
        for dialog in ('first', 'second', 'final'):
            for failure in ('absent', 'click', 'ready', 'focus', 'wait'):
                with self.subTest(dialog=dialog, failure=failure):
                    self.doCleanups()
                    self.setUp()
                    prompt = getattr(self, dialog)
                    button = {'first': self.yes_first, 'second': self.yes_second,
                              'final': self.final_ok}[dialog]
                    if failure == 'absent':
                        prompt.wait.side_effect = self.DialogAbsent()
                    elif failure == 'click':
                        button.click_input.side_effect = self.original_error
                    elif failure == 'ready':
                        prompt.wait.side_effect = [None, self.DialogAbsent()]
                    elif failure == 'focus':
                        prompt.set_focus.side_effect = self.original_error
                    else:
                        prompt.wait.side_effect = self.original_error
                    result = self.run_completion(stage='assign')
                    self.assertEqual(result['status'], 'completed' if failure == 'absent' else 'failed')
                    if failure == 'absent':
                        button.click_input.assert_not_called()
                    else:
                        self.assertEqual(result['reason'], 'complete_assignment_failed')
                        self.assertNotIn('SYNTHETIC', result['error']['message'])

    def test_silent_write_mismatch_or_unclosed_case_prompt_stops_continuation(self):
        for failure, expected_stage in (
                ('input', 'assignment input verification'),
                ('case_entry', 'Case Id verification'),
                ('case_close', 'Case Id closure')):
            with self.subTest(failure=failure):
                self.doCleanups()
                self.setUp()
                if failure == 'input':
                    self.input.iface_value.CurrentValue = ''
                elif failure == 'case_entry':
                    self.case.Edit.window_text.return_value = 'SYNTHETIC-STALE'
                else:
                    self.case.wait_not.side_effect = self.DialogAbsent()
                result = self.run_completion(stage='assign')
                self.assertEqual(result['reason'], 'complete_assignment_failed')
                self.assertIn(expected_stage, result['error']['message'])
                self.assertNotIn('SYNTHETIC', result['error']['message'])
                self.yes_second.click_input.assert_not_called()
                self.final_ok.click_input.assert_not_called()
                if failure == 'input':
                    self.input_ok.click_input.assert_not_called()
                    self.lookup.assert_not_called()
                if failure != 'case_close':
                    self.case.OK.click_input.assert_not_called()
                else:
                    self.case.OK.click_input.assert_called_once()

    def test_lookup_browser_cleanup_retains_original_failure(self):
        function = legacy_function('CheckCaseUnity', 'CMS.py')
        function.body = [node for node in function.body if not isinstance(node, ast.FunctionDef)]
        for failed_stage in (None, 'login', 'search'):
            for cleanup_fails in (False, True):
                with self.subTest(failed_stage=failed_stage, cleanup_fails=cleanup_fails):
                    driver = Mock()
                    if cleanup_fails:
                        driver.quit.side_effect = RuntimeError('synthetic cleanup failure')
                    login, search = Mock(), Mock(return_value='SYNTHETIC-CASE')
                    if failed_stage:
                        (login if failed_stage == 'login' else search).side_effect = self.original_error
                    env = {'create_driver': lambda: driver, 'loginUnity': login,
                           'NavigateSearch': search, 'sys': sys}
                    exec(compile(ast.Module(body=[function], type_ignores=[]), '<actual lookup cleanup>', 'exec'), env)
                    if failed_stage or cleanup_fails:
                        with self.assertRaises(RuntimeError) as raised:
                            env['CheckCaseUnity']('SYNTHETIC-CURRENT')
                        if failed_stage:
                            self.assertIs(raised.exception, self.original_error)
                    else:
                        self.assertEqual(env['CheckCaseUnity']('SYNTHETIC-CURRENT'), 'SYNTHETIC-CASE')
                        search.assert_called_once_with('SYNTHETIC-CURRENT', driver)
                    driver.quit.assert_called_once()

    def actual_lookup(self, *, ready=True):
        """Keep the real login and cleanup, replacing only external boundaries."""
        function = legacy_function('CheckCaseUnity', 'CMS.py')
        function.body = [node for node in function.body
                         if not isinstance(node, ast.FunctionDef)
                         or node.name == 'loginUnity']
        driver = Mock()
        driver.current_window_handle = 'synthetic-unity'
        driver.get.side_effect = lambda url: self.events.append('open Unity')
        field = Mock()
        field.is_displayed.return_value = ready
        field.is_enabled.return_value = True
        field.get_attribute.return_value = None
        driver.find_elements.return_value = [field]
        driver.find_elements.side_effect = lambda *args: self.events.append('check ready') or [field]
        search = Mock(side_effect=lambda *args: self.events.append('search') or 'SYNTHETIC-CASE')
        def until(predicate, message):
            self.events.append('wait ready')
            if not predicate(driver):
                raise self.original_error
        wait = Mock(return_value=SimpleNamespace(until=until))
        env = {'create_driver': lambda: self.events.append('create browser') or driver,
               'NavigateSearch': search, 'sys': sys, 'By': SimpleNamespace(ID='id'),
               'WebDriverWait': wait}
        exec(compile(ast.Module(body=[function], type_ignores=[]), '<actual lookup login>', 'exec'), env)
        self.lookup.side_effect = env['CheckCaseUnity']
        return driver, field, search, wait

    def test_desktop_manual_login_blocks_browser_and_rrs_until_explicit_yes(self):
        driver, field, search, wait = self.actual_lookup()
        api = backend.UnityAPI()
        api._started = time.monotonic()
        original_confirm = api.confirm
        pending = threading.Event()
        def confirm(title, message):
            if title == 'Unity manual sign-in':
                pending.set()
                return original_confirm(title, message)
            self.events.append('approval')
            return True
        api.confirm = confirm
        api.notify = Mock()
        with patch.dict(sys.modules, {'pythoncom': Mock()}):
            worker = threading.Thread(target=api._run, args=(synthetic(), 'assign', True, None))
            worker.start()
            try:
                self.assertTrue(pending.wait(2))
                # The bridge has no timeout; allow it to publish the existing prompt.
                deadline = time.monotonic() + 2
                while api.get_state()['prompt'] is None and time.monotonic() < deadline:
                    time.sleep(.001)
                prompt = api.get_state()['prompt']
                self.assertEqual(prompt['kind'], 'confirm')
                self.assertEqual(prompt['title'], 'Unity manual sign-in')
                self.assertTrue(worker.is_alive())
                self.assertEqual(self.events[-2:], ['create browser', 'open Unity'])
                driver.get.assert_called_once()
                driver.switch_to.window.assert_not_called()
                driver.find_elements.assert_not_called()
                wait.assert_not_called()
                search.assert_not_called()
                self.case.Edit.set_edit_text.assert_not_called()
                self.assertTrue(api.answer_prompt(prompt['id'], True)['ok'])
            finally:
                prompt = api.get_state()['prompt']
                if prompt:
                    api.answer_prompt(prompt['id'], False)
                worker.join(2)
            self.assertFalse(worker.is_alive())
        self.assertEqual(api.get_state()['result']['status'], 'completed')
        self.assertLess(self.events.index('wait ready'), self.events.index('search'))
        self.assertLess(self.events.index('search'), self.events.index(('Case Id', 'SYNTHETIC-CASE')))
        # Already-authenticated sessions still needed the same manual confirmation.
        driver.get.assert_called_once()
        driver.quit.assert_called_once()

    def test_cli_manual_login_wiring_reaches_actual_lookup_after_approval(self):
        driver, field, search, wait = self.actual_lookup(ready=False)
        def answer(prompt):
            if prompt.startswith('Unity manual sign-in:'):
                self.events.append('login confirmation')
                field.is_displayed.return_value = True
            return 'yes'
        with patch('builtins.input', side_effect=answer) as confirm, \
                patch.object(self.legacy, 'main', Mock(), create=True), \
                backend.redirect_stdout(backend.StringIO()):
            code = self.runner.main(['--live', '--stage', 'assign',
                                     '--claim-number', 'SYNTHETIC-001'])
        self.assertEqual(code, 0)
        self.assertEqual(confirm.call_count, 2)
        self.assertLess(self.events.index('open Unity'), self.events.index('login confirmation'))
        self.assertLess(self.events.index('login confirmation'), self.events.index('search'))
        driver.quit.assert_called_once()

    def test_login_decline_prompt_error_and_unready_page_stop_all_continuation(self):
        for failure in ('decline', 'nonboolean', 'prompt error', 'logged out',
                        'disabled', 'readonly'):
            with self.subTest(failure=failure):
                self.doCleanups()
                self.setUp()
                driver, field, search, wait = self.actual_lookup(ready=failure != 'logged out')
                if failure == 'disabled':
                    field.is_enabled.return_value = False
                if failure == 'readonly':
                    field.get_attribute.return_value = 'true'
                confirm = Mock(return_value=False if failure == 'decline' else
                               'yes' if failure == 'nonboolean' else True)
                if failure == 'prompt error':
                    confirm.side_effect = self.original_error
                with self.assertRaises(RRSCompletionError) as raised:
                    self.runner._complete_live(model.ROOT, 'SYNTHETIC-001',
                                               confirm_unity_login=confirm)
                self.assertEqual(raised.exception.stage, 'Unity case lookup')
                self.assertNotIn('SYNTHETIC', str(raised.exception))
                if failure in ('prompt error', 'logged out', 'disabled', 'readonly'):
                    self.assertIs(raised.exception.__cause__, self.original_error)
                if failure in ('decline', 'nonboolean', 'prompt error'):
                    driver.get.assert_called_once()
                    driver.find_elements.assert_not_called()
                    wait.assert_not_called()
                else:
                    self.assertEqual(driver.get.call_count, 2)
                    wait.assert_called_once_with(driver, 15)
                search.assert_not_called()
                self.case.Edit.set_edit_text.assert_not_called()
                self.case.OK.click_input.assert_not_called()
                self.yes_second.click_input.assert_not_called()
                self.final_ok.click_input.assert_not_called()
                driver.quit.assert_called_once()

    def test_post_confirmation_navigation_can_reach_search_before_bounded_readiness(self):
        driver, field, search, wait = self.actual_lookup(ready=False)
        def navigate(url):
            self.events.append('open Unity')
            if driver.get.call_count == 2:
                field.is_displayed.return_value = True
        driver.get.side_effect = navigate
        confirm = Mock(side_effect=lambda *args: self.events.append('login confirmation') or True)
        self.runner._complete_live(model.ROOT, 'SYNTHETIC-001', confirm_unity_login=confirm)
        self.assertEqual(driver.get.call_count, 2)
        self.assertEqual(self.events.count('search'), 1)
        self.assertLess(self.events.index('login confirmation'), self.events.index('check ready'))
        wait.assert_called_once_with(driver, 15)
        driver.quit.assert_called_once()

    def test_direct_legacy_default_is_attended_without_notify_fallback(self):
        for approved in (False, True, self.original_error):
            with self.subTest(approved=approved):
                self.doCleanups()
                self.setUp()
                driver, field, search, wait = self.actual_lookup()
                root = Mock()
                messagebox = SimpleNamespace(askyesno=Mock(return_value=approved))
                if isinstance(approved, Exception):
                    messagebox.askyesno.side_effect = self.original_error
                tk = SimpleNamespace(Tk=Mock(return_value=root), messagebox=messagebox)
                with patch.dict(sys.modules, {'tkinter': tk, 'tkinter.messagebox': messagebox}):
                    if approved is True:
                        self.legacy.CompleteAssignment('SYNTHETIC-001')
                        search.assert_called_once()
                    else:
                        with self.assertRaises(RRSCompletionError) as raised:
                            self.legacy.CompleteAssignment('SYNTHETIC-001')
                        if isinstance(approved, Exception):
                            self.assertIs(raised.exception.__cause__, self.original_error)
                        search.assert_not_called()
                        self.case.Edit.set_edit_text.assert_not_called()
                messagebox.askyesno.assert_called_once()
                self.assertIs(messagebox.askyesno.call_args.kwargs['parent'], root)
                root.destroy.assert_called_once()
                driver.quit.assert_called_once()

    def all_fixture(self):
        """Real frontend/controller/delegates; only external I/O is synthetic."""
        driver, field, search, wait = self.actual_lookup()
        env = {
            'create_driver': Mock(return_value=driver),
            'login': Mock(side_effect=lambda actual: self.assertIs(actual, driver)),
            'validate_and_complete_task': Mock(),
            'UnityTaskCompletionError': UnityTaskCompletionError, 'sys': sys,
        }
        exec(compile(ast.Module(body=[legacy_function('main')], type_ignores=[]),
                     '<actual borrowed Unity main>', 'exec'), env)
        unity = SimpleNamespace(**{key: env[key] for key in
            ('main', 'create_driver', 'login', 'validate_and_complete_task')})
        self.loader.side_effect = lambda name, path: (
            unity if name.endswith('legacy_unity') else self.legacy)
        config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py',
                                 FCM_SCRIPT='synthetic-fcm.py')
        self.enterContext(patch.dict(sys.modules, {'fcm_intake.config': config}))
        tracking = self.enterContext(patch.object(self.runner, '_record_live',
            side_effect=lambda *args: driver.quit.assert_not_called()))
        # Desktop loads a runner lazily; use this same real module with its I/O patched.
        self.enterContext(patch.object(backend, '_runner', return_value=self.runner))
        def finish():
            driver.quit.assert_not_called()
            self.events.append('RRS finished')
        self.final_ok.click_input.side_effect = finish
        driver.quit.side_effect = lambda: self.events.append('quit')
        return unity, driver, field, search, tracking, config

    def execute_all(self, frontend, config, answer, api=None):
        if frontend == 'desktop':
            api = api or backend.UnityAPI()
            api._started = time.monotonic()
            api.confirm, api.notify = answer, Mock()
            with patch.dict(sys.modules, {'pythoncom': Mock()}), \
                    patch.object(Path, 'glob', return_value=[Path('synthetic.pdf')]), \
                    patch.object(Path, 'is_file', return_value=True):
                api._run(synthetic(), 'all', True, None)
            return api.get_state()['result']
        with patch.object(self.runner, '_preflight_tracking'), \
                patch.object(self.runner, 'ask_yes_no', side_effect=answer), \
                patch.object(Path, 'glob', return_value=[Path('synthetic.pdf')]), \
                patch.object(Path, 'is_file', return_value=True), \
                backend.redirect_stdout(backend.StringIO()) as output, \
                patch('sys.stderr', output):
            code = self.runner.main(['--live', '--claim-number', 'SYNTHETIC-001',
                '--customer', 'Synthetic', '--claimant-first', 'Invented',
                '--claimant-last', 'Example'])
        if code == 1:
            return {'status': 'failed', 'reason': 'cli_exception', 'error': output.getvalue()}
        text = output.getvalue()
        return json.loads(text[text.index('{\n'):])

    def test_actual_all_frontends_borrow_identical_driver_until_rrs_finishes(self):
        for frontend in ('desktop', 'cli'):
            with self.subTest(frontend=frontend):
                self.doCleanups()
                self.setUp()
                unity, driver, field, search, tracking, config = self.all_fixture()
                def answer(title, message):
                    driver.quit.assert_not_called()
                    if title == 'Unity manual sign-in':
                        # No lookup navigation/probe even when already signed in.
                        self.events.append('manual Yes')
                        driver.get.assert_not_called()
                        driver.find_elements.assert_not_called()
                        search.assert_not_called()
                    else:
                        self.events.append('approval')
                    return True
                result = self.execute_all(frontend, config, answer)
                self.assertEqual(result['status'], 'completed', result)
                unity.create_driver.assert_called_once()
                unity.login.assert_called_once_with(driver)
                self.assertIs(unity.validate_and_complete_task.call_args.args[0], driver)
                self.lookup.assert_called_once()
                self.assertIs(self.lookup.call_args.kwargs['unity_driver'], driver)
                search.assert_called_once_with('SYNTHETIC-001', driver)
                tracking.assert_called_once()
                driver.get.assert_called_once()
                driver.quit.assert_called_once()
                self.assertNotIn('create browser', self.events)
                self.assertLess(self.events.index('approval'), self.events.index('manual Yes'))
                self.assertLess(self.events.index('manual Yes'), self.events.index('open Unity'))
                self.assertEqual(self.events[-2:], ['RRS finished', 'quit'])

    def test_all_stops_clean_up_without_retry_or_replacement(self):
        for frontend in ('desktop', 'cli'):
            for failure in ('approval', 'tracking', 'lookup', 'login No', 'stale',
                            'cleanup', 'lookup and cleanup', 'task'):
                with self.subTest(frontend=frontend, failure=failure):
                    self.doCleanups()
                    self.setUp()
                    unity, driver, field, search, tracking, config = self.all_fixture()
                    if failure == 'tracking':
                        tracking.side_effect = self.original_error
                    if 'lookup' in failure:
                        search.side_effect = self.original_error
                    if failure == 'stale':
                        driver.get.side_effect = self.original_error
                    if 'cleanup' in failure:
                        driver.quit.side_effect = RuntimeError('synthetic cleanup failed')
                    if failure == 'task':
                        unity.validate_and_complete_task.side_effect = UnityTaskCompletionError('click')
                    def answer(title, message):
                        if title == 'Unity manual sign-in':
                            driver.get.assert_not_called()
                            driver.find_elements.assert_not_called()
                            return failure != 'login No'
                        return failure != 'approval'
                    result = self.execute_all(frontend, config, answer)
                    self.assertEqual(result['status'], 'stopped' if failure == 'approval' else 'failed', result)
                    unity.create_driver.assert_called_once()
                    unity.login.assert_called_once()
                    unity.validate_and_complete_task.assert_called_once()
                    driver.quit.assert_called_once()
                    self.assertNotIn('create browser', self.events)
                    if failure in ('approval', 'tracking', 'task'):
                        self.lookup.assert_not_called()
                        driver.get.assert_not_called()
                    else:
                        self.lookup.assert_called_once()
                    if failure in ('approval', 'tracking', 'task', 'login No', 'stale'):
                        search.assert_not_called()
                    if failure != 'cleanup':
                        self.final_ok.click_input.assert_not_called()
                    if failure == 'lookup and cleanup':
                        self.assertIn('complete_assignment_failed', str(result))
                        self.assertIn('cleanup failed', str(result))

    def test_all_desktop_retains_only_three_manual_recovery_types(self):
        for error_type in (UnityFlagError, UnityCustomerFieldsError, UnityEmployerFieldsError):
            with self.subTest(error_type=error_type):
                self.doCleanups()
                self.setUp()
                unity, driver, field, search, tracking, config = self.all_fixture()
                error = error_type('synthetic stage')
                unity.validate_and_complete_task.side_effect = error
                result = self.execute_all('desktop', config, Mock())
                self.assertEqual(result['reason'], 'run_unity_failed')
                driver.quit.assert_not_called()
                driver.service.stop.assert_called_once()
                driver.command_executor.close.assert_called_once()
                self.assertTrue(error.browser_left_open)
                self.lookup.assert_not_called()

    def test_borrowed_lookup_decline_and_failure_never_quit_or_create(self):
        for failure in ('No', 'stale', 'search'):
            with self.subTest(failure=failure):
                self.doCleanups()
                self.setUp()
                driver, field, search, wait = self.actual_lookup()
                if failure == 'stale':
                    driver.get.side_effect = self.original_error
                if failure == 'search':
                    search.side_effect = self.original_error
                with self.assertRaises(RRSCompletionError) as caught:
                    self.runner._complete_live(model.ROOT, 'SYNTHETIC-001',
                        confirm_unity_login=lambda *args: failure != 'No', unity_driver=driver)
                if failure != 'No':
                    self.assertIs(caught.exception.__cause__, self.original_error)
                self.assertNotIn('create browser', self.events)
                driver.quit.assert_not_called()
                if failure == 'No':
                    driver.get.assert_not_called()
                    driver.find_elements.assert_not_called()
                    search.assert_not_called()

    def test_independent_all_runs_use_fresh_drivers_on_same_frontend(self):
        stable_runner = self.runner
        for frontend in ('desktop', 'cli'):
            api = backend.UnityAPI()
            previous = None
            for operation in range(2):
                with self.subTest(frontend=frontend, operation=operation):
                    self.doCleanups()
                    self.setUp()
                    self.runner = stable_runner
                    unity, driver, field, search, tracking, config = self.all_fixture()
                    result = self.execute_all(frontend, config, lambda *args: True, api)
                    self.assertEqual(result['status'], 'completed', result)
                    self.assertIsNot(driver, previous)
                    unity.create_driver.assert_called_once()
                    search.assert_called_once_with('SYNTHETIC-001', driver)
                    driver.quit.assert_called_once()
                    self.assertEqual(self.events[-2:], ['RRS finished', 'quit'])
                    previous = driver

    def test_legacy_main_owns_default_but_not_borrowed_driver_and_keeps_swallow_policy(self):
        for borrowed in (False, True):
            for failure in (None, ValueError('synthetic unrelated'),
                            UnityTaskCompletionError('synthetic completion')):
                with self.subTest(borrowed=borrowed, failure=type(failure)):
                    driver = Mock()
                    env = {'create_driver': Mock(return_value=driver), 'login': Mock(),
                           'validate_and_complete_task': Mock(side_effect=failure),
                           'UnityTaskCompletionError': UnityTaskCompletionError, 'sys': sys}
                    function = legacy_function('main')
                    exec(compile(ast.Module(body=[function], type_ignores=[]),
                                 '<actual Unity ownership>', 'exec'), env)
                    options = {'unity_driver': driver} if borrowed else {}
                    with backend.redirect_stdout(backend.StringIO()):
                        if isinstance(failure, UnityTaskCompletionError):
                            with self.assertRaises(UnityTaskCompletionError) as caught:
                                env['main']('Synthetic', 'Example', {}, **options)
                            self.assertIs(caught.exception, failure)
                        else:
                            env['main']('Synthetic', 'Example', {}, **options)
                    self.assertEqual(env['create_driver'].call_count, int(not borrowed))
                    env['login'].assert_called_once_with(driver)
                    env['validate_and_complete_task'].assert_called_once_with(driver, 'Synthetic', 'Example', {})
                    self.assertEqual(driver.quit.call_count, int(not borrowed))

    def test_actual_case_search_replaces_old_claim_with_exact_current_claim(self):
        function = next(node for node in legacy_function('CheckCaseUnity', 'CMS.py').body
                        if isinstance(node, ast.FunctionDef) and node.name == 'NavigateSearch')
        driver, field = Mock(), Mock()
        stored = ['SYNTHETIC-PREVIOUS']
        field.clear.side_effect = lambda: stored.clear()
        field.send_keys.side_effect = lambda claim: stored.append(claim)
        driver.find_element.side_effect = lambda by, locator: (
            field if locator == 'tasks-claim-number-general-input' else Mock())
        driver.find_elements.return_value = []
        env = {'WebDriverWait': Mock(), 'EC': Mock(), 'Select': Mock(),
               'By': SimpleNamespace(ID='id', XPATH='xpath'),
               'time': SimpleNamespace(sleep=lambda _: None)}
        exec(compile(ast.Module(body=[function], type_ignores=[]),
                     '<actual exact current claim search>', 'exec'), env)
        claim = 'SYNTHETIC-CURRENT/001-A'
        with backend.redirect_stdout(backend.StringIO()):
            self.assertIsNone(env['NavigateSearch'](claim, driver))
        field.clear.assert_called_once_with()
        field.send_keys.assert_called_once_with(claim)
        self.assertEqual(stored, [claim])

    def test_session_cleanup_surfaces_primary_exception_and_cleanup_failure(self):
        from AI.unity_browser import UnityBrowserSession
        primary = ValueError('SYNTHETIC-PRIVATE process failure')
        original_cause = LookupError('synthetic root cause')
        primary.__cause__ = original_cause
        cleanup = RuntimeError('SYNTHETIC-PRIVATE cleanup failure')
        driver = Mock()
        driver.quit.side_effect = cleanup
        with self.assertRaisesRegex(RuntimeError, 'cleanup failed after ValueError') as caught:
            with UnityBrowserSession() as session:
                session.acquire(lambda: driver)
                raise primary
        self.assertIs(caught.exception.__cause__, primary)
        self.assertIs(primary.__cause__, original_cause)
        self.assertIs(caught.exception.cleanup_error, cleanup)
        self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
        driver.quit.assert_called_once()
        for result in (
                None,
                {'status': 'completed', 'reason': '', 'error': {}},
                {'status': 'failed', 'reason': 'complete_assignment_failed',
                 'error': {'type': 'RRSCompletionError', 'message': 'SYNTHETIC-PRIVATE'}}):
            with self.subTest(result=result):
                driver.quit.reset_mock()
                with self.assertRaisesRegex(RuntimeError, 'Unity browser cleanup failed') as caught:
                    with UnityBrowserSession() as session:
                        session.acquire(lambda: driver)
                        session.result = result
                self.assertIs(caught.exception.__cause__, cleanup)
                self.assertIs(caught.exception.process_result, result)
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                if result and result['status'] == 'failed':
                    self.assertIn('complete_assignment_failed', str(caught.exception))
                    self.assertIn('RRSCompletionError', str(caught.exception))
                driver.quit.assert_called_once()


class ModelTests(unittest.TestCase):
    def test_contract_and_ast(self):
        audit = model.verify_legacy_keys()
        self.assertEqual(audit['count'], 39)
        self.assertEqual(audit['context_only'], [])
        self.assertEqual(audit['unexpected'], [])
        self.assertEqual(len(model.CONTEXT_KEYS), 39)
        self.assertEqual(len(model.defaults()), 35)
        consumed_derived = set(model.DERIVED_KEYS) & set(model.CONTEXT_KEYS)
        self.assertEqual(len(consumed_derived), 5)
        self.assertEqual(set(model.defaults()), set(model.CONTEXT_KEYS) - consumed_derived | {'claimNumber'})
        self.assertEqual([g['title'] for g in model.schema()], ['Claim', 'Claimant', 'Provider', 'Adjuster', 'Attorney'])

    def test_defaults_and_requirements_by_stage(self):
        self.assertFalse(any(model.defaults().values()))
        self.assertEqual(model.validate({'claimNumber': ' X '}, 'assign')['claimNumber'], 'X')
        for stage in ('all', 'unity'):
            with self.subTest(stage=stage), self.assertRaisesRegex(ValueError, 'customer, claimantFirst, claimantLast'):
                model.validate({'claimNumber': 'X'}, stage)
        for stage in ('all', 'unity', 'assign'):
            with self.subTest(stage=stage), self.assertRaisesRegex(ValueError, 'claimNumber'):
                model.validate({}, stage)
        self.assertTrue(model.warnings(model.validate(synthetic())))

    def test_invalid_supplied_values_even_on_assign(self):
        for key, value in [('dob', '02/30/2000'), ('nextApptDate', 'tomorrow'), ('nextApptTime', '25:90'),
                           ('state', 'ZZ'), ('providerState', 'California'), ('claimantAttyState', 'zz'),
                           ('zip', '123'), ('providerZip', '123456'), ('claimantAttyZip', 'oops'),
                           ('phoneNumber', '123'), ('claimantAttyPhone', 'invalid'), ('gender', 'unsupported'),
                           ('serviceType', 'Unsupported'), ('claimID', 5), ('claimID', 'x' * 1001)]:
            with self.subTest(key=key, value=value), self.assertRaises(ValueError):
                model.validate(synthetic() | {key: value}, 'assign')
        for payload, stage in [([], 'all'), ({'unknown': 'x'}, 'assign'), ({}, 'wrong')]:
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                model.validate(payload, stage)

    def test_date_time_zip_phone_normalization(self):
        result = model.validate(synthetic() | {'dob': '2000-01-02', 'nextApptTime': '14:30',
                        'zip': '00123', 'providerZip': '00123-4567', 'phoneNumber': '(555) 010-2345'})
        self.assertEqual(result['dob'], '01/02/2000')
        self.assertEqual(result['nextApptTime'], '02:30 PM')
        self.assertEqual(result['zip'], '00123')
        self.assertIn('no date', model.warnings(result)[1])

    def test_instruction_limits_and_exact_controller_preservation(self):
        for length in (1001, 32767):
            instructions = synthetic_instructions(length)
            for stage in ('all', 'unity', 'assign'):
                with self.subTest(length=length, stage=stage):
                    checked = model.validate(synthetic() | {'specialInstructions': instructions}, stage)
                    self.assertEqual(checked['specialInstructions'], instructions)
                    self.assertEqual(model.to_controller(checked)['specialInstructions'], instructions)
        for value in (synthetic_instructions(32768), 123, None):
            with self.subTest(value_type=type(value)), self.assertRaisesRegex(
                    ValueError, 'specialInstructions.*32,767 characters'):
                model.validate(synthetic() | {'specialInstructions': value})
        self.assertEqual(model.validate(synthetic() | {'specialInstructions': ' \n\t '})
                         ['specialInstructions'], ' \n\t ')
        self.assertEqual(model.validate(synthetic() | {'claimID': 'x' * 1000})['claimID'], 'x' * 1000)
        self.assertEqual(model.validate(synthetic() | {'claimID': ' X '})['claimID'], 'X')
        for key in model.defaults().keys() - {'specialInstructions'}:
            with self.subTest(key=key), self.assertRaisesRegex(ValueError, f'{key}.*1,000 characters'):
                model.validate(synthetic() | {key: 'x' * 1001})

    def test_frontend_fixture_matches_shared_length_schema(self):
        fixture = json.loads((HERE / 'frontend' / 'src' / 'schema.json').read_text(encoding='utf-8'))
        self.assertEqual(fixture['schema'], model.schema())
        fields = {field['key']: field for group in model.schema() for field in group['fields']}
        self.assertEqual(fields['specialInstructions']['maxLength'], 32767)
        self.assertTrue(all(field['maxLength'] == 1000 for key, field in fields.items()
                            if key != 'specialInstructions'))

    def test_derivation_defaults_and_state_is_not_jurisdiction(self):
        data = synthetic() | {'state': 'MI', 'providerZip': '00123'}
        before = data.copy()
        result = model.to_controller(data)
        self.assertEqual(data, before)
        self.assertEqual(result['claimStateAbbr'], '')
        self.assertEqual(result['correctZipCode'], '00123')
        self.assertEqual(result['refSource'], '')
        self.assertEqual(result['serviceType'], 'Medical Full')
        self.assertEqual(result['CaseObj'], 'Coordination of Care/Services')

    def test_continental_branch_and_explicit_semantics(self):
        data = synthetic() | {'customer': 'Synthetic Continental Tire division', 'claimType': 'Auto'}
        result = model.to_controller(data, claim_jurisdiction='MI')
        self.assertEqual((result['claimStateAbbr'], result['claimStateFull']), ('MI', 'Michigan'))
        self.assertEqual((result['ServiceType'], result['serviceType'], result['refType'], result['CaseObj']),
                         ('Medical Task', 'Medical Task', 'Task', 'MI MCCA Attendant Care'))
        explicit = model.to_controller(data | {'serviceType': 'Vocational Full'}, claim_jurisdiction='MI')
        self.assertEqual(explicit['serviceType'], 'Vocational Full')
        self.assertEqual(explicit['ServiceType'], 'Medical Task')
        self.assertEqual(model.to_controller(data, claim_jurisdiction='CA')['serviceType'], 'Medical Full')


class LegacyFlagTests(unittest.TestCase):
    def setUp(self):
        from selenium.webdriver.common.by import By
        from selenium.webdriver.remote.webelement import WebElement
        from selenium.webdriver.support import expected_conditions as EC

        self.events = []
        self.failure = None
        self.original_error = None
        self.open = False
        self.old_note_locator = (
            By.XPATH,
            './modal-template/div/div/div/div[3]/div/form/div/div[3]/text-area/form-control/div/textarea')
        self.stable_note_locator = (By.ID, 'intake-flag-for-intervention-description-text-area')
        owner = self

        class Element(WebElement):
            def __init__(self, name):
                self.name = name
                self.value = 'stale note'
                self.children = {}

            def find_element(self, by, selector):
                owner.events.append(('find', self.name, by, selector))
                child = self.children[(by, selector)]
                if owner.failure == child.name + '-locate':
                    owner.fail()
                return child

            def is_displayed(self):
                if self.name == 'modal':
                    owner.events.append(('modal-visible', owner.open))
                    return owner.open and owner.failure != 'modal-timeout'
                return True

            def is_enabled(self):
                return True

            def click(self):
                owner.events.append(('click', self.name))
                if owner.failure == self.name:
                    owner.fail()
                if self.name == 'flag':
                    owner.open = True
                elif self.name == 'reason':
                    owner.modal.children.pop(owner.old_note_locator, None)
                elif self.name == 'submit' and owner.failure != 'still-open':
                    owner.open = False

            def clear(self):
                owner.events.append(('clear', self.name))
                if owner.failure == 'note-clear':
                    owner.fail()
                self.value = ''

            def send_keys(self, value):
                owner.events.append(('type', self.name, value))
                if owner.failure == 'note-type':
                    owner.fail()
                self.value += value

            def get_attribute(self, name):
                owner.events.append(('read', self.name, name))
                if owner.failure == 'note-read':
                    owner.fail()
                return 'mismatched' if owner.failure == 'note' else self.value

        class Wait:
            def __init__(self, target, timeout):
                self.target = target

            def until(self, condition):
                for _ in range(3):
                    value = condition(self.target)
                    if value:
                        return value
                owner.fail(TimeoutError)

        self.driver, self.modal = Element('driver'), Element('modal')
        self.driver.children = {
            (By.XPATH, '/html/body/my-app/div/div/div/ng-component/div[1]/div/div[1]/span'): Element('flag'),
            (By.TAG_NAME, 'flag-for-intervention-modal'): self.modal,
        }
        note = Element('note')
        self.modal.children = {
            (By.XPATH, './modal-template/div/div/div/div[3]/div/form/div/div[1]/radio-group-input/form-control/div/label/div[1]/div/i[1]'): Element('reason'),
            self.old_note_locator: note,
            self.stable_note_locator: note,
            (By.XPATH, './modal-template/div/div/div/div[4]/div/div/div/button'): Element('submit'),
        }
        self.env = {'WebDriverWait': Wait, 'By': By, 'EC': EC, 'sys': sys,
                    'UnityFlagError': UnityFlagError,
                    'UnityCustomerFieldsError': UnityCustomerFieldsError,
                    'UnityEmployerFieldsError': UnityEmployerFieldsError,
                    'UnityTaskCompletionError': UnityTaskCompletionError,
                    'complete_task': complete_task}
        helper = legacy_function('flag_missing_appointment_time')
        exec(compile(ast.Module(body=[helper], type_ignores=[]), '<actual flag helper>', 'exec'), self.env)
        validation = legacy_function('validate_and_complete_task')
        branch, = [node for node in validation.body[0].body if isinstance(node, ast.If)
                   and any(isinstance(child, ast.Call) and isinstance(child.func, ast.Name)
                           and child.func.id == 'flag_missing_appointment_time'
                           for child in ast.walk(node))]
        flag, = [node for node in branch.body if isinstance(node, ast.If)
                 and any(isinstance(child, ast.Call) and isinstance(child.func, ast.Name)
                         and child.func.id == 'flag_missing_appointment_time'
                         for child in ast.walk(node))]
        # Keep the actual enclosing guard, flag statement and outer handler.
        # Replace unrelated portal work with sentinels at the following boundaries.
        self.assertEqual(branch.body[branch.body.index(flag) + 1].name, 'ValidateAddrPrv')
        branch = deepcopy(branch)
        branch.body = [deepcopy(flag), ast.parse('following_provider_work()').body[0]]
        branch.orelse = []
        validation.body[0].body = [branch, ast.parse('later_completion()').body[0]]
        self.env.update(following_provider_work=Mock(), later_completion=Mock())
        self.guard_code = compile(ast.fix_missing_locations(ast.Module(
            body=[branch], type_ignores=[])), '<actual flag guard>', 'exec')
        exec(compile(ast.fix_missing_locations(ast.Module(body=[validation], type_ignores=[])),
                     '<actual flag error boundary>', 'exec'), self.env)

    def fail(self, error_type=RuntimeError):
        self.original_error = error_type('SYNTHETIC-PRIVATE page content and Selenium message')
        raise self.original_error

    def prepare(self, **fields):
        data = model.to_controller(synthetic() | fields)
        self.env.update(driver=self.driver, dataRef=data,
                        providerName=data['providerName'], providerFirst=data['providerFirst'],
                        providerLast=data['providerLast'])
        return data

    def test_reason_layout_change_keeps_note_lookup_scoped_to_stable_id(self):
        self.prepare(providerName='Synthetic Facility')
        self.assertIs(self.modal.children[self.old_note_locator],
                      self.modal.children[self.stable_note_locator])
        exec(self.guard_code, self.env)
        self.assertNotIn(self.old_note_locator, self.modal.children)
        note_lookup = ('find', 'modal', *self.stable_note_locator)
        self.assertIn(note_lookup, self.events)
        self.assertLess(self.events.index(('click', 'reason')), self.events.index(note_lookup))
        self.assertNotIn(('find', 'modal', *self.old_note_locator), self.events)
        self.assertNotIn(('find', 'driver', *self.stable_note_locator), self.events)
        self.assertEqual(self.modal.children[self.stable_note_locator].value,
                         'Missing Provider Appointment Time')
        self.assertEqual(self.events.count(('click', 'submit')), 1)
        self.assertEqual(self.events[-1], ('modal-visible', False))
        self.env['following_provider_work'].assert_called_once()

    def test_actual_guard_and_successful_flag_continue(self):
        cases = [({}, False), ({'providerName': 'Synthetic Facility'}, True),
                 ({'providerFirst': 'Invented'}, True), ({'providerLast': 'Example'}, True),
                 ({'nextApptDate': '10/23/2026'}, True),
                 ({'providerName': 'Synthetic Facility', 'nextApptTime': '09:00 AM'}, False),
                 ({'nextApptDate': '10/23/2026', 'nextApptTime': '09:00 AM'}, False)]
        for fields, should_flag in cases:
            with self.subTest(fields=fields):
                self.setUp()
                self.prepare(**fields)
                exec(self.guard_code, self.env)
                clicks = [event for event in self.events if event[0] == 'click']
                self.assertEqual(clicks, [('click', name) for name in
                                         ('flag', 'reason', 'submit')] if should_flag else [])
                if should_flag:
                    self.assertEqual([event for event in self.events
                                      if event[0] in ('clear', 'type', 'read')],
                                     [('clear', 'note'), ('type', 'note', 'Missing Provider Appointment Time'),
                                      ('read', 'note', 'value')])
                    self.assertEqual(self.events[-1], ('modal-visible', False))
                    self.env['following_provider_work'].assert_called_once()
                elif fields:
                    self.env['following_provider_work'].assert_called_once()
                else:
                    self.env['following_provider_work'].assert_not_called()

    def test_flag_failures_stop_following_work_and_fail_checked_controller(self):
        stages = {
            'flag-locate': 'opening the flag dialog', 'flag': 'opening the flag dialog',
            'modal-timeout': 'waiting for the visible flag dialog',
            'reason-locate': 'selecting the missing-information reason',
            'reason': 'selecting the missing-information reason',
            'note-locate': 'locating the flag note', 'note-clear': 'clearing the flag note',
            'note-type': 'entering the flag note', 'note-read': 'verifying the flag note',
            'note': 'verifying the flag note', 'submit-locate': 'submitting the flag',
            'submit': 'submitting the flag', 'still-open': 'waiting for the flag dialog to close',
        }
        for failure, stage in stages.items():
            with self.subTest(failure=failure):
                self.setUp()
                data = self.prepare(providerName='Synthetic Facility')
                self.failure = failure
                with self.assertRaises(UnityFlagError) as caught:
                    exec(self.guard_code, self.env)
                self.assertEqual(caught.exception.stage, stage)
                self.assertIn(stage, str(caught.exception))
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                if failure == 'note':
                    self.assertIsInstance(caught.exception.__cause__, RuntimeError)
                else:
                    self.assertIs(caught.exception.__cause__, self.original_error)
                self.env['following_provider_work'].assert_not_called()
                clicks = self.events.count(('click', 'submit'))
                self.assertEqual(clicks, int(failure in ('submit', 'still-open')))
                if failure not in ('flag', 'flag-locate'):
                    self.assertTrue(self.open)
                unity = SimpleNamespace(validate_and_complete_task=self.env['validate_and_complete_task'])
                complete, confirm, record = Mock(), Mock(return_value=True), Mock()
                self.events.clear()
                result = backend._runner().load_core().run_flow(
                    data, list_attachments=lambda: [Path('synthetic.pdf')],
                    run_unity=lambda customer, claimant, payload:
                        backend.checked_validate(unity, self.driver, customer, claimant, payload),
                    record_assignment=record, complete_assignment=complete, confirm=confirm, notify=Mock())
                self.assertEqual((result['status'], result['reason']), ('failed', 'run_unity_failed'))
                self.assertEqual(result['error']['type'], 'UnityFlagError')
                self.assertIn(stage, result['error']['message'])
                self.assertNotIn('SYNTHETIC-PRIVATE', result['error']['message'])
                record.assert_called_once()
                confirm.assert_not_called()
                complete.assert_not_called()
                self.env['later_completion'].assert_not_called()
                self.env['following_provider_work'].assert_not_called()
                self.assertEqual(self.events.count(('click', 'submit')), clicks)

    def test_actual_outer_boundary_still_marks_unrelated_errors(self):
        self.prepare(providerName='Synthetic Facility')
        self.env['flag_missing_appointment_time'] = Mock(side_effect=ValueError('synthetic unrelated'))
        unity = SimpleNamespace(validate_and_complete_task=self.env['validate_and_complete_task'])
        with self.assertRaisesRegex(RuntimeError, 'Legacy Unity reported failure:.*synthetic unrelated'):
            backend.checked_validate(unity, self.driver, 'Synthetic', 'Example', self.env['dataRef'])
        self.env['following_provider_work'].assert_not_called()
        self.env['later_completion'].assert_not_called()

    def test_actual_flag_failure_through_desktop_preserves_browser_and_stops_flow(self):
        data = self.prepare(providerName='Synthetic Facility')
        self.failure = 'still-open'
        self.driver.service = Mock()
        self.driver.command_executor = Mock()
        self.driver.quit = Mock()
        unity = SimpleNamespace(create_driver=Mock(return_value=self.driver), login=Mock(),
                                validate_and_complete_task=self.env['validate_and_complete_task'])
        config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py')
        with patch.dict(sys.modules, {
                'fcm_intake.config': config,
                'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=lambda *args: unity)}):
            calls = backend.live_callbacks()
        record, confirm, complete = Mock(), Mock(return_value=True), Mock()
        result = backend._runner().load_core().run_flow(
            data, list_attachments=lambda: [Path('synthetic.pdf')], run_unity=calls['run_unity'],
            record_assignment=record, confirm=confirm, complete_assignment=complete, notify=Mock())
        self.assertEqual((result['status'], result['reason']), ('failed', 'run_unity_failed'))
        self.assertEqual(result['error']['type'], 'UnityFlagError')
        self.assertIn('waiting for the flag dialog to close', result['error']['message'])
        self.assertIn('Browser left open for manual recovery', result['error']['message'])
        self.assertNotIn('SYNTHETIC-PRIVATE', result['error']['message'])
        self.assertEqual(self.events.count(('click', 'submit')), 1)
        self.driver.service.stop.assert_called_once()
        self.driver.command_executor.close.assert_called_once()
        self.driver.quit.assert_not_called()
        self.env['following_provider_work'].assert_not_called()
        self.env['later_completion'].assert_not_called()
        record.assert_called_once()
        confirm.assert_not_called()
        complete.assert_not_called()


class CompleteTaskTests(unittest.TestCase):
    class Button:
        def __init__(self, text, *, visible=True, enabled=True, aria=None):
            self.text, self.visible, self.enabled, self.aria = text, visible, enabled, aria
            self.click = Mock()

        def is_displayed(self):
            return self.visible

        def is_enabled(self):
            return self.enabled

        def get_attribute(self, name):
            assert name == 'aria-disabled'
            return self.aria

    class Wait:
        def __init__(self, driver, timeout):
            assert timeout == 15
            self.driver = driver

        def until(self, condition):
            for _ in range(3):
                result = condition(self.driver)
                if result:
                    return result
            raise TimeoutError('SYNTHETIC-PRIVATE timeout details')

    def driver(self, buttons):
        from selenium.webdriver.common.by import By

        def find(by, selector):
            self.assertEqual((by, selector), (By.TAG_NAME, 'button'))
            self.assertTrue(all(button.click.call_count == 0 for button in buttons))
            return buttons

        return SimpleNamespace(find_elements=Mock(side_effect=find), execute_script=Mock())

    def test_role_and_order_independent_exact_rendered_label_only(self):
        for labels in (
                ['Save', 'Complete Task'],
                ['Delete Task', 'Save', 'Complete Task'],
                ['Complete Task', 'Delete Task', 'Do not Complete Task', 'Complete Tasks'],
                ['complete task', 'Complete Task now', 'Delete task'],
                ['Delete task', 'COMPLETE TASK'],
                ['Complete Task now', ' \n cOmPlEtE \t TaSk \n ', 'Delete Task']):
            with self.subTest(labels=labels):
                buttons = [self.Button(label) for label in labels]
                target, = [button for button in buttons
                           if ' '.join(button.text.split()).casefold() == 'complete task']
                driver = self.driver(buttons)
                with patch('AI.unity_completion.WebDriverWait', self.Wait):
                    complete_task(driver)
                target.click.assert_called_once_with()
                for button in buttons:
                    if button is not target:
                        button.click.assert_not_called()
                driver.execute_script.assert_called_once()
                self.assertNotIn('click', driver.execute_script.call_args.args[0].lower())

    def test_hidden_duplicates_ignored_and_wait_does_not_click(self):
        hidden = self.Button('COMPLETE TASK', visible=False)
        target = self.Button('Complete task', enabled=False)
        driver = self.driver([hidden, target, self.Button('Delete Task')])
        polls = []

        def poll(by, selector):
            polls.append(True)
            target.click.assert_not_called()
            if len(polls) == 2:
                target.enabled = True
            return [hidden, target]

        driver.find_elements.side_effect = poll
        with patch('AI.unity_completion.WebDriverWait', self.Wait):
            complete_task(driver)
        self.assertEqual(len(polls), 3)
        target.click.assert_called_once_with()
        hidden.click.assert_not_called()

    def test_missing_hidden_disabled_aria_disabled_and_duplicate_fail_closed(self):
        cases = [
            [], [self.Button('Delete task')], [self.Button('Do not complete task')],
            [self.Button('COMPLETE TASKS')], [self.Button('complete task now')],
            [self.Button('Complete Task', visible=False)],
            [self.Button('Complete Task', enabled=False)],
            [self.Button('Complete Task', aria=' TRUE ')],
            [self.Button('Complete Task'), self.Button('Complete Task')],
            [self.Button('Complete Task'), self.Button('Complete Task', enabled=False)],
            [self.Button('Complete task'), self.Button('COMPLETE TASK')],
            [self.Button('Complete task'), self.Button('COMPLETE TASK', enabled=False)],
        ]
        for buttons in cases:
            with self.subTest(buttons=[b.text for b in buttons]):
                driver = self.driver(buttons)
                with patch('AI.unity_completion.WebDriverWait', self.Wait):
                    with self.assertRaises(UnityTaskCompletionError) as caught:
                        complete_task(driver)
                self.assertEqual(caught.exception.stage,
                                 'waiting for a unique actionable Complete Task button')
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                self.assertIsInstance(caught.exception.__cause__, (RuntimeError, TimeoutError))
                driver.execute_script.assert_not_called()
                for button in buttons:
                    button.click.assert_not_called()

    def test_scroll_mutations_fail_before_click_and_never_choose_replacement(self):
        for mutation in ('relabel', 'replace', 'hide', 'disable', 'aria', 'duplicate', 'remove'):
            with self.subTest(mutation=mutation):
                target = self.Button('cOmPlEtE tAsK')
                buttons = [target, self.Button('Delete Task')]
                driver = self.driver(buttons)

                def scroll(*args):
                    if mutation == 'relabel':
                        target.text = 'Delete Task'
                    elif mutation == 'replace':
                        buttons[0] = self.Button('Complete Task')
                    elif mutation == 'hide':
                        target.visible = False
                    elif mutation == 'disable':
                        target.enabled = False
                    elif mutation == 'aria':
                        target.aria = 'true'
                    elif mutation == 'duplicate':
                        buttons.append(self.Button('COMPLETE TASK'))
                    elif mutation == 'remove':
                        buttons.remove(target)

                driver.execute_script.side_effect = scroll
                with patch('AI.unity_completion.WebDriverWait', self.Wait):
                    with self.assertRaises(UnityTaskCompletionError) as caught:
                        complete_task(driver)
                self.assertEqual(caught.exception.stage, 'revalidating the Complete Task button')
                target.click.assert_not_called()
                for button in buttons:
                    button.click.assert_not_called()

    def test_interaction_errors_keep_original_cause_and_click_is_never_retried(self):
        for failure in ('lookup', 'scroll', 'click'):
            with self.subTest(failure=failure):
                target = self.Button('Complete Task')
                driver = self.driver([target])
                error = RuntimeError('SYNTHETIC-PRIVATE driver response')
                operation = {'lookup': driver.find_elements, 'scroll': driver.execute_script,
                             'click': target.click}[failure]
                operation.side_effect = error
                with patch('AI.unity_completion.WebDriverWait', self.Wait):
                    with self.assertRaises(UnityTaskCompletionError) as caught:
                        complete_task(driver)
                self.assertIs(caught.exception.__cause__, error)
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                self.assertEqual(target.click.call_count, int(failure == 'click'))
                self.assertEqual(operation.call_count, 1)

    def test_actual_validation_desktop_controller_and_cli_stop_with_normal_cleanup(self):
        from contextlib import redirect_stdout
        from io import StringIO

        for record_on_failure in (True, False):
            with self.subTest(record_on_failure=record_on_failure):
                fixture = LegacyCustomerFieldsTests()
                fixture.setUp()
                fixture.complete.text = 'Delete Task'
                env = fixture.env
                env.update(create_driver=Mock(return_value=fixture.driver), login=Mock())
                exec(compile(ast.Module(body=[legacy_function('main')], type_ignores=[]),
                             '<actual Unity main>', 'exec'), env)
                unity = SimpleNamespace(**{name: env[name] for name in (
                    'main', 'create_driver', 'login', 'validate_and_complete_task')})
                config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py')
                with patch.dict(sys.modules, {
                        'fcm_intake.config': config,
                        'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=lambda *args: unity)}):
                    callbacks = backend.live_callbacks()
                record, confirm, finish = Mock(), Mock(return_value=True), Mock()
                with patch('AI.unity_completion.WebDriverWait', self.Wait):
                    result = backend._runner().load_core().run_flow(
                        fixture.data, list_attachments=lambda: ['synthetic.pdf'],
                        run_unity=callbacks['run_unity'], record_assignment=record,
                        confirm=confirm, complete_assignment=finish, notify=Mock(),
                        record_attempt_on_failure=record_on_failure)
                self.assertEqual((result['status'], result['reason']), ('failed', 'run_unity_failed'))
                self.assertEqual(result['error']['type'], 'UnityTaskCompletionError')
                self.assertNotIn('SYNTHETIC-PRIVATE', result['error']['message'])
                self.assertEqual(record.call_count, int(record_on_failure))
                confirm.assert_not_called()
                finish.assert_not_called()
                fixture.driver.quit.assert_called_once()
                fixture.driver.service.stop.assert_not_called()
                fixture.driver.command_executor.close.assert_not_called()
                fixture.env['notify'].assert_called_once_with(
                    'Notice', 'Stopped for Manual Validation before Completing Task.')
                fixture.env['time'].sleep.assert_any_call(2)
                self.assertNotIn(((15,), {}), fixture.env['time'].sleep.call_args_list)
                fixture.driver.quit.reset_mock()
                with patch('AI.unity_completion.WebDriverWait', self.Wait):
                    with self.assertRaises(UnityTaskCompletionError):
                        unity.main('Synthetic', 'Example', fixture.data)
                fixture.driver.quit.assert_called_once()

                # Exercise run.py's actual --live wiring with AST-extracted
                # legacy code and local synthetic inputs, never live bindings.
                runner = backend._runner()
                unity.main = lambda customer, claimant, data, **options: env['main'](
                    customer, claimant, fixture.data | data, **options)
                with tempfile.TemporaryDirectory(dir=HERE) as directory:
                    (Path(directory) / 'synthetic.pdf').touch()
                    config.ATTACHMENT_FOLDER = directory
                    fixture.driver.quit.reset_mock()
                    output = StringIO()
                    with patch.dict(sys.modules, {
                            'fcm_intake.config': config,
                            'fcm_intake.legacy_loader': SimpleNamespace(
                                load_module_from_path=lambda *args: unity)}), \
                            patch.object(runner, '_preflight_tracking'), \
                            patch.object(runner, '_record_live') as tracking, \
                            patch.object(runner, '_complete_live') as rrs, \
                            patch.object(runner, 'ask_yes_no') as approval, \
                            patch('AI.unity_completion.WebDriverWait', self.Wait), redirect_stdout(output):
                        code = runner.main(['--live', '--customer', 'Synthetic',
                            '--claim-number', 'SYNTHETIC-001', '--claimant-first', 'Synthetic',
                            '--claimant-last', 'Example', '--record-attempt-on-failure',
                            'yes' if record_on_failure else 'no'])
                    self.assertEqual(code, 2, output.getvalue())
                    self.assertIn('"reason": "run_unity_failed"', output.getvalue())
                    self.assertIn('"type": "UnityTaskCompletionError"', output.getvalue())
                    self.assertEqual(tracking.call_count, int(record_on_failure))
                    approval.assert_not_called()
                    rrs.assert_not_called()
                    fixture.driver.quit.assert_called_once()
                self.assertNotIn(('complete', 'click'), fixture.events)

    def test_actual_validation_click_failure_propagates_once_without_post_click_sleep(self):
        fixture = LegacyCustomerFieldsTests()
        fixture.setUp()
        error = RuntimeError('SYNTHETIC-PRIVATE uncertain click result')
        fixture.complete.click = Mock(side_effect=error)
        unity = SimpleNamespace(validate_and_complete_task=fixture.env['validate_and_complete_task'])
        with patch('AI.unity_completion.WebDriverWait', self.Wait):
            with self.assertRaises(UnityTaskCompletionError) as caught:
                backend.checked_validate(unity, fixture.driver, 'Synthetic', 'Example', fixture.data)
        self.assertIs(caught.exception.__cause__, error)
        self.assertEqual(caught.exception.stage, 'clicking the Complete Task button')
        self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
        fixture.complete.click.assert_called_once_with()
        self.assertNotIn(((15,), {}), fixture.env['time'].sleep.call_args_list)

    def test_actual_main_still_swallows_unrelated_legacy_errors_and_quits(self):
        from contextlib import redirect_stdout
        from io import StringIO

        driver = SimpleNamespace(quit=Mock())
        env = {'create_driver': Mock(return_value=driver), 'login': Mock(),
               'validate_and_complete_task': Mock(side_effect=ValueError('synthetic failure')),
               'UnityTaskCompletionError': UnityTaskCompletionError, 'sys': sys}
        exec(compile(ast.Module(body=[legacy_function('main')], type_ignores=[]),
                     '<actual Unity main>', 'exec'), env)
        output = StringIO()
        with redirect_stdout(output):
            self.assertIsNone(env['main']('Synthetic', 'Example', {}))
        self.assertIn('error occurred:', output.getvalue())
        driver.quit.assert_called_once_with()


class LegacyCustomerFieldsTests(unittest.TestCase):
    def setUp(self):
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support import expected_conditions as EC

        self.events = []
        self.failure = None
        self.original_error = None
        owner = self

        class Element:
            def __init__(self, name, value=''):
                self.name, self.value, self.reads = name, value, 0
                self.visible = name not in ('date', 'claim')
                self.text = 'Complete Task' if name == 'complete' else ''

            def is_displayed(self):
                self.action('visibility')
                return self.visible

            def is_enabled(self):
                return True

            def action(self, action):
                owner.events.append((self.name, action))
                if owner.failure == (self.name, action):
                    owner.original_error = ValueError('SYNTHETIC-PRIVATE Selenium/referral details')
                    raise owner.original_error

            def click(self):
                self.action('click')
                if self.name == 'group' and owner.failure != ('group', 'still-hidden'):
                    owner.date.visible = not owner.date.visible
                    owner.claim.visible = not owner.claim.visible

            def clear(self):
                self.action('clear')
                self.value = ''

            def send_keys(self, value):
                self.action('type')
                self.value += value

            def get_attribute(self, name):
                if name == 'aria-disabled':
                    return None
                self.reads += 1
                self.action('read' if self.reads == 1 else 'readback')
                if owner.failure == (self.name, 'mismatch') and self.reads > 1:
                    return 'SYNTHETIC-PRIVATE wrong value'
                if owner.failure == (self.name, 'delayed') and self.reads == 2:
                    return ''
                return self.value

        class Wait:
            def __init__(self, target, timeout):
                self.target = target

            def until(self, condition):
                for _ in range(3):
                    result = condition(self.target)
                    if result:
                        return result
                owner.original_error = TimeoutError('SYNTHETIC-PRIVATE wait')
                raise owner.original_error

        self.group, self.date, self.claim, self.complete = [
            Element(name) for name in ('group', 'date', 'claim', 'complete')]
        locators = {
            (By.XPATH, "//button[contains(@class,'btn') and contains(@class,'btn-link') and "
             "contains(normalize-space(.), 'Liberty Mutual Commercial Market') and "
             "contains(normalize-space(.), '20 fields')]"): self.group,
            (By.XPATH, "//field-label[normalize-space()='Date of First Appointment']"
             "/following::input[@type='text'][1]"): self.date,
            (By.ID, 'intake-customer-specific-field-fields--Claim Identifier-text-general-input'): self.claim,
        }

        def find(by, selector):
            field = locators.get((by, selector))
            if field is None:
                self.events.append(('unexpected-lookup', selector))
                raise AssertionError('Unexpected later portal work')
            field.action('locate')
            return field

        def find_all(by, selector):
            if (by, selector) == (By.TAG_NAME, 'button'):
                return [self.complete]
            field = find(by, selector)
            return [field, field] if self.failure == (field.name, 'ambiguous') else [field]

        self.driver = SimpleNamespace(find_element=find, find_elements=find_all, execute_script=Mock(),
                                      service=Mock(), command_executor=Mock(), quit=Mock())
        self.env = {'WebDriverWait': Wait, 'By': By, 'EC': EC, 'sys': sys,
                    'time': Mock(), 'notify': Mock(), 'UnityFlagError': UnityFlagError,
                    'UnityCustomerFieldsError': UnityCustomerFieldsError,
                    'UnityEmployerFieldsError': UnityEmployerFieldsError,
                    'UnityTaskCompletionError': UnityTaskCompletionError,
                    'complete_task': complete_task}
        helper = legacy_function('fill_customer_defined_fields')
        validation = legacy_function('validate_and_complete_task')
        statements = validation.body[0].body
        index, = [i for i, node in enumerate(statements)
                  if isinstance(node, ast.Expr) and isinstance(node.value, ast.Call)
                  and isinstance(node.value.func, ast.Name)
                  and node.value.func.id == 'fill_customer_defined_fields']
        # Execute the actual customer call, employer branch, manual gate,
        # Complete Task block and outer handlers, excluding earlier portal work.
        validation.body[0].body = statements[index:]
        exec(compile(ast.fix_missing_locations(ast.Module(
            body=[helper, validation], type_ignores=[])), '<actual customer fields>', 'exec'), self.env)
        self.data = model.to_controller(synthetic() | {
            'nextApptDate': '10/23/2026', 'claimID': 'SYNTHETIC-ID'})

    def test_actual_helper_fills_blanks_preserves_populated_and_skips_missing_sources(self):
        for existing in ('', ' \t ', 'Already populated', '  Already populated \t '):
            with self.subTest(existing=existing):
                self.setUp()
                self.date.value = self.claim.value = existing
                self.env['validate_and_complete_task'](
                    self.driver, 'Synthetic', 'Example', self.data)
                for field, key in ((self.date, 'nextApptDate'), (self.claim, 'claimID')):
                    expected = existing if existing.strip() else self.data[key]
                    self.assertEqual(field.value, expected)
                    self.assertEqual(field.reads, 1 if existing.strip() else 2)
                    self.assertEqual(self.events.count((field.name, 'type')), int(not existing.strip()))
                self.env['notify'].assert_called_once_with(
                    'Notice', 'Stopped for Manual Validation before Completing Task.')
                self.assertEqual(self.events.count(('complete', 'click')), 1)
                self.env['time'].sleep.assert_any_call(15)
        for source in ({}, {'nextApptDate': '', 'claimID': None},
                       {'nextApptDate': ' \t ', 'claimID': ' \t '}):
            with self.subTest(source=source):
                self.setUp()
                self.date.value = self.claim.value = ' \t '
                self.env['fill_customer_defined_fields'](
                    self.driver, next_appointment_date=source.get('nextApptDate'),
                    claim_identifier=source.get('claimID'))
                self.assertEqual((self.date.value, self.claim.value), (' \t ', ' \t '))
                self.assertFalse(any(action in ('clear', 'type') for _, action in self.events))
        for delayed in ('date', 'claim'):
            with self.subTest(delayed_readback=delayed):
                self.setUp()
                self.failure = (delayed, 'delayed')
                self.env['fill_customer_defined_fields'](
                    self.driver, next_appointment_date='10/23/2026', claim_identifier='SYNTHETIC-ID')
                self.assertEqual(self.date.value, '10/23/2026')
                self.assertEqual(self.claim.value, 'SYNTHETIC-ID')
                field = self.date if delayed == 'date' else self.claim
                self.assertEqual(field.reads, 3)
                self.assertEqual(self.events.count((delayed, 'type')), 1)
                self.assertEqual(self.events.count(('group', 'click')), 1)

    def test_actual_helper_toggles_only_collapsed_group_and_requires_visible_unambiguous_fields(self):
        for expanded in (False, True):
            with self.subTest(expanded=expanded):
                self.setUp()
                self.date.visible = self.claim.visible = expanded
                self.date.value = self.claim.value = 'Already populated'
                self.env['fill_customer_defined_fields'](
                    self.driver, next_appointment_date='10/23/2026', claim_identifier='SYNTHETIC-ID')
                self.assertEqual(self.events.count(('group', 'click')), int(not expanded))
                self.assertTrue(self.date.visible and self.claim.visible)
                self.assertEqual((self.date.value, self.claim.value), ('Already populated', 'Already populated'))
                self.assertFalse(any(action in ('clear', 'type') for _, action in self.events))
        cases = [
            (('group', 'still-hidden'), 'waiting for customer fields to become visible', 1),
            (('date', 'visibility'), 'inspecting customer field visibility', 0),
            (('date', 'ambiguous'), 'locating Date of First Appointment', 0),
            (('claim', 'ambiguous'), 'locating Claim Identifier', 0),
            (('group', 'mixed'), 'inspecting customer field visibility', 0),
        ]
        for failure, stage, clicks in cases:
            with self.subTest(failure=failure):
                self.setUp()
                self.failure = failure
                if failure == ('group', 'mixed'):
                    self.date.visible = True
                with self.assertRaises(UnityCustomerFieldsError) as caught:
                    self.env['validate_and_complete_task'](
                        self.driver, 'Synthetic', 'Example', self.data)
                self.assertEqual(caught.exception.stage, stage)
                if self.original_error is not None:
                    self.assertIs(caught.exception.__cause__, self.original_error)
                else:
                    self.assertIsInstance(caught.exception.__cause__, RuntimeError)
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                self.assertEqual(self.events.count(('group', 'click')), clicks)
                self.assertFalse(any(action in ('clear', 'type') for _, action in self.events))
                self.env['notify'].assert_not_called()
                self.assertNotIn(('complete', 'click'), self.events)

    def test_actual_helper_errors_are_safe_chained_and_stop_later_source_work(self):
        stages = {'locate': 'locating', 'read': 'reading', 'click': 'focusing',
                  'clear': 'clearing', 'type': 'entering',
                  'readback': 'verifying', 'mismatch': 'verifying'}
        cases = [(('group', 'locate'), 'locating the customer field group'),
                 (('group', 'click'), 'opening the customer field group')]
        cases += [((field, action), f'{verb} {label}')
                  for field, label in (('date', 'Date of First Appointment'), ('claim', 'Claim Identifier'))
                  for action, verb in stages.items()]
        for failure, stage in cases:
            with self.subTest(failure=failure):
                self.setUp()
                self.failure = failure
                self.data['customer'] = 'GOODYEAR TIRE & RUBBER CO'
                unity = SimpleNamespace(validate_and_complete_task=self.env['validate_and_complete_task'])
                with self.assertRaises(UnityCustomerFieldsError) as caught:
                    backend.checked_validate(unity, self.driver, 'Synthetic', 'Example', self.data)
                self.assertEqual(caught.exception.stage, stage)
                self.assertIn(stage, str(caught.exception))
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                self.assertIs(caught.exception.__cause__, self.original_error)
                self.env['notify'].assert_not_called()
                self.assertFalse(any(name in ('unexpected-lookup', 'complete') for name, _ in self.events))
                if failure[0] != 'claim':
                    self.assertNotIn(('claim', 'read'), self.events)

    def test_actual_customer_failure_desktop_recovery_stops_gate_rrs_and_records_cleanup_failures(self):
        for release_failure in (None, 'service', 'connection', 'both'):
            with self.subTest(release_failure=release_failure):
                self.setUp()
                self.failure = ('claim', 'readback')
                if release_failure in ('service', 'both'):
                    self.driver.service.stop.side_effect = RuntimeError('SYNTHETIC-PRIVATE service')
                if release_failure in ('connection', 'both'):
                    self.driver.command_executor.close.side_effect = RuntimeError('SYNTHETIC-PRIVATE connection')
                raised = []

                def validate(*args):
                    try:
                        self.env['validate_and_complete_task'](*args)
                    except UnityCustomerFieldsError as error:
                        raised.append(error)
                        raise

                unity = SimpleNamespace(create_driver=Mock(return_value=self.driver), login=Mock(),
                                        validate_and_complete_task=validate)
                config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py')
                with patch.dict(sys.modules, {
                        'fcm_intake.config': config,
                        'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=lambda *args: unity)}):
                    calls = backend.live_callbacks()
                record, confirm, complete = Mock(), Mock(return_value=True), Mock()
                result = backend._runner().load_core().run_flow(
                    self.data, list_attachments=lambda: [Path('synthetic.pdf')],
                    run_unity=calls['run_unity'], record_assignment=record,
                    confirm=confirm, complete_assignment=complete, notify=Mock())
                self.assertEqual((result['status'], result['reason']), ('failed', 'run_unity_failed'))
                self.assertEqual(result['error']['type'], 'UnityCustomerFieldsError')
                self.assertIn('verifying Claim Identifier', result['error']['message'])
                self.assertIn('Browser left open for manual recovery', result['error']['message'])
                self.assertNotIn('SYNTHETIC-PRIVATE', result['error']['message'])
                self.assertEqual(len(raised), 1)
                self.assertIs(raised[0].__cause__, self.original_error)
                self.assertTrue(raised[0].browser_left_open)
                self.assertIs(raised[0].service_release_failed, release_failure is not None)
                self.assertEqual('release could not be confirmed' in result['error']['message'],
                                 release_failure is not None)
                unity.create_driver.assert_called_once_with(detach=True)
                self.driver.service.stop.assert_called_once()
                self.driver.command_executor.close.assert_called_once()
                self.driver.quit.assert_not_called()
                self.env['notify'].assert_not_called()
                self.assertNotIn(('complete', 'click'), self.events)
                record.assert_called_once()
                confirm.assert_not_called()
                complete.assert_not_called()


class LegacyEmployerFieldsTests(unittest.TestCase):
    def setUp(self):
        LegacyCustomerFieldsTests.setUp(self)
        owner = self
        self.date.value = self.claim.value = 'Already populated'
        self.data.update(customer='GOODYEAR TIRE & RUBBER CO',
                         specialInstructions='', providerCity='Lawton')
        self.selected_index = 0
        self.selection_reads = 0
        self.written = False

        class Element:
            def __init__(self, name, text='', value='', enabled=True):
                self.name, self.text, self.value = name, text, value
                self.enabled = enabled
                self.visible = name != 'facility'

            def action(self, action):
                owner.events.append((self.name, action))
                if owner.failure == (self.name, action):
                    owner.original_error = ValueError('SYNTHETIC-PRIVATE employer details')
                    raise owner.original_error

            def click(self):
                self.action('click')
                if owner.failure != ('employer', 'still-hidden'):
                    owner.facility.visible = not owner.facility.visible

            def is_displayed(self):
                self.action('visibility')
                return self.visible

            def is_enabled(self):
                self.action('enabled')
                return self.enabled

            def get_attribute(self, name):
                self.action('read')
                return self.value

        self.employer, self.facility = Element('employer'), Element('facility')
        self.options = [Element('placeholder', 'Select...', '0: null'),
                        Element('lawton', 'PL0153 Lawton', '1: facility-a'),
                        Element('topeka', 'PL0133 Topeka', '2: facility-b')]

        class Select:
            def __init__(self, field):
                self.field = field

            @property
            def first_selected_option(self):
                self.field.action('readback' if owner.written else 'read')
                owner.selection_reads += 1
                if owner.written:
                    if owner.failure == ('facility', 'mismatch'):
                        return owner.options[0]
                    if owner.failure == ('facility', 'wrong-value'):
                        return Element('wrong', owner.options[1].text, 'SYNTHETIC-wrong-value')
                    if owner.failure == ('facility', 'wrong-text'):
                        return Element('wrong', 'SYNTHETIC-wrong-text', owner.options[1].value)
                    if owner.failure == ('facility', 'delayed') and owner.selection_reads == 2:
                        return owner.options[0]
                return owner.options[owner.selected_index]

            @property
            def options(self):
                self.field.action('options')
                return owner.options

            def select_by_visible_text(self, text):
                self.field.action('select')
                owner.written = True
                owner.selected_index, = [i for i, option in enumerate(owner.options) if option.text == text]

        original_find = self.driver.find_element
        By = self.env['By']

        def find(by, selector):
            if 'GOODYEAR TIRE & RUBBER CO' in selector:
                self.employer.action('locate')
                return self.employer
            if selector == 'intake-customer-specific-field-fields--Facility-select-input':
                self.facility.action('locate')
                return self.facility
            return original_find(by, selector)

        def find_all(by, selector):
            if (by, selector) == (By.TAG_NAME, 'button'):
                return [self.complete]
            field = find(by, selector)
            if field is self.facility and self.failure == ('facility', 'ambiguous'):
                return [field, field]
            return [field]

        self.driver.find_element, self.driver.find_elements = find, find_all
        self.env.update(Select=Select, re=re)
        exec(compile(ast.fix_missing_locations(ast.Module(
            body=[legacy_function('fill_employer_defined_fields')], type_ignores=[])),
            '<actual employer fields>', 'exec'), self.env)

    def helper(self, instructions='', city='Lawton'):
        return self.env['fill_employer_defined_fields'](
            self.driver, special_instructions=instructions, provider_city=city)

    def test_actual_helper_opens_once_only_when_hidden_and_preserves_populated(self):
        for visible in (True, False):
            for populated in (True, False):
                with self.subTest(visible=visible, populated=populated):
                    self.setUp()
                    self.facility.visible = visible
                    self.selected_index = 2 if populated else 0
                    result = self.helper()
                    self.assertEqual(result, 'PL0133 Topeka' if populated else 'PL0153 Lawton')
                    self.assertEqual(self.events.count(('employer', 'click')), int(not visible))
                    self.assertEqual(self.events.count(('facility', 'select')), int(not populated))
                    if populated:
                        self.assertNotIn(('facility', 'options'), self.events)

    def test_actual_header_xpath_matches_mixed_case_but_keeps_group_scope_and_uniqueness(self):
        from xml.etree import ElementTree
        from win32com.client import Dispatch

        cases = [
            ('Goodyear Tire & Rubber Co', '4 fields', 1, True, True),
            ('goodyear  tire &\n rubber co', '4 fields', 1, True, True),
            ('GOODYEAR TIRE & RUBBER CO', '4 fields', 1, True, True),
            ('Other Synthetic Employer', '4 fields', 1, True, False),
            ('Goodyear Tire & Rubber Company', '3 fields', 1, True, False),
            ('Goodyear Tire & Rubber Co', '3 fields', 1, True, False),
            ('Goodyear Tire & Rubber Co', '14 fields', 1, True, False),
            ('Goodyear Tire & Rubber Co', '24 fields', 1, True, False),
            ('Goodyear Tire & Rubber Co', '4 fields', 2, True, False),
            ('Goodyear Tire & Rubber Co', '4 fields', 2, False, True),
        ]
        for name, count, number, second_visible, succeeds in cases:
            with self.subTest(name=name, count=count, number=number, second_visible=second_visible):
                self.setUp()
                root = ElementTree.Element('div')
                headers = []
                for index in range(number):
                    node = ElementTree.SubElement(root, 'button', attrib={
                        'class': 'btn btn-link', 'data-index': str(index)})
                    node.text = f' \n {name} \t {count} '
                    header = self.employer if index == 0 else deepcopy(self.employer)
                    header.visible = index == 0 or second_visible
                    headers.append(header)
                document = Dispatch('Msxml2.DOMDocument.6.0')
                self.assertTrue(document.loadXML(ElementTree.tostring(root, encoding='unicode')))
                original_find_all = self.driver.find_elements

                def find_all(by, selector):
                    if 'GOODYEAR TIRE & RUBBER CO' not in selector:
                        return original_find_all(by, selector)
                    self.employer.action('locate')
                    return [headers[int(node.getAttribute('data-index'))]
                            for node in document.selectNodes(selector)]

                self.driver.find_elements = find_all
                if succeeds:
                    self.assertEqual(self.helper(), 'PL0153 Lawton')
                    self.assertEqual(self.events.count(('employer', 'click')), 1)
                else:
                    with self.assertRaises(UnityEmployerFieldsError) as caught:
                        self.helper()
                    self.assertEqual(caught.exception.stage, 'locating the employer field group')
                    self.assertNotIn(('employer', 'click'), self.events)
                    self.assertNotIn(('facility', 'select'), self.events)
        for state in ('visible', 'enabled'):
            with self.subTest(unavailable_header_state=state):
                self.setUp()
                setattr(self.employer, state, False)
                with self.assertRaises(UnityEmployerFieldsError) as caught:
                    self.helper()
                self.assertEqual(caught.exception.stage, 'locating the employer field group')
                self.assertNotIn(('employer', 'click'), self.events)
                self.assertNotIn(('facility', 'select'), self.events)

    def test_actual_helper_unique_matching_and_instructions_precede_city(self):
        for instructions, city, expected in (
                ('', 'lawton', 'PL0153 Lawton'),
                ('Use PL0133 Topeka', 'Lawton', 'PL0133 Topeka'),
                ('Unrecognized synthetic instruction', 'Lawton', 'PL0153 Lawton'),
                ('Use PL0153 Lawton', '', 'PL0153 Lawton')):
            with self.subTest(instructions=instructions, city=city):
                self.setUp()
                self.assertEqual(self.helper(instructions, city), expected)
                self.assertEqual(self.events.count(('facility', 'select')), 1)

    def test_actual_helper_blank_and_angular_placeholders_are_not_populated_facilities(self):
        for text, value in (('', ''), ('Select...', ''), ('Select...', 'null'),
                            ('Select...', 'undefined'), ('Select...', '0: null')):
            with self.subTest(text=text, value=value):
                self.setUp()
                self.options[0].text, self.options[0].value = text, value
                self.assertEqual(self.helper(), 'PL0153 Lawton')
                self.assertEqual(self.events.count(('facility', 'select')), 1)
        self.setUp()
        self.options[0].value = None
        with self.assertRaises(UnityEmployerFieldsError) as caught:
            self.helper()
        self.assertEqual(caught.exception.stage, 'reading Facility')
        self.assertIsInstance(caught.exception.__cause__, RuntimeError)
        self.assertNotIn(('facility', 'select'), self.events)

    def test_actual_helper_no_or_ambiguous_matches_leave_unchanged_for_manual_gate(self):
        cases = ('blank-city', 'whitespace-city', 'no-match', 'two-tokens',
                 'recognized-no-option', 'ambiguous-city', 'ambiguous-token',
                 'disabled-option', 'blank-option', 'null-option', 'angular-null-option')
        for case in cases:
            with self.subTest(case=case):
                self.setUp()
                if case == 'blank-city':
                    self.data['providerCity'] = ''
                elif case == 'whitespace-city':
                    self.data['providerCity'] = ' \t '
                elif case == 'no-match':
                    self.data['providerCity'] = 'Invented City'
                elif case == 'two-tokens':
                    self.data['specialInstructions'] = 'PL0153 Lawton or PL0133 Topeka'
                elif case == 'recognized-no-option':
                    self.data['specialInstructions'] = 'Use PL0126 Danville'
                elif case in ('ambiguous-city', 'ambiguous-token'):
                    duplicate = deepcopy(self.options[1])
                    duplicate.text += ' Annex'
                    self.options.append(duplicate)
                    if case == 'ambiguous-token':
                        self.data['specialInstructions'] = 'Use PL0153 Lawton'
                elif case == 'disabled-option':
                    self.options[1].enabled = False
                else:
                    self.options[1].value = {
                        'blank-option': '', 'null-option': 'null',
                        'angular-null-option': '7: null'}[case]
                self.env['validate_and_complete_task'](
                    self.driver, 'Synthetic', 'Example', self.data)
                self.assertEqual(self.selected_index, 0)
                self.assertNotIn(('facility', 'select'), self.events)
                self.assertEqual(self.env['notify'].call_args_list[0].args, (
                    'Goodyear Employer Defined Field - Facility', 'No matching Facility from Dropdown.'))
                self.assertEqual(self.env['notify'].call_args_list[1].args, (
                    'Notice', 'Stopped for Manual Validation before Completing Task.'))
                self.assertEqual(self.events.count(('complete', 'click')), 1)

    def test_actual_helper_polling_is_read_only_and_requires_exact_value_and_text(self):
        self.failure = ('facility', 'delayed')
        self.assertEqual(self.helper(), 'PL0153 Lawton')
        self.assertEqual(self.selection_reads, 3)
        self.assertEqual(self.events.count(('facility', 'select')), 1)
        for failure in ('mismatch', 'wrong-value', 'wrong-text'):
            with self.subTest(failure=failure):
                self.setUp()
                self.failure = ('facility', failure)
                with self.assertRaises(UnityEmployerFieldsError) as caught:
                    self.helper()
                self.assertEqual(caught.exception.stage, 'verifying Facility')
                self.assertIs(caught.exception.__cause__, self.original_error)
                self.assertEqual(self.events.count(('facility', 'select')), 1)

    def test_actual_helper_failures_are_typed_chained_safe_and_stop_later_work(self):
        cases = [
            (('employer', 'locate'), 'locating the employer field group'),
            (('employer', 'visibility'), 'locating the employer field group'),
            (('employer', 'enabled'), 'locating the employer field group'),
            (('facility', 'locate'), 'locating Facility'),
            (('facility', 'ambiguous'), 'locating Facility'),
            (('facility', 'visibility'), 'inspecting employer field visibility'),
            (('employer', 'click'), 'opening the employer field group'),
            (('employer', 'still-hidden'), 'waiting for Facility to become visible'),
            (('facility', 'read'), 'reading Facility'),
            (('placeholder', 'read'), 'reading Facility'),
            (('facility', 'options'), 'matching Facility'),
            (('lawton', 'read'), 'matching Facility'),
            (('lawton', 'enabled'), 'matching Facility'),
            (('facility', 'enabled'), 'selecting Facility'),
            (('facility', 'select'), 'selecting Facility'),
            (('facility', 'readback'), 'verifying Facility'),
        ]
        for failure, stage in cases:
            with self.subTest(failure=failure):
                self.setUp()
                self.failure = failure
                unity = SimpleNamespace(validate_and_complete_task=self.env['validate_and_complete_task'])
                with self.assertRaises(UnityEmployerFieldsError) as caught:
                    backend.checked_validate(unity, self.driver, 'Synthetic', 'Example', self.data)
                self.assertEqual(caught.exception.stage, stage)
                if self.original_error is None:
                    self.assertIsInstance(caught.exception.__cause__, RuntimeError)
                else:
                    self.assertIs(caught.exception.__cause__, self.original_error)
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                self.env['notify'].assert_not_called()
                self.assertNotIn(('complete', 'click'), self.events)

    def test_actual_employer_failure_desktop_recovery_stops_approval_and_rrs(self):
        for release_failure in (None, 'service', 'connection', 'both'):
            with self.subTest(release_failure=release_failure):
                self.setUp()
                self.failure = ('facility', 'readback')
                if release_failure in ('service', 'both'):
                    self.driver.service.stop.side_effect = RuntimeError('SYNTHETIC-PRIVATE service')
                if release_failure in ('connection', 'both'):
                    self.driver.command_executor.close.side_effect = RuntimeError('SYNTHETIC-PRIVATE connection')
                raised = []

                def validate(*args):
                    try:
                        self.env['validate_and_complete_task'](*args)
                    except UnityEmployerFieldsError as error:
                        raised.append(error)
                        raise

                unity = SimpleNamespace(create_driver=Mock(return_value=self.driver), login=Mock(),
                                        validate_and_complete_task=validate)
                config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py')
                with patch.dict(sys.modules, {
                        'fcm_intake.config': config,
                        'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=lambda *args: unity)}):
                    calls = backend.live_callbacks()
                record, confirm, complete = Mock(), Mock(return_value=True), Mock()
                result = backend._runner().load_core().run_flow(
                    self.data, list_attachments=lambda: [Path('synthetic.pdf')],
                    run_unity=calls['run_unity'], record_assignment=record,
                    confirm=confirm, complete_assignment=complete, notify=Mock())
                self.assertEqual((result['status'], result['reason']), ('failed', 'run_unity_failed'))
                self.assertEqual(result['error']['type'], 'UnityEmployerFieldsError')
                self.assertIn('verifying Facility', result['error']['message'])
                self.assertIn('Browser left open for manual recovery', result['error']['message'])
                self.assertNotIn('SYNTHETIC-PRIVATE', result['error']['message'])
                self.assertEqual(len(raised), 1)
                self.assertIs(raised[0].__cause__, self.original_error)
                self.assertTrue(raised[0].browser_left_open)
                self.assertIs(raised[0].service_release_failed, release_failure is not None)
                self.assertEqual('release could not be confirmed' in result['error']['message'],
                                 release_failure is not None)
                unity.create_driver.assert_called_once_with(detach=True)
                self.driver.service.stop.assert_called_once()
                self.driver.command_executor.close.assert_called_once()
                self.driver.quit.assert_not_called()
                self.env['notify'].assert_not_called()
                self.assertNotIn(('complete', 'click'), self.events)
                record.assert_called_once()
                confirm.assert_not_called()
                complete.assert_not_called()


class LegacyProviderAddressTests(unittest.TestCase):
    def setUp(self):
        from selenium.webdriver.common.by import By
        from selenium.webdriver.remote.webelement import WebElement
        from selenium.webdriver.support import expected_conditions as EC

        self.events = []
        self.failure = None
        self.open = False
        self.closing = False
        self.closure_polls = 0
        owner = self

        class Element(WebElement):
            def __init__(self, name, classes='', visible=True, enabled=True):
                self.name = name
                self.classes = classes
                self.visible = visible
                self.enabled = enabled
                self.children = []
                self.synthetic_text = 'Invalid address'

            @property
            def text(self):
                return self.synthetic_text

            def find_element(self, by, selector):
                if self.name == 'section':
                    return owner.provider
                if self.name == 'provider' and selector == ".//button[normalize-space()='Validate Address']":
                    return owner.validate
                raise AssertionError(f'Unexpected single lookup: {self.name}, {selector}')

            def find_elements(self, by, selector):
                if self.name == 'driver':
                    owner.events.append(('modal-lookup',))
                    if owner.failure == 'wait':
                        raise RuntimeError('synthetic wait failure')
                    if by != By.CSS_SELECTOR:
                        raise AssertionError('Expected token-based modal selector')
                    selectors = [part.strip().split('.') for part in selector.split(',')]
                    return [element for element in self.children
                            if any(parts[0] == 'div' and set(parts[1:]) <= set(element.classes.split())
                                   for parts in selectors)]
                if by == By.XPATH and selector == ".//button[normalize-space()='Use']":
                    return self.children
                raise AssertionError(f'Unexpected child lookup: {self.name}, {selector}')

            def is_displayed(self):
                if self.name == 'modal':
                    if owner.closing:
                        owner.closure_polls += 1
                        if owner.closure_polls >= 2 and owner.failure != 'still-open':
                            owner.open = False
                    owner.events.append(('modal-visible', owner.open))
                    return owner.open
                return self.visible

            def is_enabled(self):
                return self.enabled

            def click(self):
                owner.events.append(('click', self.name))
                if owner.failure == self.name:
                    raise RuntimeError('synthetic click failure')
                if self.name == 'validate':
                    owner.open = owner.failure != 'modal-timeout'
                elif self.name == 'use':
                    owner.closing = True

        class Wait:
            def __init__(self, target, timeout):
                self.target = target

            def until(self, condition):
                for _ in range(3):
                    value = condition(self.target)
                    if value:
                        return value
                raise TimeoutError('synthetic address timeout')

        self.section = Element('section')
        self.provider = Element('provider')
        self.validate = Element('validate')
        self.modal = Element('modal', classes='modal fade show')
        self.use = Element('use')
        self.modal.children = [self.use]
        self.driver = Element('driver')
        self.driver.children = [self.modal]
        self.driver.execute_script = Mock()
        self.element = Element
        validation = legacy_function('validate_and_complete_task')
        helper, = [node for node in ast.walk(validation)
                   if isinstance(node, ast.FunctionDef) and node.name == 'ValidateAddrPrv']
        env = {'driver': self.driver, 'By': By, 'EC': EC, 'WebDriverWait': Wait}
        exec(compile(ast.Module(body=[helper], type_ignores=[]), '<actual provider address helper>', 'exec'), env)
        self.run_helper = env['ValidateAddrPrv']

    def test_current_and_legacy_modal_tokens_accept_only_after_closure(self):
        for classes in ('modal fade show', 'show extra modal fade', 'modal fade in', 'in fade modal'):
            with self.subTest(classes=classes):
                self.setUp()
                self.modal.classes = classes
                self.assertIs(self.run_helper(self.section), True)
                self.assertEqual([event for event in self.events if event[0] == 'click'],
                                 [('click', 'validate'), ('click', 'use')])
                self.assertEqual(self.events[-1], ('modal-visible', False))
                self.assertEqual(self.closure_polls, 2)

    def test_hidden_and_unrelated_use_controls_are_not_clicked(self):
        hidden = self.element('hidden-modal', classes='modal fade show', visible=False)
        hidden.children = [self.element('hidden-modal-use')]
        hidden_button = self.element('hidden-use', visible=False)
        disabled_button = self.element('disabled-use', enabled=False)
        unrelated = self.element('unrelated', classes='modal-dialog fade show')
        unrelated.children = [self.element('unrelated-use')]
        self.driver.children = [hidden, unrelated, self.element('outside-use'), self.modal]
        self.modal.children = [hidden_button, disabled_button, self.use]
        self.assertIs(self.run_helper(self.section), True)
        self.assertEqual([event for event in self.events if event[0] == 'click'],
                         [('click', 'validate'), ('click', 'use')])

    def test_missing_hidden_disabled_or_ambiguous_confirmation_returns_false(self):
        for case in ('hidden-modal', 'hidden-use', 'disabled-use', 'wrong-class', 'ambiguous'):
            with self.subTest(case=case):
                self.setUp()
                if case == 'hidden-modal':
                    self.modal = self.element('hidden-modal', classes='modal fade show', visible=False)
                    self.modal.children = [self.use]
                    self.driver.children = [self.modal]
                elif case == 'hidden-use':
                    self.use.visible = False
                elif case == 'disabled-use':
                    self.use.enabled = False
                elif case == 'wrong-class':
                    self.modal.classes = 'modal fade showing'
                else:
                    self.modal.children.append(self.element('other-use'))
                self.assertIs(self.run_helper(self.section), False)
                self.assertEqual([event for event in self.events if event[0] == 'click'],
                                 [('click', 'validate')])

    def test_click_wait_and_closure_failures_never_report_acceptance(self):
        for failure in ('validate', 'modal-timeout', 'wait', 'use', 'still-open'):
            with self.subTest(failure=failure):
                self.setUp()
                self.failure = failure
                self.assertIs(self.run_helper(self.section), False)
                self.assertEqual(self.events.count(('click', 'validate')), 1)
                self.assertEqual(self.events.count(('click', 'use')),
                                 int(failure in ('use', 'still-open')))

    def test_already_validated_provider_preserves_existing_fast_path(self):
        self.provider.synthetic_text = 'Validated'
        self.assertIs(self.run_helper(self.section), True)
        self.assertEqual(self.events, [])


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=HERE)
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'synthetic.xlsx'
        excel.create_template(self.path)

    def fill(self, data, row=2):
        book = load_workbook(self.path)
        sheet = book[excel.SHEET]
        for index, (key, _, _) in enumerate(excel.COLUMNS, 1):
            if key in data:
                sheet.cell(row, index, data[key])
        book.save(self.path)
        book.close()

    def test_template_styles_and_no_sample(self):
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        sheet = book[excel.SHEET]
        self.assertEqual(sheet.max_column, 35)
        self.assertEqual(sheet.freeze_panes, 'C2')
        self.assertTrue(sheet.auto_filter.ref)
        self.assertEqual(sheet['A1'].fill.fgColor.rgb, '002563EB')
        self.assertTrue(sheet['A1'].comment)
        self.assertEqual(sheet['A2'].number_format, '@')
        self.assertFalse(any(cell.value for row in sheet.iter_rows(min_row=2) for cell in row))
        self.assertEqual(len(sheet.data_validations.dataValidation), 5)
        index = next(i for i, (key, _, _) in enumerate(excel.COLUMNS, 1) if key == 'specialInstructions')
        self.assertIn('32,767 characters', sheet.cell(1, index).comment.text)
        notes = '\n'.join(str(row[0].value) for row in book['Instructions'])
        self.assertIn('Special instructions: up to 32,767 characters', notes)
        self.assertIn('All other fields: up to 1,000 characters', notes)

    def test_synthetic_roundtrip_and_row_numbers(self):
        self.fill(synthetic() | {'zip': '00123', 'dob': '2000-01-02'})
        self.fill({'claimNumber': 'SYNTHETIC-002'}, 4)
        rows = excel.read_workbook(self.path)
        self.assertEqual([r['row'] for r in rows], [2, 4])
        self.assertEqual(rows[0]['data']['zip'], '00123')
        self.assertIn('customer', rows[1]['warning'])

    def test_maximum_multiline_instructions_roundtrip_and_bridge_selection(self):
        instructions = synthetic_instructions(32767)
        self.assertEqual(len(instructions), 32767)
        self.fill(synthetic() | {'specialInstructions': instructions, 'claimID': ' ID '})
        self.fill(synthetic() | {'claimNumber': 'SYNTHETIC-004', 'specialInstructions': ' \n\t '}, 4)
        rows = excel.read_workbook(self.path)
        self.assertEqual([row['row'] for row in rows], [2, 4])
        self.assertEqual(rows[0]['data']['specialInstructions'], instructions)
        self.assertEqual(rows[0]['data']['claimID'], 'ID')
        self.assertEqual(rows[1]['data']['specialInstructions'], ' \n\t ')
        api = backend.UnityAPI(lambda: self.fail('Loading must not create controller callbacks.'))
        with patch.object(api, '_dialog', return_value='OPEN'):
            pdf = Path(self.temp.name) / 'synthetic.pdf'
            pdf.write_bytes(b'%PDF-1.4\n')
            api._window = Mock()
            api._window.create_file_dialog.return_value = [str(pdf)]
            self.assertTrue(api.choose_referral_document()['ok'])
        self.assertTrue(api.load_excel(self.path)['ok'])
        selected = api.select_excel_row(2)['data']['data']
        self.assertEqual(selected['specialInstructions'], instructions)
        checked = api.validate_fields(selected)
        self.assertTrue(checked['ok'], checked)
        self.assertEqual(checked['data']['fields']['specialInstructions'], instructions)
        self.assertEqual(api.select_excel_row(4)['data']['data']['specialInstructions'], ' \n\t ')
        self.assertEqual(api.get_state()['attachment'], {'source': 'local', 'filename': pdf.name})
        self.assertIsNone(api._thread)

    def test_ordinary_oversized_field_loads_for_correction_without_losing_instructions(self):
        instructions = synthetic_instructions(1001)
        self.fill(synthetic() | {'claimID': 'x' * 1001, 'specialInstructions': instructions})
        record = excel.read_workbook(self.path)[0]
        self.assertIn('claimID. Use text up to 1,000 characters', record['warning'])
        self.assertEqual(record['data']['claimID'], 'x' * 1001)
        self.assertEqual(record['data']['specialInstructions'], instructions)

    def test_excel_dates_and_zero_formatted_identifiers(self):
        from datetime import datetime, time as clocktime
        self.fill(synthetic() | {'dob': datetime(2000, 1, 2), 'nextApptTime': clocktime(14, 30)})
        book = load_workbook(self.path)
        index = next(i for i, (key, _, _) in enumerate(excel.COLUMNS, 1) if key == 'zip')
        cell = book[excel.SHEET].cell(2, index, 123)
        cell.number_format = '00000'
        book.save(self.path)
        book.close()
        data = excel.read_workbook(self.path)[0]['data']
        self.assertEqual((data['dob'], data['nextApptTime'], data['zip']), ('01/02/2000', '02:30 PM', '00123'))

    def test_formulas_and_errors_rejected(self):
        for value in ('=1+1', '#DIV/0!'):
            self.fill(synthetic() | {'claimID': value})
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'plain values'):
                excel.read_workbook(self.path)

    def test_bad_schema_empty_extension_duplicate(self):
        with self.assertRaisesRegex(ValueError, 'No filled rows'):
            excel.read_workbook(self.path)
        with self.assertRaisesRegex(ValueError, '.xlsx'):
            excel.read_workbook(self.path.with_suffix('.csv'))
        book = load_workbook(self.path)
        sheet = book[excel.SHEET]
        sheet['B1'] = sheet['A1'].value
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Duplicate'):
            excel.read_workbook(self.path)
        book = load_workbook(self.path)
        book[excel.SHEET].title = 'Wrong sheet'
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Missing'):
            excel.read_workbook(self.path)

    def test_malformed_optional_row_loads_with_warning(self):
        self.fill(synthetic() | {'zip': 'wrong'})
        record = excel.read_workbook(self.path)[0]
        self.assertIn('five-digit ZIP', record['warning'])
        self.assertEqual(record['data']['zip'], 'wrong')


class BridgeTests(unittest.TestCase):
    def setUp(self):
        self.calls = {'list_attachments': Mock(return_value=['synthetic-referral.pdf']),
                      'run_unity': Mock(), 'record_assignment': Mock(), 'complete_assignment': Mock()}
        self.api = backend.UnityAPI(lambda: self.calls)

    def wait(self, predicate):
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            state = self.api.get_state()
            if predicate(state):
                return state
            time.sleep(.005)
        self.fail('Synthetic worker did not reach expected state.')

    def finish(self):
        state = self.wait(lambda s: not s['busy'])
        self.api._thread.join(1)
        return state

    def test_unity_stage_and_injected_derived_data(self):
        data = synthetic() | {'customer': 'Synthetic Continental Tire', 'claimType': 'Auto'}
        self.assertTrue(self.api.start(data, 'unity', True, 'MI')['ok'])
        state = self.finish()
        self.assertEqual(state['status'], 'completed')
        self.assertEqual(self.calls['run_unity'].call_args.args[2]['serviceType'], 'Medical Task')
        self.calls['complete_assignment'].assert_not_called()
        self.assertIsNone(state['prompt'])

    def test_long_instructions_reach_callbacks_exactly_and_keep_approval_gate(self):
        for length in (1001, 32767):
            with self.subTest(length=length):
                instructions = synthetic_instructions(length)
                data = synthetic() | {'specialInstructions': instructions}
                checked = self.api.validate_fields(data)
                self.assertTrue(checked['ok'], checked)
                self.assertEqual(checked['data']['fields']['specialInstructions'], instructions)
                self.assertTrue(self.api.start(data)['ok'])
                state = self.wait(lambda s: s['prompt'])
                passed = self.calls['run_unity'].call_args.args[2]
                self.assertEqual(passed['specialInstructions'], instructions)
                self.assertEqual(passed['claimNumber'], data['claimNumber'])
                self.calls['complete_assignment'].assert_not_called()
                self.assertTrue(self.api.answer_prompt(state['prompt']['id'], True)['ok'])
                self.assertEqual(self.finish()['status'], 'completed')
                self.calls['complete_assignment'].assert_called_once_with(data['claimNumber'])
                for call in self.calls.values():
                    call.reset_mock()

    def test_oversized_instructions_rejected_before_any_side_effect(self):
        factory = Mock(return_value=self.calls)
        self.api = backend.UnityAPI(factory)
        data = synthetic() | {'specialInstructions': synthetic_instructions(32768)}
        for stage in ('all', 'unity', 'assign'):
            with self.subTest(stage=stage):
                for reply in (self.api.validate_fields(data, stage), self.api.start(data, stage)):
                    self.assertFalse(reply['ok'])
                    self.assertIn('specialInstructions. Use text up to 32,767 characters', reply['error'])
                self.assertEqual(self.api.get_state()['status'], 'idle')
                self.assertIsNone(self.api._thread)
        factory.assert_not_called()
        for callback in self.calls.values():
            callback.assert_not_called()

    def test_blocking_confirm_polling_and_stale_answer(self):
        self.assertTrue(self.api.start(synthetic())['ok'])
        state = self.wait(lambda s: s['prompt'])
        identifier = state['prompt']['id']
        for _ in range(30):
            self.assertTrue(self.api.get_state()['busy'])
        self.assertFalse(self.api.answer_prompt('stale', True)['ok'])
        self.assertFalse(self.api.answer_prompt(identifier, 'yes')['ok'])
        self.assertFalse(self.api.start(synthetic())['ok'])
        self.assertTrue(self.api.answer_prompt(identifier, True)['ok'])
        self.assertFalse(self.api.answer_prompt(identifier, False)['ok'])
        self.assertEqual(self.finish()['status'], 'completed')
        self.calls['complete_assignment'].assert_called_once_with('SYNTHETIC-001')

    def test_assign_decline(self):
        self.api.start({'claimNumber': 'SYNTHETIC-001'}, 'assign')
        state = self.wait(lambda s: s['prompt'])
        self.api.answer_prompt(state['prompt']['id'], False)
        state = self.finish()
        self.assertEqual((state['status'], state['result']['reason']), ('stopped', 'icasemanager_declined'))
        for call in self.calls.values():
            call.assert_not_called()

    def test_attachment_stops_acknowledged(self):
        for files, reason in [([], 'attachment_missing'), (['synthetic-a.pdf', 'synthetic-b.pdf'], 'attachment_ambiguous')]:
            with self.subTest(reason=reason):
                self.setUp()
                self.calls['list_attachments'].return_value = files
                self.api.start(synthetic())
                state = self.wait(lambda s: s['prompt'])
                self.assertEqual(state['prompt']['kind'], 'notify')
                self.assertFalse(self.api.answer_prompt(state['prompt']['id'], False)['ok'])
                self.api.answer_prompt(state['prompt']['id'], True)
                state = self.finish()
                self.assertEqual((state['status'], state['result']['reason']), ('stopped', reason))
                self.calls['run_unity'].assert_not_called()

    def test_failures_and_tracking_preference(self):
        for tracking in (True, False):
            self.setUp()
            self.calls['run_unity'].side_effect = SystemExit(7)
            self.api.start(synthetic(), 'unity', tracking)
            state = self.finish()
            self.assertEqual(state['result']['error']['exit_code'], 7)
            self.assertEqual(state['status'], 'failed')
            self.assertEqual(self.calls['record_assignment'].call_count, int(tracking))

    def test_preflight_no_side_effects(self):
        for data, stage, jurisdiction in [({}, 'all', ''), (synthetic(), 'invalid', ''), (synthetic(), 'all', 'ZZ')]:
            self.assertFalse(self.api.start(data, stage, True, jurisdiction)['ok'])
        self.assertFalse(self.api.start(synthetic(), 'all', 'yes')['ok'])
        self.assertEqual(self.api.get_state()['status'], 'idle')
        for call in self.calls.values():
            call.assert_not_called()

    def test_worker_factory_error_and_new_run_resets_result(self):
        self.api._factory = Mock(side_effect=RuntimeError('Synthetic connection failure'))
        self.api.start(synthetic(), 'unity')
        self.assertEqual(self.finish()['result']['reason'], 'desktop_worker_failed')
        self.api._factory = lambda: self.calls
        self.api.start(synthetic(), 'unity')
        self.assertEqual(self.finish()['status'], 'completed')
        state = self.api.get_state()
        state['result']['steps'].clear()
        self.assertTrue(self.api.get_state()['result']['steps'])

    def test_default_factory_com_lifecycle(self):
        com = Mock()
        self.api = backend.UnityAPI()
        with patch.dict(os.environ, {'FCM_RRS_DB_CONNECTION': 'SYNTHETIC-TRACKING'}), \
                patch.dict(sys.modules, {'pythoncom': com}), \
                patch.object(backend, 'live_callbacks', return_value=self.calls) as live:
            self.api.start(synthetic(), 'unity')
            self.assertEqual(self.finish()['status'], 'completed')
            live.assert_called_once()
            com.CoInitialize.assert_called_once()
            com.CoUninitialize.assert_called_once()

    def test_swallowed_legacy_failure_marker_is_failed(self):
        unity = SimpleNamespace(validate_and_complete_task=lambda *args: print('error occurred: synthetic failure (line 123)'))
        with self.assertRaisesRegex(RuntimeError, 'Legacy Unity reported failure'):
            backend.checked_validate(unity, Mock(), 'Synthetic Customer', 'Invented Example', {})
        unity.validate_and_complete_task = lambda *args: print('Synthetic successful operation')
        backend.checked_validate(unity, Mock(), 'Synthetic Customer', 'Invented Example', {})

    def test_checked_flag_failure_keeps_exception_identity_and_original_type(self):
        from selenium.common.exceptions import TimeoutException

        original = TimeoutException('SYNTHETIC-PRIVATE page details')
        error = UnityFlagError('waiting for the visible flag dialog')

        def fail(*args):
            raise error from original

        unity = SimpleNamespace(validate_and_complete_task=fail)
        with self.assertRaises(UnityFlagError) as caught:
            backend.checked_validate(unity, Mock(), 'Synthetic', 'Example', {})
        self.assertIs(caught.exception, error)
        self.assertIs(caught.exception.__cause__, original)
        self.assertIs(type(caught.exception.__cause__), TimeoutException)
        self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))

    def test_live_portal_is_lazy_and_closes_driver_on_unrelated_failure(self):
        unity = Mock()
        loader = Mock(return_value=unity)
        config = SimpleNamespace(ATTACHMENT_FOLDER=str(HERE / 'nonexistent-synthetic-folder'), UNITY_SCRIPT='synthetic-script.py')
        with patch.dict(sys.modules, {'fcm_intake.config': config,
                                     'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=loader)}):
            calls = backend.live_callbacks()
            loader.assert_not_called()
            self.assertEqual(calls['list_attachments'](), [])
            unity.validate_and_complete_task.side_effect = SystemExit(7)
            with self.assertRaises(SystemExit):
                calls['run_unity']('Synthetic Customer', 'Invented Example', synthetic())
            unity.validate_and_complete_task.assert_called_once_with(
                unity.create_driver.return_value, 'Synthetic Customer', 'Invented Example', synthetic())
            unity.create_driver.return_value.quit.assert_called_once()
            unity.create_driver.assert_called_once_with(detach=True)

    def test_live_flag_failure_releases_service_but_preserves_browser_and_chain(self):
        for release_failure in (None, 'service', 'connection'):
            with self.subTest(release_failure=release_failure):
                unity = Mock()
                driver = unity.create_driver.return_value
                original = ValueError('SYNTHETIC-PRIVATE Selenium details')
                flag_error = UnityFlagError('submitting the flag')

                def fail(*args):
                    raise flag_error from original

                unity.validate_and_complete_task.side_effect = fail
                if release_failure == 'service':
                    driver.service.stop.side_effect = RuntimeError('SYNTHETIC-PRIVATE service')
                if release_failure == 'connection':
                    driver.command_executor.close.side_effect = RuntimeError('SYNTHETIC-PRIVATE connection')
                config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py')
                with patch.dict(sys.modules, {
                        'fcm_intake.config': config,
                        'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=lambda *args: unity)}):
                    calls = backend.live_callbacks()
                    with self.assertRaises(UnityFlagError) as caught:
                        calls['run_unity']('Synthetic', 'Example', synthetic())
                self.assertIs(caught.exception, flag_error)
                self.assertIs(caught.exception.__cause__, original)
                self.assertIn('Browser left open for manual recovery', str(caught.exception))
                self.assertIn('do not blindly retry submission', str(caught.exception))
                self.assertNotIn('SYNTHETIC-PRIVATE', str(caught.exception))
                self.assertIs(caught.exception.service_release_failed, release_failure is not None)
                unity.create_driver.assert_called_once_with(detach=True)
                unity.validate_and_complete_task.assert_called_once()
                driver.service.stop.assert_called_once()
                driver.command_executor.close.assert_called_once()
                driver.quit.assert_not_called()

    def test_live_success_and_nonflag_failures_still_quit(self):
        for failure in (None, ValueError('synthetic unrelated'), SystemExit(7)):
            with self.subTest(failure=failure):
                unity = Mock()
                unity.validate_and_complete_task.side_effect = failure
                config = SimpleNamespace(ATTACHMENT_FOLDER='synthetic', UNITY_SCRIPT='synthetic.py')
                with patch.dict(sys.modules, {
                        'fcm_intake.config': config,
                        'fcm_intake.legacy_loader': SimpleNamespace(load_module_from_path=lambda *args: unity)}):
                    calls = backend.live_callbacks()
                    if failure is None:
                        calls['run_unity']('Synthetic', 'Example', synthetic())
                    else:
                        with self.assertRaises(type(failure)):
                            calls['run_unity']('Synthetic', 'Example', synthetic())
                driver = unity.create_driver.return_value
                driver.quit.assert_called_once()
                driver.service.stop.assert_not_called()

    def test_excel_bridge_and_busy_guard(self):
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            path = Path(directory) / 'synthetic.xlsx'
            self.assertTrue(self.api.generate_workbook(str(path))['ok'])
            book = load_workbook(path)
            for i, (key, _, _) in enumerate(excel.COLUMNS, 1):
                book[excel.SHEET].cell(2, i, synthetic()[key])
            book.save(path)
            book.close()
            self.assertTrue(self.api.load_excel(str(path))['ok'])
            self.assertEqual(self.api.select_excel_row(2)['data']['data'], synthetic())
            self.assertFalse(self.api.select_excel_row(3)['ok'])
            self.api.start(synthetic())
            state = self.wait(lambda s: s['prompt'])
            self.assertFalse(self.api.select_excel_row(2)['ok'])
            self.assertFalse(self.api.load_excel(str(path))['ok'])
            self.assertFalse(self.api.generate_workbook(str(path))['ok'])
            self.api.answer_prompt(state['prompt']['id'], False)
            self.finish()

    def pick(self, paths):
        self.api._window = Mock()
        self.api._window.create_file_dialog.return_value = paths
        with patch.object(self.api, '_dialog', return_value='OPEN'):
            return self.api.choose_referral_document()

    def test_document_picker_validation_cancel_and_session_state(self):
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            from webview.util import parse_file_type
            for extension in ('.pDf', '.MsG', '.eMl'):
                document = Path(directory) / ('synthetic' + extension)
                document.touch()
                self.assertTrue(self.pick([str(document)])['ok'])
                state = self.api.initialize()['state']['attachment']
                self.assertEqual(state, {'source': 'local', 'filename': document.name})
                self.assertEqual(self.api._attachment_path, document.resolve())
                self.api._window.create_file_dialog.assert_called_once_with(
                    'OPEN', allow_multiple=False,
                    file_types=('Referral document (*.pdf;*.msg;*.eml;*.PDF;*.MSG;*.EML)',))
                _, patterns = parse_file_type(self.api._window.create_file_dialog.call_args.kwargs['file_types'][0])
                self.assertEqual(set(patterns.split(';')),
                                 {'*.pdf', '*.msg', '*.eml', '*.PDF', '*.MSG', '*.EML'})
                self.assertEqual(self.pick(None), {'ok': True, 'data': None})
                for suffix in ('.txt', '.zip'):
                    wrong = document.with_suffix(suffix)
                    wrong.touch()
                    self.assertFalse(self.pick([str(wrong)])['ok'])
                folder = Path(directory) / ('folder' + extension)
                folder.mkdir()
                for paths in ([str(document.with_name('missing' + extension))],
                              [str(folder)], [str(document), str(document)]):
                    with self.subTest(paths=paths):
                        self.assertFalse(self.pick(paths)['ok'])
                        self.assertEqual(self.api.get_state()['attachment'], state)
                self.assertTrue(self.api.validate_fields(synthetic())['ok'])
                self.assertEqual(self.api.get_state()['attachment'], state)
            self.assertTrue(self.api.use_rrs_export_folder()['ok'])
            self.assertIsNone(self.api._attachment_path)
            self.assertEqual(self.api.get_state()['attachment'], {'source': 'rrs', 'filename': ''})
            self.api.start(synthetic(), 'unity')
            self.assertEqual(self.finish()['result']['attachment'], 'synthetic-referral.pdf')
            self.calls['list_attachments'].assert_called_once()

    def test_document_picker_busy_and_dialog_race(self):
        self.api._state['busy'] = True
        self.assertFalse(self.pick(None)['ok'])
        self.api._window.create_file_dialog.assert_not_called()
        self.assertFalse(self.api.use_rrs_export_folder()['ok'])
        self.api._state['busy'] = False
        self.api._window.create_file_dialog.side_effect = lambda *a, **kw: self.api._state.update(busy=True)
        with patch.object(self.api, '_dialog', return_value='OPEN'):
            self.assertFalse(self.api.choose_referral_document()['ok'])
        self.assertIsNone(self.api._attachment_path)

    def test_selected_document_overrides_exports_failure_policy_and_approval(self):
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            for extension in ('.PDF', '.mSg', '.EmL'):
                document = Path(directory) / ('synthetic' + extension)
                document.touch()
                for exports in ([], ['old-a.pdf', 'old-b.pdf']):
                    for tracking in (True, False):
                        self.calls['list_attachments'].return_value = exports
                        self.assertTrue(self.pick(str(document))['ok'])
                        self.calls['run_unity'].side_effect = RuntimeError('synthetic upload failure')
                        self.assertTrue(self.api.start(synthetic(), 'all', tracking)['ok'])
                        state = self.finish()
                        self.assertEqual(state['result']['attachment'], str(document.resolve()))
                        self.assertEqual(state['result']['reason'], 'run_unity_failed')
                        self.assertEqual(self.calls['record_assignment'].call_count, int(tracking))
                        self.calls['list_attachments'].assert_not_called()
                        self.calls['complete_assignment'].assert_not_called()
                        for call in self.calls.values():
                            call.reset_mock()
            self.calls['run_unity'].side_effect = None
            self.assertTrue(self.api.start(synthetic(), 'all')['ok'])
            state = self.wait(lambda s: s['prompt'])
            self.assertFalse(self.api.use_rrs_export_folder()['ok'])
            self.calls['complete_assignment'].assert_not_called()
            self.api.answer_prompt(state['prompt']['id'], False)
            self.assertEqual(self.finish()['result']['reason'], 'icasemanager_declined')

    def test_deleted_selected_document_never_falls_back_and_assign_ignores_it(self):
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            for extension in ('.PDF', '.MsG', '.EmL'):
                self.setUp()
                document = Path(directory) / ('synthetic' + extension)
                document.touch()
                self.assertTrue(self.pick([str(document)])['ok'])
                document.unlink()
                for stage in ('unity', 'all'):
                    self.assertFalse(self.api.start(synthetic(), stage)['ok'])
                    self.assertIsNone(self.api._thread)
                for call in self.calls.values():
                    call.assert_not_called()
                self.assertTrue(self.api.start({'claimNumber': 'SYNTHETIC-001'}, 'assign')['ok'])
                state = self.wait(lambda s: s['prompt'])
                self.api.answer_prompt(state['prompt']['id'], True)
                self.assertEqual(self.finish()['status'], 'completed')
                self.calls['list_attachments'].assert_not_called()
                self.calls['run_unity'].assert_not_called()
                self.calls['record_assignment'].assert_not_called()
                self.calls['complete_assignment'].assert_called_once()

    def test_live_selected_document_is_exact_and_checked_before_upload(self):
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            for name in ('old-a.pdf', 'old-b.pdf'):
                (Path(directory) / name).touch()
            for extension in ('.pDf', '.mSg', '.eML'):
                document = Path(directory) / ('synthetic' + extension)
                document.touch()
                unity = Mock()
                config = SimpleNamespace(ATTACHMENT_FOLDER=directory, UNITY_SCRIPT='synthetic-script.py')
                loader = SimpleNamespace(load_module_from_path=Mock(return_value=unity))
                with patch.dict(sys.modules, {'fcm_intake.config': config, 'fcm_intake.legacy_loader': loader}):
                    calls = backend.live_callbacks(document.resolve())
                    self.assertEqual(calls['list_attachments'](), [document.resolve()])
                    payload = synthetic()
                    calls['run_unity']('Synthetic Customer', 'Invented Example', payload)
                    unity.validate_and_complete_task.assert_called_once_with(
                        unity.create_driver.return_value, 'Synthetic Customer', 'Invented Example', payload,
                        attachment_path=str(document.resolve()))
                    unity.create_driver.return_value.quit.assert_called_once()
                    document.unlink()
                    with self.assertRaisesRegex(ValueError, 'existing referral document'):
                        calls['list_attachments']()
                    unity.create_driver.reset_mock()
                    with self.assertRaisesRegex(ValueError, 'existing referral document'):
                        calls['run_unity']('Synthetic', 'Example', {})
                    unity.create_driver.assert_not_called()
                    with self.assertRaisesRegex(ValueError, 'existing referral document'):
                        backend.checked_validate(unity, Mock(), 'Synthetic', 'Example', {}, attachment_path=document)
                    self.assertEqual(unity.validate_and_complete_task.call_count, 1)

    def test_document_disappearing_after_start_fails_controller_preflight(self):
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            for extension in ('.PDF', '.MsG', '.eMl'):
                self.setUp()
                document = Path(directory) / ('synthetic' + extension)
                document.touch()
                self.assertTrue(self.pick([str(document)])['ok'])
                def factory():
                    document.unlink()
                    return self.calls
                self.api._factory = factory
                self.assertTrue(self.api.start(synthetic(), 'unity')['ok'])
                self.assertEqual(self.finish()['result']['reason'], 'resolve_attachment_failed')
                for call in self.calls.values():
                    call.assert_not_called()

    def test_actual_legacy_upload_ast_binds_manual_path_and_propagates_failure(self):
        function = legacy_function('validate_and_complete_task')
        upload_block, = [node for node in function.body[0].body if isinstance(node, ast.Try)
                         and any(isinstance(child, ast.ImportFrom)
                                 and child.module == 'AI.referral_attachments'
                                 for child in ast.walk(node))]
        tree = ast.Module(body=[upload_block], type_ignores=[])
        code = compile(tree, '<actual legacy upload block>', 'exec')
        with tempfile.TemporaryDirectory(dir=HERE) as directory:
            upload, globber = Mock(), Mock()
            env = {'os': os, 'glob': globber,
                   'ATTACHMENT_FOLDER': directory, 'driver': Mock(), 'time': Mock(),
                   'WebDriverWait': Mock(return_value=SimpleNamespace(until=Mock(return_value=upload))),
                   'EC': Mock(), 'By': SimpleNamespace(XPATH='xpath')}
            for extension in ('.pDf', '.MsG', '.eMl'):
                document = Path(directory) / ('synthetic' + extension)
                document.touch()
                env['attachment_path'] = str(document.resolve())
                upload.reset_mock(side_effect=True)
                exec(code, env)
                upload.send_keys.assert_called_once_with(str(document.resolve()))
                globber.glob.assert_not_called()
                upload.send_keys.side_effect = RuntimeError('synthetic selenium upload failure')
                with self.assertRaisesRegex(RuntimeError, 'synthetic selenium'):
                    exec(code, env)
                document.unlink()
                with self.assertRaisesRegex(ValueError, 'existing referral document'):
                    exec(code, env)
            for suffix in ('.txt', '.zip', '.msg'):
                invalid = Path(directory) / ('invalid' + suffix)
                if suffix == '.msg':
                    invalid.mkdir()
                else:
                    invalid.touch()
                env['attachment_path'] = str(invalid)
                with self.assertRaises(ValueError):
                    exec(code, env)
            globber.glob.assert_not_called()
            # Original four-argument/export mode retains its legacy skip behavior.
            env['attachment_path'] = None
            upload.reset_mock(side_effect=True)
            globber.glob.return_value = ['synthetic-first.pdf', 'synthetic-second.msg']
            with patch('builtins.print'):
                exec(code, env)
            upload.send_keys.assert_called_once_with('synthetic-first.pdf')
            upload.send_keys.side_effect = RuntimeError('synthetic export upload failure')
            with patch('builtins.print'):
                exec(code, env)
            globber.glob.return_value = []
            with patch('builtins.print'):
                exec(code, env)


if __name__ == '__main__':
    unittest.main()
