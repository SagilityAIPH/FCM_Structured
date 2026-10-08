"""Offline native-window test. Never touches RRS or any real referral."""
import argparse
from pathlib import Path
import sys
import time
import traceback


class FakeAdapter:
    def __init__(self, check):
        self.check = check
        self.created = False

    def open(self): pass
    def identity(self, data): pass
    def lookup(self, kind, data): pass
    def ensure_lookup_closed(self): pass
    def details(self, data): pass
    def create(self, data): self.created = True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--self-test', action='store_true')
    parser.add_argument('--report', type=Path, required=True)
    args = parser.parse_args()
    args.report.parent.mkdir(parents=True, exist_ok=True)
    import webview
    from backend import ManualAPI
    from manual_flow import defaults
    from AI.workbench.window_layout import constrain_window
    adapter = FakeAdapter(lambda: None)
    api = ManualAPI(lambda check: adapter)
    index = Path(getattr(sys, '_MEIPASS', Path(__file__).parent)) / 'frontend/dist/index.html'
    window = webview.create_window('SubjectLineBuilder offline smoke', str(index), js_api=api,
                                  hidden=True, width=480, height=900, min_size=(480, 600))
    window.events.before_show += lambda: constrain_window(window)
    passed = []

    def verify():
        try:
            # Verify packaged UIA/COM dependencies without touching any window.
            import pythoncom
            pythoncom.CoInitialize()
            try:
                from rrs_ui import RRSAdapter
                assert RRSAdapter().desktop is not None
            finally:
                pythoncom.CoUninitialize()
            window.events.loaded.wait(20)
            deadline = time.monotonic() + 30
            while not window.evaluate_js("Boolean(document.querySelector('input[name=claimNumber]'))"):
                if time.monotonic() > deadline:
                    raise RuntimeError('Manual form did not load')
                time.sleep(.2)
            assert window.evaluate_js("document.querySelectorAll('.field-section').length") == 5
            assert window.evaluate_js("Boolean(document.querySelector('progress'))")
            from System import Action
            from System.Drawing import Size
            def check_size():
                form = window.native
                width = round(480 * form.DeviceDpi / 96)
                assert form.MinimumSize.Width == form.MaximumSize.Width == width
                original = form.Height
                form.Size = Size(width + 100, form.MinimumSize.Height)
                assert form.Width == width
                form.Size = Size(width, original)
            window.native.Invoke(Action(check_size))
            payload = defaults() | {'customer': 'Synthetic Co', 'claimNumber': 'TEST-001',
                'claimantFirst': 'Test', 'claimantLast': 'Example', 'providerName': 'Synthetic Clinic'}
            import tempfile
            from excel_input import create_template, COLUMNS, SHEET
            from openpyxl import load_workbook
            with tempfile.TemporaryDirectory() as folder:
                excel_path = Path(folder) / 'smoke.xlsx'
                create_template(excel_path)
                book = load_workbook(excel_path)
                for index, (key, _, _) in enumerate(COLUMNS, 1):
                    book[SHEET].cell(2, index, payload[key])
                book.save(excel_path)
                book.close()
                assert api._load_excel(excel_path)['ok']
                imported = api.select_excel_row(2)
                assert imported['ok'] and imported['data']['claimNumber'] == 'TEST-001'
                assert not api.get_state()['busy']
            assert api.start(imported['data'])['ok']
            deadline = time.monotonic() + 20
            reviews = 0
            while api.get_state()['busy']:
                if time.monotonic() > deadline:
                    raise RuntimeError('Offline workflow timed out')
                state = api.get_state()
                if state['status'] in ('lookup', 'review'):
                    assert not adapter.created
                    assert api.proceed(state['review_id'], state['status'] == 'review')['ok']
                    reviews += 1
                time.sleep(.1)
            assert adapter.created and reviews == 4
            assert api.get_state()['status'] == 'completed'
            passed.append(True)
            args.report.write_text('PASS: native WebView2, fixed 480 width, vertical resize, 5 manual groups, Excel template generation/import/row selection, bridge, progress, guided lookup and explicit Create gates. Synthetic adapter only; no live RRS, CMS or PDFs.', encoding='utf8')
        except Exception:
            args.report.write_text(traceback.format_exc(), encoding='utf8')
        finally:
            api.stop()
            window.destroy()
    webview.start(verify, gui='edgechromium', http_server=True, private_mode=True)
    if not passed:
        raise SystemExit(args.report.read_text())


if __name__ == '__main__':
    main()
