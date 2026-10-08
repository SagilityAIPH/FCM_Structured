import sys
from pathlib import Path
import tempfile
import unittest
from datetime import datetime, time

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'Deploy-Ready'))
from openpyxl import load_workbook
from excel_input import COLUMNS, SHEET, create_template, read_workbook
from backend import ManualAPI


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'input.xlsx'
        create_template(self.path)

    def fill(self, rows):
        book = load_workbook(self.path)
        for row_number, fields in enumerate(rows, 2):
            for index, (key, _, _) in enumerate(COLUMNS, 1):
                if key in fields:
                    book[SHEET].cell(row_number, index, fields[key])
        book.save(self.path)
        book.close()

    def test_dates_identifiers_and_incomplete_rows(self):
        self.fill([{'claimNumber': '000123', 'claimantFirst': 'Synthetic', 'zip': '00123',
                    'dob': datetime(1980, 1, 2), 'nextApptTime': time(8, 30)}])
        record = read_workbook(self.path)[0]
        self.assertEqual(record['row'], 2)
        self.assertEqual(record['data']['claimNumber'], '000123')
        self.assertEqual(record['data']['zip'], '00123')
        self.assertEqual(record['data']['dob'], '01/02/1980')
        self.assertEqual(record['data']['nextApptTime'], '08:30 AM')
        self.assertIn('Required:', record['warning'])

    def test_blank_template_is_not_a_referral(self):
        with self.assertRaisesRegex(ValueError, 'No filled rows'):
            read_workbook(self.path)

    def test_formula_is_rejected(self):
        self.fill([{'claimNumber': '=1+1'}])
        with self.assertRaisesRegex(ValueError, 'plain values'):
            read_workbook(self.path)

    def test_duplicate_headers_rejected(self):
        book = load_workbook(self.path)
        book[SHEET].cell(1, 2, book[SHEET].cell(1, 1).value)
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Duplicate column'):
            read_workbook(self.path)

    def test_row_switch_clears_previous_optional_fields_and_never_starts(self):
        self.fill([{'claimNumber': 'FIRST', 'attorneyFirst': 'Synthetic'}, {'claimNumber': 'SECOND'}])
        api = ManualAPI()
        self.assertTrue(api._load_excel(self.path)['ok'])
        self.assertEqual(api.select_excel_row(2)['data']['attorneyFirst'], 'Synthetic')
        self.assertEqual(api.select_excel_row(3)['data']['attorneyFirst'], '')
        self.assertFalse(api.get_state()['busy'])
        self.assertFalse(api.select_excel_row(99)['ok'])
        api._update(busy=True)
        self.assertFalse(api.select_excel_row(3)['ok'])
        self.assertFalse(api._load_excel(self.path)['ok'])


if __name__ == '__main__':
    unittest.main()
