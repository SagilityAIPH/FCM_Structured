"""Synthetic referral ICD-10-CM validation regressions (official CMS codes only)."""
import json
from unittest.mock import patch

import pytest

from AI import field_review, icd10
from AI.daily_output import read_record, save_daily_output
from openpyxl import load_workbook
from AI.field_review import review_formats
from AI.referral_schema import FIELD_ONLY_PROMPT, _TEMPLATE, export_payload, force_exact_field_output


@pytest.mark.parametrize(('code', 'status'), [
    ('S62.101A', 'valid'), ('s62101a', 'valid'), ('M54', 'incomplete'),
    ('S62.101', 'incomplete'), ('X99.9', 'incomplete'), ('T12', 'invalid'),
    ('S99.999X', 'invalid'), ('M54.5', 'incomplete'), ('X99.99', 'invalid'),
])
def test_membership(code, status):
    assert icd10.check(code) == status


def test_incomplete_codes_do_not_flag_by_default():
    assert field_review.FLAG_INCOMPLETE_ICD10 is False
    for code in ('M54', 'M54.5', 'M48.06', 'S32.314', 'S62.101'):
        assert icd10.check(code) == 'incomplete'
        for key in ('Diagnosis Code', 'Additional Diagnosis Codes'):
            assert review_formats({key: code}) == []
    assert review_formats({'Diagnosis Code': 'M54.5',
                           'Additional Diagnosis Codes': 'M48.06; S32.314'}) == []


def test_incomplete_flags_restore_with_switch(monkeypatch):
    monkeypatch.setattr(field_review, 'FLAG_INCOMPLETE_ICD10', True)
    flags = review_formats({'Diagnosis Code': 'M54',
                            'Additional Diagnosis Codes': 'M54.5; S32.314; S62.101'})
    assert [(f['field'], f['reason']) for f in flags] == [
        ('Diagnosis Code', "'M54' is an ICD-10-CM category, not a billable code -- more characters needed (FY2027)"),
        ('Additional Diagnosis Codes', "'M54.5' is an incomplete ICD-10-CM code, not billable -- more characters needed (FY2027)"),
        ('Additional Diagnosis Codes', "'S32.314' is an incomplete ICD-10-CM code, not billable -- more characters needed (FY2027)"),
        ('Additional Diagnosis Codes', "'S62.101' is an incomplete ICD-10-CM code, not billable -- more characters needed (FY2027)"),
    ]


def test_invalid_codes_still_flag_with_switch_off():
    assert field_review.FLAG_INCOMPLETE_ICD10 is False
    flags = review_formats({'Diagnosis Code': 'T12',
                            'Additional Diagnosis Codes': 'M54.5; 854; S62.101A'})
    assert [(f['field'], f['reason']) for f in flags] == [
        ('Diagnosis Code', "'T12' is not a valid ICD-10-CM code (FY2027)"),
        ('Additional Diagnosis Codes', "'854' is not a valid ICD-10-CM code (FY2027)"),
    ]


def test_missing_list_is_unchecked_without_crash():
    assert field_review.FLAG_INCOMPLETE_ICD10 is False
    icd10._codes.cache_clear()
    with patch.object(icd10, '_list_path', side_effect=FileNotFoundError):
        assert icd10.check('M54') == 'unchecked'
        flags = review_formats({'Diagnosis Code': 'M54',
                                'Additional Diagnosis Codes': 'M54.5; T12'})
        assert [(flag['field'], flag['reason']) for flag in flags] == [
            ('Diagnosis Code', "ICD-10-CM validation could not run for 'M54' (FY2027); verify the code manually"),
            ('Additional Diagnosis Codes', "ICD-10-CM validation could not run for 'M54.5' (FY2027); verify the code manually"),
            ('Additional Diagnosis Codes', "ICD-10-CM validation could not run for 'T12' (FY2027); verify the code manually"),
        ]
    icd10._codes.cache_clear()


def test_final_values_flags_recheck_and_workbook(tmp_path):
    raw = dict(_TEMPLATE)
    raw.update({'Diagnosis Code': 'M54', 'Additional Diagnosis Codes': 'S62.101A; X99.99',
                'Provider Information': []})
    text, fields = force_exact_field_output(json.dumps(raw), source_text='')
    assert fields['Diagnosis Code'] == 'M54'
    assert fields['Additional Diagnosis Codes'] == 'S62.101A; X99.99'
    reasons = [f['reason'] for f in fields.review_flags]
    assert not any("'M54'" in reason for reason in reasons)
    assert any("'X99.99'" in reason and 'not a valid' in reason for reason in reasons)
    assert not any('S62.101A' in reason for reason in reasons)
    assert 'Diagnosis Code: M54' in text
    assert review_formats(fields, fields.review_flags) == fields.review_flags
    assert export_payload(fields)['Review Flags'] == fields.review_flags
    path, record_id = save_daily_output(fields, 'synthetic.pdf', directory=tmp_path)
    book = load_workbook(path, read_only=True)
    try:
        sheet = book['Review Flags']
        saved = [json.loads(row[1]) for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert saved == fields.review_flags
    finally:
        book.close()
    restored, _ = read_record(path, record_id)
    assert restored.review_flags == fields.review_flags
    assert restored['Diagnosis Code'] == 'M54'


def test_edits_recheck_instead_of_retaining_stale_flags():
    assert not review_formats({'Diagnosis Code': 'S62.101A'},
                              review_formats({'Diagnosis Code': 'M54'}))
    assert 'Transcribe diagnosis codes exactly as written; never correct, complete or invent a code; Python validates them against ICD-10-CM.' in FIELD_ONLY_PROMPT
