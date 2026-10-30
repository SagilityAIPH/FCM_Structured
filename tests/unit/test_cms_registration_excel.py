"""Synthetic-only tests for local CMS workbook import and conversion."""

from datetime import datetime, time
from copy import copy
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from AI import cms_registration_excel as excel


@pytest.fixture
def workbook_path():
    # Keep test artifacts inside the repository, not an OS temporary directory.
    root = Path("tests") / ".cms-excel-test-artifacts"
    root.mkdir(exist_ok=True)
    paths = []

    def make(name="input.xlsx", *, headers=None, rows=None, sheet=excel.SHEET):
        path = root / f"{uuid4().hex}-{name}"
        paths.append(path)
        book = Workbook()
        page = book.active
        page.title = sheet
        page.append(list(excel.LEGACY_HEADERS if headers is None else headers))
        for row in rows or []:
            page.append(row)
        book.save(path)
        book.close()
        return path

    yield make
    for path in paths:
        if path.exists():
            path.unlink()
    if not any(root.iterdir()):
        root.rmdir()


def _row(values, headers=excel.LEGACY_HEADERS):
    return [values.get(header) for header in headers]


def _edit(path, update):
    book = load_workbook(path)
    try:
        update(book.active)
        book.save(path)
    finally:
        book.close()


def test_conversion_preserves_source_values_formats_and_adds_only_blanks(workbook_path):
    values = {header: f"SYN-{index}" for index, header in enumerate(excel.SOURCE_HEADERS)}
    values["Claim - Claim number"] = 73
    values["Provider - Appointment date"] = datetime(2030, 4, 5)
    values["Claim - Referral type"] = "Full Case Management"
    source = workbook_path("source.xlsx", headers=excel.SOURCE_HEADERS,
                           rows=[_row(values, excel.SOURCE_HEADERS)], sheet=excel.SOURCE_SHEET)

    def format_cells(sheet):
        sheet["C2"].number_format = "000000"
        sheet["C2"].font = Font(bold=True, color="AABBCC")
        sheet["AA2"].number_format = "mm/dd/yyyy"
        sheet["H2"] = "=literal text"
        sheet["H2"].data_type = "s"

    _edit(source, format_cells)
    destination = workbook_path("converted.xlsx")
    destination.unlink()
    before = sha256(source.read_bytes()).digest()
    assert excel.convert_subject_line_workbook(source, destination) == str(destination)
    assert sha256(source.read_bytes()).digest() == before
    original = load_workbook(source)
    converted = load_workbook(destination)
    try:
        assert converted.sheetnames == [excel.SHEET, "Instructions"]
        sheet = converted[excel.SHEET]
        assert sheet.max_column == 54
        assert sheet.freeze_panes == "C2"
        for row in original.active:
            for cell in row:
                copied = sheet.cell(cell.row, cell.column)
                assert copied.value == cell.value
                assert copied.number_format == cell.number_format
                assert copy(copied.font) == copy(cell.font)
                assert copy(copied.fill) == copy(cell.fill)
                assert copy(copied.border) == copy(cell.border)
                assert copy(copied.alignment) == copy(cell.alignment)
                assert copy(copied.protection) == copy(cell.protection)
                assert copied.data_type == cell.data_type
        assert all(cell.value is None for row in sheet.iter_rows(
            min_row=2, min_col=44) for cell in row)
        assert "no rrs" in " ".join(str(row[0].value) for row in converted["Instructions"]).lower()
    finally:
        original.close()
        converted.close()
    records = excel.read_workbook(destination)
    assert records[0]["form"]["claim_number"] == "000073"
    assert records[0]["form"]["referral_type"] == ""


def test_request_and_setup_mapping_preserves_separate_metadata(workbook_path):
    values = {
        "Claim - Claim number": "001234",
        "Claimant - First name": "Synthetic", "Claimant - Last name": "Example",
        "Claim - Employer / customer": "Synthetic Employer",
        "Claim - Body part": "Hand", "Claim - Referral type": "Full Case Management",
        "Provider - Facility name": "Synthetic Facility",
        "Provider - Doctor first name": "Synthetic", "Provider - Doctor last name": "Doctor",
        "Provider - Phone": "0010203040", "Adjuster - Referral source": "Adjuster",
        "Adjuster - Email": "adjuster@example.invalid",
        "Claimant - Address line 1": "Synthetic claimant address",
        "Attorney - Firm": "Synthetic Firm", "Attorney - Address line 1": "Synthetic address",
        "Attorney - Address line 2": "Synthetic suite", "Attorney - City": "Example",
        "Attorney - State": "NY", "Attorney - ZIP": "00123", "Attorney - Phone": "000123",
        "Provider - Appointment date": datetime(2030, 4, 5),
        "Provider - Appointment time": time(8, 30),
        "CMS - Case ID": "00005", "CMS - Referral type": "Vocational",
        "CMS - Referral number": "00009", "CMS - CEM employer request": "yes",
        "CMS - Diagnosis code": "SYN.1", "CMS - Office number": "00002",
        "CMS - Special instructions": "  Synthetic\ninstructions  ",
        "CMS - Claim Identifier": "Synthetic identifier",
        "CMS - Other Notes": " Synthetic notes ",
        "CMS - Customer email": "employer@example.invalid",
        "CMS - Insured address line 1": "Synthetic insured address",
    }
    path = workbook_path(rows=[_row(values)])
    result = excel.read_workbook(path)[0]
    assert result["form"] == {
        "claim_number": "001234", "claimant_first_name": "Synthetic",
        "claimant_last_name": "Example", "referral_type": "Vocational",
        "cms_case_id": "00005", "referral_number": "00009", "cem_employer_request": "yes",
    }
    data = result["data"]
    expected = {
        "dxCode": "SYN.1", "bodyParts": "Hand", "officeNumber": "00002",
        "xlCustomerName": "Synthetic Employer", "xlReferralSource": "Adjuster",
        "xlProvider": "Synthetic Facility", "provType": "facility",
        "xlProviderPhone": "0010203040", "xlCustomerEmail": "employer@example.invalid",
        "xlInsuredAddr1": "Synthetic insured address", "xlAttorney": "Synthetic Firm",
        "xlAttorneyAdd1": "Synthetic address", "xlAttorneyAdd2": "Synthetic suite",
        "xlAttorneyCity": "Example", "xlAttorneyState": "NY", "xlAttorneyZip": "00123",
        "xlAttorneyPhone": "000123", "apptDate": "04/05/2030", "refType": "Vocational",
        "Special_Instructions": "  Synthetic\ninstructions  ", "cem_employer_request": True,
        "reviewClaimIdentifier": "Synthetic identifier", "reviewOtherNotes": " Synthetic notes ",
        "reviewAppointmentDate": "04/05/2030", "reviewAppointmentTime": "08:30 AM",
    }
    assert {key: data[key] for key in expected} == expected
    assert len(data["review_fields"]) == 43
    assert data["review_fields"]["Provider - Doctor last name"] == "Doctor"
    assert "Synthetic" not in result["warning"]
    assert "001234" not in result["warning"]


def test_unknowns_stay_blank_without_office_referral_or_cem_inference(workbook_path):
    path = workbook_path(rows=[_row({
        "Claim - Claim number": "000-987-12345",
        "Claim - Referral type": "Full Case Management",
        "Adjuster - Referral source": "Adjuster",
        "Adjuster - Email": "adjuster@example.invalid",
        "Claimant - Address line 1": "Synthetic claimant address",
    })])
    record = excel.read_workbook(path)[0]
    assert record["form"]["referral_type"] == ""
    assert record["form"]["referral_number"] == ""
    assert record["form"]["cem_employer_request"] == "auto"
    assert "cem_employer_request" not in record["data"]
    for key in ("officeNumber", "dxCode", "xlCustomerEmail", "xlInsuredAddr1", "refType"):
        assert record["data"][key] == ""
    assert "Select CMS referral type" in record["warning"]
    assert "unknown" in record["warning"]


def test_partial_headers_row_numbers_native_dates_and_zero_padding(workbook_path):
    headers = ["Claim - Claim number", "Claim - Referral type", "Provider - Appointment date",
               "Provider - Appointment time", "CMS - CEM employer request"]
    path = workbook_path(headers=headers, rows=[
        [None] * 5, [42, "medical", datetime(2030, 1, 2), time(14, 5), "no"],
    ])
    _edit(path, lambda sheet: setattr(sheet["A3"], "number_format", "000000"))
    record = excel.read_workbook(path)[0]
    assert record["row"] == 3
    assert record["form"]["claim_number"] == "000042"
    assert record["form"]["referral_type"] == "Medical"
    assert record["form"]["claimant_first_name"] == ""
    assert record["data"]["apptDate"] == "01/02/2030"
    assert record["data"]["reviewAppointmentTime"] == "02:05 PM"
    assert record["data"]["cem_employer_request"] is False
    assert "Claimant first name is required." in record["warning"]


def test_doctor_and_attorney_names_used_when_no_facility(workbook_path):
    path = workbook_path(rows=[_row({
        "Provider - Doctor first name": "Synthetic", "Provider - Doctor last name": "Doctor",
        "Attorney - First name": "Synthetic", "Attorney - Last name": "Lawyer",
        "Attorney - Firm": "Synthetic Firm", "CMS - CEM employer request": "auto",
    })])
    data = excel.read_workbook(path)[0]["data"]
    assert data["xlProvider"] == "Synthetic Doctor"
    assert data["provType"] == ""
    assert data["xlAttorney"] == "Synthetic Lawyer"
    assert data["review_fields"]["Attorney - Firm"] == "Synthetic Firm"
    assert "cem_employer_request" not in data


@pytest.mark.parametrize("header,value,message", [
    ("CMS - Referral type", "Full Case Management", "CMS referral type must be"),
    ("CMS - CEM employer request", "SYN-secret-invalid", "CEM employer request must be"),
])
def test_invalid_explicit_choices_are_safe_errors(workbook_path, header, value, message):
    path = workbook_path(headers=[header], rows=[[value]])
    with pytest.raises(excel.ExcelValidationError) as error:
        excel.read_workbook(path)
    assert "A2" in str(error.value)
    assert message in str(error.value)
    assert value not in str(error.value)


@pytest.mark.parametrize("value", ['="SYN-secret-formula"', "#VALUE!"])
def test_rejects_formulas_and_excel_errors_even_in_unknown_columns(workbook_path, value):
    path = workbook_path(headers=["Claim - Claim number", "Unrecognized"],
                         rows=[["SYN-001", value]])
    with pytest.raises(excel.ExcelValidationError, match="B2: replace formulas or Excel errors") as error:
        excel.read_workbook(path)
    assert value not in str(error.value)


@pytest.mark.parametrize("headers,rows,sheet,message", [
    (["Claim - Claim number", "claim - claim NUMBER"], [["SYN", "SYN"]], excel.SHEET, "duplicate"),
    (["Claim - Claim number"], [["SYN"]], "Wrong Sheet", "Missing"),
    (["Unrecognized"], [["SYN"]], excel.SHEET, "No recognized"),
    (["Claim - Claim number"], [], excel.SHEET, "No filled recognized"),
    (["Claim - Claim number", "Unrecognized"], [[None, "SYN"]], excel.SHEET, "No filled recognized"),
])
def test_bad_schema_and_empty_workbooks_fail_explicitly(
        workbook_path, headers, rows, sheet, message):
    path = workbook_path(headers=headers, rows=rows, sheet=sheet)
    with pytest.raises(ValueError, match=message):
        excel.read_workbook(path)


def test_template_blank_and_existing_files_never_overwritten(workbook_path):
    path = workbook_path()
    before = path.read_bytes()
    with pytest.raises(ValueError, match="already exists"):
        excel.create_template(path)
    assert path.read_bytes() == before
    path.unlink()
    assert excel.create_template(path) == str(path)
    book = load_workbook(path)
    try:
        assert [c.value for c in book.active[1]] == list(excel.HEADERS)
        assert all(c.value is None for row in book.active.iter_rows(min_row=2) for c in row)
        assert len(book.active.data_validations.dataValidation) == 2
    finally:
        book.close()


def test_conversion_refuses_wrong_source_overwrite_and_formula(workbook_path):
    source = workbook_path("source.xlsx", headers=excel.SOURCE_HEADERS,
                           rows=[_row({"Claim - Claim number": "SYN"}, excel.SOURCE_HEADERS)],
                           sheet=excel.SOURCE_SHEET)
    destination = workbook_path("destination.xlsx")
    before = destination.read_bytes()
    with pytest.raises(ValueError, match="already exists"):
        excel.convert_subject_line_workbook(source, destination)
    assert destination.read_bytes() == before
    with pytest.raises(ValueError, match="separate"):
        excel.convert_subject_line_workbook(source, source)
    destination.unlink()
    _edit(source, lambda sheet: setattr(sheet["C2"], "value", '="SYN-secret-formula"'))
    with pytest.raises(ValueError, match="C2: replace formulas"):
        excel.convert_subject_line_workbook(source, destination)
    assert not destination.exists()
    wrong = workbook_path("wrong.xlsx", headers=["Claim - Claim number"],
                          sheet=excel.SOURCE_SHEET)
    with pytest.raises(ValueError, match="original 43"):
        excel.convert_subject_line_workbook(wrong, destination)
    assert not destination.exists()


def _ai_source(workbook_path):
    rows = []
    for index in range(7):
        values = {header: "Not found" for header in excel.AI_SOURCE_HEADERS}
        values.update({
            "Accurate": "Synthetic review", "Record ID": f"SYN-record-{index}",
            "Claim Number": f"000{index}", "Claim ID": f"SYN-not-cms-{index}",
            "First Name": "Synthetic", "Last Name": "Example",
            "Referral Type": "Full Case Management", "Diagnosis Code": "SYN.1",
            "Claims Office Number": "00002", "Employer Name": "Synthetic Employer",
            "Attorney Name": "Synthetic Attorney", "Attorney Address": "Synthetic line 1",
            "Attorney City": "Example", "Attorney State": "NY", "Attorney Zip": "00123",
            "Referral Instructions": "Separate source metadata",
            "Special Instructions": "  Synthetic\nspecial instructions  ",
            "Compensable Body/Part(s)": "Synthetic raw description",
            "Claims Case Manager E-mail Address": "manager@example.invalid",
            "Address-line-1": "Synthetic claimant address",
            "Appointment Date": datetime(2030, 4, 5),
        })
        rows.append(_row(values, excel.AI_SOURCE_HEADERS))
        if index == 0:
            rows.append([None] * len(excel.AI_SOURCE_HEADERS))
    path = workbook_path("ai.xlsx", headers=excel.AI_SOURCE_HEADERS,
                         rows=rows, sheet="Referrals")
    book = load_workbook(path)
    try:
        page = book["Referrals"]
        page.cell(2, excel.AI_SOURCE_HEADERS.index("Claim Number") + 1).font = Font(bold=True)
        providers = book.create_sheet("Providers")
        providers.append(["Record ID", "Provider Index", "Provider Name (First Name / Last Name)",
                          "Determining if Doctor or Provider Name", "Provider Phone",
                          "Provider Address", "Doctor First Name", "Doctor Last Name"])
        for index in range(7):
            providers.append([f"SYN-record-{index}", "1", "Synthetic Facility", "Facility",
                              "000123", "Synthetic provider address", None, None])
        providers.append(["SYN-record-0", "2", "Synthetic Doctor", "Doctor", "000456",
                          "Synthetic doctor address", "Synthetic", "Doctor"])
        classification = book.create_sheet("RRS Classification")
        classification.append(["Record ID", "Body Part", "Review"])
        for index in range(7):
            classification.append([f"SYN-record-{index}", "Hand", "Synthetic review"])
        review = book.create_sheet("Name Inference")
        review.append(["Record ID", "Review JSON"])
        review.append(["SYN-record-0", '{"synthetic": true}'])
        review.append(["SYN-record-6", '{"unselected": true}'])
        book.create_sheet("Schema").append(["Version", "3"])
        book.save(path)
    finally:
        book.close()
    return path


def test_ai_first_five_conversion_preserves_raw_columns_and_related_rows(workbook_path):
    source = _ai_source(workbook_path)
    destination = workbook_path("ai-cms.xlsx")
    destination.unlink()
    before = sha256(source.read_bytes()).digest()
    excel.convert_ai_workbook(source, destination)
    assert sha256(source.read_bytes()).digest() == before
    original = load_workbook(source)
    converted = load_workbook(destination)
    try:
        target = converted[excel.SHEET]
        source_column_count = len(excel.AI_SOURCE_HEADERS)
        assert [cell.value for cell in target[1]][:source_column_count] == list(excel.AI_SOURCE_HEADERS)
        for source_number, target_number in zip((2, 4, 5, 6, 7), range(2, 7)):
            for column in range(1, source_column_count + 1):
                old = original["Referrals"].cell(source_number, column)
                new = target.cell(target_number, column)
                assert (new.value, new.number_format, new.data_type) == (
                    old.value, old.number_format, old.data_type)
                assert copy(new.font) == copy(old.font)
        assert all(cell.value is None for row in target.iter_rows(
            min_row=2, min_col=source_column_count + 1) for cell in row)
        assert converted["Providers"].max_row == 7
        assert converted["RRS Classification"].max_row == 6
        assert converted["Name Inference"].max_row == 2
        assert converted["Schema"]["B1"].value == "3"
    finally:
        original.close()
        converted.close()
    records = excel.read_workbook(destination)
    assert len(records) == 5
    assert [record["form"]["claim_number"] for record in records] == [
        "0000", "0001", "0002", "0003", "0004"]
    record = records[0]
    assert record["form"]["cms_case_id"] == record["form"]["referral_number"] == ""
    assert record["form"]["referral_type"] == ""
    assert record["form"]["cem_employer_request"] == "auto"
    data = record["data"]
    assert data["dxCode"] == "SYN.1"
    assert data["bodyParts"] == "Hand"
    assert data["officeNumber"] == "00002"
    assert data["xlCustomerName"] == "Synthetic Employer"
    assert data["xlAttorney"] == "Synthetic Attorney"
    assert data["xlAttorneyAdd1"] == "Synthetic line 1"
    assert data["xlAttorneyCity"] == "Example"
    assert data["Special_Instructions"] == "  Synthetic\nspecial instructions  "
    assert data["apptDate"] == data["reviewAppointmentDate"] == "04/05/2030"
    assert data["xlCustomerEmail"] == data["xlInsuredAddr1"] == ""
    assert data["review_fields"]["Referral Instructions"] == "Separate source metadata"
    assert data["review_fields"]["Provider Phone"] == "Not found"
    assert data["xlProviderPhone"] == ""
    assert len(data["providers"]) == 2
    assert data["providers"][0]["type"] == "facility"
    assert data["providers"][1]["type"] == "doctor"
    assert data["providers"][1]["first_name"] == "Synthetic"
    assert data["providers"][0]["address"] == "Synthetic provider address"
    with pytest.raises(excel.ExcelValidationError, match="already exists"):
        excel.convert_ai_workbook(source, destination)


def test_ai_explicit_overrides_unknowns_and_unmapped_body_are_not_inferred(workbook_path):
    values = {
        "Record ID": "SYN-record", "Claim Number": "SYN-claim",
        "First Name": "Synthetic", "Last Name": "Example", "Claim ID": "SYN-not-cms",
        "Referral Type": "Medical", "Diagnosis Code": "Not found",
        "Compensable Body/Part(s)": "Raw description",
        "CMS - Diagnosis code": "SYN.override", "CMS - CEM employer request": "no",
        "CMS - Customer email": "employer@example.invalid",
        "CMS - Insured address line 1": "Synthetic insured address",
        "CMS - Special instructions": "Explicit override",
        "Special Instructions": "Source instructions",
    }
    path = workbook_path(headers=excel.HEADERS, rows=[_row(values, excel.HEADERS)])
    result = excel.read_workbook(path)[0]
    assert result["form"]["referral_type"] == "Medical"
    assert result["data"]["cem_employer_request"] is False
    assert result["data"]["dxCode"] == "SYN.override"
    assert result["data"]["bodyParts"] == ""
    assert result["data"]["Special_Instructions"] == "Explicit override"
    assert result["data"]["xlCustomerEmail"] == "employer@example.invalid"
    assert result["data"]["xlInsuredAddr1"] == "Synthetic insured address"
    assert "no mapped CMS body part" in result["warning"]
    assert result["form"]["cms_case_id"] == ""
    assert len(result["data"]["review_fields"]) == len(excel.HEADERS)


@pytest.mark.parametrize("change,message", [
    ("duplicate", "Repeated referral Record ID"),
    ("missing", "no Record ID"),
    ("classification", "Repeated RRS Classification"),
    ("insufficient", "fewer populated"),
])
def test_ai_ambiguous_identifiers_and_insufficient_rows_fail_safely(
        workbook_path, change, message):
    source = _ai_source(workbook_path)
    book = load_workbook(source)
    try:
        if change == "duplicate":
            book["Referrals"]["B4"] = book["Referrals"]["B2"].value
        elif change == "missing":
            book["Referrals"]["B2"] = None
        elif change == "classification":
            book["RRS Classification"].append(["SYN-record-0", "Hand"])
        else:
            book["Referrals"].delete_rows(7, 10)
        book.save(source)
    finally:
        book.close()
    destination = workbook_path("bad-conversion.xlsx")
    destination.unlink()
    with pytest.raises(excel.ExcelValidationError, match=message) as error:
        excel.convert_ai_workbook(source, destination)
    assert "SYN-record" not in str(error.value)
    assert not destination.exists()


def test_setup_edits_allow_partial_clears_and_return_independent_provider_data():
    values = {"dxCode": "", "apptDate": "04/05/2030",
              "providers": [{"name": "Synthetic", "type": "doctor", "phone": ""}]}
    edited = excel.validate_setup_edits(values)
    assert edited == values
    edited["providers"][0]["name"] = "Changed"
    assert values["providers"][0]["name"] == "Synthetic"
    assert excel.validate_setup_edits({}) == {}
    keys = [field["key"] for field in excel.EDITABLE_SETUP_FIELDS]
    assert "apptDate" in keys and "reviewAppointmentDate" not in keys
    assert len(keys) == len(set(keys))


@pytest.mark.parametrize("values", [
    None, [], {"dxCode": 12}, {"dxCode": None}, {"dxCode": True},
    {"review_fields": {}}, {"source_reviews": {}}, {"claim_number": "SYN"},
    {"SYN-secret-key": "SYN-secret-value"}, {"provType": "SYN-secret-value"},
    {"providers": {}}, {"providers": []}, {"providers": ["SYN-secret-value"]},
    {"providers": [{"name": 1}]}, {"providers": [{"type": "SYN-secret-value"}]},
    {"providers": [{"Record ID": "SYN-secret-value"}]},
])
def test_setup_edits_reject_types_metadata_and_unsupported_choices_safely(values):
    with pytest.raises(excel.ExcelValidationError) as error:
        excel.validate_setup_edits(values)
    assert "SYN-secret" not in str(error.value)
