from copy import deepcopy
from datetime import date
import json

import pytest
from openpyxl import load_workbook

from AI import referral_schema as schema
from AI.appointment_rules import appointment_checks, confirm_appointments, us_holidays
from AI.daily_output import save_daily_output, read_record, LEGACY_HEADERS, LEGACY_PROVIDER_HEADERS, V2_HEADERS
from tests.unit.test_referral_schema import complete_payload, payload, provider


def parse(data):
    return schema.parse_field_block(json.dumps(data))


def test_matrix_fields_unique_and_source_priority():
    assert len(schema.REQUIRED_FIELDS) == len(set(schema.REQUIRED_FIELDS)) == 63
    assert sum(map(len, schema.FIELD_GROUPS.values())) == 59
    data = payload()
    data.update({'Attorney Name': 'Section Attorney', 'Accident Description': 'Section description',
                 'Employer Contact Name': 'Customer contact', 'Employer First Name': 'Section',
                 'Diagnosis Code': 'S01', 'Special Instructions': 'Please arrange visit. Referrer Name: Excluded Person'})
    data['Special Instruction Fields'] = {'Attorney Name': 'Instruction Attorney',
        'Accident Description': 'Other description', 'Employer First Name': 'Instruction',
        'Employer Contact Name': 'Forbidden', 'Diagnosis Code': 'S02',
        'Compensable Body/Part(s)': 'Left knee', 'Additional Diagnosis Codes': 'S01; S02; S02'}
    fields = parse(data)
    assert fields['Attorney Name'] == 'Instruction Attorney'
    assert fields['Accident Description'] == 'Section description'
    assert fields['Employer First Name'] == 'Instruction'
    assert fields['Employer Contact Name'] == 'Customer contact'
    assert fields['Diagnosis Code'] == 'S01'
    assert fields['Additional Diagnosis Codes'] == 'S02'
    assert fields['Compensable Body/Part(s)'] == 'Left knee'
    assert fields['Special Instructions'] == 'Please arrange visit.'


def test_doctor_identity_and_date_only_phone_optional():
    data = complete_payload()
    record = data['Provider Information'][0]
    record.update({'Provider Name (First Name / Last Name)': 'Not found', 'Doctor First Name': 'Jane',
                   'Doctor Last Name': 'Doe', 'Provider Phone': 'Not found', 'Appointment Time': 'Not found'})
    fields = parse(data)
    assert schema.completeness(fields)['status'] == 'Passed'
    assert schema.export_payload(fields)['Provider Information'][0]['Date Only'] is True
    assert schema.export_payload(fields)['Provider Information'][0]['Provider / Facility'] == 'Not found'
    record['Doctor Last Name'] = 'Not found'
    assert schema.completeness(parse(data))['status'] == 'Failed'
    record['Provider Name (First Name / Last Name)'] = 'Example Clinic'
    assert schema.completeness(parse(data))['status'] == 'Passed'
    record['Appointment Date'] = 'Not found'
    assert schema.completeness(parse(data))['status'] == 'Failed'


def test_special_then_referral_provider_fallback_without_cross_provider_mix():
    data = payload()
    identity = {'Provider Name (First Name / Last Name)': 'Clinic A'}
    data['Special Instruction Fields'] = {'Provider Information': [provider(**identity, **{'Provider Phone': '111'})]}
    data['Referral Instruction Providers'] = [provider(**identity, **{'Provider Address': '1 Main St',
        'Appointment Date': '2026-10-06', 'Appointment Time': '2 PM'})]
    data['Provider Information'] = [provider(**identity, **{'Provider Phone': '222', 'Provider City': 'Boston'}),
                                    provider(**{'Provider Name (First Name / Last Name)': 'Clinic B', 'Provider Zip': '01234'})]
    fields = parse(data)
    assert len(fields.providers) == 2
    first = fields.providers[0]
    assert first['Provider Phone'] == '111'
    assert first['Provider Address'] == '1 Main St'
    assert first['Appointment Date'] == '2026-10-06'
    assert first['Provider Zip'] == 'Not found'


def test_one_line_addresses_split_without_guessing():
    data = payload()
    data.update({'Address-line-1': '123 Summer St Apt 2, Worcester MA 01608'})
    data['Provider Information'] = [provider(**{'Provider Address': '123 Summer St Suite 5, Worcester MA 01608'})]
    fields = parse(data)
    assert [fields[k] for k in ('Address-line-1', 'Address-line-2', 'City', 'State', 'Zip')] == [
        '123 Summer St', 'Apt 2', 'Worcester', 'MA', '01608']
    assert fields.providers[0]['Provider Address'] == '123 Summer St'
    assert fields.providers[0]['Provider Address Line 2'] == 'Suite 5'
    assert fields.providers[0]['Provider Zip'] == '01608'


@pytest.mark.parametrize('email,first,last', [('jane.doe@example.com', 'Jane', 'Doe'),
    ('info@example.com', 'Not found', 'Not found'), ('human.resources@example.com', 'Not found', 'Not found'),
    ('jdoe@example.com', 'Not found', 'Not found')])
def test_email_names_only_when_recognizable(email, first, last):
    data = payload()
    data['Special Instruction Fields'] = {'Employer Contact Email': email}
    fields = parse(data)
    assert fields['Employer First Name'] == first
    assert fields['Employer Last Name'] == last
    assert fields['Employer Contact Name'] == 'Not found'
    assert bool(fields.name_inference) == (first != 'Not found')


def test_holidays_weekend_past_confirmation_and_invalidation(tmp_path):
    data = complete_payload()
    data['Provider Information'][0]['Appointment Date'] = '2026-07-04'
    fields = parse(data)
    checks = appointment_checks(fields, date(2026, 7, 5))
    assert len(checks[0]['reasons']) == 3
    assert schema.completeness(fields, date(2026, 7, 5))['status'] == 'Failed'
    confirmed = confirm_appointments(fields, [1], date(2026, 7, 5))
    assert schema.completeness(confirmed, date(2026, 7, 5))['status'] == 'Passed'
    assert not getattr(fields, 'appointment_review', [])
    path, rid = save_daily_output(confirmed, 'synthetic.pdf', directory=tmp_path)
    loaded, status = read_record(path, rid)
    assert status['status'] == 'Passed'
    assert loaded.appointment_review[0]['confirmed_at']
    loaded.providers[0]['Appointment Date'] = '2026-07-05'
    assert schema.completeness(loaded, date(2026, 7, 5))['status'] == 'Failed'
    assert date(2026, 7, 3) in us_holidays(2026)
    assert date(2021, 12, 31) in us_holidays(2021)  # next year's New Year observed
    assert date(2026, 11, 26) in us_holidays(2026)


def test_invalid_date_cannot_be_confirmed():
    data = complete_payload()
    data['Provider Information'][0]['Appointment Date'] = '02/30/2026'
    fields = parse(data)
    assert schema.completeness(fields)['status'] == 'Failed'
    with pytest.raises(ValueError):
        confirm_appointments(fields, [1])


def test_confirmation_updates_same_daily_row(tmp_path):
    data = complete_payload()
    data['Provider Information'][0]['Appointment Date'] = '2020-07-04'
    fields = parse(data)
    path, rid = save_daily_output(fields, 'synthetic.pdf', directory=tmp_path)
    assert read_record(path, rid)[1]['status'] == 'Failed'
    confirmed = confirm_appointments(fields, [1])
    save_daily_output(confirmed, 'synthetic.pdf', directory=tmp_path, record_id=rid, update_existing=True)
    assert read_record(path, rid)[1]['status'] == 'Passed'
    book = load_workbook(path)
    assert book['Referrals'].max_row == 2
    assert book['Appointment Review'].max_row == 2
    book.close()


def test_v1_workbooks_remain_readable_and_upgrade_without_losing_records(tmp_path):
    data = parse(complete_payload())
    path, rid = save_daily_output(data, 'old.pdf', directory=tmp_path, record_id='old')
    book = load_workbook(path)
    for name, headers in [('Referrals', LEGACY_HEADERS), ('Providers', LEGACY_PROVIDER_HEADERS)]:
        sheet = book[name]
        for i in range(sheet.max_column, 0, -1):
            if sheet.cell(1, i).value not in headers:
                sheet.delete_cols(i)
    book['Schema']['B1'] = '1'
    book.save(path)
    book.close()
    loaded, status = read_record(path, rid)
    assert loaded['Special Instructions'] == 'Not found'
    assert status['status'] == 'Failed'  # new required field cannot be invented
    save_daily_output(data, 'new.pdf', directory=tmp_path, record_id='new')
    assert read_record(path, 'old')[0]['Claim Number'] == data['Claim Number']
    assert read_record(path, 'new')[1]['status'] == 'Passed'


def test_matrix_sections_and_required_manager_names():
    data = complete_payload()
    data['Claims Case Manager First Name'] = data['Claims Case Manager Last Name'] = 'Not found'
    data['Claims Case Manager Name'] = 'Saylor, Jessica'
    data['Special Instruction Fields'] = {'Claims Case Manager Name': 'Other, Person'}
    fields = parse(data)
    assert fields['Claims Case Manager First Name'] == 'Jessica'
    assert fields['Claims Case Manager Last Name'] == 'Saylor'
    output = schema.export_payload(fields)
    assert list(output['Referral Information']) == ['Referral Type', 'Referral Priority']
    assert list(output['NCM Information']) == ['NCM', 'Nurse Case Manager E-mail Address']
    assert 'Claims Case Manager Name' not in output['Case Manager Information']
    assert 'Referral Instructions' not in output['Attorney Information']
    fields['Claims Case Manager Last Name'] = 'Not found'
    assert 'Case Manager Information / Claims Case Manager Last Name' in schema.completeness(fields)['missing_fields']


def test_manager_names_fall_back_to_special_only_when_section_missing():
    data = payload()
    data['Special Instruction Fields'] = {'Claims Case Manager First Name': 'Jessica',
                                         'Claims Case Manager Last Name': 'Saylor'}
    fields = parse(data)
    assert fields['Claims Case Manager First Name'] == 'Jessica'
    assert fields['Claims Case Manager Last Name'] == 'Saylor'
    assert fields['Claims Case Manager Name'] == 'Jessica Saylor'


def test_employer_email_overrides_conflicting_names_with_provenance():
    data = payload()
    data['Special Instruction Fields'] = {'Employer First Name': 'Different', 'Employer Last Name': 'Person',
                                         'Employer Contact Email': 'jane.doe@example.com'}
    fields = parse(data)
    assert (fields['Employer First Name'], fields['Employer Last Name']) == ('Jane', 'Doe')
    assert [item['previous_value'] for item in fields.name_inference] == ['Different', 'Person']
    data['Special Instruction Fields']['Employer Contact Email'] = 'claims.team@example.com'
    fields = parse(data)
    assert (fields['Employer First Name'], fields['Employer Last Name']) == ('Different', 'Person')
    assert not fields.name_inference


def test_ncm_special_priority_fallback_and_email_name():
    data = payload()
    data['Nurse Case Manager E-mail Address'] = 'section.nurse@example.com'
    data['Special Instruction Fields'] = {'Nurse Case Manager E-mail Address': 'jane.doe@example.com'}
    fields = parse(data)
    assert fields['NCM'] == 'Jane Doe'
    assert fields['Nurse Case Manager E-mail Address'] == 'jane.doe@example.com'
    assert fields.name_inference[0]['source'] == 'Nurse Case Manager E-mail Address'
    data['Special Instruction Fields']['NCM'] = 'Documented Nurse'
    assert parse(data)['NCM'] == 'Documented Nurse'
    data['Special Instruction Fields'] = {}
    assert parse(data)['Nurse Case Manager E-mail Address'] == 'section.nurse@example.com'
    assert parse(data)['NCM'] == 'Not found'  # generic mailbox name must not be guessed


def test_v2_upgrade_preserves_manager_names_and_appointment_confirmation(tmp_path):
    data = complete_payload()
    data.update({'Claims Case Manager Name': 'Saylor, Jessica',
                 'Claims Case Manager First Name': 'Jessica', 'Claims Case Manager Last Name': 'Saylor'})
    data['Provider Information'][0]['Appointment Date'] = '2020-07-04'
    fields = confirm_appointments(parse(data), [1])
    path, _ = save_daily_output(fields, 'old.pdf', directory=tmp_path, record_id='old')
    book = load_workbook(path)
    sheet = book['Referrals']
    for i in range(sheet.max_column, 0, -1):
        if sheet.cell(1, i).value not in V2_HEADERS:
            sheet.delete_cols(i)
    book['Schema']['B1'] = '2'
    book.save(path)
    book.close()
    old, status = read_record(path, 'old')
    assert old['Claims Case Manager First Name'] == 'Jessica'
    assert status['status'] == 'Passed'
    save_daily_output(fields, 'new.pdf', directory=tmp_path, record_id='new')
    old, status = read_record(path, 'old')
    assert old['Claims Case Manager Last Name'] == 'Saylor'
    assert status['status'] == 'Passed'
    assert old.appointment_review == fields.appointment_review
    book = load_workbook(path)
    assert book['Schema']['B1'].value == '3'
    assert book['Referrals']['D2'].value == 'Passed'
    book.close()
