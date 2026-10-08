import json
from copy import deepcopy
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from AI.referral_schema import REQUIRED_FIELDS, PROVIDER_FIELDS, parse_field_block, completeness, export_payload
from AI.address_enrichment import suggest_addresses, apply_suggestions
from AI.daily_output import save_daily_output, read_record


def fields():
    data = {key: "Not found" for key in REQUIRED_FIELDS if key not in PROVIDER_FIELDS}
    data["Address-line-1"] = "123 Summer St"
    data["City"] = "Worcester"
    data["Provider Information"] = []
    return parse_field_block(json.dumps(data))


def match():
    return {"matchedAddress": "123 SUMMER ST, WORCESTER, MA, 01608",
            "addressComponents": {"city": "WORCESTER", "state": "MA", "zip": "01608"}}


def test_suggestions_do_not_mutate_and_acceptance_fills_multiple_address_fields():
    original = fields()
    before = deepcopy(original)
    report = suggest_addresses(original, lambda query: [match()])
    assert original == before
    assert completeness(original) == completeness(before)
    assert report[0]["status"] == "suggested"
    assert set(report[0]["changes"]) == {"State", "Zip"}
    accepted = apply_suggestions(original, report, [0])
    assert accepted["State"] == "MA"
    assert accepted["Zip"] == "01608"
    assert accepted["City"] == "Worcester"
    assert accepted["Address-line-2"] == "Not found"
    assert accepted.address_review[0]["status"] == "accepted"
    assert original["State"] == "Not found"
    assert "Address Review" in export_payload(accepted)


@pytest.mark.parametrize("matches, reason", [([], "No match"), ([match(), match()], "Multiple matches")])
def test_unmatched_and_ambiguous_leave_missing_fields(matches, reason):
    report = suggest_addresses(fields(), lambda _: matches)
    assert report[0]["changes"] == {}
    assert reason in report[0]["reason"]


def test_conflict_and_no_address_do_not_invent_values():
    data = fields()
    data["State"] = "CA"
    report = suggest_addresses(data, lambda _: [match()])
    assert "conflicts" in report[0]["reason"]
    data["Address-line-1"] = "Not found"
    def unexpected(_):
        pytest.fail("Cannot look up a street from city/state alone")
    report = suggest_addresses(data, unexpected)
    assert "Insufficient" in report[0]["reason"]


def test_lookup_errors_do_not_fail_extraction():
    def timeout(_):
        raise TimeoutError()
    report = suggest_addresses(fields(), timeout)
    assert report[0]["status"] == "unavailable"
    assert not report[0]["changes"]


def test_city_lookup_and_provider_isolation():
    data = fields()
    data["City"] = "Not found"
    data["State"] = "Massachusetts"
    data["Zip"] = "01608"
    record = {key: "Not found" for key in PROVIDER_FIELDS}
    record.update({"Provider Address": "123 Summer St", "Provider City": "Worcester"})
    data.providers = [record, {key: "Not found" for key in PROVIDER_FIELDS}]
    report = suggest_addresses(data, lambda _: [match()])
    assert "City" in report[0]["changes"]
    assert report[1]["target"] == "Provider 1"
    accepted = apply_suggestions(data, report, [1])
    assert accepted.providers[0]["Provider Zip"] == "01608"
    assert accepted.providers[1]["Provider Zip"] == "Not found"
    assert accepted["Provider Zip"] == "01608"
    assert accepted["City"] == "Not found"


def test_existing_fields_cannot_be_overwritten_after_lookup():
    data = fields()
    report = suggest_addresses(data, lambda _: [match()])
    data["State"] = "CA"
    with pytest.raises(ValueError, match="cannot be overwritten"):
        apply_suggestions(data, report, [0])


def test_accepted_review_updates_same_daily_row_and_persists_provenance(tmp_path):
    data = fields()
    date = datetime(2026, 9, 30, tzinfo=timezone.utc)
    path, _ = save_daily_output(data, "sample.pdf", directory=tmp_path, record_id="same", extracted_at=date)
    accepted = apply_suggestions(data, suggest_addresses(data, lambda _: [match()]), [0])
    for _ in range(2):
        save_daily_output(accepted, "sample.pdf", directory=tmp_path, record_id="same", extracted_at=date, update_existing=True)
    book = load_workbook(path)
    assert book["Referrals"].max_row == 2
    assert book["Address Review"].max_row == 2
    book.close()
    loaded, _ = read_record(path, "same")
    assert loaded["Zip"] == "01608"
    assert loaded.address_review[0]["changes"]["Zip"]["original"] == "Not found"
