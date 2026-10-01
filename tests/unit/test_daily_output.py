from datetime import datetime, timezone
import json
import importlib.util
from pathlib import Path

from openpyxl import load_workbook
import pytest

from AI import referral_schema as schema
from AI.daily_output import save_daily_output, read_record, reopen_input
from tests.unit.test_referral_schema import future_workday


def fields():
    values = {key: "Documented value" for key in schema.REQUIRED_FIELDS if key not in schema.PROVIDER_FIELDS}
    for key in schema.OPTIONAL_FIELDS:
        values[key] = "Not found"
    values.update({"Claim Number": "000123", "Claim ID": "000456", "First Name": "Jane", "Last Name": "Doe",
                   "Employer Name": "=Example", "Social Security Number": "001-02-0003"})
    values["Provider Information"] = [{key: "Provider one" for key in schema.PROVIDER_FIELDS},
                                      {key: "Provider two" for key in schema.PROVIDER_FIELDS}]
    for provider in values['Provider Information']:
        provider['Appointment Date'] = future_workday()
    return schema.parse_field_block(json.dumps(values))


def test_new_optional_fields_do_not_fail():
    data = fields()
    assert schema.completeness(data)["status"] == "Passed"
    for key in ["Employer Contact Name", "Employer Contact Mobile", "Diagnosis Code"]:
        assert data[key] == "Not found"


def test_daily_append_retry_rollover_and_read(tmp_path):
    data = fields()
    day = datetime(2026, 9, 29, 10, tzinfo=timezone.utc)
    path, _ = save_daily_output(data, "test.pdf", directory=tmp_path, record_id="one", extracted_at=day)
    save_daily_output(data, "test.pdf", directory=tmp_path, record_id="one", extracted_at=day)
    failed = fields()
    failed["Claim Number"] = "Not found"
    save_daily_output(failed, "failed.pdf", directory=tmp_path, record_id="two", extracted_at=day)
    book = load_workbook(path)
    assert book["Referrals"].max_row == 3
    assert book["Providers"].max_row == 5
    assert book["Referrals"]["D2"].value == "Passed"
    assert book["Referrals"]["D3"].value == "Failed"
    assert not any(cell.data_type == "f" for row in book["Referrals"] for cell in row)
    book.close()
    loaded, assessment = read_record(path, "one")
    assert assessment["status"] == "Passed"
    assert loaded["Claim Number"] == "000123"
    assert loaded["Social Security Number"] == "001-02-0003"
    assert len(loaded.providers) == 2
    assert reopen_input(loaded)["claimantFull"] == "Jane Doe"
    assert reopen_input(loaded)["customer"] == "=Example"
    tomorrow = day.replace(day=30)
    next_path, _ = save_daily_output(data, "test.pdf", directory=tmp_path, extracted_at=tomorrow)
    assert next_path != path
    assert len(list(tmp_path.glob("*.xlsx"))) == 2


def test_save_failure_preserves_workbook_and_can_retry(tmp_path, monkeypatch):
    import AI.daily_output as output
    data = fields()
    path, _ = save_daily_output(data, "test.pdf", directory=tmp_path, record_id="one")
    original = path.read_bytes()
    def locked(*_):
        raise PermissionError("Workbook open in Excel")
    with monkeypatch.context() as patch:
        patch.setattr(output.os, "replace", locked)
        with pytest.raises(PermissionError):
            save_daily_output(data, "test.pdf", directory=tmp_path, record_id="two")
    assert path.read_bytes() == original
    assert not list(tmp_path.glob("*.lock"))
    save_daily_output(data, "test.pdf", directory=tmp_path, record_id="two")
    assert read_record(path, "two")[1]["status"] == "Passed"


def test_reader_recomputes_status_and_rejects_formulas(tmp_path):
    path, _ = save_daily_output(fields(), "test.pdf", directory=tmp_path, record_id="one")
    book = load_workbook(path)
    headers = [c.value for c in book["Referrals"][1]]
    book["Referrals"].cell(2, headers.index("Claim Number") + 1, "Not found")
    book.save(path)
    assert read_record(path, "one")[1]["status"] == "Failed"
    book["Referrals"].cell(2, headers.index("Claim Number") + 1, "=1+1")
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="Formula"):
        read_record(path, "one")


def test_standalone_previews_excel_input_without_cms(tmp_path, capsys):
    path, _ = save_daily_output(fields(), "test.pdf", directory=tmp_path, record_id="one")
    runner_path = Path(__file__).resolve().parents[2] / "Processes/Reopen-Check/Stand-Alone/run.py"
    spec = importlib.util.spec_from_file_location("excel_standalone_test", runner_path)
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    assert runner.main(["--excel", str(path), "--record-id", "one"]) == 0
    output = json.loads(capsys.readouterr().out)
    assert output["input"]["claimNumber"] == "000123"
    assert output["input"]["claimantFull"] == "Jane Doe"
    failed = fields()
    failed["Claim Number"] = "Not found"
    save_daily_output(failed, "test.pdf", directory=tmp_path, record_id="failed")
    assert runner.main(["--excel", str(path), "--record-id", "failed", "--live"]) == 1
    assert "missing required information" in capsys.readouterr().err
