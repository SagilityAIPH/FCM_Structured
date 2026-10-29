"""Synthetic-only page-one referral identifier contract."""
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from AI.daily_output import HEADERS, read_record, save_daily_output
from AI.referral_number import read_referral_number
from AI.referral_schema import (FIELD_ONLY_PROMPT, FIELD_GROUPS, PROVIDER_FIELDS,
                                REQUIRED_FIELDS, export_payload, force_exact_field_output)

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'Final Code'))
from finalcode.field_adapter import to_all_data


def model(number='Not found'):
    values = {key: 'Not found' for key in REQUIRED_FIELDS if key not in PROVIDER_FIELDS}
    values['Referral Number'] = number
    values['Provider Information'] = [{key: 'Not found' for key in PROVIDER_FIELDS}]
    return json.dumps(values)


@pytest.mark.parametrize(('text', 'expected'), [
    ('Date Submitted to vendor:\n03/09/2026\nReferral Time (Eastern Time Zone) :\n16:47:11\nReferral Number:\n\n  10000001  \nClaimant Information', '10000001'),
    ('ReferralNumber: 01000001\nClaimant Information', '01000001'),
    ('referral #: 10000002\nClaimant Information', '10000002'),
    ('REFERRAL No.: 10000003\nClaimant Information', '10000003'),
    ('Referral Number: 10000004\nClaimant Information\nReferral Number: 10000005', '10000004'),
    ('Referral Number: invalid\nClaimant Information\nReferral Number: 10000006', '10000006'),
    ('Referral Number: 10000007\n===== PAGE 2 =====\nReferral Number: 10000008', '10000007'),
    ('Claimant Information\nReferral Number: 10000009\n===== PAGE 2 =====\nReferral Number: 10000010', '10000009'),
    ('Date Submitted to vendor:\n03/09/2026\nReferral Time:\n16:47:11\nClaim Number: 10000011\nCase ID: 10000012', 'Not found'),
    ('Referral Number:\nClaim Number: 10000013\nClaimant Information', 'Not found'),
    ('Referral Number: 10000014-1\nClaimant Information', 'Not found'),
    ('===== PAGE 1 =====\nClaimant Information\n===== PAGE 2 =====\nReferral Number: 10000015', 'Not found'),
])
def test_page_one_reader(text, expected):
    assert read_referral_number(text) == expected


def test_authoritative_value_and_exports(tmp_path):
    source = ('===== PAGE 1 =====\nDate Submitted to vendor:\n03/09/2026\n'
              'Referral Time (Eastern Time Zone) :\n16:47:11\n'
              'Referral Number:\n01000001\nClaimant Information\n')
    display, fields = force_exact_field_output(model('10000002'), source)
    assert 'Referral Number: 01000001' in display
    assert fields['Referral Number'] == '01000001'
    assert export_payload(fields)['Claim Information']['Referral Number'] == '01000001'
    assert any(flag['field'] == 'Referral Number' and flag['kind'] == 'disagreement'
               and flag['reader_value'] == '01000001' for flag in fields.review_flags)
    assert to_all_data(fields)['referralNumber'] == '01000001'
    path, ident = save_daily_output(fields, 'synthetic.pdf', directory=tmp_path)
    book = load_workbook(path)
    try:
        assert [cell.value for cell in book['Referrals'][1]] == HEADERS
        assert HEADERS[-1] == 'Referral Number'
        cell = book['Referrals'].cell(2, len(HEADERS))
        assert cell.value == '01000001' and cell.data_type == 's' and cell.number_format == '@'
    finally:
        book.close()
    restored, _ = read_record(path, ident)
    assert restored['Referral Number'] == '01000001'
    from AI.referral_json import _fields
    assert _fields(fields)['Referral Number'] == '01000001'


def test_old_v3_workbook_migrates_without_shifting_existing_columns(tmp_path):
    _, fields = force_exact_field_output(model(), 'Referral Number: 01000031')
    path, ident = save_daily_output(fields, 'synthetic.pdf', directory=tmp_path, record_id='old')
    book = load_workbook(path)
    try:
        book['Referrals'].delete_cols(len(HEADERS))
        book['Schema']['B1'] = '3'
        book.save(path)
    finally:
        book.close()
    _, updated = force_exact_field_output(model(), 'Referral Number: 01000032')
    save_daily_output(updated, 'synthetic.pdf', directory=tmp_path, record_id='new')
    book = load_workbook(path)
    try:
        assert book['Schema']['B1'].value == '4'
        assert [c.value for c in book['Referrals'][1]] == HEADERS
        assert book['Referrals'].cell(2, len(HEADERS)).value == 'Not found'
        assert book['Referrals'].cell(3, len(HEADERS)).value == '01000032'
    finally:
        book.close()
    assert read_record(path, ident)[0]['Referral Number'] == 'Not found'


def test_missing_label_rejects_model_guess_and_is_optional():
    _, fields = force_exact_field_output(model('10000020'), 'Claim Number: 10000020\nReferral Time: 16:47:11')
    assert fields['Referral Number'] == 'Not found'
    assert export_payload(fields)['Claim Information']['Referral Number'] == 'Not found'
    assert to_all_data(fields)['referralNumber'] == ''
    assert not any(flag['field'] == 'Referral Number' for flag in fields.review_flags)


def test_prompt_and_schema_contract():
    assert 'Referral Number' in FIELD_GROUPS['Claim Information']
    for phrase in ('header at the top of page 1', 'Date Submitted to vendor', 'Referral Time',
                   'copy digits only, exactly as printed', 'Never use the claim number',
                   'a phone or a case/office ID', 'Return Not found if absent'):
        assert phrase in FIELD_ONLY_PROMPT
