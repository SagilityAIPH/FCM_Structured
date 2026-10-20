from contextlib import nullcontext
import json
import sys
import threading
import time
from types import SimpleNamespace

import fitz
from openpyxl import load_workbook
import pytest

from AI.workbench.backend import WorkbenchAPI
from AI.workbench import backend
from AI import address_enrichment
from AI.daily_output import read_record
from tests.unit.test_referral_schema import complete_payload


@pytest.fixture
def workbench(tmp_path, monkeypatch):
    api = WorkbenchAPI(tmp_path / 'output')
    api._key = 'test-secret-not-real'
    monkeypatch.setattr(backend.core, 'create_bedrock_client', lambda *args: nullcontext(None))
    monkeypatch.setattr(address_enrichment, 'census_lookup', lambda query: [])
    path = tmp_path / 'synthetic.pdf'
    with fitz.open() as pdf:
        pdf.new_page().insert_text((40, 60), 'Synthetic referral for offline bridge testing.')
        pdf.save(path)
    identifier = api._import_paths([path])[0]
    return api, identifier, path


def extract(monkeypatch, data):
    result, fields = backend.core.force_exact_field_output(json.dumps(data))
    monkeypatch.setattr(backend.core, 'run_reasoning', lambda *args: (result, fields, json.dumps(data), .1))


def finish(api):
    api._thread.join(10)
    assert not api._thread.is_alive()
    return api.get_state()['value']


def test_prepared_capture_counts_full_text_and_reader_limit_remains_visible(workbench, monkeypatch):
    reader, file_id, _ = workbench
    extract(monkeypatch, complete_payload())
    reader._settings['document_limit'] = 12
    sent = []
    original = backend.core.run_reasoning
    def observed(*args):
        sent.append(args[2])
        return original(*args)
    monkeypatch.setattr(backend.core, 'run_reasoning', observed)
    assert reader.start_processing([file_id])['ok']
    assert finish(reader)['stage'] == 'Queue finished. Review validation and Excel save status.'
    detail = reader.get_document(file_id)['value']
    assert len(sent[0]) == detail['sent_characters'] == 12
    assert detail['total_characters'] > detail['sent_characters'] > 0

    memory = WorkbenchAPI(memory_only=True)
    assert memory.get_state()['value']['stage'] == 'Waiting for RRS capture'
    assert WorkbenchAPI().get_state()['value']['stage'] == 'Ready to import documents'
    capture_id = memory._import_text('synthetic capture', 'Synthetic referral captured from RRS.')
    fields = backend.core.force_exact_field_output(json.dumps(complete_payload()))[1]
    memory._settings['document_limit'] = 12  # Prepared extraction does not apply this limit.
    assert memory.start_processing([capture_id], prepared_fields=fields)['ok']
    captured = memory.get_document(capture_id)['value']
    assert captured['sent_characters'] == captured['total_characters'] == len(captured['text'])
    assert captured['sent_characters'] > memory._settings['document_limit']
    assert captured['workbook'] == ''


def test_memory_only_queue_stage_never_claims_excel(workbench, monkeypatch):
    reader, _, _ = workbench
    memory = WorkbenchAPI(memory_only=True)
    memory._key = reader._key
    monkeypatch.setattr(memory, '_enrich', lambda fields: (fields, []))
    identifier = memory._import_text('synthetic capture', 'Synthetic captured referral.')
    extract(monkeypatch, complete_payload())
    assert memory.start_processing([identifier])['ok']
    assert finish(memory)['stage'] == 'Queue finished. Review validation.'
    assert memory.get_document(identifier)['value']['workbook'] == ''


def test_import_auto_extracts_enriches_and_saves_daily(workbench, monkeypatch):
    api, identifier, path = workbench
    data = complete_payload()
    data.update({'Address-line-1': '123 Summer St', 'City': 'Worcester', 'State': 'Not found', 'Zip': 'Not found'})
    extract(monkeypatch, data)
    monkeypatch.setattr(address_enrichment, 'census_lookup', lambda query: [{
        'matchedAddress': '123 SUMMER ST, WORCESTER, MA, 01608',
        'addressComponents': {'city': 'WORCESTER', 'state': 'MA', 'zip': '01608'},
    }])
    api._window = type('Window', (), {'create_file_dialog': lambda *_args, **_kw: (str(path),)})()
    monkeypatch.setitem(sys.modules, 'webview', SimpleNamespace(FileDialog=SimpleNamespace(OPEN=1)))
    assert api.choose_documents()['value'] == [identifier]  # existing imports aren't duplicated
    state = finish(api)
    assert len(state['documents']) == 1
    doc = api._documents[identifier]
    assert doc['status'] == 'Completed'
    assert doc['fields']['State'] == 'MA'
    assert doc['fields']['Zip'] == '01608'
    saved, checked = read_record(doc['workbook'], doc['record_id'])
    assert checked['status'] == 'Passed'
    assert saved.address_review[0]['applied_automatically']
    assert 'test-secret' not in json.dumps(state)
    assert api.render_page(identifier, 0, .5)['value'].startswith('data:image/png;base64,')
    assert not api.render_page(identifier, 7, .5)['ok']


def test_conflicting_address_is_not_applied_and_failed_result_is_saved(workbench, monkeypatch):
    api, identifier, _ = workbench
    data = complete_payload()
    data.update({'Address-line-1': '123 Summer St', 'City': 'Worcester', 'State': 'MA', 'Zip': 'Not found'})
    extract(monkeypatch, data)
    monkeypatch.setattr(address_enrichment, 'census_lookup', lambda query: [{
        'matchedAddress': '123 OTHER ST, WORCESTER, MA, 01608',
        'addressComponents': {'city': 'WORCESTER', 'state': 'MA', 'zip': '01608'},
    }])
    assert api.start_processing([identifier])['ok']
    finish(api)
    doc = api._documents[identifier]
    assert doc['status'] == 'Needs Review'
    saved, checked = read_record(doc['workbook'], doc['record_id'])
    assert saved['Zip'] == 'Not found'
    assert checked['status'] == 'Failed'
    assert not getattr(saved, 'address_review', [])


def test_edit_preserves_provider_records_revalidates_and_updates_same_excel_row(workbench, monkeypatch):
    api, identifier, _ = workbench
    data = complete_payload()
    data['Claim ID'] = 'Not found'
    extract(monkeypatch, data)
    api.start_processing([identifier]); finish(api)
    doc = api._documents[identifier]
    record_id = doc['record_id']
    revision = doc['revision']
    edits = [{'key': 'Claim ID', 'provider_index': None, 'value': '=literal-claim'},
             {'key': 'Provider City', 'provider_index': 0, 'value': 'Worcester'}]
    assert api.save_fields(identifier, edits, revision)['ok']
    finish(api)
    assert doc['record_id'] == record_id
    assert doc['status'] == 'Completed'
    saved, _ = read_record(doc['workbook'], record_id)
    assert saved['Claim ID'] == '=literal-claim'
    assert saved.providers[0]['Provider City'] == 'Worcester'
    book = load_workbook(doc['workbook'])
    assert book['Referrals'].max_row == 2
    book.close()
    assert not api.save_fields(identifier, edits, revision)['ok']
    assert api.get_document(identifier)['value']['groups'][0]['fields'][0]['confidence'] is None


def test_doctor_display_is_separate_from_editable_facility(workbench, monkeypatch):
    from tests.unit.test_provider_merge_decisions import cases
    api, identifier, _ = workbench
    data = complete_payload()
    data['Provider Information'] = cases()['C']
    extract(monkeypatch, data)
    api.start_processing([identifier]); finish(api)
    group = next(group for group in api.get_document(identifier)['value']['groups'] if group['name'] == 'Provider 1')
    row = next(row for row in group['fields'] if row['label'] == 'Doctor / Facility')
    assert row['display_value'] == 'Alex Sample'
    assert row['value'] == 'Synthetic Clinic'
    assert row['key'] == 'Provider Name (First Name / Last Name)'
    assert row['conditional']


def test_token_failure_does_not_reuse_old_success(workbench, monkeypatch):
    api, identifier, _ = workbench
    extract(monkeypatch, complete_payload())
    api.start_processing([identifier]); finish(api)
    def fail(*args):
        raise RuntimeError('Token budget exhausted test-secret-not-real')
    monkeypatch.setattr(backend.core, 'run_reasoning', fail)
    api.start_processing([identifier]); finish(api)
    result = api.get_document(identifier)['value']
    assert result['status'] == 'Failed'
    assert result['groups'] == [] and result['next_step'] is None
    assert result['error'] == 'Token budget exhausted [redacted]'


def test_settings_unblock_waiting_imports_and_busy_rejects_edits(workbench, monkeypatch):
    api, identifier, _ = workbench
    extract(monkeypatch, complete_payload())
    gate = threading.Event()
    run = backend.core.run_reasoning
    monkeypatch.setattr(backend.core, 'run_reasoning', lambda *args: (gate.wait(5), run(*args))[1])
    api._key = ''
    api._queue_pending()
    assert not api._busy
    settings = api.get_state()['value']['settings']
    assert settings['max_tokens'] == 32768
    assert api.save_settings(settings, 'test-secret-not-real')['ok']
    assert api._busy
    assert not api.save_settings(settings, '')['ok']
    assert not api.start_processing([identifier])['ok']
    gate.set(); finish(api)


def test_workbook_save_error_is_visible_and_retry_is_idempotent(workbench, monkeypatch):
    api, identifier, _ = workbench
    extract(monkeypatch, complete_payload())
    real_save = backend.save_daily_output
    def locked(*args, **kwargs):
        raise PermissionError('Workbook open in Excel')
    monkeypatch.setattr(backend, 'save_daily_output', locked)
    api.start_processing([identifier]); finish(api)
    assert api.get_document(identifier)['value']['save_error']
    monkeypatch.setattr(backend, 'save_daily_output', real_save)
    assert api.retry_save(identifier)['ok']
    assert api.retry_save(identifier)['ok']
    doc = api._documents[identifier]
    assert not doc['save_error']
    book = load_workbook(doc['workbook'])
    assert book['Referrals'].max_row == 2
    book.close()


def test_empty_provider_can_be_completed_through_form(workbench, monkeypatch):
    api, identifier, _ = workbench
    data = complete_payload()
    provider = data['Provider Information'][0]
    data['Provider Information'] = []
    extract(monkeypatch, data)
    api.start_processing([identifier]); finish(api)
    doc = api.get_document(identifier)['value']
    group = next(g for g in doc['groups'] if g['name'] == 'Provider 1')
    edits = [{'key': f['key'], 'provider_index': 0, 'value': provider[f['key']]} for f in group['fields']]
    assert api.save_fields(identifier, edits, doc['revision'])['ok']
    finish(api)
    assert api.get_document(identifier)['value']['next_step'] == 'Passed'
    assert len(api._documents[identifier]['fields'].providers) == 1


def test_editing_appointment_invalidates_previous_confirmation(workbench, monkeypatch):
    api, identifier, _ = workbench
    data = complete_payload()
    data['Provider Information'][0]['Appointment Date'] = '2020-07-04'
    extract(monkeypatch, data)
    api.start_processing([identifier]); finish(api)
    doc = api.get_document(identifier)['value']
    assert doc['validation']['confirmation_required']
    assert api.confirm_appointment(identifier, 1)['ok']
    doc = api.get_document(identifier)['value']
    assert doc['next_step'] == 'Passed'
    assert api.save_fields(identifier, [{'key': 'Appointment Date', 'provider_index': 0, 'value': '2020-07-05'}], doc['revision'])['ok']
    finish(api)
    assert api.get_document(identifier)['value']['next_step'] == 'Failed'
    assert api.get_document(identifier)['value']['validation']['confirmation_required']


def test_unsaved_state_is_pushed_and_close_handler_never_queries_page():
    api = WorkbenchAPI()
    assert api.set_unsaved(True) == {'ok': True, 'value': True} and api._unsaved is True
    assert api.set_unsaved(False)['ok'] and api._unsaved is False
    # pywebview runs `closing` on the UI thread; evaluate_js there deadlocks the window.
    from pathlib import Path
    launcher = (Path(backend.__file__).parent / 'launcher.py').read_text(encoding='utf-8')
    assert 'evaluate_js' not in launcher


def test_import_text_matches_file_shape_and_pipeline(workbench):
    api, file_id, _ = workbench
    text = '  Synthetic referral — text capture.  '
    identifier = api._import_text('synthetic capture.pdf', text)
    doc = api._documents[identifier]
    assert doc.keys() == api._documents[file_id].keys()
    assert doc['path'] is None and doc['content'] == b''
    assert doc['name'] == 'synthetic capture.pdf'
    pages = backend.core.extract_pages_from_txt(text.encode('utf-8'))
    assert doc['text'] == backend.core.join_pages(pages)
    assert doc['pages'] == len(pages)
    detail = api.get_document(identifier)['value']
    assert not detail['is_pdf']
    assert detail['preview_message'] == 'No preview — referral text captured from RRS'


@pytest.mark.parametrize('text', ['', ' \t\n\r '])
def test_import_text_rejects_empty_content(workbench, text):
    api, _, _ = workbench
    before = set(api._documents)
    with pytest.raises(ValueError, match='No selectable text found'):
        api._import_text('empty capture', text)
    reply = api.import_text('empty capture', text)
    assert not reply['ok'] and 'No selectable text found' in reply['error']
    assert set(api._documents) == before


def test_text_only_render_page_has_successful_no_preview_response(workbench):
    api, _, _ = workbench
    identifier = api._import_text('synthetic.pdf', 'Synthetic referral capture.')
    assert api.render_page(identifier, 0) == {
        'ok': True, 'value': {'available': False, 'message': 'No preview — referral text captured from RRS'},
    }


def test_text_captures_do_not_collide_and_bridge_respects_busy(workbench):
    api, file_id, path = workbench
    first = api.import_text('same label', 'Synthetic capture one.')
    second = api.import_text('same label', 'Synthetic capture two.')
    assert first['ok'] and second['ok'] and first['value'] != second['value']
    assert len(api._documents) == 3
    assert api.get_state()['value']['stage'] == 'Imported referral text'
    assert api._import_paths([path]) == [file_id]
    api._busy = True
    assert not api.import_text('busy capture', 'Synthetic text.')['ok']
    assert len(api._documents) == 3


def test_text_capture_can_extract_save_fields_validate_and_export(workbench, monkeypatch, tmp_path):
    api, _, _ = workbench
    identifier = api._import_text('synthetic capture.pdf', 'Synthetic referral capture.')
    data = complete_payload()
    data['Claim ID'] = 'Not found'
    extract(monkeypatch, data)
    assert api.start_processing([identifier])['ok']
    finish(api)
    doc = api._documents[identifier]
    assert api.get_document(identifier)['value']['next_step'] == 'Failed'
    record_id = doc['record_id']
    assert api.save_fields(identifier, [{'key': 'Claim ID', 'provider_index': None,
                                       'value': 'SYNTHETIC-TEXT-ONLY'}], doc['revision'])['ok']
    finish(api)
    assert api.get_document(identifier)['value']['next_step'] == 'Passed'
    assert doc['record_id'] == record_id and not doc['save_error']
    saved, checked = read_record(doc['workbook'], record_id)
    assert saved['Claim ID'] == 'SYNTHETIC-TEXT-ONLY' and checked['status'] == 'Passed'
    target = tmp_path / 'synthetic-export.json'
    def dialog(*args, **kwargs):
        assert kwargs['save_filename'] == 'synthetic capture.json'
        return str(target)
    api._window = SimpleNamespace(create_file_dialog=dialog)
    monkeypatch.setitem(sys.modules, 'webview', SimpleNamespace(FileDialog=SimpleNamespace(SAVE=1)))
    assert api.export_document(identifier, 'json')['ok']
    assert 'SYNTHETIC-TEXT-ONLY' in target.read_text(encoding='utf-8')


@pytest.mark.parametrize("final_code", [False, True])
def test_launch_window_preserves_topmost_policy_through_native_layouts(monkeypatch, final_code):
    from functools import partial
    from pathlib import Path
    from unittest.mock import Mock
    from AI.workbench import launcher, window_layout

    class Size:
        def __init__(self, width, height):
            self.Width, self.Height = width, height

    class Event:
        def __init__(self):
            self.handler = None
        def __iadd__(self, handler):
            self.handler = handler
            return self

    class Form:
        DeviceDpi = 96
        InvokeRequired = False
        WindowState = 'Normal'
        Left = 78
        Top = 78
        MaximizeBox = True
        def __init__(self):
            self.Size = Size(480, 900)
            self.MinimumSize = Size(0, 0)
            self.MaximumSize = Size(0, 0)
            self.DpiChanged = Event()
            self.ResizeEnd = Event()
        @property
        def Height(self):
            return self.Size.Height
        def BeginInvoke(self, callback):
            callback()

    form = Form()
    screen = SimpleNamespace(FromControl=lambda _: SimpleNamespace(
        WorkingArea=SimpleNamespace(Height=1392, Left=0, Top=0, Right=1920, Bottom=1392)))
    monkeypatch.setitem(sys.modules, 'System', SimpleNamespace(Action=lambda callback: callback))
    monkeypatch.setitem(sys.modules, 'System.Drawing', SimpleNamespace(Size=Size))
    monkeypatch.setitem(sys.modules, 'System.Windows.Forms', SimpleNamespace(
        FormWindowState=SimpleNamespace(Minimized='Minimized'), Screen=screen))
    window = SimpleNamespace(native=form, events=SimpleNamespace(
        before_show=Event(), closing=Event(), loaded=Event()))

    def create_window(*args, **kwargs):
        form.TopMost = kwargs["on_top"]
        return window

    create = Mock(side_effect=create_window)
    start = Mock(side_effect=lambda **kwargs: window.events.before_show.handler())
    monkeypatch.setitem(sys.modules, "webview", SimpleNamespace(create_window=create, start=start))
    monkeypatch.setattr(launcher, "load_dotenv", Mock())
    # Use an existing synthetic-test source as the mock page, not generated frontend output.
    monkeypatch.setattr(launcher, "launch_window", partial(launcher.launch_window, index=__file__))
    if final_code:
        monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[2] / "Final Code"))
        from finalcode import app
        api = app.launch()
    else:
        api = SimpleNamespace(_busy=False, _workflow_active=False, _unsaved=False)
        launcher.launch_window(api=api)
    assert create.call_args.kwargs["on_top"] is (not final_code)
    assert api._window is window
    start.assert_called_once_with(gui="edgechromium", http_server=True, private_mode=True)
    assert form.TopMost is (not final_code)
    assert (form.Size.Width, form.Size.Height) == (480, 900)
    form.WindowState = 'Minimized'
    form.Size = Size(160, 28)
    form.Left = form.Top = -32000
    assert form.ResizeEnd.handler is not None
    form.ResizeEnd.handler()
    assert (form.Size.Width, form.Size.Height, form.Left, form.Top) == (160, 28, -32000, -32000)
    window._set_compact_layout(True)
    assert form.Size.Height == 28
    form.WindowState = 'Normal'
    form.Size = Size(480, 900)  # WinForms applies its intact RestoreBounds.
    form.Left = form.Top = 78
    window._set_compact_layout(False)
    assert (form.Size.Width, form.Size.Height) == (480, 900)
    window._set_compact_layout(True)
    assert form.Size.Height == window_layout.COMPACT_HEIGHT
    assert form.TopMost is (not final_code)
    window._set_compact_layout(False)
    assert form.Size.Height == 900
    form.DpiChanged.handler()
    form.ResizeEnd.handler()
    assert form.TopMost is (not final_code)


def test_compact_state_headless_and_injected_resize(workbench):
    api, _, _ = workbench
    assert api.get_state()['value']['compact'] is False
    assert api.set_compact(True) == {'ok': True, 'value': True}
    assert api.get_state()['value']['compact'] is True
    calls = []
    def resize(value):
        # Simulate a UI callback that needs the state lock: must not deadlock.
        polled = []
        thread = threading.Thread(target=lambda: polled.append(api.get_state()))
        thread.start(); thread.join(2)
        assert not thread.is_alive()
        assert polled[0]['ok']
        calls.append(value)
    api._window = SimpleNamespace(_set_compact_layout=resize)
    assert api.set_compact(False)['ok']
    assert calls == [False]
    assert api.get_state()['value']['compact'] is False


@pytest.mark.parametrize('value', [None, 0, 1, 'false', [], {}])
def test_compact_rejects_non_boolean(workbench, value):
    api, _, _ = workbench
    assert not api.set_compact(value)['ok']
    assert api.get_state()['value']['compact'] is False


def test_compact_resize_failure_does_not_publish_success(workbench):
    api, _, _ = workbench
    api._window = SimpleNamespace()
    assert not api.set_compact(True)['ok']
    assert api.get_state()['value']['compact'] is False


def test_external_stage_idle_and_progress_copy(workbench):
    api, _, _ = workbench
    assert api.set_stage('Searching CMS cases…', 1, 3)['ok']
    state = api.get_state()['value']
    assert state['stage'] == 'Searching CMS cases…' and not state['busy']
    assert state['progress'] == {'done': 1, 'total': 3}
    state['progress']['done'] = 99
    assert api.get_state()['value']['progress']['done'] == 1
    assert api.set_stage('Validating customer…')['ok']
    assert api.get_state()['value']['progress'] == api._progress


@pytest.mark.parametrize('fails', [False, True])
def test_external_stage_wins_running_job_until_next_start(workbench, fails):
    api, _, _ = workbench
    gate = threading.Event()
    def work():
        assert gate.wait(5)
        with api._lock:
            api._stage = 'Internal completion'
            api._progress = {'done': 2, 'total': 2}
        if fails:
            raise RuntimeError('Synthetic job failure')
    api._start('Internal start', work)
    try:
        assert api.set_stage('Filling Subject Line Builder…', 1, 4)['ok']
        with pytest.raises(ValueError):
            api._start('Rejected concurrent start', lambda: None)
        assert api.get_state()['value']['stage'] == 'Filling Subject Line Builder…'
    finally:
        gate.set()
    state = finish(api)
    assert state['stage'] == 'Filling Subject Line Builder…'
    assert state['progress'] == {'done': 1, 'total': 4}
    api._start('Next job', lambda: None)
    state = finish(api)
    assert state['stage'] == 'Next job'
    assert state['progress'] == {'done': 2, 'total': 2}


@pytest.mark.parametrize('args', [(None,), ('',), ('  ',), (4,), ('Stage', 1),
                                  ('Stage', None, 2), ('Stage', -1, 2),
                                  ('Stage', 3, 2), ('Stage', True, 2),
                                  ('Stage', 1, 2.5), ('Stage', '1', 2)])
def test_external_stage_invalid_input_is_atomic(workbench, args):
    api, _, _ = workbench
    api.set_stage('Valid synthetic stage', 0, 2)
    assert not api.set_stage(*args)['ok']
    state = api.get_state()['value']
    assert state['stage'] == 'Valid synthetic stage'
    assert state['progress'] == {'done': 0, 'total': 2}


@pytest.mark.parametrize('kind,answer', [('confirm', True), ('confirm', False), ('notify', True)])
def test_live_prompt_roundtrip_without_deadlock(workbench, monkeypatch, kind, answer):
    from AI.workbench import window_layout
    api, _, _ = workbench
    api._compact = True
    ready = threading.Event()
    def resize(window, compact):
        # A native UI Invoke may poll from a different thread. No state lock held.
        poll = threading.Thread(target=lambda: (api.get_state(), ready.set()), daemon=True)
        poll.start()
        assert ready.wait(2), 'Resize deadlocked against get_state'
        assert compact is False
    monkeypatch.setattr(window_layout, 'set_compact', resize)
    returned = []
    worker = threading.Thread(target=lambda: returned.append(getattr(api, kind)('Synthetic prompt', 'Synthetic instruction')), daemon=True)
    worker.start()
    assert ready.wait(2)
    prompt = api.get_state()['value']['pending_prompt']
    try:
        assert prompt['kind'] == kind
        assert prompt['title'] == 'Synthetic prompt'
        assert worker.is_alive()
        assert not api.answer_prompt('stale-id', True)['ok']
        assert not api.answer_prompt(prompt['id'], 'yes')['ok']
        assert not api.set_compact(True)['ok']
        with pytest.raises(RuntimeError, match='already pending'):
            api.confirm('Other prompt', 'Must not overwrite')
        if kind == 'notify':
            assert not api.answer_prompt(prompt['id'], False)['ok']
        assert api.answer_prompt(prompt['id'], answer)['ok']
    finally:
        api.answer_prompt(prompt['id'], True)
    worker.join(2)
    assert not worker.is_alive()
    assert returned == ([answer] if kind == 'confirm' else [None])
    assert api.get_state()['value']['pending_prompt'] is None
    assert api.get_state()['value']['compact'] is False
    assert not api.answer_prompt(prompt['id'], answer)['ok']


def test_answer_without_pending_prompt_is_rejected(workbench):
    api, _, _ = workbench
    assert api.get_state()['value']['pending_prompt'] is None
    assert api.answer_prompt('not-pending', True) == {'ok': False, 'error': 'Prompt is no longer pending.'}


def test_prompt_resize_failure_cleans_up(workbench, monkeypatch):
    from AI.workbench import window_layout
    api, _, _ = workbench
    def fail(*args):
        raise RuntimeError('Synthetic resize failure')
    monkeypatch.setattr(window_layout, 'set_compact', fail)
    with pytest.raises(RuntimeError, match='Synthetic resize failure'):
        api.confirm('Synthetic prompt', 'Synthetic instruction')
    assert api.get_state()['value']['pending_prompt'] is None


def test_workflow_result_snapshot_and_clear(workbench):
    api, _, _ = workbench
    result = {'outcome': 'stopped', 'reason': 'tire_customer_open_tcm', 'customer': 'SYNTHETIC EMPLOYER',
              'cases': [{'cms_caseNum': 'SYNTHETIC-CASE'}], 'missing_fields': [], 'stages': ['CMS'], 'payloads': {}}
    assert api.set_workflow_result(result)['ok']
    result['cases'].clear()
    state = api.get_state()['value']
    assert state['workflow_result']['cases'] == [{'cms_caseNum': 'SYNTHETIC-CASE'}]
    state['workflow_result']['cases'].clear()
    assert api.get_state()['value']['workflow_result']['cases']
    assert not api.set_workflow_result({'outcome': 'unknown'})['ok']
    assert api.set_workflow_result(None)['ok']
    assert api.get_state()['value']['workflow_result'] is None


def _await_review(api, timeout=5):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        review = api.get_state()['value']['pending_review']
        if review:
            return review
        time.sleep(0.01)
    raise AssertionError('No review became pending.')


def _extracted(api, identifier, monkeypatch):
    """Give a document fields, which wait_for_review requires."""
    extract(monkeypatch, complete_payload())
    api.start_processing([identifier])
    finish(api)


def test_pending_review_is_visible_and_approval_returns_edited_fields(workbench, monkeypatch):
    api, identifier, _ = workbench
    _extracted(api, identifier, monkeypatch)
    answer = {}

    def worker():
        answer['fields'] = api.wait_for_review(identifier)

    thread = threading.Thread(target=worker, daemon=True)
    thread.start()
    review = _await_review(api)
    assert review['document_id'] == identifier
    assert review['id'] != identifier
    assert set(review) == {'id', 'document_id'}
    review['document_id'] = 'mutated snapshot'
    assert api.get_state()['value']['pending_review']['document_id'] == identifier
    assert api.answer_review(review['id'], True)['ok']
    thread.join(5)
    assert not thread.is_alive()
    assert answer['fields']['Claim Number'] == complete_payload()['Claim Number']
    assert answer['fields'] is not api._documents[identifier]['fields']
    assert api.get_state()['value']['pending_review'] is None


def test_approval_returns_saved_fields_from_memory_only_review(monkeypatch):
    from AI.workbench import window_layout
    api = WorkbenchAPI(memory_only=True)
    monkeypatch.setattr(window_layout, 'set_compact', lambda *_: None)
    monkeypatch.setattr(address_enrichment, 'census_lookup', lambda _: [])
    identifier = api._import_text('synthetic.txt', 'Synthetic captured text.')
    fields = backend.core.force_exact_field_output(json.dumps(complete_payload()))[1]
    assert api.start_processing([identifier], prepared_fields=fields)['ok']
    returned = {}
    thread = threading.Thread(target=lambda: returned.update(fields=api.wait_for_review(identifier)), daemon=True)
    thread.start()
    review = _await_review(api)
    try:
        detail = api.get_document(identifier)['value']
        assert api.save_fields(identifier, [{'key': 'Accident Description', 'provider_index': None,
                                            'value': 'Synthetic correction'}], detail['revision'])['ok']
        finish(api)
        assert api.set_unsaved(False)['ok']
        assert api.answer_review(review['id'], True)['ok']
    finally:
        api.answer_review(review['id'], False)
    thread.join(5)
    assert not thread.is_alive()
    assert returned['fields']['Accident Description'] == 'Synthetic correction'
    assert api.get_state()['value']['pending_review'] is None


def test_review_cancel_returns_none(workbench, monkeypatch):
    api, identifier, _ = workbench
    _extracted(api, identifier, monkeypatch)
    answer = {}
    thread = threading.Thread(target=lambda: answer.update(fields=api.wait_for_review(identifier)),
                              daemon=True)
    thread.start()
    review = _await_review(api)
    assert api.answer_review(review['id'], False)['ok']
    thread.join(5)
    assert not thread.is_alive()
    assert answer['fields'] is None
    assert api.get_state()['value']['pending_review'] is None


def test_unsaved_review_approval_is_rejected_until_saved(workbench, monkeypatch):
    api, identifier, _ = workbench
    _extracted(api, identifier, monkeypatch)
    thread = threading.Thread(target=lambda: api.wait_for_review(identifier), daemon=True)
    thread.start()
    review = _await_review(api)
    try:
        assert api.set_unsaved(True)['ok']
        assert api.answer_review(review['id'], True) == {
            'ok': False, 'error': 'Save field changes before approving.'}
        assert thread.is_alive()
        assert api.get_state()['value']['pending_review'] == review
        assert api.set_unsaved(False)['ok']
        assert api.answer_review(review['id'], True)['ok']
    finally:
        api.answer_review(review['id'], False)
    thread.join(5)
    assert not thread.is_alive()


def test_polling_and_editing_survive_a_parked_review_worker(workbench, monkeypatch):
    """The worker blocks with the state lock released, or the operator could
    never reach the button that unblocks it."""
    api, identifier, _ = workbench
    _extracted(api, identifier, monkeypatch)
    thread = threading.Thread(target=lambda: api.wait_for_review(identifier), daemon=True)
    thread.start()
    review = _await_review(api)
    for _ in range(5):
        assert api.get_state()['ok']
    assert api.get_state()['value']['documents']
    assert api.answer_review(review['id'], True)['ok']
    thread.join(5)
    assert not thread.is_alive()


@pytest.mark.parametrize('answer', [1, 'yes', None, 0])
def test_review_answer_must_be_a_real_boolean(workbench, monkeypatch, answer):
    api, identifier, _ = workbench
    _extracted(api, identifier, monkeypatch)
    thread = threading.Thread(target=lambda: api.wait_for_review(identifier), daemon=True)
    thread.start()
    review = _await_review(api)
    assert not api.answer_review(review['id'], answer)['ok']
    assert api.answer_review(review['id'], True)['ok']
    thread.join(5)


def test_stale_or_unknown_review_identifier_is_rejected(workbench, monkeypatch):
    api, identifier, _ = workbench
    _extracted(api, identifier, monkeypatch)
    assert not api.answer_review('no-such-review', True)['ok']
    thread = threading.Thread(target=lambda: api.wait_for_review(identifier), daemon=True)
    thread.start()
    review = _await_review(api)
    assert not api.answer_review(identifier, True)['ok']  # Document id is not the review UUID.
    assert not api.answer_review('wrong-review', False)['ok']
    assert thread.is_alive()
    assert api.answer_review(review['id'], True)['ok']
    thread.join(5)
    assert not thread.is_alive()
    assert not api.answer_review(review['id'], True)['ok']


def test_review_requires_an_extracted_document(workbench):
    api, identifier, _ = workbench
    with pytest.raises(ValueError, match='Extract this document first'):
        api.wait_for_review(identifier)


def test_cms_credential_prompt_is_keyed_private_and_lock_free(workbench, monkeypatch):
    from AI.workbench import window_layout
    api, _, _ = workbench
    api._memory_only = True
    api._compact = True
    polled = threading.Event()

    def resize(window, compact):
        reader = threading.Thread(target=lambda: (api.get_state(), polled.set()), daemon=True)
        reader.start()
        assert polled.wait(2), 'Resize held the state lock'
        assert compact is False

    monkeypatch.setattr(window_layout, 'set_compact', resize)
    returned = []
    worker = threading.Thread(target=lambda: returned.append(api.request_cms_credentials()), daemon=True)
    worker.start()
    assert polled.wait(2)
    prompt = api.get_state()['value']['pending_prompt']
    assert prompt['kind'] == 'credentials'
    assert worker.is_alive()
    assert not api.answer_prompt('stale-id', {'username': 'synthetic-user', 'password': 'synthetic-secret'})['ok']
    with pytest.raises(RuntimeError, match='already pending'):
        api.confirm('Concurrent prompt', 'Synthetic instruction')
    wrong = api.answer_prompt(prompt['id'], {'username': '', 'password': 'synthetic-secret'})
    assert not wrong['ok'] and 'synthetic-secret' not in wrong['error']
    assert not api.answer_prompt(prompt['id'], True)['ok']
    assert not api.set_compact(True)['ok']
    # An independent thread polls while the worker is parked on its Event.
    seen = []
    poll = threading.Thread(target=lambda: seen.append(api.get_state()['value']), daemon=True)
    poll.start(); poll.join(2)
    assert seen and seen[0]['pending_prompt']['id'] == prompt['id']
    assert 'synthetic-secret' not in json.dumps(seen)
    assert api.answer_prompt(prompt['id'], {'username': 'synthetic-user', 'password': 'synthetic-secret', 'remember': False}) == {'ok': True, 'value': True}
    worker.join(2)
    assert not worker.is_alive()
    assert returned == [('synthetic-user', 'synthetic-secret')]
    assert 'synthetic-secret' not in json.dumps(api.get_state())
    stale = api.answer_prompt(prompt['id'], {'username': 'synthetic-user', 'password': 'synthetic-secret'})
    assert not stale['ok'] and 'synthetic-secret' not in str(stale)
    assert api.get_state()['value']['compact'] is False


def test_cms_credential_cancel_and_reader_isolation(workbench):
    reader, _, _ = workbench
    with pytest.raises(RuntimeError, match='attended'):
        reader.request_cms_credentials()
    assert reader.get_state()['value']['pending_prompt'] is None
    api = WorkbenchAPI(memory_only=True)
    answers = []
    worker = threading.Thread(target=lambda: answers.append(api.request_cms_credentials()), daemon=True)
    worker.start()
    deadline = time.monotonic() + 2
    while not api.get_state()['value']['pending_prompt'] and time.monotonic() < deadline:
        time.sleep(.01)
    prompt = api.get_state()['value']['pending_prompt']
    assert prompt and api.answer_prompt(prompt['id'], None)['ok']
    worker.join(2)
    assert answers == [None]


def test_final_code_cms_wiring_requests_credentials_before_search(monkeypatch):
    import importlib.util
    from pathlib import Path
    import sys
    from unittest.mock import Mock
    path = Path(__file__).resolve().parents[2] / 'Final Code' / 'finalcode' / 'live_wiring.py'
    spec = importlib.util.spec_from_file_location('synthetic_cms_wiring', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    make_cms = module.make_cms
    calls = []
    session = SimpleNamespace(set_credentials=lambda user, secret: calls.append(('set', bool(user), bool(secret))), CmsAuthenticationRejected=type('CmsAuthenticationRejected', (RuntimeError,), {}))
    runner = SimpleNamespace(run_live=lambda data, **kwargs: (calls.append(('search',)), {'status': 'completed'})[1])
    monkeypatch.setitem(sys.modules, 'fcm_intake.cms.session', session)
    monkeypatch.setitem(sys.modules, 'fcm_intake.workflows.reopen_flow', runner)
    api = WorkbenchAPI(memory_only=True)
    monkeypatch.setattr(api, 'request_cms_credentials', lambda: ('synthetic-user', 'synthetic-secret'))
    cms = make_cms(notify=api.notify, confirm=api.confirm)
    assert cms({'customer': 'SYNTHETIC'})['status'] == 'completed'
    assert calls == [('set', True, True), ('search',), ('set', False, False)]
    calls.clear()
    monkeypatch.setattr(api, 'request_cms_credentials', lambda: None)
    assert cms({'customer': 'SYNTHETIC'})['reason'] == 'cms_sign_in_cancelled'
    assert not calls
    monkeypatch.setattr(api, 'request_cms_credentials', lambda: ('synthetic-user', 'synthetic-secret'))
    runner.run_live = Mock(side_effect=RuntimeError('synthetic-secret'))
    failing_cms = make_cms(notify=api.notify, confirm=api.confirm)
    with pytest.raises(RuntimeError) as caught:
        failing_cms({'customer': 'SYNTHETIC'})
    assert 'synthetic-secret' not in str(caught.value)
    assert calls[-1] == ('set', False, False)
