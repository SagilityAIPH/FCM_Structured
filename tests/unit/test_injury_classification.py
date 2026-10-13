"""RRS Body Part / Injury Type / Injury Cause classification in the AI PDF Reader (synthetic text only)."""
import json

import pytest
from openpyxl import load_workbook

from AI import injury_classification as ic
from AI import referral_schema as schema
from AI.daily_output import CLASSIFICATION_SHEET, check_layout, read_record, save_daily_output
from tests.unit.test_referral_schema import complete_payload
from tests.unit.test_workbench import extract, finish, workbench  # noqa: F401  (pytest fixture)

REAL_LOAD = ic.load_options
LISTS = {**ic.provisional_options(), 'provisional': False}


@pytest.fixture(autouse=True)
def fixed_lists(monkeypatch):
    # Independent of whether a live AI/rrs_dropdowns.json has been captured on this machine.
    def fixed(path=None):
        return LISTS
    fixed.cache_clear = lambda: None
    monkeypatch.setattr(ic, 'load_options', fixed)


def texts(accident='Not found', injury='Not found', special='Not found'):
    return {'Accident Description': accident, 'Injury Description': injury, 'Special Instructions': special}


def by_field(detail):
    return {item['field']: item for item in detail}


def test_rules_fallback_uses_legacy_inputs():
    fields = texts('Employee lifted a heavy box', 'left knee strain')
    detail = by_field(ic.classify(fields, {}))
    assert [fields[f] for f in ic.FIELDS] == ['Knee', 'Strain', 'Lifting']
    assert {item['source'] for item in detail.values()} == {'Rules'}
    assert not any(item['review'] for item in detail.values())


def test_body_part_blends_special_instructions_like_legacy():
    fields = texts('Employee lifted a heavy box', 'left knee strain', 'Please arrange a face-to-face visit.')
    ic.classify(fields, {})
    assert fields['Body Part'] == 'Multiple Body Parts'  # owner-chosen legacy rule; AI is the guard


def test_type_and_cause_use_special_instructions_only_as_fallback():
    fields = texts('Employee was hurt at work', special='Lifted boxes; lower back strain')
    ic.classify(fields, {})
    assert (fields['Injury Type'], fields['Injury Cause']) == ('Strain', 'Lifting')
    fields = texts('Cut hand on a blade', special='Prior fracture noted')
    ic.classify(fields, {})
    assert fields['Injury Type'] == 'Laceration'


def test_ai_choice_wins_and_disagreement_is_flagged():
    fields = texts('Employee lifted a heavy box', 'left knee strain', 'Please arrange a face-to-face visit.')
    detail = by_field(ic.classify(fields, {'Body Part': 'Knee', 'Injury Type': 'strain', 'Injury Cause': 'LIFTING'}))
    assert [fields[f] for f in ic.FIELDS] == ['Knee', 'Strain', 'Lifting']
    assert detail['Body Part']['source'] == 'AI'
    assert 'Multiple Body Parts' in detail['Body Part']['review']
    assert not detail['Injury Type']['review'] and not detail['Injury Cause']['review']


def test_ai_answer_outside_the_list_falls_back_to_rules():
    fields = texts(injury='left knee strain')
    detail = by_field(ic.classify(fields, {'Body Part': 'Kneecap area'}))
    assert fields['Body Part'] == 'Knee'
    assert detail['Body Part']['source'] == 'Rules'
    assert 'not an item' in detail['Body Part']['review']


def test_rule_default_is_not_reported_as_a_disagreement():
    fields = texts('Something happened at work')
    detail = by_field(ic.classify(fields, {'Injury Type': 'Sprain or Tear'}))
    assert fields['Injury Type'] == 'Sprain or Tear'
    assert detail['Injury Type']['rule_value'] == ic.TYPE_DEFAULT
    assert not detail['Injury Type']['review']


def test_manual_choice_wins_and_non_list_value_is_flagged():
    fields = texts(injury='left knee strain')
    detail = by_field(ic.classify(fields, {'Body Part': 'Knee'}, manual={'Body Part': 'Hand'}))
    assert (fields['Body Part'], detail['Body Part']['source'], detail['Body Part']['review']) == ('Hand', 'Manual', '')
    detail = by_field(ic.classify(fields, {}, manual={'Body Part': 'Pinky'}))
    assert fields['Body Part'] == 'Pinky' and 'not an item' in detail['Body Part']['review']


def test_no_injury_text_is_not_found():
    fields = texts()
    detail = by_field(ic.classify(fields, {}))
    assert [fields[f] for f in ic.FIELDS] == ['Not found'] * 3
    assert {item['source'] for item in detail.values()} == {'None'}


def test_provisional_lists_hold_every_rule_value_under_its_rrs_name():
    lists = ic.provisional_options()
    assert 'Skull' in lists['Body Part'] and 'Head' not in lists['Body Part']
    assert 'Holding or Carrying' in lists['Injury Cause']
    assert ic.CAUSE_DEFAULT in lists['Injury Cause'] and ic.TYPE_DEFAULT in lists['Injury Type']
    for pattern_value in [v for _, v in ic.CAUSE_RULES] + ic.FALL_VALUES:
        assert ic.list_item(ic.CAUSE_ALIASES.get(pattern_value, pattern_value), lists['Injury Cause'])


def test_live_lists_are_loaded_and_partial_files_fall_back(tmp_path):
    REAL_LOAD.cache_clear()
    assert REAL_LOAD(tmp_path / 'missing.json')['provisional'] is True
    live = tmp_path / 'rrs.json'
    live.write_text(json.dumps({'Body Part': ['None', 'Knee', ''], 'Injury Type': ['Strain'],
                                'Injury Cause': ['Lifting']}), encoding='utf-8')
    loaded = REAL_LOAD(live)
    assert loaded == {'Body Part': ['Knee'], 'Injury Type': ['Strain'], 'Injury Cause': ['Lifting'], 'provisional': False}
    partial = tmp_path / 'partial.json'
    partial.write_text(json.dumps({'Body Part': ['Knee']}), encoding='utf-8')
    assert REAL_LOAD(partial)['provisional'] is True
    REAL_LOAD.cache_clear()


def test_prompt_asks_for_list_items_but_canonical_keys_are_unchanged():
    assert '"RRS Classification"' in schema.FIELD_ONLY_PROMPT
    for field in ic.FIELDS:
        assert f'Allowed {field} items:' in schema.FIELD_ONLY_PROMPT
        assert field not in schema.REQUIRED_FIELDS and field in schema.OPTIONAL_FIELDS
    assert schema.FIELD_GROUPS['RRS Classification'] == list(ic.FIELDS)


def classified_payload(ai=True):
    data = complete_payload()
    data.update({'Accident Description': 'Employee lifted a heavy box', 'Injury Description': 'left knee strain'})
    if ai:
        data['RRS Classification'] = {'Body Part': 'Knee', 'Injury Type': 'Strain', 'Injury Cause': 'Lifting'}
    return data


def test_parse_classifies_after_extraction_without_changing_next_step():
    fields = schema.parse_field_block(json.dumps(classified_payload()))
    assert [fields[f] for f in ic.FIELDS] == ['Knee', 'Strain', 'Lifting']
    assert {item['source'] for item in fields.classification} == {'AI'}
    exported = schema.export_payload(fields)
    assert exported['RRS Classification'] == {'Body Part': 'Knee', 'Injury Type': 'Strain', 'Injury Cause': 'Lifting'}
    assert exported['RRS Classification Detail'] == fields.classification
    assert exported['NEXT STEP']['status'] == schema.completeness(fields)['status'] == 'Passed'
    # Older model output without the object still parses; the legacy rules fill the values.
    older = schema.parse_field_block(json.dumps(classified_payload(ai=False)))
    assert [older[f] for f in ic.FIELDS] == ['Knee', 'Strain', 'Lifting']
    assert {item['source'] for item in older.classification} == {'Rules'}


def test_workbook_keeps_referrals_layout_and_round_trips(tmp_path):
    fields = schema.parse_field_block(json.dumps(classified_payload()))
    path, record_id = save_daily_output(fields, 'synthetic.pdf', directory=tmp_path)
    save_daily_output(fields, 'synthetic.pdf', directory=tmp_path, record_id=record_id, update_existing=True)
    book = load_workbook(path)
    check_layout(book)  # existing readers (CMSCustomerSearch) accept the workbook unchanged
    assert 'Body Part' not in [cell.value for cell in book['Referrals'][1]]
    assert book[CLASSIFICATION_SHEET].max_row == 2  # updating the same record never duplicates
    book.close()
    saved, _ = read_record(path, record_id)
    assert [saved[f] for f in ic.FIELDS] == ['Knee', 'Strain', 'Lifting']
    assert saved.classification[0]['source'] == 'AI'


def test_workbench_shows_fields_recomputes_on_edit_and_keeps_manual_choice(workbench, monkeypatch):
    api, identifier, _ = workbench
    extract(monkeypatch, classified_payload(ai=False))
    api.start_processing([identifier]); finish(api)

    def shown():
        detail = api.get_document(identifier)['value']
        group = next(g for g in detail['groups'] if g['name'] == 'RRS Classification')
        return {f['key']: (f['value'], f['source']) for f in group['fields']}

    assert shown() == {'Body Part': ('Knee', 'Rules'), 'Injury Type': ('Strain', 'Rules'), 'Injury Cause': ('Lifting', 'Rules')}
    doc = api._documents[identifier]
    edit = {'key': 'Injury Description', 'provider_index': None, 'value': 'right shoulder sprain'}
    assert api.save_fields(identifier, [edit], doc['revision'])['ok']; finish(api)
    assert shown()['Body Part'] == ('Shoulder(s)', 'Rules') and shown()['Injury Type'] == ('Sprain or Tear', 'Rules')
    assert api.save_fields(identifier, [{'key': 'Body Part', 'provider_index': None, 'value': 'Hand'}], doc['revision'])['ok']
    finish(api)
    edit = {'key': 'Injury Description', 'provider_index': None, 'value': 'left ankle contusion'}
    assert api.save_fields(identifier, [edit], doc['revision'])['ok']; finish(api)
    assert shown()['Body Part'] == ('Hand', 'Manual') and shown()['Injury Type'] == ('Contusion', 'Rules')
    saved, _ = read_record(doc['workbook'], doc['record_id'])
    assert (saved['Body Part'], saved['Injury Type']) == ('Hand', 'Contusion')


def test_export_script_reads_three_dropdowns_without_selecting(tmp_path):
    from AI import export_rrs_dropdowns as export

    class Combo:
        def __init__(self, items): self.items = items

    class FakeAdapter:
        def unique(self, pattern):
            assert pattern == export.BUILDER_TITLE
            return 'builder'
        def control(self, parent, ident, kind):
            assert (parent, kind) == ('builder', 'ComboBox')
            return Combo({'cbBodyPart': ['None', 'Knee', 'Knee', 'Skull'], 'cbInjuryType': ['Strain'],
                          'cbInjuryCause': ['Lifting', '']}[ident])
        def _combo_options(self, combo): return combo.items

    lists = export.capture(FakeAdapter())
    assert lists == {'Body Part': ['Knee', 'Skull'], 'Injury Type': ['Strain'], 'Injury Cause': ['Lifting']}
    assert 'Hand' in export.unmapped_rule_values(lists)['Body Part']
    path = export.write(lists, tmp_path / 'rrs.json')
    assert REAL_LOAD(path)['Body Part'] == ['Knee', 'Skull']
    REAL_LOAD.cache_clear()
