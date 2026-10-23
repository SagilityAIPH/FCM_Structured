"""Synthetic application persistence hooks, independent of live integrations."""
from contextlib import nullcontext
from io import BytesIO
import json
from pathlib import Path
import threading

import fitz
from openpyxl import load_workbook
import pytest

from AI.workbench import backend
from AI.workbench.backend import WorkbenchAPI
from AI.referral_json import ReferralRecord
from tests.unit.test_referral_schema import complete_payload


def ready(tmp_path, monkeypatch):
    api = WorkbenchAPI(tmp_path / "workbooks")
    api._key = "SYNTHETIC-SECRET"
    monkeypatch.setattr(api, "_client", lambda: nullcontext(None))
    def reasoning(*args):
        payload = complete_payload()
        payload.update({"First Name": "Synthetic", "Last Name": "Original"})
        fields = backend.core.force_exact_field_output(json.dumps(payload))[1]
        return "Synthetic formatted output", fields, "SYNTHETIC MODEL RESPONSE", .1
    monkeypatch.setattr(backend.core, "run_reasoning", reasoning)
    monkeypatch.setattr(api, "_enrich", lambda fields: (fields, []))
    original = ReferralRecord.save
    def unlocked(record, *args, **kwargs):
        assert not api._lock._is_owned(), "referral I/O while holding API lock"
        return original(record, *args, **kwargs)
    monkeypatch.setattr(ReferralRecord, "save", unlocked)
    return api


def finish(api):
    api._thread.join(10)
    assert not api._thread.is_alive()


def record(api, identifier):
    detail = api.get_document(identifier)["value"]
    path = Path(detail["referral_path"])
    return path, json.loads(path.read_text(encoding="utf-8"))


def test_reader_multiple_documents_edit_renames_and_workbook_contract(tmp_path, monkeypatch):
    api = ready(tmp_path, monkeypatch)
    ids = [api._import_text(f"synthetic-{index}.txt", f"Synthetic selectable document {index} only.") for index in range(2)]
    assert api.start_processing(ids)["ok"]
    finish(api)
    first_path, first = record(api, ids[0])
    second_path, second = record(api, ids[1])
    assert first_path != second_path
    assert first["record_id"] != second["record_id"]
    assert first["source_file_name"] == "synthetic-0.txt"
    assert first["source_text"] == api._documents[ids[0]]["text"]
    assert second["source_text"] == api._documents[ids[1]]["text"]
    assert first["source_text"] != second["source_text"]
    assert "sent_text" not in first and "sent_text" not in second
    assert first["record_id"] == api._documents[ids[0]]["record_id"]
    assert first["status"] in ("Completed", "Needs Review")
    assert [s["reason"] for s in first["saves"]] == ["after_extraction", "address_check"]
    workbook = Path(api._documents[ids[0]]["workbook"])
    before = load_workbook(workbook)
    rows = {sheet.title: sheet.max_row for sheet in before}
    before.close()
    revision = api.get_document(ids[0])["value"]["revision"]
    api._documents[ids[0]]["text"] = "Synthetic later displayed text"
    api._documents[ids[0]]["model_sent_text"] = "Synthetic later sent text"
    api._settings["document_limit"] = 3
    assert api.save_fields(ids[0], [{"key": "Last Name", "provider_index": None, "value": "Revised"}], revision)["ok"]
    finish(api)
    new_path, updated = record(api, ids[0])
    assert not first_path.exists() and new_path != first_path
    assert updated["record_id"] == first["record_id"]
    assert updated["ai_extracted"]["Last Name"] == "Original"
    assert updated["current"]["Last Name"] == "Revised"
    assert updated["source_text"] == first["source_text"]
    assert "sent_text" not in updated
    assert [s["reason"] for s in updated["saves"]] == ["after_extraction", "address_check", "save_changes", "address_check"]
    assert len(list(new_path.parent.glob("*.json"))) == 2
    assert record(api, ids[1])[1] == second
    after = load_workbook(workbook)
    assert {sheet.title: sheet.max_row for sheet in after} == rows
    after.close()
    serialized = json.dumps(updated)
    assert "SYNTHETIC MODEL RESPONSE" not in serialized and "SYNTHETIC-SECRET" not in serialized
    assert api.get_state()["value"]["referral_notice"].startswith("Saved")


def test_reader_extraction_survives_address_service_failure(tmp_path, monkeypatch):
    api = ready(tmp_path, monkeypatch)
    monkeypatch.setattr(api, "_enrich", lambda fields: (_ for _ in ()).throw(RuntimeError("Synthetic service failure")))
    identifier = api._import_text("synthetic.txt", "Synthetic text only.")
    assert api.start_processing([identifier])["ok"]
    finish(api)
    path, data = record(api, identifier)
    assert path.exists()
    assert [s["reason"] for s in data["saves"]] == ["after_extraction", "processing_failed"]
    assert data["status"] == "Failed"
    assert data["ai_extracted"]["First Name"] == "Synthetic"
    assert data["source_text"] == api._documents[identifier]["text"]
    assert api.get_document(identifier)["value"]["status"] == "Failed"


def test_reader_full_source_and_actual_truncated_input_saved_before_enrichment(tmp_path, monkeypatch):
    api = ready(tmp_path, monkeypatch)
    api._settings["document_limit"] = 18
    identifier = api._import_text("synthetic-truncated.txt", "Synthetic full document source SYNTHETIC-SECRET beyond the model limit.")
    full_source = api._documents[identifier]["text"]
    seen = []
    original = backend.core.run_reasoning
    def reasoning(*args):
        seen.append((args[2], args[3]))
        api._settings["document_limit"] = 500
        return original(*args)
    monkeypatch.setattr(backend.core, "run_reasoning", reasoning)
    def enrich(fields):
        _, initial = record(api, identifier)
        assert [item["reason"] for item in initial["saves"]] == ["after_extraction"]
        assert initial["source_text"] == full_source.replace("SYNTHETIC-SECRET", "[redacted]")
        assert initial["sent_text"] == seen[0][0]
        return fields, []
    monkeypatch.setattr(api, "_enrich", enrich)
    assert api.start_processing([identifier])["ok"]
    finish(api)
    _, data = record(api, identifier)
    assert seen == [(full_source[:18], full_source)]
    assert data["sent_text"] == full_source[:18]
    assert data["source_text"] == full_source.replace("SYNTHETIC-SECRET", "[redacted]")
    assert "SYNTHETIC-SECRET" not in json.dumps(data)
    assert "SYNTHETIC MODEL RESPONSE" not in json.dumps(data)


def test_reader_uploaded_synthetic_pdf_preserves_joined_page_source(tmp_path, monkeypatch):
    api = ready(tmp_path, monkeypatch)
    path = tmp_path / "synthetic-two-page.pdf"
    with fitz.open() as pdf:
        for text in ("Synthetic first page referral text.", "Synthetic second page referral text."):
            pdf.new_page().insert_text((40, 60), text)
        pdf.save(path)
    stream = BytesIO(path.read_bytes())
    stream.name = path.name
    expected = backend.core.join_pages(backend.core.extract_document_pages(stream))
    identifier = api._import_paths([path])[0]
    assert api.start_processing([identifier])["ok"]
    finish(api)
    _, data = record(api, identifier)
    assert data["source_text"] == expected == api._documents[identifier]["text"]
    assert "Synthetic first page referral text." in data["source_text"]
    assert "Synthetic second page referral text." in data["source_text"]
    assert data["source_file_name"] == path.name
    assert "sent_text" not in data
    assert "SYNTHETIC MODEL RESPONSE" not in json.dumps(data)


def test_reader_stored_sent_text_matches_actual_core_model_prompt(tmp_path, monkeypatch):
    from AI import bedrock_core
    api = ready(tmp_path, monkeypatch)
    api._settings["document_limit"] = 20
    identifier = api._import_text("synthetic-prompt.txt",
                                  "Synthetic leading text. SYNTHETIC EXCLUDED SOURCE TAIL")
    source = api._documents[identifier]["text"]
    prompts = []
    def llm(**kwargs):
        prompts.append(kwargs["prompt"])
        return json.dumps(complete_payload())
    def reasoning(*args):
        return bedrock_core.run_reasoning(*args, llm_call=llm)
    monkeypatch.setattr(backend.core, "run_reasoning", reasoning)
    assert api.start_processing([identifier])["ok"]
    finish(api)
    _, data = record(api, identifier)
    assert data["source_text"] == source
    assert data["sent_text"] == source[:20]
    assert len(prompts) == 1
    assert data["sent_text"] in prompts[0]
    assert "SYNTHETIC EXCLUDED SOURCE TAIL" not in prompts[0]


def test_reader_reextract_starts_new_referral_id_matching_workbook(tmp_path, monkeypatch):
    api = ready(tmp_path, monkeypatch)
    identifier = api._import_text("synthetic.txt", "Synthetic text only.")
    api.start_processing([identifier]); finish(api)
    old_path, old = record(api, identifier)
    api.start_processing([identifier]); finish(api)
    new_path, new = record(api, identifier)
    assert old_path.exists() and new_path.exists() and old_path != new_path
    assert new["record_id"] != old["record_id"]
    assert new["record_id"] == api._documents[identifier]["record_id"]


@pytest.mark.parametrize("memory_only", [False, True])
def test_saved_edits_survive_address_failure_in_both_apps(tmp_path, monkeypatch, memory_only):
    api = ready(tmp_path, monkeypatch)
    api._memory_only = memory_only
    identifier = api._import_text("synthetic.txt", "Synthetic selectable text.")
    assert api.start_processing([identifier])["ok"]
    finish(api)
    original_path, original = record(api, identifier)
    monkeypatch.setattr(api, "_enrich", lambda fields: (_ for _ in ()).throw(RuntimeError("Synthetic address outage")))
    revision = api.get_document(identifier)["value"]["revision"]
    assert api.save_fields(identifier, [{"key": "Last Name", "provider_index": None, "value": "Corrected"}], revision)["ok"]
    finish(api)
    new_path, saved = record(api, identifier)
    assert not original_path.exists() and new_path.exists()
    assert saved["record_id"] == original["record_id"]
    assert saved["ai_extracted"]["Last Name"] == "Original"
    assert saved["current"]["Last Name"] == "Corrected"
    assert saved["saves"][-1]["reason"] == "save_changes"


@pytest.mark.parametrize("target", ["json", "workbook"])
def test_polling_works_during_blocked_persistence(tmp_path, monkeypatch, target):
    api = ready(tmp_path, monkeypatch)
    entered, release, polled = threading.Event(), threading.Event(), threading.Event()
    original = ReferralRecord.save if target == "json" else backend.save_daily_output
    def blocked(*args, **kwargs):
        entered.set()
        assert release.wait(5)
        return original(*args, **kwargs)
    monkeypatch.setattr(ReferralRecord if target == "json" else backend,
                        "save" if target == "json" else "save_daily_output", blocked)
    identifier = api._import_text("synthetic.txt", "Synthetic selectable text.")
    assert api.start_processing([identifier])["ok"]
    assert entered.wait(5)
    def poll():
        assert api.get_state()["ok"]
        polled.set()
    poller = threading.Thread(target=poll, daemon=True)
    poller.start()
    try:
        assert polled.wait(2), "polling blocked behind persistence"
    finally:
        release.set()
        poller.join(5)
        finish(api)
